"""Deterministic exact-date driver defect reports; source workbooks are never modified."""
from __future__ import annotations
import html,io,re
from pathlib import Path
import openpyxl
from tools.fleet.overflow.report import normalize,validate_date
from integrations.hermes.function_domain.scoped import ScopedReader
from integrations.hermes.function_domain.extract import zip_safe

SOURCE_NAME='گزارش روزانه رانندگان2.xlsx'
SECTIONS={'mechanical':'شرح معایب مکانیکی','metalwork':'شرح معایب آهنگری'}
EMPTY='شرح خرابی برای این تاریخ ثبت نشده است'
STALE='گزارش این تاریخ به‌روز نشده است؛ شرح خرابی برای این تاریخ ثبت نشده است'

def key(value):return normalize(value).replace('آ','ا').replace(' ','')

def load_driver(date,reader=None,*,data=None):
    date=validate_date(date);reader=reader or ScopedReader()
    data=reader.snapshot(SOURCE_NAME) if data is None else data;zip_safe(data)
    book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
    try:
        matches=[];dates=[]
        for sheet in book:
            found=set()
            for row in sheet.iter_rows(max_row=1,max_col=100,values_only=True):
                for cell in row:
                    for raw in re.findall(r'(?<!\d)1[34]\d{2}\s*[/.-]\s*\d{1,2}\s*[/.-]\s*\d{1,2}(?!\d)',normalize(cell)):
                        found.add(validate_date(re.sub(r'\s+','',raw)))
            if len(found)>1:
                if date in found:raise ValueError('Ambiguous driver sheet header dates')
                continue
            if found:
                actual=found.pop();dates.append(actual)
                if actual==date:matches.append(sheet)
        if len(matches)>1:raise ValueError('Multiple driver sheets have the requested date')
        base={'date':date,'latest_date':max(dates) if dates else None,'sections':{}}
        if not matches:
            for name,label in SECTIONS.items():base['sections'][name]={'title':label,'rows':[],'status':'date_missing','message':STALE}
            return base
        ws=matches[0];base['sheet']=ws.title
        if ws.max_row>10000 or ws.max_column>128:raise ValueError('Driver sheet dimensions exceed safe reporting limit')
        headers=next(ws.iter_rows(min_row=2,max_row=2,values_only=True),())
        labels=['ردیف','نوع دستگاه','کد جدید',*SECTIONS.values()]
        indices=[]
        for label in labels:
            found=[i for i,v in enumerate(headers) if key(v)==key(label)]
            if len(found)!=1:raise ValueError('Missing or ambiguous driver report column: '+label)
            indices.append(found[0])
        sections={name:[] for name in SECTIONS}
        for row in ws.iter_rows(min_row=3,values_only=True):
            values=[row[i] if i<len(row) else None for i in indices]
            if not any(values[:3]):continue
            for index,name in enumerate(SECTIONS,3):
                if normalize(values[index]):sections[name].append(values[:3]+[str(values[index])])
        for name,label in SECTIONS.items():
            rows=sections[name];base['sections'][name]={'title':label,'rows':rows,'status':'ready' if rows else 'no_defects','message':'' if rows else EMPTY}
        return base
    finally:book.close()

def render_driver(report,output_dir=None):
    import fitz
    out=Path(output_dir) if output_dir is not None else None
    if out is not None:out.mkdir(parents=True,exist_ok=True)
    paths=[]
    for section,content in report['sections'].items():
        title=content['title']+' — '+report['date']
        headers=['ردیف','نوع دستگاه','کد دستگاه',content['title']]
        def cells(values,tag):return ''.join(f'<{tag} dir="auto">{html.escape(str(v if v is not None else "—"))}</{tag}>' for v in reversed(values))
        table='<tr>'+cells(headers,'th')+'</tr>'
        for i,row in enumerate(content['rows']):
            color='#f0f5f8' if i%2==0 else '#ffffff'
            table+=f'<tr style="background:{color}">'+cells(row,'td')+'</tr>'
        body=f'<html><body><h2 dir="rtl">{html.escape(title)}</h2><p dir="rtl">تاریخ گزارش: {report["date"]}</p>'
        if content['message']:body+='<h3 dir="rtl">'+html.escape(content['message'])+'</h3>'
        else:body+='<table>'+table+'</table>'
        body+='</body></html>'
        css='body{font-family:sans-serif;font-size:19px;color:#173448} h2{font-size:32px} h3{font-size:26px} table{border-collapse:collapse;width:100%} td,th{border:1px solid #becdd6;padding:10px 6px;text-align:center} th{background:#173448;color:white} td:first-child{width:65%;text-align:right} p{font-size:18px}'
        height=1200
        while True:
            with fitz.open() as doc:
                page=doc.new_page(width=1450,height=height)
                spare,_=page.insert_htmlbox(fitz.Rect(25,20,1425,height-25),body,css=css,scale_low=1)
                if spare>=0:
                    image=page.get_pixmap(clip=fitz.Rect(0,0,1450,min(height,height-25-spare+40))).tobytes('png')
                    if len(image)>9*1024*1024:raise ValueError('Driver photo exceeds transport limit')
                    if out is None:paths.append(image)
                    else:
                        path=out/f'driver-{section}-{report["date"].replace("/","-")}.png'
                        path.write_bytes(image);paths.append(str(path))
                    break
            height+=1000
            if height>8200:raise ValueError('Full driver section exceeds two-image layout; no data truncated')
    if len(paths)!=2:raise ValueError('Exactly two driver images required')
    return paths

def build_driver(date,output_dir,reader=None):
    report=load_driver(date,reader)
    return {'ok':True,'report':report,'images':render_driver(report,output_dir)}


def render_pdf_notice(report, section, output):
    """Searchable vector PDF for a missing day/empty section; never a raster wrapper."""
    import fitz
    content=report['sections'][section]
    body=f'<html><body><h2 dir="rtl">{html.escape(content["title"])}</h2><p dir="rtl">تاریخ گزارش: {report["date"]}</p><h3 dir="rtl">{html.escape(content["message"])}</h3></body></html>'
    css='body{font-family:sans-serif;font-size:16px;color:#173448} h2{font-size:22px} h3{font-size:18px}'
    with fitz.open() as doc:
        page=doc.new_page(width=595,height=842)
        spare,_=page.insert_htmlbox(fitz.Rect(35,35,560,800),body,css=css,scale_low=1)
        if spare<0:raise ValueError('PDF notice does not fit; no truncation')
        doc.subset_fonts();doc.save(output,garbage=4,deflate=True)


def build_driver_pdf(date, output_dir, reader=None):
    """Two native Excel PDFs from one immutable snapshot of the exact requested day."""
    from tempfile import TemporaryDirectory
    from tools.fleet.repairs.report import export_pdf, section_columns
    reader=reader or ScopedReader()
    data=reader.snapshot(SOURCE_NAME)
    report=load_driver(date,reader,data=data)
    out=Path(output_dir).resolve();out.mkdir(parents=True,exist_ok=True)
    documents=[]
    with TemporaryDirectory(prefix='driver-excel-',dir=out) as directory:
        work=Path(directory);snapshot=work/'source.xlsx'
        # Only this private temporary copy is opened by Excel. Source is never edited.
        if any(c['status']=='ready' for c in report['sections'].values()):snapshot.write_bytes(data)
        for section in SECTIONS:
            path=out/f'driver-{section}-{report["date"].replace("/","-")}.pdf'
            content=report['sections'][section]
            if content['status']=='ready':
                columns=section_columns(snapshot,report['sheet'],section)
                export_pdf(snapshot,report['sheet'],path,work,columns)
            else:
                render_pdf_notice(report,section,path)
            import fitz
            with fitz.open(path) as pdf:
                if not pdf.page_count or not any(p.get_text().strip() for p in pdf):raise ValueError('Empty driver PDF')
            if path.stat().st_size>45*1024*1024:raise ValueError('Driver PDF exceeds document transport limit')
            documents.append(str(path))
    if len(documents)!=2:raise ValueError('Exactly two driver PDFs required')
    return {'ok':True,'report':report,'documents':documents,'format':'pdf'}

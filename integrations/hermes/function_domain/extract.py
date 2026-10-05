"""Safe, paged extraction. Never evaluate macros, formulas, external links or binaries."""
from __future__ import annotations
import csv,io,json,zipfile
from pathlib import PurePosixPath
from .scoped import ScopedReader

TEXT_FORMATS={'.txt','.csv','.tsv','.log','.md','.json','.yaml','.yml','.xml','.html','.htm','.ini','.cfg','.rtf'}

def decode(data):
    if data.startswith((b'\xff\xfe',b'\xfe\xff')):return data.decode('utf-16')
    for encoding in ['utf-8-sig','cp1256']:
        try:return data.decode(encoding)
        except UnicodeDecodeError:pass
    return data.decode('utf-8',errors='replace')

def zip_safe(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        items=z.infolist()
        if len(items)>20000 or sum(i.file_size for i in items)>256*1024*1024:
            raise ValueError('Compressed document exceeds safe extraction limit')
        if any(i.file_size>64*1024*1024 for i in items):raise ValueError('Compressed member too large')

def extract(reader,relative,*,sheet=None,offset=0,limit=100):
    if type(offset)!=int or offset<0 or type(limit)!=int or not 1<=limit<=500:raise ValueError('Invalid page')
    suffix=PurePosixPath(relative.replace('\\','/')).suffix.lower()
    if suffix not in TEXT_FORMATS | {'.xlsx','.xlsm','.xls','.xlsb','.pdf','.docx','.odt'}:
        return {**reader.metadata(relative),'content_supported':False,'instruction':'Unknown binary: metadata and existing-file attachment only; execution unavailable.'}
    if PurePosixPath(relative.replace('\\','/')).name.startswith('~$'):
        return {**reader.metadata(relative),'content_supported':False,'instruction':'Excel owner/lock file, not workbook content. Metadata and existing-file attachment remain available.'}
    data=reader.snapshot(relative)
    if suffix in {'.xlsx','.xlsm','.xlsb','.docx','.odt'}:
        if not data.startswith(b'PK'):
            return {**reader.metadata(relative),'content_supported':False,'instruction':'Filename has a document extension but valid package content was not recognized. Do not infer data; metadata and existing-file attachment remain available.'}
        zip_safe(data)
    result={'path':relative,'format':suffix,'offset':offset}
    if suffix in {'.xlsx','.xlsm'}:
        import openpyxl
        book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
        try:
            if sheet is None:return {**result,'sheets':book.sheetnames,'instruction':'Select an exact sheet; no latest-sheet selection.'}
            ws=book[sheet];rows=list(ws.iter_rows(min_row=offset+1,max_row=offset+limit,values_only=True))
            return {**result,'sheet':sheet,'rows':rows,'total_rows':ws.max_row,'next_offset':offset+len(rows) if offset+len(rows)<ws.max_row else None}
        finally:book.close()
    if suffix=='.xls':
        import xlrd
        book=xlrd.open_workbook(file_contents=data,on_demand=True)
        try:
            if sheet is None:return {**result,'sheets':book.sheet_names()}
            ws=book.sheet_by_name(sheet);end=min(ws.nrows,offset+limit)
            return {**result,'sheet':sheet,'rows':[ws.row_values(i) for i in range(offset,end)],'total_rows':ws.nrows,'next_offset':end if end<ws.nrows else None}
        finally:book.release_resources()
    if suffix=='.xlsb':
        import pyxlsb
        with pyxlsb.open_workbook(io.BytesIO(data)) as book:
            if sheet is None:return {**result,'sheets':book.sheets}
            with book.get_sheet(sheet) as ws:
                from itertools import islice
                rows=[[c.v for c in row] for row in islice(ws.rows(),offset,offset+limit+1)]
            return {**result,'sheet':sheet,'rows':rows[:limit],'next_offset':offset+limit if len(rows)>limit else None}
    if suffix=='.pdf':
        import fitz
        with fitz.open(stream=data,filetype='pdf') as doc:
            end=min(doc.page_count,offset+min(limit,20))
            return {**result,'pages':[{'page':i+1,'text':doc[i].get_text()} for i in range(offset,end)],'total_pages':doc.page_count,'next_offset':end if end<doc.page_count else None}
    if suffix in {'.docx','.odt'}:
        from defusedxml.ElementTree import fromstring
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            tree=fromstring(z.read('word/document.xml' if suffix=='.docx' else 'content.xml'))
        lines=[''.join(n.itertext()) for n in tree.iter() if n.tag.split('}')[-1] in {'p','h'}]
    elif suffix in TEXT_FORMATS:
        text=decode(data)
        if suffix in {'.csv','.tsv'}:
            rows=list(csv.reader(io.StringIO(text),delimiter='\t' if suffix=='.tsv' else ','))
            end=min(len(rows),offset+limit)
            return {**result,'rows':rows[offset:end],'total_rows':len(rows),'next_offset':end if end<len(rows) else None}
        lines=text.splitlines()
    else:
        return {**reader.metadata(relative),'content_supported':False,'instruction':'Unknown binary: metadata and existing-file attachment only; execution unavailable.'}
    end=min(len(lines),offset+limit)
    return {**result,'lines':lines[offset:end],'total_lines':len(lines),'next_offset':end if end<len(lines) else None}

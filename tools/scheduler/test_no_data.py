"""Exact-day, independent artifact receipts, and bounded late delivery fixtures."""
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
import io, sqlite3, tempfile, unittest
from contextlib import closing
import openpyxl, yaml
from zoneinfo import ZoneInfo
from .runner import tick, connect, _save_receipt, _closed, STATE
from .tasks import ROOT, TransportUnavailable, TransportUncertain
from .no_data import notice_text
from tools.fleet.overflow.report import load_report, ReportError
from integrations.hermes.function_domain.driver_report import load_driver

TZ=ZoneInfo('Asia/Tehran')
DUE=datetime(2026,10,7,9,0,tzinfo=TZ)
TARGET='1405/07/14'

class NoDataTests(unittest.TestCase):
    def setUp(self):
        root=ROOT/'runtime/scheduler-tests';root.mkdir(parents=True,exist_ok=True)
        self.temp=tempfile.TemporaryDirectory(dir=root);self.addCleanup(self.temp.cleanup)
        self.dir=Path(self.temp.name);self.config=self.dir/'config.yaml';self.state=self.dir/'state.db'
        self.users=['A','B'];self.present=False;self.latest='1405/07/12';self.delivery=[]
        self.auth=Mock()
        self.auth.has_capability.return_value=True
        self.auth.resolve_active_recipients.side_effect=lambda _: SimpleNamespace(status='ready',recipients=tuple(self.users))
        self.job=dict(id='test',enabled=True,task='mechanical_overflow',recipient_capability='reports.overflow.mechanical_daily_receive',timezone='Asia/Tehran',misfire=dict(policy='replay',horizon_days=7),trigger=dict(type='cron',hour=9,minute=0),params={})
        self.write()
        self.sender=Mock()
        self.sender.message.side_effect=lambda u,t:self.send('notice',u,t)
        self.sender.photo.side_effect=lambda u,p,t:self.send('report',u,t)
        self.worker_patch=patch('tools.scheduler.tasks.build_report',self.worker);self.worker_patch.start();self.addCleanup(self.worker_patch.stop)
        self.sender_patch=patch('tools.scheduler.tasks.BaleSender',return_value=self.sender);self.sender_patch.start();self.addCleanup(self.sender_patch.stop)
    def write(self):
        self.config.write_text(yaml.safe_dump(dict(version=1,timezone='Asia/Tehran',schedules=[self.job])),encoding='utf-8')
    def send(self,kind,user,text):
        self.delivery.append((kind,user,text));return str(len(self.delivery))
    async def worker(self,date,directory):
        self.assertEqual(date,TARGET)
        if not self.present:return dict(ok=False,reason='date_missing',message='missing',latest_date=self.latest)
        p=Path(directory)/'test.png';p.write_bytes(b'test')
        return dict(ok=True,images=[p],report=dict(date=TARGET))
    def run_at(self,seconds=0):return tick(self.config,self.state,DUE+timedelta(seconds=seconds),authorization=self.auth)
    def row(self):
        with closing(sqlite3.connect(self.state)) as c:
            c.row_factory=sqlite3.Row
            return dict(c.execute('select * from runs where due=?',('2026-10-07T05:30:00+00:00',)).fetchone())
    def receipts(self):
        with closing(sqlite3.connect(self.state)) as c:return c.execute('select recipient_id,delivery_kind,status,message_id from receipts').fetchall()
    def test_present_at_due_report_once_no_notice(self):
        self.present=True;self.run_at();self.run_at(61)
        self.assertEqual([x[:2] for x in self.delivery],[('report','A'),('report','B')]);self.assertEqual(self.row()['status'],'succeeded')
    def test_missing_old_date_notice_contains_target_latest_waits(self):
        self.run_at();self.assertEqual(self.row()['status'],'waiting_for_data')
        self.assertTrue(all(TARGET in x[2] and self.latest in x[2] for x in self.delivery));self.assertEqual(self.row()['data_state'],'missing')
    def test_repeated_ticks_notice_once_independent_recipients(self):
        self.run_at();self.run_at(61);self.run_at(182)
        self.assertEqual([x[:2] for x in self.delivery],[('notice','A'),('notice','B')]);self.assertEqual(sum(r[1]=='notice' for r in self.receipts()),2)
    def test_target_appears_after_notice_report_once_and_succeeds(self):
        self.run_at();self.present=True;self.run_at(61);self.run_at(182)
        self.assertEqual([x[:2] for x in self.delivery],[('notice','A'),('notice','B'),('report','A'),('report','B')]);self.assertEqual(len(self.receipts()),4);self.assertEqual(self.row()['status'],'succeeded')
    def test_revoked_before_later_report_never_sent(self):
        self.run_at();self.users.remove('B');self.present=True;self.run_at(61)
        self.assertEqual([x[1] for x in self.delivery if x[0]=='report'],['A'])
    def test_no_valid_date_truthful_notice(self):
        self.latest=None;self.run_at()
        self.assertNotIn('1405/07/12',self.delivery[0][2]);self.assertIn('تاریخ معتبر دیگری',self.delivery[0][2])
    def test_old_date_never_becomes_actual_report(self):
        self.run_at();self.run_at(61)
        self.sender.photo.assert_not_called();self.assertEqual(self.row()['status'],'waiting_for_data')
    def test_notice_connection_down_retry_safe(self):
        self.sender.check_connection.side_effect=TransportUnavailable();self.run_at()
        self.assertEqual(self.row()['notice_delivery_state'],'retry_wait');self.assertEqual(self.delivery,[])
        self.sender.check_connection.side_effect=None;self.run_at(61)
        self.assertEqual(len(self.delivery),2)
    def test_report_connection_down_delivery_retry(self):
        self.present=True;self.sender.check_connection.side_effect=TransportUnavailable();self.run_at()
        self.assertEqual(self.row()['status'],'retry_wait');self.assertEqual(self.delivery,[])
        self.sender.check_connection.side_effect=None;self.run_at(61);self.assertEqual(self.row()['status'],'succeeded')
    def test_uncertain_notice_no_blind_duplicate_but_later_report_delivered(self):
        self.sender.message.side_effect=TransportUncertain();self.run_at();self.run_at(61)
        self.assertEqual(self.sender.message.call_count,2);self.assertEqual(self.row()['notice_delivery_state'],'uncertain')
        self.present=True;self.run_at(182);self.assertEqual(self.row()['status'],'succeeded');self.assertEqual(len(self.delivery),2)
    def test_uncertain_report_no_blind_duplicate(self):
        self.present=True;self.sender.photo.side_effect=TransportUncertain();self.run_at();self.run_at(61)
        self.assertEqual(self.sender.photo.call_count,2);self.assertEqual(self.row()['status'],'uncertain')
    def test_horizon_expiry_finalizes_without_repeating_notice(self):
        self.run_at();self.run_at(7*86400+1)
        self.assertEqual(self.row()['status'],'no_data_final')
        self.assertEqual(sum(x[0]=='notice' for x in self.delivery if TARGET in x[2]),2)
    def test_late_within_horizon_is_delivered(self):
        self.run_at();self.present=True;self.run_at(6*86400)
        self.assertEqual(self.row()['status'],'succeeded')
    def test_disabled_job_untouched(self):
        self.job['enabled']=False;self.write();self.run_at()
        self.assertEqual(self.delivery,[]);self.assertEqual(self.receipts(),[])
    def test_sent_recipient_excluded_from_notice_and_report_retry(self):
        self.run_at();c=connect(self.state)
        _save_receipt(c,('test','2026-10-07T05:30:00+00:00'),'A','sent','actual-A')
        c.close();self.present=True;self.run_at(61)
        self.assertEqual([x[1] for x in self.delivery if x[0]=='report'],['B'])
    def test_failed_report_recipient_revoked_does_not_keep_retrying(self):
        self.present=True;self.sender.photo.side_effect=lambda u,p,t: (_ for _ in ()).throw(TransportUnavailable()) if u=='B' else self.send('report',u,t)
        self.run_at();self.assertEqual(self.row()['status'],'partial');self.users.remove('B');self.run_at(61)
        self.assertEqual(self.row()['status'],'succeeded')
    def test_stale_sending_notice_becomes_uncertain_and_source_still_checked(self):
        self.run_at();c=connect(self.state)
        c.execute("update receipts set status='sending' where delivery_kind='notice'")
        c.execute("update runs set status='running',next_attempt_at=NULL");c.commit();c.close()
        self.run_at(901);self.assertEqual(self.sender.message.call_count,2);self.assertEqual(self.row()['status'],'waiting_for_data')
        self.present=True;self.run_at(1022);self.assertEqual(self.row()['status'],'succeeded')

    def test_crash_after_first_report_does_not_close_unsent_recipient(self):
        self.present=True;self.run_at()
        with closing(connect(self.state)) as c:
            c.execute("update receipts set status='pending',message_id=NULL where recipient_id='B' and delivery_kind='report'")
            c.execute("update runs set status='running',next_attempt_at=NULL");c.commit()
        self.delivery=[];self.run_at(901);self.run_at(962)
        self.assertEqual([x[:2] for x in self.delivery],[('report','B')]);self.assertEqual(self.row()['status'],'succeeded')

class ParserTests(unittest.TestCase):
    def test_driver_invalid_and_ambiguous_dates_not_latest(self):
        b=openpyxl.Workbook();b.remove(b.active)
        for title,header in [('old','1405/07/12'),('invalid','1405/13/40'),('ambiguous','1405/07/13 1405/07/11')]:b.create_sheet(title).cell(1,1,header)
        data=io.BytesIO();b.save(data);b.close()
        result=load_driver(TARGET,data=data.getvalue())
        self.assertEqual(result['latest_date'],'1405/07/12')
        self.assertTrue(all(x['status']=='date_missing' for x in result['sections'].values()))
    def test_truthful_text_without_latest_and_invalid_latest_rejected(self):
        self.assertIn(TARGET,notice_text('گزارش',TARGET));self.assertNotIn('آخرین اطلاعات معتبر',notice_text('گزارش',TARGET))
        with self.assertRaises(ValueError):notice_text('گزارش',TARGET,'invalid')

class TransportTests(unittest.TestCase):
    def test_post_connection_reset_is_uncertain_but_new_connection_failure_retryable(self):
        import requests
        from urllib3.exceptions import NewConnectionError
        from .tasks import BaleSender
        with patch.dict('os.environ',{'BALE_BOT_TOKEN':'fixture-token'}):sender=BaleSender()
        self.addCleanup(sender.close)
        with patch.object(sender.session,'post',side_effect=requests.exceptions.ConnectionError('reset after upload')):
            with self.assertRaises(TransportUncertain):sender.message('A','fixture')
        with patch.object(sender.session,'post',side_effect=requests.exceptions.ConnectionError(NewConnectionError(None,'connection not established'))):
            with self.assertRaises(TransportUnavailable):sender.message('A','fixture')

class FamilyTests(unittest.TestCase):
    def test_all_enabled_families_return_waiting_and_never_upload_old_files(self):
        from .tasks import TASKS
        auth=Mock();auth.has_capability.return_value=True
        auth.resolve_daily_recipient.return_value=SimpleNamespace(status='ready',recipient='A',holder_count=1)
        auth.resolve_active_recipients.return_value=SimpleNamespace(status='ready',recipients=('A',))
        from .runner import load_config
        jobs=[j for j in load_config()[2] if j.get('enabled',True)]
        self.assertEqual(len(jobs),6)
        report={'date':TARGET,'latest_date':'1405/07/12','sections':{k:{'status':'date_missing'} for k in ('mechanical','metalwork')}}
        missing={'ok':False,'report':report,'documents':[]}
        from unittest.mock import AsyncMock
        for job in jobs:
            with self.subTest(task=job['task']), patch('tools.scheduler.tasks.BaleSender') as sender, patch('tools.scheduler.tasks.build_report',AsyncMock(return_value=dict(ok=False,reason='date_missing',latest_date='1405/07/12'))), patch('integrations.hermes.function_domain.driver_report.build_driver_pdf',return_value=missing), patch('tools.fleet.maintenance_daily.delivery.build',return_value={'status':'skipped','reason':'no_target_date_records','report':dict(date=TARGET,available_dates=['1405/07/12'],source_sha256='fixture',sheets_scanned=1,devices_matched=0,rows_matched=0)}):
                events=[]
                result=TASKS[job['task']][1]('A',{'date':TARGET},authorization=auth,on_receipt=lambda *args:events.append(args))
                self.assertEqual(result['status'],'waiting_for_data')
                self.assertEqual([e[3] for e in events],['notice','notice'])
                sender.return_value.message.assert_called_once()
                sender.return_value.document.assert_not_called();sender.return_value.photo.assert_not_called()

    def test_overflow_latest_ignores_duplicate_or_malformed_dates(self):
        root=ROOT/'runtime/scheduler-tests';root.mkdir(parents=True,exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as directory:
            book=openpyxl.Workbook();book.remove(book.active)
            for name,day in [('old','1405/07/12'),('dupe1','1405/07/13'),('dupe2','1405/07/13'),('invalid','1405/13/40')]:
                sheet=book.create_sheet(name);sheet.append([day]);sheet.append(['ردیف','نام دستگاه','کد دستگاه']);sheet.append([1,'fixture','A'])
            path=Path(directory)/'fixture.xlsx';book.save(path);book.close()
            with self.assertRaises(ReportError) as caught:load_report(TARGET,path)
            self.assertEqual(caught.exception.reason,'date_missing');self.assertEqual(caught.exception.latest_date,'1405/07/12')
            book=openpyxl.Workbook();book.active.append(['invalid']);book.active.append(['ردیف']);book.save(path);book.close()
            with self.assertRaises(ReportError) as caught:load_report(TARGET,path)
            self.assertIsNone(caught.exception.latest_date);self.assertEqual(caught.exception.reason,'date_missing')

    def test_existing_receipt_migration_preserves_sent_report(self):
        root=ROOT/'runtime/scheduler-tests';root.mkdir(parents=True,exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as directory:
            path=Path(directory)/'state.db'
            with closing(sqlite3.connect(path)) as conn:
                conn.execute('CREATE TABLE receipts(schedule_id TEXT,due TEXT,recipient_id TEXT,status TEXT,message_id TEXT,updated TEXT,PRIMARY KEY(schedule_id,due,recipient_id))')
                conn.execute("INSERT INTO receipts VALUES('x','due','A','sent','123','now')");conn.commit()
            with closing(connect(path)) as conn:
                row=conn.execute('select delivery_kind,status,message_id from receipts').fetchone()
                self.assertEqual(tuple(row),('report','sent','123'))
                _save_receipt(conn,('x','due'),'A','sent','notice-id','notice')
                self.assertEqual(_closed(conn,('x','due'))['sent'],{'A'})
                self.assertEqual(_closed(conn,('x','due'))['notice']['sent'],{'A'})

if __name__=='__main__':unittest.main()

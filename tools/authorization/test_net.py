"""NET-1 contracts and real W/R pipelines: temporary stores/files, fake transport."""
from contextlib import closing
import asyncio
import importlib
import json
from pathlib import Path
import sqlite3
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch
import uuid

import openpyxl
import yaml
from tools.authorization.store import AuthorizationStore, ROOT, FUNCTION_READ
from tools.authorization.net import (
    NET_MANAGER, NET_DEPUTY, NET_CAPABILITIES, WO_CAPABILITIES, REPAIRS_CAPABILITIES,
    WO_DOMAIN, REPAIRS_DOMAIN, RESOURCES, domain_decision, migrate_net,
)
from tools.fleet.work_orders.core import db, service, permissions, staff_dispatch, delivery, review
from tools.fleet.work_orders.channels.bale import create_worker, staff_flow
from tools.fleet.repairs import entry_service, entry_bale
from tools.bale_ui.runtime import _main_menu


class NetFixture(unittest.TestCase):
    def setUp(self):
        parent = ROOT/'runtime/net-tests';parent.mkdir(exist_ok=True)
        self.tmp = tempfile.TemporaryDirectory(dir=parent);self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name);self.auth_path = self.root/'auth.db'
        with closing(sqlite3.connect(self.auth_path)) as c,c:
            c.execute('CREATE TABLE channel_users (platform TEXT,user_id TEXT,chat_id TEXT,registration_status TEXT,PRIMARY KEY(platform,user_id))')
            c.executemany('INSERT INTO channel_users VALUES (?,?,?,?)',
                [('bale',u,u,'approved') for u in ['101','202','404','505']] + [('bale','303','303','pending_approval')])
        self.auth = AuthorizationStore(self.auth_path)
        self.auth.migrate(self.root/'backups');self.auth.migrate_admin(self.root/'backups')
        migrate_net(self.auth,self.root/'backups')
        self.fleet = self.root/'fleet.db'
        with closing(sqlite3.connect(self.fleet)) as c,c:
            c.execute('CREATE TABLE machines (id INTEGER PRIMARY KEY,canonical_code TEXT,machine_type_hint TEXT)')
            c.execute("INSERT INTO machines VALUES (1,'HD714','دامپتراک')")
            importlib.import_module('tools.fleet.work_orders.migrations.001_create_work_order_schema_v1').create_schema(c)
            importlib.import_module('tools.fleet.work_orders.migrations.002_create_work_order_permissions').create_schema(c)
            staff_dispatch.create_schema(c)
            c.executemany('INSERT INTO service_work_order_users(bale_id,role,active) VALUES (?,?,?)',
                [('101','MAINTENANCE_MANAGER',1),('202','MAINTENANCE_MANAGER',0),('404','MAINTENANCE_MANAGER',1)])
            c.execute("INSERT INTO service_staff(id,display_name,bale_id,service_role,active) VALUES (1,'staff','202','GENERAL',1)")
            c.execute("INSERT INTO service_staff_roster VALUES (1,'staff',1,1)")
        self.source=self.root/'drivers.xlsx';self.template=self.root/'blank.xlsx'
        wb=openpyxl.Workbook();s=wb.active
        s.title='گزارش روزانه 1405.07.14'
        s.append([None,'گزارش روزانه','تاریخ: 1405/07/14'])
        s.append(['ردیف','نوع دستگاه','کد جدید','شرح معایب مکانیکی','شرح معایب آهنگری'])
        s.append([1,'دامپتراک','HD714','old','old-metal'])
        s.append([2,'دامپتراک','HD710',None,None])
        wb.save(self.source);wb.close()
        entry_service.create_blank_template(self.source,self.template)
        self.config=self.root/'repairs.json'
        self.config.write_text(json.dumps(dict(allowed_users=['101','202','404'],source=str(self.source),template=str(self.template))),encoding='utf-8')
        self.kw=dict(config=self.config,fleet_db=self.root/'missing.db',day='1405/07/15')
        self.pdf=self.root/'preview.pdf';self.pdf.write_bytes(b'%PDF-fixture')
        for p in [
            patch('tools.authorization.net.AuthorizationStore',return_value=self.auth),
            patch.object(db,'DB_PATH',self.fleet),patch.object(permissions,'DB_PATH',self.fleet),
            patch.object(service,'WORK_ORDER_OUTPUT_ROOT',self.root/'orders'),
            patch.object(entry_bale,'CONFIG',self.config),
            patch.object(delivery,'export_staff_pdf',return_value=self.pdf),
        ]:
            p.start();self.addCleanup(p.stop)

    def sql(self, statement, args=()):
        with closing(sqlite3.connect(self.auth_path)) as c,c:return c.execute(statement,args).fetchall()

    def assign(self):
        migrate_net(self.auth,self.root/'backups',assignments=[('101',NET_MANAGER),('202',NET_DEPUTY)],actor='fixture')

    def request(self, actor='101', section='mechanical', code='HD714', text='new'):
        return dict(entry_service.preview(actor,code,section,**self.kw),description=text,operation=uuid.uuid4().hex)

    def build(self, *, output_path, **kw):
        output_path.parent.mkdir(parents=True,exist_ok=True)
        wb=openpyxl.Workbook();wb.active['A1']='fixture';wb.save(output_path);wb.close()

    def create(self, actor, kind):
        items=[dict(machine_code='HD714',machine_name='دامپتراک',action_code=kind,action_text='fixture')]
        builder=SimpleNamespace(get_items=lambda *a,**k:items,build_document=self.build)
        with patch.object(service,'_load_builder',return_value=builder):
            result=create_worker.execute_request(dict(action='create',bale_id=actor,work_order_type=kind,
                jalali_date='1405/07/15',shift='صبح',machine_codes=['HD714']))
        self.assertTrue(result['ok'],result)
        return service.get_work_order(result['order']['work_order_no'])


class NetContracts(NetFixture):
    def test_profile_alone_no_operations_and_equal_exact_roles(self):
        self.assertEqual(self.sql('SELECT * FROM auth_user_roles'),[])
        self.assertFalse(permissions.check_work_order_permission('505').allowed)
        self.assertFalse(entry_bale.permitted('505'))
        self.assertFalse(self.auth.function_scope('505'))
        for role in [NET_MANAGER,NET_DEPUTY]:
            self.assertEqual(set(self.auth.capabilities_for_role(role)),set(NET_CAPABILITIES))
        self.assertEqual(RESOURCES[FUNCTION_READ],r'E:\Function')
        self.assertFalse(any(x.startswith('maintenance.') for x in NET_CAPABILITIES))
        cfg=yaml.safe_load((ROOT/'integrations/hermes/profiles/net/config.template.yaml').read_text(encoding='utf-8'))
        for selected in cfg['platform_toolsets'].values():
            self.assertEqual(set(selected),{'komatso_technical_docs','web','komatso_public_browser','delegation','skills_readonly','no_mcp'})
        self.assertFalse(cfg['skills']['inline_shell'])
        self.assertFalse(cfg['bale']['enabled'])
        self.assertFalse(cfg['telegram']['enabled'])
        self.assertFalse({'terminal','file','code_execution','skills','mcp','process'} & set(cfg['platform_toolsets']['bale']))

    def test_legacy_only_W_R_still_authoritative(self):
        self.assertTrue(permissions.check_work_order_permission('404').allowed)
        self.assertEqual(permissions.check_work_order_permission('404').source,'legacy')
        self.assertTrue(entry_bale.permitted('404'))
        self.assertFalse(permissions.check_work_order_permission('202').allowed)
        # A capability without migration marker must not replace legacy authority.
        self.auth.assign_role('202',NET_DEPUTY,actor='fixture')
        self.assertFalse(permissions.check_work_order_permission('202').allowed)

    def test_central_roles_routes_and_atomic_idempotent_markers(self):
        self.assign()
        for u in ['101','202']:
            self.assertEqual(self.auth.resolve_profile(u,u),'net')
            self.assertTrue(self.auth.function_scope(u).all)
            self.assertTrue(permissions.check_work_order_permission(u).allowed)
            self.assertEqual(permissions.check_work_order_permission(u).source,'central')
            self.assertTrue(entry_bale.permitted(u))
        self.assertEqual(len(self.sql('SELECT * FROM auth_domain_authority')),4)
        before={t:self.sql('SELECT * FROM '+t) for t in ['auth_user_roles','auth_events','auth_domain_authority','auth_extensions']}
        self.assign()
        self.assertEqual(before,{t:self.sql('SELECT * FROM '+t) for t in before})
        self.assertEqual(self.sql('SELECT MAX(version) FROM auth_migrations'),[(2,)])
        with closing(sqlite3.connect(self.fleet)) as c:
            self.assertEqual(c.execute("SELECT active FROM service_work_order_users WHERE bale_id='202'").fetchone(),(0,))
            self.assertEqual(c.execute("SELECT service_role,active FROM service_staff WHERE bale_id='202'").fetchone(),('GENERAL',1))
        self.assertFalse(self.auth.resolve_profile('101','202'))

    def test_unapproved_assignment_rolls_back_and_profile_conflict_denies(self):
        before=self.sql('SELECT * FROM auth_user_roles')
        with self.assertRaises(ValueError):
            migrate_net(self.auth,self.root/'backups',assignments=[('101',NET_MANAGER),('303',NET_DEPUTY)])
        self.assertEqual(before,self.sql('SELECT * FROM auth_user_roles'))
        self.auth.assign_role('101','office_supervisor',actor='fixture')
        with self.assertRaises(ValueError):
            migrate_net(self.auth,self.root/'backups',assignments=[('101',NET_MANAGER)])

    def test_central_revoke_beats_active_legacy_W_R_and_stale_menu(self):
        self.assign()
        self.assertTrue(_main_menu('101',bale_approved=True).command_for('📋 حکم کار'))
        self.sql("UPDATE auth_user_roles SET active=0 WHERE user_id='101'")
        self.assertFalse(permissions.check_work_order_permission('101').allowed)
        self.assertEqual(permissions.check_work_order_permission('101').source,'central')
        self.assertFalse(entry_bale.permitted('101'))
        with self.assertRaises(PermissionError):entry_service.preview('101','HD714','mechanical',**self.kw)
        menu=_main_menu('101',bale_approved=True)
        self.assertIsNone(menu.command_for('📋 حکم کار'))
        self.assertIsNone(menu.command_for('🛠 شرح خرابی'))
        self.assertEqual(len(self.sql("SELECT * FROM auth_domain_authority WHERE user_id='101'")),2)
        self.assertFalse(self.auth.function_scope('101'))

    def test_every_capability_revoke_resource_corruption_and_store_failure_denies(self):
        self.assign()
        for domain,caps in [(WO_DOMAIN,WO_CAPABILITIES),(REPAIRS_DOMAIN,REPAIRS_CAPABILITIES)]:
            for cap in caps:
                self.sql('DELETE FROM auth_role_capabilities WHERE role=? AND capability=?',(NET_MANAGER,cap))
                self.assertFalse(domain_decision('101',domain,[cap]).allowed,cap)
                self.assertTrue(domain_decision('202',domain,[cap]).allowed,cap)
                self.sql('INSERT INTO auth_role_capabilities VALUES (?,?)',(NET_MANAGER,cap))
        self.sql("UPDATE auth_capabilities SET resource='C:/' WHERE capability='work_orders.create'")
        self.assertFalse(domain_decision('101',WO_DOMAIN,['work_orders.create']).allowed)
        missing=AuthorizationStore(self.root/'does-not-exist.db')
        self.assertEqual(domain_decision('101',WO_DOMAIN,store=missing).source,'denied')
        self.assertFalse((self.root/'does-not-exist.db').exists())
        self.sql('DROP TABLE auth_domain_authority')
        self.assertEqual(domain_decision('101',WO_DOMAIN).source,'denied')

    def test_approved_private_identity_is_rechecked(self):
        self.assign()
        for status in ['revoked','pending_approval','rejected']:
            self.sql("UPDATE channel_users SET registration_status=? WHERE user_id='101'",(status,))
            self.assertFalse(permissions.check_work_order_permission('101').allowed)
            self.assertFalse(entry_bale.permitted('101'))
        self.sql("UPDATE channel_users SET registration_status='approved',chat_id='202' WHERE user_id='101'")
        self.assertFalse(permissions.check_work_order_permission('101').allowed)


class NetWorkflows(NetFixture):
    def test_both_own_AF_OC_GR_create_preview_review_assign_approve_send_archives(self):
        self.assign()
        with closing(sqlite3.connect(self.fleet)) as c:staff_before=c.execute('SELECT * FROM service_staff').fetchall()
        for actor in ['101','202']:
            for kind in ['AIR_FILTER','OIL_CHANGE','GREASING']:
                order=self.create(actor,kind);number=order['work_order_no']
                self.assertEqual(order['status'],'FILE_READY')
                self.assertTrue(create_worker.execute_request(dict(action='preview',bale_id=actor,work_order_no=number))['ok'])
                other='202' if actor=='101' else '101'
                self.assertFalse(create_worker.execute_request(dict(action='preview',bale_id=other,work_order_no=number))['ok'])
                self.assertTrue(create_worker.execute_request(dict(action='confirm_review',bale_id=actor,work_order_no=number))['ok'])
                with self.assertRaises(ValueError):staff_dispatch.prepare_dispatch(number,other,other,1,1)
                staff_dispatch.prepare_dispatch(number,actor,actor,1,1)
                self.assertEqual(service.get_work_order(number)['status'],'APPROVED')
                self.assertEqual(service.get_work_order(number)['approved_by'],f'bale:{actor}')
                self.assertTrue(staff_dispatch.claim_send(number))
                self.assertFalse(staff_dispatch.claim_send(number))
                sender=Mock()
                with patch.object(delivery,'archive_delivered_order') as archive:
                    self.assertEqual(delivery.send_work_order(work_order_no=number,sender=sender,actor=actor)['status'],'SENT')
                    delivery.send_work_order(work_order_no=number,sender=sender,actor=actor)
                    self.assertEqual(archive.call_args.args[0]['work_order_type'],kind)
                sender.send_document.assert_called_once()
                self.assertEqual(service.get_work_order(number)['status'],'SENT')
                staff_dispatch.finish_send(number,'SENT')
                self.assertTrue(any(o['work_order_no']==number for o in staff_dispatch.recipient_orders('202')))
                staff_dispatch.acknowledge(number,'202')
                with self.assertRaises(ValueError):staff_dispatch.acknowledge(number,'101')
        with closing(sqlite3.connect(self.fleet)) as c:self.assertEqual(c.execute('SELECT * FROM service_staff').fetchall(),staff_before)

    def test_create_revoke_during_render_rolls_back_and_cleans_output(self):
        self.assign()
        original=self.build
        def render(**kw):
            original(**kw);self.sql("UPDATE auth_user_roles SET active=0 WHERE user_id='101'")
        self.build=render
        items=[dict(machine_code='HD714',machine_name='fixture',action_code='AIR_FILTER_OUTER',action_text='fixture')]
        builder=SimpleNamespace(get_items=lambda *a,**k:items,build_document=render)
        with patch.object(service,'_load_builder',return_value=builder):
            result=create_worker.execute_request(dict(action='create',bale_id='101',work_order_type='AIR_FILTER',jalali_date='1405/07/15',shift='صبح',machine_codes=['714']))
        self.assertEqual(result['error'],'DENIED')
        with closing(sqlite3.connect(self.fleet)) as c:self.assertEqual(c.execute('SELECT count(*) FROM service_work_orders').fetchone(),(0,))
        self.assertEqual(list((self.root/'orders').rglob('*.xlsx')),[])

    def test_send_revoke_during_export_and_cross_owner_denied(self):
        self.assign();order=self.create('101','AIR_FILTER');number=order['work_order_no']
        review.confirm_document_review(number,'101');staff_dispatch.prepare_dispatch(number,'101','101',1,1)
        sender=Mock()
        def export(*a):
            self.sql("UPDATE auth_user_roles SET active=0 WHERE user_id='101'");return self.pdf
        with patch.object(delivery,'export_staff_pdf',side_effect=export),patch.object(delivery,'archive_delivered_order') as archive:
            with self.assertRaises(PermissionError):delivery.send_work_order(work_order_no=number,sender=sender,actor='202')
            with self.assertRaises(PermissionError):delivery.send_work_order(work_order_no=number,sender=sender,actor='101')
        sender.send_document.assert_not_called();archive.assert_not_called()
        self.assertEqual(service.get_work_order(number)['status'],'APPROVED')

    def test_AF_GR_fixed_archive_adapter_and_OC_no_archive(self):
        from tools.fleet.work_orders.core import daily_archive
        for kind in ['AIR_FILTER','GREASING','OIL_CHANGE']:
            with patch.object(daily_archive.subprocess,'run',return_value=SimpleNamespace(returncode=0)) as run:
                daily_archive.archive_delivered_order(dict(work_order_type=kind,excel_path=str(self.root/'order.xlsx'),work_order_no='FIXTURE'))
            if kind=='OIL_CHANGE':run.assert_not_called()
            else:
                args=run.call_args.args[0]
                self.assertEqual(args[args.index('-TargetPath')+1],str(daily_archive.ARCHIVES[kind]))
                self.assertIn('archive_daily.ps1',args[args.index('-File')+1])

    def test_both_repairs_mechanical_metalwork_replace_clear_append_daily_sheet_retry(self):
        self.assign()
        for actor in ['101','202']:
            for section in ['mechanical','metalwork']:
                req=self.request(actor,section,text=actor+section)
                entry_service.commit(actor,req,runtime=self.root/'journal',**self.kw)
                # Existing succeeded operation replay returns the same result.
                entry_service.commit(actor,req,runtime=self.root/'journal',**self.kw)
                req=self.request(actor,section,text='replace')
                entry_service.commit(actor,req,runtime=self.root/'journal',**self.kw)
                req=self.request(actor,section,text='')
                entry_service.commit(actor,req,runtime=self.root/'journal',**self.kw)
        req=self.request('202',code='HD710',text='append missing')
        entry_service.commit('202',req,runtime=self.root/'journal',**self.kw)
        wb=openpyxl.load_workbook(self.source)
        try:
            self.assertEqual(len(wb.worksheets),2)
            today,old=wb.worksheets
            self.assertEqual(today['C3'].value,'HD710')
            self.assertEqual(today['D3'].value,'append missing')
            self.assertEqual(old['D3'].value,'old');self.assertEqual(old['E3'].value,'old-metal')
        finally:wb.close()

    def test_repairs_revoke_before_replace_denies_and_preserves_file(self):
        self.assign();req=self.request();before=self.source.read_bytes()
        original=openpyxl.workbook.workbook.Workbook.save
        def save(book, destination):
            original(book,destination);self.sql("UPDATE auth_user_roles SET active=0 WHERE user_id='101'")
        with patch.object(openpyxl.workbook.workbook.Workbook,'save',save),patch.object(entry_service,'replace_source') as replace:
            with self.assertRaises(PermissionError):entry_service.commit('101',req,runtime=self.root/'journal',**self.kw)
        replace.assert_not_called();self.assertEqual(self.source.read_bytes(),before)

    def test_reply_menu_no_maintenance_and_staff_inbox_coexists(self):
        self.assign()
        for actor in ['101','202']:
            menu=_main_menu(actor,bale_approved=True)
            self.assertEqual(menu.command_for('📋 حکم کار'),'حکم کار')
            self.assertEqual(menu.command_for('🛠 شرح خرابی'),'شرح خرابی')
            self.assertEqual(menu.command_for('🔄 شروع گفتگوی جدید'),'/new')
            self.assertIsNone(menu.command_for('تعمیرات'))
        event=SimpleNamespace(text='حکم کار',source=SimpleNamespace(platform='bale',chat_type='dm',user_id='202',chat_id='202'))
        sent=[]
        result=staff_flow.handle_staff_receipt(event,None,send=lambda g,c,t:sent.append(t))
        self.assertEqual(result['reason'],'staff-orders-empty')
        event.text='📋 حکم کار'
        self.assertIsNone(staff_flow.handle_staff_receipt(event,None,send=lambda *a:None))

    def test_function_fixture_surface_read_and_path_escapes_denied(self):
        from integrations.hermes.function_domain.worker import run
        from integrations.hermes.function_domain.scoped import ScopedReader
        self.assign();reader=ScopedReader(self.root)
        for actor in ['101','202']:
            identity=dict(platform='bale',chat_type='dm',user_id=actor,chat_id=actor)
            for op,args in [('list',{}),('search',{'query':'drivers'}),('metadata',{'path':'drivers.xlsx'}),('read',{'path':'drivers.xlsx','limit':2})]:
                result=run(dict(identity=identity,operation=op,args=args),store=self.auth,reader=reader,cache=self.root/'cache')
                self.assertTrue(result['read_only'])
            for path in ['../escape','C:/Windows/win.ini',r'\\server\share\file',r'..\escape',r'drivers.xlsx:stream']:
                with self.assertRaises((ValueError,PermissionError,OSError)):
                    run(dict(identity=identity,operation='read',args={'path':path}),store=self.auth,reader=reader)
            with self.assertRaises(ValueError):
                run(dict(identity=identity,operation='write',args={'path':'drivers.xlsx'}),store=self.auth,reader=reader)


class NetInlineContracts(NetFixture):
    def test_all_fixed_inline_families_equal_roles_and_revoke_denies(self):
        from tools.fleet.work_orders.channels.bale import keyboards as kb
        self.assign()
        families=[
            ('MENU',kb.entry_keyboard()),('PROPOSAL',kb.proposal_keyboard),
            ('PROPOSAL',kb.oil_proposal_keyboard),('EDIT_CODE',kb.oil_edit_code_keyboard),
            ('EDIT_INTERVAL',kb.oil_interval_keyboard),('EDIT_CONFIRM',kb.oil_edit_confirm_keyboard),
            ('SHIFT',kb.shift_keyboard),('REVIEW',kb.review_keyboard),
            ('STAFF',kb.staff_keyboard([dict(display_name='staff')])),
        ]
        for actor in ['101','202']:
            permission=permissions.check_work_order_permission(actor)
            for stage,builder in families:
                markup=builder.build('fixture',stage=stage,role=permission.role,
                    permits=lambda _:permissions.check_work_order_permission(actor).allowed)
                actual={b['callback_data'].split(':')[-1] for row in markup['inline_keyboard'] for b in row}
                self.assertEqual(actual,set(builder.actions))
                for name in actual:
                    data='ik:work_order:fixture:'+name
                    self.assertEqual(builder.resolve(data,'fixture',stage=stage,role=permission.role,
                        permits=lambda _:permissions.check_work_order_permission(actor).allowed),name)
                    with self.assertRaises(ValueError):
                        builder.resolve(data,'old',stage=stage,role=permission.role,permits=lambda _:True)
            for stage in ['SECTION','CONFIRM','DESCRIPTION','RESULT']:
                builder=entry_bale.keyboard(stage)
                markup=builder.build('fixture',stage=stage,role='',permits=lambda _:entry_bale.permitted(actor))
                self.assertEqual({b['callback_data'].split(':')[-1] for row in markup['inline_keyboard'] for b in row},set(builder.actions))
        self.sql("UPDATE auth_user_roles SET active=0 WHERE user_id='101'")
        for stage,builder in families:
            with self.assertRaises(PermissionError):
                builder.resolve('ik:work_order:fixture:'+next(iter(builder.actions)), 'fixture',stage=stage,
                                role='MAINTENANCE_MANAGER',permits=lambda _:permissions.check_work_order_permission('101').allowed)

    def test_stale_real_handler_callbacks_after_revoke_and_actor_chat_spoof_denied(self):
        from tools.fleet.work_orders.channels.bale.message_handler import WorkOrderMenuHandler,FormSession
        self.assign();handler=WorkOrderMenuHandler(db_path=self.fleet)
        session=FormSession(stage='MENU',expires=10**12,keyboard_revision='fixture')
        handler.pending[('bale','101','101')]=session
        self.sql("UPDATE auth_user_roles SET active=0 WHERE user_id='101'")
        event=SimpleNamespace(text='',message_id='fixture',raw_message=dict(bale_inline_callback=True,data='ik:work_order:fixture:air_filter',origin_message_id='1'),
            source=SimpleNamespace(platform='bale',chat_type='dm',user_id='101',chat_id='101'))
        result=handler.handle(event,None,send=lambda *a:None)
        self.assertEqual(result['action'],'skip');self.assertIn(result['reason'],{'work-order-permission-denied','inline-rejected'})
        self.assertEqual(session.stage,'MENU')
        rh=entry_bale.RepairsEntryHandler()
        event.raw_message['data']='ik:repairs_entry:fixture:mechanical'
        self.assertEqual(rh.handle(event,None,send=lambda *a:None)['reason'],'repairs-entry-denied')
        event.source.user_id='202'
        result=handler.handle(event,None,send=lambda *a:None)
        self.assertEqual(result['reason'],'inline-rejected')

if __name__=='__main__':unittest.main()

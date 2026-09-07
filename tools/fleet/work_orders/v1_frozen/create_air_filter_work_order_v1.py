import argparse
import re
import sqlite3
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


DB_PATH = Path(
    r"E:\KomatsoAI\data\fleet\db\fleet_ops.db"
)

OUTPUT_ROOT = Path(
    r"E:\KomatsoAI\work_orders\air_filter"
)

TEMPLATE_SHEET = "Sheet1 (486)"

DATA_START_ROW = 3
TEMPLATE_DATA_ROWS = 6

WORK_ORDER_TYPE = "AIR_FILTER"
TEMPLATE_KEY = "AIR_FILTER_DAY9_V1"

DEFAULT_ACTION_CODE = "AIR_FILTER_OUTER"
DEFAULT_ACTION_TEXT = "تعویض هواکش بیرونی"


def clean_code(value):
    return str(value).strip().upper()


def validate_date(value):
    m = re.fullmatch(
        r"1405/(\d{2})/(\d{2})",
        value.strip()
    )

    if not m:
        raise ValueError(
            "تاریخ باید مانند 1405/06/10 باشد"
        )

    month = int(m.group(1))
    day = int(m.group(2))

    if not 1 <= month <= 12:
        raise ValueError(
            "ماه نامعتبر است"
        )

    max_day = 31 if month <= 6 else 30

    if month == 12:
        max_day = 29

    if not 1 <= day <= max_day:
        raise ValueError(
            "روز نامعتبر است"
        )


def machine_name_from_code(code):

    c = clean_code(code)

    if c.isdigit():

        n = int(c)

        if 461 <= n <= 469:
            return "دامپتراک"

        if 701 <= n <= 716:
            return "دامپتراک"

        if n == 231:
            return "بیل"

    if c.startswith("MZ"):
        return "مزدا"

    if c in {"S1", "S3"}:
        return "کامیون سهند زرد"

    if c == "TA1":
        return "کامیون آب پاش"

    if c == "TR1":
        return "خاور"

    if c == "DG1":
        return "ژنراتور"

    if c in {"PR2", "PR3"}:
        return "پیکاپ ریچ"

    raise ValueError(
        f"کد دستگاه برای حکم هواکش شناخته نشده: {code}"
    )


def find_machine_id(con, raw_code):

    code = clean_code(raw_code)

    candidates = [
        code
    ]

    if code.isdigit():

        n = int(code)

        if 461 <= n <= 469:
            candidates.append(
                f"HD{n}"
            )

        if 701 <= n <= 716:
            candidates.append(
                f"HD{n}"
            )

        if n == 231:
            candidates.append(
                "EX231"
            )

    placeholders = ",".join(
        "?"
        for _ in candidates
    )

    row = con.execute(
        f"""
        SELECT id, canonical_code
        FROM machines
        WHERE UPPER(canonical_code)
              IN ({placeholders})
        ORDER BY id
        LIMIT 1
        """,
        candidates
    ).fetchone()

    if row:
        return row[0]

    return None


def copy_row_style(
    ws,
    source_row,
    target_row
):

    for col in range(1, 7):

        src = ws.cell(
            source_row,
            col
        )

        dst = ws.cell(
            target_row,
            col
        )

        dst._style = copy(
            src._style
        )

        dst.font = copy(
            src.font
        )

        dst.fill = copy(
            src.fill
        )

        dst.border = copy(
            src.border
        )

        dst.alignment = copy(
            src.alignment
        )

        dst.protection = copy(
            src.protection
        )

        dst.number_format = (
            src.number_format
        )

    ws.row_dimensions[
        target_row
    ].height = (
        ws.row_dimensions[
            source_row
        ].height
    )


def next_work_order_no(
    con,
    jalali_date
):

    date_part = (
        jalali_date
        .replace("/", "-")
    )

    prefix = (
        f"AF-{date_part}-"
    )

    rows = con.execute(
        """
        SELECT work_order_no
        FROM service_work_orders
        WHERE work_order_no LIKE ?
        """,
        (prefix + "%",)
    ).fetchall()

    highest = 0

    for row in rows:

        number = row[0]

        try:
            seq = int(
                number.rsplit(
                    "-",
                    1
                )[1]
            )

            highest = max(
                highest,
                seq
            )

        except Exception:
            continue

    return (
        f"{prefix}"
        f"{highest + 1:03d}"
    )


def build_excel(
    template_path,
    output_path,
    jalali_date,
    codes
):

    wb = load_workbook(
        template_path
    )

    if TEMPLATE_SHEET not in wb.sheetnames:

        wb.close()

        raise RuntimeError(
            f"شیت تمپلیت پیدا نشد: {TEMPLATE_SHEET}"
        )

    ws = wb[
        TEMPLATE_SHEET
    ]

    # فقط همان تمپلیت روز 9
    for other_ws in list(
        wb.worksheets
    ):

        if (
            other_ws.title
            != TEMPLATE_SHEET
        ):
            wb.remove(
                other_ws
            )

    count = len(codes)

    if count > TEMPLATE_DATA_ROWS:

        extra = (
            count
            - TEMPLATE_DATA_ROWS
        )

        insert_at = (
            DATA_START_ROW
            + TEMPLATE_DATA_ROWS
        )

        ws.insert_rows(
            insert_at,
            amount=extra
        )

        source_style_row = (
            DATA_START_ROW
            + TEMPLATE_DATA_ROWS
            - 1
        )

        for row in range(
            insert_at,
            DATA_START_ROW
            + count
        ):

            copy_row_style(
                ws,
                source_style_row,
                row
            )

    elif count < TEMPLATE_DATA_ROWS:

        remove_count = (
            TEMPLATE_DATA_ROWS
            - count
        )

        ws.delete_rows(
            DATA_START_ROW
            + count,
            amount=remove_count
        )

    # تاریخ
    ws["F1"] = jalali_date

    # دستگاه‌ها
    for index, code in enumerate(
        codes,
        start=1
    ):

        row = (
            DATA_START_ROW
            + index
            - 1
        )

        ws.cell(
            row,
            1
        ).value = index

        ws.cell(
            row,
            2
        ).value = (
            machine_name_from_code(
                code
            )
        )

        ws.cell(
            row,
            3
        ).value = code

        ws.cell(
            row,
            4
        ).value = (
            DEFAULT_ACTION_TEXT
        )

        ws.cell(
            row,
            5
        ).value = None

        ws.cell(
            row,
            6
        ).value = None

    signature_row = (
        DATA_START_ROW
        + count
    )

    ws.print_area = (
        f"A1:F{signature_row}"
    )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    wb.save(
        output_path
    )

    wb.close()


def main():

    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--template",
        required=True
    )

    parser.add_argument(
        "--date",
        required=True
    )

    parser.add_argument(
        "--shift",
        default="صبح"
    )

    parser.add_argument(
        "--codes",
        nargs="+",
        required=True
    )

    args = parser.parse_args()

    template_path = Path(
        args.template
    )

    if not DB_PATH.exists():

        raise SystemExit(
            f"DB NOT FOUND: {DB_PATH}"
        )

    if not template_path.exists():

        raise SystemExit(
            f"TEMPLATE NOT FOUND: {template_path}"
        )

    try:
        validate_date(
            args.date
        )

    except ValueError as exc:

        raise SystemExit(
            f"DATE ERROR: {exc}"
        )

    codes = [
        clean_code(code)
        for code in args.codes
    ]

    if not codes:

        raise SystemExit(
            "NO MACHINE CODES"
        )

    # Duplicate protection داخل خود حکم
    duplicates = sorted({
        code
        for code in codes
        if codes.count(code) > 1
    })

    if duplicates:

        raise SystemExit(
            "DUPLICATE MACHINE CODES: "
            + ", ".join(duplicates)
        )

    # قبل از هر تغییر DB،
    # همه کدها اعتبارسنجی شوند.
    machine_names = {}

    for code in codes:

        try:

            machine_names[
                code
            ] = (
                machine_name_from_code(
                    code
                )
            )

        except ValueError as exc:

            raise SystemExit(
                f"VALIDATION ERROR: {exc}"
            )

    con = sqlite3.connect(
        str(DB_PATH)
    )

    con.execute(
        "PRAGMA foreign_keys = ON"
    )

    output_path = None

    try:

        con.execute(
            "BEGIN IMMEDIATE"
        )

        work_order_no = (
            next_work_order_no(
                con,
                args.date
            )
        )

        date_folder = (
            args.date[:7]
            .replace("/", "-")
        )

        output_dir = (
            OUTPUT_ROOT
            / date_folder
        )

        output_path = (
            output_dir
            / f"{work_order_no}.xlsx"
        )

        if output_path.exists():

            raise RuntimeError(
                f"OUTPUT ALREADY EXISTS: {output_path}"
            )

        # --------------------------------------------
        # ابتدا DRAFT در DB
        # --------------------------------------------

        cur = con.execute(
            """
            INSERT INTO service_work_orders (
                work_order_no,
                work_order_type,
                jalali_date,
                shift,
                status,
                assigned_staff_id,
                template_key,
                excel_path,
                pdf_path,
                notes,
                created_by
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                work_order_no,
                WORK_ORDER_TYPE,
                args.date,
                args.shift,
                "DRAFT",
                None,
                TEMPLATE_KEY,
                None,
                None,
                None,
                "LOCAL_V1",
            )
        )

        work_order_id = (
            cur.lastrowid
        )

        # --------------------------------------------
        # آیتم‌ها
        # --------------------------------------------

        for item_no, code in enumerate(
            codes,
            start=1
        ):

            machine_id = (
                find_machine_id(
                    con,
                    code
                )
            )

            con.execute(
                """
                INSERT INTO service_work_order_items (
                    work_order_id,
                    item_no,
                    machine_id,
                    machine_code,
                    machine_name,
                    action_code,
                    action_text,
                    item_status
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    work_order_id,
                    item_no,
                    machine_id,
                    code,
                    machine_names[
                        code
                    ],
                    DEFAULT_ACTION_CODE,
                    DEFAULT_ACTION_TEXT,
                    "PENDING",
                )
            )

        # --------------------------------------------
        # Excel
        # --------------------------------------------

        build_excel(
            template_path,
            output_path,
            args.date,
            codes
        )

        if not output_path.exists():

            raise RuntimeError(
                "Excel file was not created"
            )

        # --------------------------------------------
        # فقط بعد از ساخته‌شدن فایل:
        # FILE_READY
        # --------------------------------------------

        con.execute(
            """
            UPDATE service_work_orders
            SET
                status = 'FILE_READY',
                excel_path = ?,
                finalized_at = CURRENT_TIMESTAMP,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (
                str(output_path),
                work_order_id
            )
        )

        con.commit()

        # --------------------------------------------
        # Verify
        # --------------------------------------------

        order = con.execute(
            """
            SELECT
                id,
                work_order_no,
                work_order_type,
                jalali_date,
                shift,
                status,
                excel_path
            FROM service_work_orders
            WHERE id = ?
            """,
            (
                work_order_id,
            )
        ).fetchone()

        items = con.execute(
            """
            SELECT
                item_no,
                machine_code,
                machine_name,
                action_text,
                machine_id
            FROM service_work_order_items
            WHERE work_order_id = ?
            ORDER BY item_no
            """,
            (
                work_order_id,
            )
        ).fetchall()

        print("=" * 78)
        print("AIR FILTER WORK ORDER CREATED")
        print("=" * 78)

        print(
            f"ID:            {order[0]}"
        )

        print(
            f"NUMBER:        {order[1]}"
        )

        print(
            f"TYPE:          {order[2]}"
        )

        print(
            f"DATE:          {order[3]}"
        )

        print(
            f"SHIFT:         {order[4]}"
        )

        print(
            f"STATUS:        {order[5]}"
        )

        print(
            f"EXCEL:         {order[6]}"
        )

        print()
        print(
            f"ITEMS ({len(items)})"
        )
        print("-" * 78)

        for row in items:

            machine_id_text = (
                row[4]
                if row[4] is not None
                else "-"
            )

            print(
                f"{row[0]}. "
                f"{row[1]} | "
                f"{row[2]} | "
                f"{row[3]} | "
                f"machine_id={machine_id_text}"
            )

        print()
        print("=" * 78)
        print("RESULT")
        print("=" * 78)

        print(
            "DB + EXCEL TRANSACTION PASS"
        )

        print(
            "WORK ORDER STATUS: FILE_READY"
        )

        print("=" * 78)

    except Exception as exc:

        con.rollback()

        # اگر Excel نیمه‌کاره ساخته شده،
        # پاک شود.
        if (
            output_path is not None
            and output_path.exists()
        ):

            try:
                output_path.unlink()

            except Exception:
                pass

        raise SystemExit(
            f"CREATE WORK ORDER FAILED: {exc}"
        )

    finally:

        con.close()


if __name__ == "__main__":
    main()
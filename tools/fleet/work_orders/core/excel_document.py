"""Shared six-column work-order Excel layout; business rules stay in type builders."""
from copy import copy
import math
import unicodedata
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.page import PageMargins

DATA_START_ROW = 3
TEMPLATE_DATA_ROWS = 6


def _copy_row_style(
    ws,
    source_row: int,
    target_row: int,
) -> None:

    for col in range(1, 7):

        src = ws.cell(
            source_row,
            col
        )

        dst = ws.cell(
            target_row,
            col
        )

        dst._style = copy(
            src._style
        )

        dst.font = copy(
            src.font
        )

        dst.fill = copy(
            src.fill
        )

        dst.border = copy(
            src.border
        )

        dst.alignment = copy(
            src.alignment
        )

        dst.protection = copy(
            src.protection
        )

        dst.number_format = (
            src.number_format
        )

    ws.row_dimensions[
        target_row
    ].height = (
        ws.row_dimensions[
            source_row
        ].height
    )


def build_document(
    *,
    output_path: Path,
    jalali_date: str,
    items: list[dict],
    template_path: Path,
    template_sheet: str,
    title: str,
    readable_names: bool = False,
    work_order=None,
) -> Path:
    TEMPLATE_PATH = template_path
    TEMPLATE_SHEET = template_sheet

    if not TEMPLATE_PATH.exists():

        raise RuntimeError(
            f"Work Order template "
            f"not found: {TEMPLATE_PATH}"
        )

    if not items:

        raise ValueError(
            "Work Order Work Order "
            "has no items"
        )

    wb = load_workbook(
        TEMPLATE_PATH
    )

    try:

        if (
            TEMPLATE_SHEET
            not in wb.sheetnames
        ):

            raise RuntimeError(
                "Work Order template "
                "sheet not found: "
                f"{TEMPLATE_SHEET}"
            )

        ws = wb[
            TEMPLATE_SHEET
        ]

        for other_ws in list(
            wb.worksheets
        ):

            if (
                other_ws.title
                != TEMPLATE_SHEET
            ):

                wb.remove(
                    other_ws
                )

        count = len(items)
        # openpyxl moves cells, but does not translate merged ranges or row dimensions.
        footer_start = DATA_START_ROW + TEMPLATE_DATA_ROWS
        merges = list(ws.merged_cells.ranges)
        footer_heights = {r:d.height for r,d in ws.row_dimensions.items() if r >= footer_start}
        for merged in merges:
            ws.unmerge_cells(str(merged))

        if count > TEMPLATE_DATA_ROWS:

            extra = (
                count
                - TEMPLATE_DATA_ROWS
            )

            insert_at = (
                DATA_START_ROW
                + TEMPLATE_DATA_ROWS
            )

            ws.insert_rows(
                insert_at,
                amount=extra
            )

            source_style_row = (
                DATA_START_ROW
                + TEMPLATE_DATA_ROWS
                - 1
            )

            for row in range(
                insert_at,
                DATA_START_ROW + count
            ):

                _copy_row_style(
                    ws,
                    source_style_row,
                    row,
                )

        elif count < TEMPLATE_DATA_ROWS:

            remove_count = (
                TEMPLATE_DATA_ROWS
                - count
            )

            ws.delete_rows(
                DATA_START_ROW + count,
                amount=remove_count
            )

        delta = count - TEMPLATE_DATA_ROWS
        for merged in merges:
            if merged.min_row >= footer_start:
                merged.shift(row_shift=delta)
            elif merged.max_row >= DATA_START_ROW:
                raise ValueError('Unexpected merged data rows in work-order template')
            ws.merge_cells(str(merged))
        for row,height in footer_heights.items():
            ws.row_dimensions[row+delta].height = height

        ws["F1"] = work_order['jalali_date'] if work_order else jalali_date
        ws["A1"] = title
        if work_order and work_order.get('revision_of'):
            ws['A1'] = title + '\nاصلاحیه حکم: ' + work_order['revision_of']
        if work_order and work_order.get('work_order_no'):
            ws.oddHeader.center.text = 'شماره حکم: ' + work_order['work_order_no']
            if work_order.get('created_at'):
                from datetime import datetime, timezone
                from zoneinfo import ZoneInfo
                import jdatetime
                issued = datetime.fromisoformat(work_order['created_at'])
                if issued.tzinfo is None:
                    issued = issued.replace(tzinfo=timezone.utc)
                issued_date = jdatetime.date.fromgregorian(date=issued.astimezone(ZoneInfo('Asia/Tehran')).date())
                ws.oddHeader.center.text += '\nتاریخ صدور: ' + issued_date.strftime('%Y/%m/%d')
            ws.oddHeader.center.size = 11
            ws.oddHeader.center.font = 'Arial'
        if readable_names:
            for col,width in {'A':6,'B':36,'C':14,'D':22,'E':16,'F':16}.items():
                ws.column_dimensions[col].width = width
            for row in (1,2,DATA_START_ROW+count):
                ws.row_dimensions[row].height = 40 if row != 2 else 32
                for cell in ws[row][:6]:
                    font=copy(cell.font)
                    font.sz = 20 if row==1 else 14
                    cell.font=font
            date_font=copy(ws['F1'].font)
            date_font.sz=12
            ws['F1'].font=date_font
            if work_order and work_order.get('revision_of'):
                alignment=copy(ws['A1'].alignment)
                alignment.wrap_text=True
                ws['A1'].alignment=alignment
                ws.row_dimensions[1].height=60
            ws.page_margins = PageMargins(left=.25,right=.25,top=.55 if work_order else .25,bottom=.25,header=.15,footer=0)
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.orientation = 'portrait'
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.print_title_rows = '1:2'

        for index, item in enumerate(
            items,
            start=1,
        ):

            row = (
                DATA_START_ROW
                + index
                - 1
            )

            ws.cell(
                row,
                1
            ).value = index

            ws.cell(
                row,
                2
            ).value = item[
                "machine_name"
            ]

            ws.cell(
                row,
                3
            ).value = item[
                "machine_code"
            ]

            ws.cell(
                row,
                4
            ).value = item[
                "action_text"
            ]

            ws.cell(
                row,
                5
            ).value = None

            ws.cell(
                row,
                6
            ).value = None
            if readable_names:
                for col in range(1,7):
                    font=copy(ws.cell(row,col).font)
                    font.sz=14
                    ws.cell(row,col).font=font
                    alignment = copy(ws.cell(row,col).alignment)
                    alignment.wrap_text = True
                    alignment.vertical = 'center'
                    ws.cell(row,col).alignment = alignment
                # Allow whole names to wrap, keeping the template's readable font.
                name = str(item['machine_name'])
                width = sum(1.2 if unicodedata.bidirectional(c) in {'R','AL'} else 1 for c in name)
                action = str(item['action_text'])
                action_width = sum(1.2 if unicodedata.bidirectional(c) in {'R','AL'} else 1 for c in action)
                lines = max(1,math.ceil(width/28),name.count('\n')+1,
                            math.ceil(action_width/17),action.count('\n')+1)
                ws.row_dimensions[row].height = max(32,lines*20+10)

        signature_row = (
            DATA_START_ROW
            + count
        )

        ws.print_area = (
            f"A1:F{signature_row}"
        )

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True
        )

        wb.save(
            output_path
        )

    finally:

        wb.close()

    if not output_path.exists():

        raise RuntimeError(
            "Work Order Excel "
            "was not created"
        )

    return output_path

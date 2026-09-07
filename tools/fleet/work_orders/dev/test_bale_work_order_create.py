from __future__ import annotations

import importlib
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from tools.fleet.work_orders.core import service
from tools.fleet.work_orders.dev.test_permissions import PermissionDatabaseTestCase, test_database


class WorkOrderCreateTests(PermissionDatabaseTestCase):
    def setUp(self):
        super().setUp()
        self.output_root = self.db_path.parent / "orders"
        schema = importlib.import_module("tools.fleet.work_orders.migrations.001_create_work_order_schema_v1")
        with test_database(self.db_path) as con:
            con.execute("CREATE TABLE machines (id INTEGER PRIMARY KEY, canonical_code TEXT)")
            con.executemany("INSERT INTO machines (canonical_code) VALUES (?)", [(f"HD{code}",) for code in (465,710,711,712,713,714)])
            schema.create_schema(con)
        for target, value in (
            ("tools.fleet.work_orders.core.db.DB_PATH", self.db_path),
            ("tools.fleet.work_orders.core.service.WORK_ORDER_OUTPUT_ROOT", self.output_root),
        ):
            patcher = patch(target, value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def create(self):
        return service.create_work_order(work_order_type="AIR_FILTER", jalali_date="1405/06/15", shift="صبح", machine_codes=["465","710","711","712","713","714"], created_by="bale:455740857")

    def test_real_builder_produces_six_machine_excel_and_file_ready_record(self):
        order = self.create()
        self.assertEqual(order["status"], "FILE_READY")
        self.assertEqual(order["created_by"], "bale:455740857")
        self.assertIsNone(order["assigned_staff_id"])
        self.assertIsNone(order["approved_at"])
        self.assertIsNone(order["sent_at"])
        self.assertEqual(len(order["items"]), 6)
        wb = load_workbook(order["excel_path"])
        try:
            ws = wb.active
            self.assertEqual(ws["F1"].value, "1405/06/15")
            self.assertEqual([str(ws.cell(row,3).value) for row in range(3,9)], ["465","710","711","712","713","714"])
            self.assertEqual([ws.cell(row,4).value for row in range(3,9)], ["تعویض هواکش بیرونی"]*6)
        finally:
            wb.close()

    def test_existing_output_is_preserved_on_number_collision(self):
        path = self.output_root / "air_filter/1405-06/AF-1405-06-15-001.xlsx"
        path.parent.mkdir(parents=True)
        path.write_bytes(b"EXISTING DOCUMENT")
        with self.assertRaisesRegex(RuntimeError, "OUTPUT ALREADY EXISTS"):
            self.create()
        self.assertEqual(path.read_bytes(), b"EXISTING DOCUMENT")

    def test_committed_excel_survives_result_lookup_failure(self):
        with patch.object(service, "get_work_order", side_effect=RuntimeError("TEST LOOKUP FAILURE")):
            with self.assertRaisesRegex(RuntimeError, "TEST LOOKUP FAILURE"):
                self.create()
        with test_database(self.db_path) as con:
            row = con.execute("SELECT status, excel_path FROM service_work_orders").fetchone()
        self.assertEqual(row[0], "FILE_READY")
        self.assertTrue(Path(row[1]).is_file())

    def test_failed_builder_rolls_back_rows_and_removes_partial_file(self):
        def fail(*, output_path, **kwargs):
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_bytes(b"PARTIAL")
            raise RuntimeError("TEST BUILD FAILURE")
        with patch("tools.fleet.work_orders.types.air_filter.builder.build_document", side_effect=fail):
            with self.assertRaisesRegex(RuntimeError, "TEST BUILD FAILURE"):
                self.create()
        with test_database(self.db_path) as con:
            self.assertEqual(con.execute("SELECT COUNT(*) FROM service_work_orders").fetchone()[0], 0)
            self.assertEqual(con.execute("SELECT COUNT(*) FROM service_work_order_items").fetchone()[0], 0)
        self.assertEqual(list(self.output_root.rglob("*.xlsx")), [])

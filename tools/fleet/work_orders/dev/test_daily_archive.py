import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from tools.fleet.work_orders.core.daily_archive import archive_delivered_order
from tools.fleet.work_orders.dev.test_staff_pdf import StaffPdfTests
from tools.fleet.work_orders.core.delivery import send_work_order


class ArchiveRetryTests(StaffPdfTests):
    def test_archive_failure_retry_does_not_resend(self):
        order = self.prepared_order()
        sender = Mock()
        with patch('tools.fleet.work_orders.core.delivery.export_staff_pdf', return_value=Path('test.pdf')), patch(
            'tools.fleet.work_orders.core.delivery.archive_delivered_order',
            side_effect=[RuntimeError('locked'), None],
        ) as archive:
            with self.assertRaisesRegex(RuntimeError, 'locked'):
                send_work_order(work_order_no=order['work_order_no'], sender=sender)
            result = send_work_order(work_order_no=order['work_order_no'], sender=sender)
        self.assertEqual(result['status'], 'ALREADY_SENT')
        self.assertEqual(archive.call_count, 2)
        sender.send_document.assert_called_once()


class NativeArchiveTests(unittest.TestCase):
    def test_native_copy_numbering_format_and_idempotency(self):
        with tempfile.TemporaryDirectory(dir='E:/KomatsoAI/runtime') as folder:
            target = Path(folder) / 'daily.xlsx'
            source = Path(folder) / 'approved.xlsx'
            wb = Workbook()
            wb.active.title = 'Sheet1 (494)'
            wb.active['A1'] = 'previous'
            wb.create_sheet('Sheet1 (2)')
            wb.create_sheet('Sheet1')
            wb.save(target)
            wb.close()
            wb = Workbook()
            ws = wb.active
            ws['A1'] = 'approved'
            ws['A1'].font = Font(name='Arial', bold=True, size=18)
            ws.merge_cells('A1:D1')
            ws.column_dimensions['A'].width = 24
            ws.print_area = 'A1:F9'
            wb.save(source)
            wb.close()
            order = dict(work_order_type='AIR_FILTER', excel_path=str(source), work_order_no='AF-TEST-001')
            with patch('tools.fleet.work_orders.core.daily_archive.ARCHIVES', {'AIR_FILTER': target}):
                archive_delivered_order(order)
                archive_delivered_order(order)
            wb = load_workbook(target)
            self.assertEqual(wb.sheetnames, ['Sheet1 (494)', 'Sheet1 (2)', 'Sheet1', 'Sheet1 (495)'])
            self.assertEqual(wb.worksheets[0]['A1'].value, 'previous')
            ws = wb.worksheets[-1]
            self.assertEqual(ws['A1'].value, 'approved')
            self.assertTrue(ws['A1'].font.bold)
            self.assertEqual(ws['A1'].font.sz, 18)
            self.assertIn('A1:D1', ws.merged_cells)
            self.assertAlmostEqual(ws.column_dimensions['A'].width, 24, delta=1)
            self.assertIn('$A$1:$F$9', str(ws.print_area))
            wb.close()
            self.assertEqual(len(list((Path(folder) / 'work_order_backups').glob('*.xlsx'))), 1)

    def test_oil_is_not_archived(self):
        with patch('tools.fleet.work_orders.core.daily_archive.subprocess.run') as run:
            archive_delivered_order({'work_order_type': 'OIL_CHANGE'})
        run.assert_not_called()

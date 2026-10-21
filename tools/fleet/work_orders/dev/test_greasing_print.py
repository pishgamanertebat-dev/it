import tempfile
from pathlib import Path
import unittest
from openpyxl import load_workbook
from tools.fleet.work_orders.types.greasing.builder import build_document


class GreasingPrintTests(unittest.TestCase):
    def test_variable_rows_preserve_independent_result_cells_and_signature(self):
        root=Path(__file__).resolve().parents[4]/'runtime/greasing-print-tests'
        root.mkdir(parents=True,exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as tmp:
            for count in (1,2,5,6,7,20,40,100):
                with self.subTest(count=count):
                    path=Path(tmp)/f'GR-{count}.xlsx'
                    items=[dict(machine_name='دامپتراک کوماتسو با نام بلند 785 مدل 7',machine_code=f'HD{i}',action_text='گریسکاری کامل') for i in range(count)]
                    build_document(output_path=path,jalali_date='1405/07/25',items=items)
                    wb=load_workbook(path);ws=wb.active
                    try:
                        for row in range(3,count+3):
                            self.assertTrue(ws.cell(row,2).alignment.wrap_text)
                            for col in (5,6):
                                self.assertFalse(any(ws.cell(row,col).coordinate in m for m in ws.merged_cells.ranges))
                            ws.cell(row,5).value='انجام شد'
                            ws.cell(row,6).value='نشد'
                        self.assertIn(f'E{count+3}:F{count+3}',{str(m) for m in ws.merged_cells.ranges})
                        self.assertEqual(str(ws.page_setup.paperSize),ws.PAPERSIZE_A4)
                        self.assertEqual(ws.page_setup.fitToHeight,0)
                        self.assertIn(f'F${count+3}',str(ws.print_area))
                        self.assertTrue(ws.cell(count+3,5).value)
                    finally:wb.close()

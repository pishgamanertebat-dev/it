"""Unified oil editing: real handler/core, isolated DBs, fake Bale transport."""
from contextlib import closing
from copy import deepcopy
import asyncio
import hashlib
import json
import os
from pathlib import Path
import shutil
import sqlite3
import unittest
from unittest.mock import patch

from tools.fleet.work_orders.dev import test_inline_menus as menus
from tools.fleet.work_orders.dev.test_oil_service_edit import OilServiceEditTests
from tools.fleet.oil_change.test_proposal import sample_source
from tools.fleet.work_orders.channels.bale import create_worker
from tools.fleet.work_orders.channels.bale.keyboards import keyboard_for
from tools.fleet.work_orders.channels.bale.proposal_form import (
    creation_metadata, oil_edit_items, apply_oil_edits,
)


class OilUnifiedTests(menus.InlineMenusTests):
    async def click(self, action):
        for index in range(len(self.sent)-1,-1,-1):
            try:
                data=self.button(self.sent[index],action)
                break
            except (KeyError,StopIteration):
                continue
        else:
            self.fail('Missing inline action: ' + action)
        self.message('', data=data, origin=index+1)
        await self.settle()

    async def enter(self, text):
        result = self.message(text)
        await self.settle()
        return result

    async def start(self, source=None, add_code=None):
        session = await self.oil_edit_step(source, add_code=add_code)
        await self.click('edit_excel')
        return session

    def request(self, session):
        return dict(action='create', bale_id=self.key[1], work_order_type='OIL_CHANGE',
                    jalali_date=session.jalali_date, shift='روزانه',
                    machine_codes=[i['machine_code'] for i in session.proposal['items']],
                    item_actions={i['machine_code']:i['action_code'] for i in session.proposal['items']},
                    proposal=creation_metadata(session.proposal))

    @staticmethod
    def combined(count=3, remaining=100):
        source = sample_source('HD465-7R', 'HD800', last=1600, remaining=remaining)
        for n in range(801, 800 + count):
            extra = sample_source('HD465-7R', f'HD{n}', last=1600, remaining=remaining)
            source['machines'] += extra['machines']
            source['plans'] += extra['plans']
        return source

    async def test_only_unified_button_and_legacy_add_cannot_start_old_flow(self):
        s = await self.oil_edit_step(add_code=None)
        self.assertEqual(set(keyboard_for(s).actions), {'edit_excel','remove','confirm','cancel'})
        self.assertEqual(keyboard_for(s).actions['edit_excel'].label, '➕ افزودن/ویرایش')
        forged = self.button(self.sent[-1],'edit_excel').rsplit(':',1)[0] + ':add'
        self.assertEqual(self.message('',data=forged,origin=len(self.sent))['reason'],'inline-rejected')
        await self.settle()
        await self.enter('اضافه')
        self.assertEqual(s.stage,'EDIT_CODE')
        self.assertFalse(self.requests)
        result = create_worker.execute_request(dict(action='validate_add',bale_id=self.key[1],
            work_order_type='OIL_CHANGE',machine_codes=['EX1252']),db_path=self.db_path)
        self.assertEqual(result['error'],'INVALID_INPUT')

    async def test_add_default_is_confirmed_without_false_override(self):
        s = await self.start()
        before = deepcopy(s.proposal)
        await self.enter('۱۲۵۲')
        await self.click('interval_keep')
        self.assertEqual(s.proposal,before)
        self.assertIn('افزودن با نوبت پیشنهادی 2000',self.sent[-1]['text'])
        await self.click('edit_apply')
        self.assertEqual(s.stage,'PROPOSAL')
        self.assertEqual(len(s.proposal['items']),1)
        self.assertFalse(s.proposal.get('service_overrides'))
        self.assertNotIn('از سرویس',s.proposal['items'][0]['manual_note'])
        self.assertEqual(s.proposal['evaluations'],before['evaluations'])
        self.assertFalse(self.requests)

    async def test_existing_same_interval_is_a_true_noop(self):
        s = await self.start(add_code='1252')
        before = deepcopy(s.proposal)
        await self.enter('ex1252');await self.enter('2000')
        self.assertIn('بدون تغییر',self.sent[-1]['text'])
        await self.click('edit_apply')
        self.assertEqual(s.proposal,before)

    async def test_changed_interval_applies_only_after_confirmation(self):
        s = await self.start(add_code='1252')
        before = deepcopy(s.proposal)
        await self.enter('EX1252');await self.enter('1800')
        self.assertEqual(s.proposal,before)
        await self.click('edit_apply')
        self.assertEqual(s.proposal['service_overrides'],{'EX1252':1800})
        for key in ('remaining','current_meter','target_meter','last_interval','last_service'):
            self.assertEqual(s.proposal['items'][0]['components']['oil_change'][key],before['items'][0]['components']['oil_change'][key])

    async def test_batch_defaults_optional_change_and_atomic_confirmation(self):
        s = await self.start(self.combined())
        before = deepcopy(s.proposal)
        await self.click('edit_batch')
        await self.enter('HD800, ۸۰۱، HD802')
        self.assertEqual(s.stage,'EDIT_CONFIRM')
        self.assertEqual(s.proposal,before)
        await self.click('edit_pending_1')
        await self.enter('1600')
        self.assertEqual(s.proposal,before)
        self.handler = self.new_handler()
        s = self.handler.pending[self.key]
        await self.click('edit_apply')
        self.assertEqual([i['machine_code'] for i in s.proposal['items']],['HD800','HD801','HD802'])
        self.assertEqual(s.proposal['service_overrides'],{'HD801':1600})
        self.assertEqual([i['components']['oil_change']['next_interval'] for i in s.proposal['items']],[1800,1600,1800])

    async def test_batch_duplicates_unknown_codes_and_existing_units(self):
        s = await self.start(self.combined(),add_code='800')
        before = deepcopy(s.proposal)
        for codes in ('HD800 800','HD800 UNKNOWN'):
            result = await self.enter(codes)
            self.assertEqual(result['reason'],'work-order-input-rejected')
            self.assertEqual(s.proposal,before)
        await self.enter('HD800 HD801')
        self.assertIn('از قبل در حکم',self.sent[-1]['text'])
        await self.click('edit_apply')
        self.assertEqual([i['machine_code'] for i in s.proposal['items']],['HD800','HD801'])

    async def test_batch_back_never_copies_another_units_pending_interval(self):
        source=self.combined(2)
        source['plans'][1]['last_interval']=200
        s=await self.start(source)
        before=deepcopy(s.proposal)
        await self.enter('HD800 HD801')
        await self.click('edit_pending_0');await self.enter('1600')
        await self.click('edit_pending_1');await self.click('edit_back')
        self.assertEqual(s.edit_intervals,{'HD800':1600,'HD801':400})
        await self.click('edit_back')
        self.assertEqual(s.stage,'EDIT_CODE')
        self.assertEqual(s.proposal,before)

    async def test_batch_failure_is_atomic_and_shared_oil_add_uses_snapshot(self):
        from tools.fleet.oil_change.proposal import resolve_items
        from tools.fleet.work_orders.channels.bale.proposal_form import add_items
        source=self.combined(2)
        s=await self.start(source)
        before=deepcopy(s.proposal)
        with self.assertRaises(ValueError):
            apply_oil_edits(s.proposal,['HD800','HD801'],{'HD800':1600,'HD801':1801})
        self.assertEqual(s.proposal,before)
        fresh=deepcopy(source)
        fresh['plans'][0]['last_interval']=1800
        fresh['machines'][0]['current_meter']=3100
        add_items(s.proposal,resolve_items(['HD800'],fresh),'')
        item=s.proposal['items'][0]
        self.assertEqual(item['action_code'],'OIL_CHANGE_HD465-7R_1800')
        self.assertEqual(item['components']['oil_change']['next_interval'],1800)
        self.assertEqual(item['components']['oil_change']['current_meter'],1002)

    async def test_batch_limit_and_confirmation_pagination(self):
        s = await self.start(self.combined(100))
        codes = [f'HD{n}' for n in range(800,900)]
        result = await self.enter(' '.join(codes + ['HD900']))
        self.assertEqual(result['reason'],'work-order-input-rejected')
        await self.enter(' '.join(codes))
        self.assertEqual(len(s.edit_codes),100)
        old = self.button(self.sent[-1],'edit_pending_0')
        await self.click('edit_confirm_next')
        self.assertEqual(s.edit_confirm_page,1)
        self.assertEqual(self.message('',data=old,origin=len(self.sent))['reason'],'inline-rejected')
        await self.settle()
        # A rejected old callback sends a notice; get the latest active card.
        self.handler._send_reply(self.gateway,self.key[2],'batch',lambda *a:None,key=self.key)
        await self.settle()
        await self.click('edit_pending_16')
        self.assertEqual(s.edit_machine_code,'HD816')
        await self.click('interval_keep')
        await self.click('edit_apply')
        self.assertEqual(len(s.proposal['items']),100)

    async def test_aliases_and_ambiguous_numeric_codes_use_snapshot(self):
        source = self.combined(1)
        extra = sample_source('PC600-8','EX800',remaining=100)
        source['machines'] += extra['machines'];source['plans'] += extra['plans']
        s = await self.start(source)
        result = await self.enter('۸۰۰')
        self.assertEqual(result['reason'],'work-order-input-rejected')
        self.assertEqual(s.stage,'EDIT_CODE')
        await self.enter('hd۸۰۰');self.assertEqual(s.edit_machine_code,'HD800')
        source = sample_source('D155A-6','W151',remaining=100)
        source['machines'][0].update(code='D155',legacy_code='151')
        s = await self.start(source)
        await self.enter('W۱۵۱');self.assertEqual(s.edit_machine_code,'D151')

    async def test_cancel_only_current_edit_preserves_original_proposal(self):
        s = await self.start(add_code='1252')
        before = deepcopy(s.proposal)
        await self.enter('EX1252');await self.enter('1600')
        await self.click('edit_cancel')
        self.assertEqual(s.stage,'PROPOSAL')
        self.assertIs(self.handler.pending[self.key],s)
        self.assertEqual(s.proposal,before)
        self.assertFalse(s.edit_codes)

    async def test_step_back_preserves_page_and_discards_pending_interval(self):
        s = await self.start(self.combined(20))
        await self.click('edit_next')
        await self.click('edit_unit_16')
        await self.enter('1600')
        await self.click('edit_back')
        self.assertEqual(s.stage,'EDIT_INTERVAL')
        self.assertEqual(s.edit_intervals['HD816'],1800)
        await self.click('edit_back')
        self.assertEqual(s.stage,'EDIT_CODE')
        self.assertEqual(s.edit_code_page,1)
        self.assertFalse(s.proposal['items'])
        await self.click('edit_back')
        self.assertEqual(s.stage,'PROPOSAL')

    async def test_near_due_warning_returns_after_removal(self):
        s = await self.start(self.combined(1,remaining=14))
        original = deepcopy(s.proposal['evaluations'])
        self.assertEqual(len(s.proposal['warnings']),1)
        await self.enter('HD800');await self.click('interval_keep');await self.click('edit_apply')
        self.assertFalse(s.proposal['warnings'])
        await self.click('remove');await self.enter('1');await self.enter('تایید حذف')
        self.assertEqual(len(s.proposal['warnings']),1)
        self.assertIn('14 ساعت',s.proposal['warnings'][0])
        self.assertEqual(s.proposal['evaluations'],original)

    async def test_editor_never_rereads_source_and_creation_requires_refresh(self):
        s = await self.start()
        snapshot = deepcopy(s.proposal['evaluations'])
        with patch('tools.fleet.oil_change.proposal.read_source',side_effect=AssertionError('snapshot only')):
            await self.enter('EX1252');await self.enter('1800');await self.click('edit_apply')
        self.assertEqual(s.proposal['evaluations'],snapshot)
        with patch('tools.fleet.oil_change.source.source_hash',return_value='changed'),patch('tools.fleet.work_orders.core.service.create_work_order') as create:
            self.result = create_worker.execute_request(self.request(s),db_path=self.db_path)
            create.assert_not_called()
        await self.click('confirm')
        self.assertEqual(s.stage,'OIL_RECOVERY')
        self.assertEqual(s.oil_recovery,'refresh')
        self.assertIn('اطلاعات مبنا تغییر کرده',self.sent[-1]['text'])
        self.assertEqual(set(keyboard_for(s).actions),{'proposal_refresh','cancel'})
        fresh = deepcopy(s.proposal);fresh['source_sha256']='changed'
        self.result = {'ok':True,'proposal':fresh}
        await self.click('proposal_refresh')
        self.assertEqual(self.requests[-1]['action'],'propose')
        self.assertEqual(s.stage,'PROPOSAL')
        self.assertEqual(s.proposal['source_sha256'],'changed')

    async def test_correctable_error_returns_without_automatic_creation(self):
        s = await self.oil_edit_step()
        before = deepcopy(s.proposal)
        self.result = {'ok':False,'error':'INVALID_INPUT','message':'correctable test error'}
        await self.click('confirm')
        self.assertEqual(s.stage,'OIL_RECOVERY')
        requests = len(self.requests)
        await self.click('proposal_back')
        self.assertEqual(s.stage,'PROPOSAL')
        self.assertEqual(s.proposal,before)
        self.assertEqual(len(self.requests),requests)

    async def test_uncertain_build_error_offers_inspection_without_retry(self):
        s = await self.oil_edit_step()
        self.result = {'ok':False,'error':'CREATE_FAILED','message':'uncertain build'}
        await self.click('confirm')
        self.assertEqual(s.oil_recovery,'inspect')
        self.assertIn('proposal_inspect',keyboard_for(s).actions)
        self.result = {'ok':False,'error':'INVALID_INPUT','message':'no file found'}
        await self.click('proposal_inspect')
        self.assertEqual(self.requests[-1]['action'],'preview_latest')
        self.assertEqual(s.stage,'OIL_RECOVERY')
        count = len(self.requests)
        await self.click('proposal_back')
        self.assertEqual(s.stage,'PROPOSAL')
        self.assertEqual(len(self.requests),count)

    async def test_source_refresh_of_revision_keeps_parent_identity(self):
        s = await self.oil_edit_step()
        s.proposal.update(revision_of='OC-1405-06-23-001',revision_machine_code='EX1252')
        self.result = {'ok':False,'error':'INVALID_INPUT','recovery':'refresh','message':'changed source'}
        await self.click('confirm')
        self.result = {'ok':True,'work_order_type':'OIL_CHANGE','proposal':deepcopy(s.proposal)}
        await self.click('proposal_refresh')
        self.assertEqual(self.requests[-1],{'action':'edit','work_order_no':'OC-1405-06-23-001','bale_id':self.key[1]})
        self.assertEqual(s.proposal['revision_of'],'OC-1405-06-23-001')

    async def test_revision_parent_rejected_at_selection_other_unit_removable(self):
        s = await self.oil_edit_step(self.combined(2,remaining=-2),add_code=None)
        s.proposal.update(revision_of='OC-1405-06-23-001',revision_machine_code='HD800')
        await self.click('remove')
        await self.click('pick_0')
        self.assertEqual(s.removal_selection['selected'],[])
        self.assertIn('دستگاه اصلی حکم',self.sent[-1]['text'])
        result = await self.enter('1')
        self.assertEqual(result['reason'],'work-order-input-rejected')
        self.assertEqual(s.removal_selection['selected'],[])
        await self.click('pick_1');await self.click('select_confirm')
        self.assertEqual([i['machine_code'] for i in s.proposal['items']],['HD800'])

    async def test_legacy_session_is_migrated_without_losing_confirmed_draft(self):
        s = await self.oil_edit_step()
        before = deepcopy(s.proposal)
        s.stage='ADD_CODES';s.oil_editor_version=0
        self.handler._persist()
        self.handler = self.new_handler()
        restored = self.handler.pending[self.key]
        self.assertEqual(restored.stage,'PROPOSAL')
        self.assertEqual(restored.proposal,before)
        self.assertEqual(restored.keyboard_revision,'')

    async def test_net_manager_deputy_and_unauthorized_actor(self):
        from tools.authorization.store import AuthorizationStore
        from tools.authorization.net import migrate_net, NET_MANAGER, NET_DEPUTY
        path = self.db_path.parent/'central.db'
        users = [('641220453',NET_MANAGER),('1294822197',NET_DEPUTY)]
        with closing(sqlite3.connect(path)) as con,con:
            con.execute('CREATE TABLE channel_users(platform TEXT,user_id TEXT,chat_id TEXT,registration_status TEXT,PRIMARY KEY(platform,user_id))')
            con.executemany('INSERT INTO channel_users VALUES(?,?,?,?)',[('bale',u,u,'approved') for u,_ in users])
        auth=AuthorizationStore(path)
        auth.migrate(self.db_path.parent/'backups');auth.migrate_admin(self.db_path.parent/'backups')
        migrate_net(auth,self.db_path.parent/'backups',assignments=users,actor='ISOLATED TEST')
        with patch('tools.authorization.net.AuthorizationStore',return_value=auth):
            for user,_role in users:
                self.key=('bale',user,user)
                s=await self.start()
                await self.enter('EX1252');await self.enter('1800')
                with closing(sqlite3.connect(path)) as con,con:
                    con.execute('UPDATE auth_user_roles SET active=0 WHERE user_id=?',(user,))
                await self.click('edit_apply')
                self.assertFalse(s.proposal['items'])
            self.key=('bale','900009','900009')
            self.assertEqual((await self.enter('حکم کار'))['reason'],'work-order-permission-denied')

    async def test_real_excel_pdf_review_and_revision_from_net_menu(self):
        from tools.bale_ui import StateStore
        from tools.fleet.work_orders.channels.bale.message_handler import send_manager_excel
        from tools.fleet.work_orders.core import service, review, delivery
        from openpyxl import load_workbook
        import fitz
        fixture=OilServiceEditTests('test_repeated_override_retains_original_evidence_for_all_cycles')
        fixture.setUp();self.addCleanup(fixture.doCleanups)
        # Exercise HD465 end to end; PC1250's three native print contracts are
        # checked separately below, including every unchanged technical cell.
        code='HD464'
        fixture.source=sample_source('HD465-7R',code,last=1600,remaining=90.5)
        self.db_path=fixture.db_path;self.store=StateStore(self.db_path.parent/'real-ui.json')
        permission_patch=patch('tools.fleet.work_orders.core.permissions.DB_PATH',self.db_path)
        permission_patch.start();self.addCleanup(permission_patch.stop)
        documents=[]
        document_errors=[]
        async def send_document(**kwargs):
            documents.append((kwargs['filename'],kwargs['document'].read()))
        self.gateway.adapters['bale']._bot.send_document=send_document
        async def worker(request):
            self.requests.append(deepcopy(request))
            return create_worker.execute_request(request,db_path=self.db_path)
        async def document_sender(gateway,chat,order):
            try:
                await send_manager_excel(gateway,chat,order)
            except Exception:
                import traceback
                document_errors.append(traceback.format_exc())
                raise
        import subprocess
        real_run=subprocess.run
        def diagnostic_run(*args,**kwargs):
            result=real_run(*args,**kwargs)
            if result.returncode:
                document_errors.append(result.stderr.decode('utf-8',errors='replace'))
            return result
        process_patch=patch('tools.fleet.work_orders.core.pdf_document.subprocess.run',diagnostic_run)
        process_patch.start();self.addCleanup(process_patch.stop)
        self.worker=worker;self.document_sender=document_sender;self.handler=self.new_handler()
        await self.enter('حکم کار');await self.click('oil')
        await self.click('edit_excel');await self.enter(code);await self.enter('1800')
        await self.click('edit_apply');await self.click('confirm')
        self.assertEqual(len(documents),1,document_errors)
        number=self.handler.pending[self.key].order_no
        order=service.get_work_order(number)
        book=load_workbook(order['excel_path'])
        try:
            self.assertEqual(book.sheetnames,['1800'])
            self.assertEqual(book.active['F3'].value,code)
            self.assertEqual(book.active['C2'].value,1800)
        finally:book.close()
        with fitz.open(stream=documents[0][1],filetype='pdf') as pdf:
            text='\n'.join(p.get_text() for p in pdf)
            compact=''.join(text.split())
            self.assertIn(code,compact);self.assertIn('1800',compact)
        await self.click('review_confirm')
        reviewed=service.get_work_order(number)
        identity=json.loads(next(line.split(': ',1)[1] for line in reviewed['notes'].splitlines() if line.startswith('MANAGER_DOCUMENT_REVIEW: ')))
        self.assertEqual(identity['pdf_sha256'],hashlib.sha256(documents[0][1]).hexdigest())
        self.assertEqual(identity['sha256'],hashlib.sha256(Path(order['excel_path']).read_bytes()).hexdigest())
        self.assertEqual(review.approved_pdf_bytes(reviewed),documents[0][1])
        # Back from staff selection, reopen this exact reviewed order, revise it.
        await self.click('staff_back');await self.click('review_edit')
        s=self.handler.pending[self.key]
        self.assertEqual(s.proposal['revision_machine_code'],code)
        await self.click('edit_excel');await self.enter(code);await self.enter('1600')
        await self.click('edit_apply');await self.click('confirm')
        revised=self.handler.pending[self.key].order_no
        self.assertNotEqual(revised,number)
        self.assertEqual(service.get_work_order(number)['superseded_by'],revised)
        self.assertEqual(service.get_work_order(revised)['revision_of'],number)
        with self.assertRaisesRegex(ValueError,'منسوخ'):
            review.confirm_document_review(number,self.key[1])
        with self.assertRaisesRegex(ValueError,'منسوخ'):
            delivery.send_work_order(work_order_no=number,sender=None)
        evidence=os.environ.get('KOMATSO_OIL_EVIDENCE')
        if evidence:
            destination=Path(evidence);destination.mkdir(parents=True,exist_ok=True)
            for result in (order,service.get_work_order(revised)):
                for path in [Path(result['excel_path']),Path(result['pdf_path'])]:
                    shutil.copy2(path,destination/path.name)
            (destination/'e2e.json').write_text(json.dumps(dict(success=True,transport='SIMULATED',pdf='NATIVE EXCEL',
                original=number,revision=revised,review_hash_verified=True,old_dispatch_blocked=True),indent=2),encoding='utf-8')

    async def test_pc1250_real_outputs_1600_1800_2000_change_only_cycle_title(self):
        import fitz
        from openpyxl import load_workbook
        from zipfile import ZipFile
        from tools.fleet.work_orders.core import service
        from tools.fleet.work_orders.core.pdf_document import export_staff_pdf
        from tools.fleet.work_orders.types.oil_change.builder import TEMPLATES, action_for
        fixture=OilServiceEditTests('test_repeated_override_retains_original_evidence_for_all_cycles')
        fixture.setUp();self.addCleanup(fixture.doCleanups)
        source=TEMPLATES['PC1250-8'][0]
        before=hashlib.sha256(source.read_bytes()).hexdigest()
        original=load_workbook(source)
        self.addCleanup(original.close)
        results=[]
        for interval in (1600,1800,2000):
            order=service.create_work_order(work_order_type='OIL_CHANGE',jalali_date=f'1405/06/{interval//200:02d}',
                shift='روزانه',machine_codes=['EX1252'],item_actions={'EX1252':action_for('PC1250-8',interval)},
                created_by='bale:455740857')
            workbook=load_workbook(order['excel_path'])
            try:
                self.assertEqual(workbook.sheetnames,[str(interval)])
                expected,actual=original[str(interval)],workbook.active
                self.assertEqual(actual['C2'].value,str(interval)+'ساعت')
                self.assertEqual(actual['C2']._style,expected['C2']._style)
                for row in expected:
                    for cell in row:
                        if cell.coordinate not in {'C2','F3','E2','J2'}:
                            self.assertEqual(actual[cell.coordinate].value,cell.value,cell.coordinate)
                        self.assertEqual(actual[cell.coordinate]._style,cell._style,cell.coordinate)
                self.assertEqual(str(actual.merged_cells),str(expected.merged_cells))
            finally:workbook.close()
            with ZipFile(source) as old,ZipFile(order['excel_path']) as new:
                changed=[n for n in old.namelist() if old.read(n)!=new.read(n)]
                self.assertEqual(len(changed),2)
            pdf=await asyncio.to_thread(export_staff_pdf,Path(order['excel_path']))
            with fitz.open(pdf) as document:
                compact=''.join(''.join(p.get_text() for p in document).split())
                self.assertIn('EX1252',compact)
                self.assertIn(str(interval)+'ساعت',compact)
            evidence=os.environ.get('KOMATSO_OIL_EVIDENCE')
            if evidence:
                directory=Path(evidence);directory.mkdir(parents=True,exist_ok=True)
                shutil.copy2(order['excel_path'],directory/Path(order['excel_path']).name)
                shutil.copy2(pdf,directory/pdf.name)
            results.append(dict(interval=interval,excel_sha256=hashlib.sha256(Path(order['excel_path']).read_bytes()).hexdigest(),
                                pdf_sha256=hashlib.sha256(pdf.read_bytes()).hexdigest(),technical_cells_preserved=True))
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(),before)
        if os.environ.get('KOMATSO_OIL_EVIDENCE'):
            (Path(os.environ['KOMATSO_OIL_EVIDENCE'])/'pc1250-native.json').write_text(
                json.dumps(dict(success=True,source_unchanged=True,outputs=results),indent=2),encoding='utf-8')


def load_tests(loader, tests, pattern):
    # Reuse the transport fixture without counting inherited baseline tests twice.
    return unittest.TestSuite(OilUnifiedTests(name) for name in OilUnifiedTests.__dict__ if name.startswith('test_'))

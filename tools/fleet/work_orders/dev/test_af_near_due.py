"""AF forecast boundaries and explicit selection using disposable source/DB/UI."""
from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tools.fleet.air_filter import proposal as af
from tools.fleet.air_filter.rules import BOTH, OUTER
from tools.fleet.air_filter.test_proposal import day
from tools.fleet.work_orders.channels.bale import create_worker, proposal_form
from tools.fleet.work_orders.channels.bale.message_handler import FormSession
from tools.fleet.work_orders.core import service
from tools.fleet.work_orders.dev import test_inline_menus as inline
from tools.fleet.work_orders.dev import test_oil_p1_safety as shared


def source_file(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1 (2)'
    ws.append(['ردیف', 'نام دستگاه', 'کد دستگاه'])
    for index, date in enumerate(('06.20', '06.21', '06.22')):
        for offset, field in enumerate(('کارکرد', 'درونی', 'بیرونی')):
            col = 4 + index * 3 + offset
            ws.cell(1, col, date)
            ws.cell(2, col, field)
    for row, code, hours in [(3, '710', 6), (4, '719', 5)]:
        ws.cell(row, 1, row - 2)
        ws.cell(row, 2, 'دامپتراک')
        ws.cell(row, 3, code)
        ws.cell(row, 4, 24)  # Service-day hours remain excluded.
        ws.cell(row, 5, 'صبح')
        ws.cell(row, 7, 14)
        ws.cell(row, 10, hours)
    wb.save(path)
    wb.close()
    return path


class ForecastTests(unittest.TestCase):
    def evaluate(self, hours, horizon=3, daily=None, code='710', name='دامپتراک'):
        return af.evaluate_machine(dict(code=code, name=name, daily=daily if daily is not None else
                                   [day((6, 20), 24, inner='صبح'), day((6, 22), hours)]),
                                   (6, 22), (6, 23), forecast_horizon_hours=horizon)

    def test_zero_exact_due_and_one_three_four_six_remaining(self):
        for horizon in (0, 3, 6):
            for remaining in (-1, 0, 1, 3, 4, 6, 7):
                with self.subTest(horizon=horizon, remaining=remaining):
                    r = self.evaluate(20 - remaining, horizon)
                    c = r['components']['outer']
                    expected = 'DUE' if remaining <= 0 else 'NEAR_DUE' if remaining <= horizon else 'OK'
                    self.assertEqual(c['state'], expected)
                    self.assertEqual(c['remaining_hours'], remaining)
                    self.assertEqual(c['value'], 20 - remaining)
                    self.assertEqual(r['action_code'], OUTER if remaining <= 0 else None)

    def test_inner_and_outer_independent_and_mixed_states_keep_due_action(self):
        # 97 inner / 20 outer: the outer is due, inner is forecast-only.
        daily = [day((6, 1), 24, inner='صبح'), day((6, 2), 77),
                 day((6, 20), 0, outer='شب'), day((6, 22), 20)]
        r = self.evaluate(None, daily=daily)
        self.assertEqual((r['components']['inner']['state'], r['components']['outer']['state']), ('NEAR_DUE', 'DUE'))
        self.assertEqual(r['action_code'], OUTER)
        # 100 inner / 17 outer: preserve the existing BOTH action.
        daily[1]['hours'] = 83
        daily[-1]['hours'] = 17
        r = self.evaluate(None, daily=daily)
        self.assertEqual((r['components']['inner']['state'], r['components']['outer']['state']), ('DUE', 'NEAR_DUE'))
        self.assertEqual(r['action_code'], BOTH)

    def test_missing_baseline_and_untrustworthy_hours_never_forecast(self):
        variants = [[day((6, 22), 19)],
                    [day((6, 20), inner='صبح'), day((6, 22), 'نامشخص')],
                    [day((6, 20), inner='صبح'), day((6, 22), 19, outer='نامشخص')]]
        for daily in variants:
            with self.subTest(daily=daily):
                r = self.evaluate(None, daily=daily)
                self.assertEqual(r['components']['outer']['state'], 'NEEDS_REVIEW')
                self.assertTrue(r['issues'])
                self.assertIsNone(r['action_code'])
        # A blank post-service sum must not become a 6h forecast for a 5h interval.
        for blank in (None, '', '-'):
            r = self.evaluate(None, 6, code='EX801', name='بیل مکانیکی',
                              daily=[day((6, 20), 24, inner='صبح'), day((6, 22), blank)])
            self.assertEqual(r['components']['outer']['state'], 'NEEDS_REVIEW')
            self.assertTrue(r['issues'])
        r = self.evaluate(0, 6, code='EX801', name='بیل مکانیکی')
        self.assertEqual(r['components']['outer']['state'], 'NEAR_DUE')  # Explicit recorded zero.

    def test_five_recorded_hours_once_and_forecast_never_changes_input(self):
        daily = [day((6, 20), 24, inner='صبح'), day((6, 21), 14), day((6, 22), 5)]
        before = deepcopy(daily)
        for horizon in (0, 3, 6):
            r = self.evaluate(None, horizon, daily)
            self.assertEqual(r['components']['outer']['value'], 19)
            self.assertEqual(r['components']['outer']['last_service'], '1405/06/20')
            self.assertIsNone(r['action_code'])
        self.assertEqual(daily, before)

    def test_forecast_does_not_change_calendar_rules_or_dates(self):
        for horizon in (0, 3, 6):
            r = self.evaluate(None, horizon, [day((6, 22), outer='صبح')], 'MZ1', 'مزدا')
            self.assertEqual(r['components']['outer']['state'], 'OK')
            self.assertNotIn('remaining_hours', r['components']['outer'])

    def test_invalid_horizons_rejected(self):
        for horizon in (-1, True, float('nan'), float('inf'), '3'):
            with self.assertRaises(ValueError):
                self.evaluate(19, horizon)

    def test_workbook_immutable_due_set_stable_and_warning_is_not_an_item(self):
        root = Path('runtime/af_near_due_tests')
        root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as temp:
            path = source_file(Path(temp)/'source.xlsx')
            before = path.read_bytes(), path.stat().st_mtime_ns
            proposals = [af.build_proposal(path, forecast_horizon_hours=h) for h in (0, 3, 6)]
            for p in proposals:
                self.assertEqual([i['machine_code'] for i in p['items']], ['710'])
                self.assertEqual((p['cutoff'], p['plan_date']), ('1405/06/22', '1405/06/23'))
            self.assertFalse(proposals[0]['warnings'])
            self.assertEqual(proposals[1]['warnings'], ['719 : بیرونی: 1 ساعت باقی‌مانده'])
            self.assertEqual(before, (path.read_bytes(), path.stat().st_mtime_ns))

    def test_mixed_machine_one_card_with_both_component_states(self):
        r = self.evaluate(None, daily=[day((6, 1), inner='صبح'), day((6, 2), 77),
                                      day((6, 20), outer='صبح'), day((6, 22), 20)])
        with patch.object(af, 'read_source', return_value=([{}], (6, 22), (6, 23), [], 'test')), \
                patch.object(af, 'evaluate_machine', return_value=r):
            p = af.build_proposal()
        self.assertEqual(len(p['items']), 1)
        self.assertFalse(p['warnings'])
        text = proposal_form.render(p)
        self.assertEqual(text.count('دستگاه 710'), 1)
        self.assertIn('موعدرسیده', text)
        self.assertIn('نزدیک موعد — پیش‌بینی', text)
        self.assertIn('3 ساعت تا موعد سرویس', text)
        self.assertEqual(p['items'][0]['action_code'], OUTER)


class ForecastOrderTests(unittest.TestCase):
    setUp = shared.OilSafetyTests.setUp

    def proposal(self):
        path = source_file(self.db_path.parent/'af.xlsx')
        machines = af.read_source(path)
        for patcher in (patch.object(af, 'SOURCE', path), patch.object(af, 'read_source', return_value=machines)):
            patcher.start()
            self.addCleanup(patcher.stop)
        return path, af.build_proposal(path)

    def create(self, p):
        result = create_worker.execute_request(dict(action='create', bale_id=shared.MANAGER,
            work_order_type='AIR_FILTER', jalali_date=self.fixture.test_today, shift='صبح',
            machine_codes=[i['machine_code'] for i in p['items']],
            item_actions={i['machine_code']:i['action_code'] for i in p['items']},
            proposal=proposal_form.creation_metadata(p)), db_path=self.db_path)
        self.assertTrue(result['ok'], result)
        return service.get_work_order(result['order']['work_order_no'])

    def assert_excel(self, order, expected):
        wb = load_workbook(order['excel_path'])
        try:
            codes = [str(wb.active.cell(row, 3).value) for row in range(3, 3+len(expected))]
            self.assertEqual(codes, expected)
        finally:
            wb.close()
        self.assertEqual([i['machine_code'] for i in order['items']], expected)
        self.assertIsNone(order['sent_at'])

    def test_unselected_forecast_not_created_or_exported(self):
        path, p = self.proposal()
        before = path.read_bytes()
        order = self.create(p)
        self.assert_excel(order, ['710'])
        self.assertEqual(path.read_bytes(), before)

    def test_manual_near_due_uses_existing_validation_action_and_confirmation(self):
        path, p = self.proposal()
        before = path.read_bytes()
        r = create_worker.execute_request(dict(action='validate_add', bale_id=shared.MANAGER,
            work_order_type='AIR_FILTER', machine_codes=['719']), db_path=self.db_path)
        self.assertTrue(r['ok'], r)
        proposal_form.add_items(p, r['items'], OUTER)
        self.assertFalse(p['warnings'])
        self.assertIn('با انتخاب مسئول نت', p['items'][-1]['manual_note'])
        self.assertEqual(p['items'][-1]['components']['outer']['value'], 19)
        self.assert_excel(self.create(p), ['710', '719'])
        self.assertEqual(path.read_bytes(), before)

    def test_permission_source_change_and_duplicate_guards_preserved(self):
        path, p = self.proposal()
        for actor, codes in [('1003', ['719']), (shared.MANAGER, ['719', '719']), (shared.MANAGER, ['UNKNOWN'])]:
            r = create_worker.execute_request(dict(action='validate_add', bale_id=actor,
                work_order_type='AIR_FILTER', machine_codes=codes), db_path=self.db_path)
            self.assertFalse(r['ok'], r)
        path.write_bytes(path.read_bytes()+b'changed')
        req = dict(action='create', bale_id=shared.MANAGER, work_order_type='AIR_FILTER',
                   machine_codes=['710'], item_actions={'710':OUTER}, jalali_date=self.fixture.test_today,
                   shift='صبح', proposal=proposal_form.creation_metadata(p))
        self.assertFalse(create_worker.execute_request(req, db_path=self.db_path)['ok'])


class ForecastUITests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.h = inline.InlineMenusTests()
        self.h.setUp()
        self.addCleanup(self.h.doCleanups)
        h = self.h
        path = source_file(h.db_path.parent/'af.xlsx')
        self.p = af.build_proposal(path)
        self.p['work_order_type'] = 'AIR_FILTER'
        h.test_today = self.p['plan_date']
        h.handler.pending[h.key] = FormSession(expires=h.handler.clock()+600, stage='PROPOSAL',
            work_order_type='AIR_FILTER', jalali_date=h.test_today, proposal=self.p)

    async def test_confirm_without_add_sends_only_due_to_existing_worker(self):
        h = self.h
        text = proposal_form.render(self.p)
        self.assertIn('نزدیک موعد — پیش‌بینی تا 3 ساعت کارکرد دیگر', text)
        self.assertIn('به حکم اضافه نشده‌اند', text)
        h.message('تایید')
        await h.settle()
        self.assertFalse(h.requests)
        self.assertEqual(h.handler.pending[h.key].stage, 'SHIFT')
        h.message('', data=h.button(h.sent[-1], 'shift_morning'), origin=len(h.sent))
        await h.settle()
        self.assertEqual(h.requests[-1]['machine_codes'], ['710'])

    async def test_manual_add_back_cancel_and_confirm_keep_existing_stages(self):
        h = self.h
        h.message('اضافه')
        await h.settle()
        h.result = {'ok':True, 'items':[dict(machine_code='719', machine_name='دامپتراک')]}
        h.message('719')
        await h.settle()
        self.assertEqual(h.handler.pending[h.key].stage, 'ADD_ACTION')
        self.assertEqual([i['machine_code'] for i in self.p['items']], ['710'])
        h.message('برگشت')
        await h.settle()
        self.assertEqual([i['machine_code'] for i in self.p['items']], ['710'])
        h.message('اضافه')
        await h.settle()
        h.message('719')
        await h.settle()
        h.message('1')
        await h.settle()
        self.assertEqual([i['machine_code'] for i in self.p['items']], ['710', '719'])
        self.assertFalse(any(r['action'] == 'create' for r in h.requests))
        h.result = {'ok':False, 'message':'TEST creation request captured locally'}
        h.message('تایید')
        await h.settle()
        h.message('', data=h.button(h.sent[-1], 'shift_night'), origin=len(h.sent))
        await h.settle()
        self.assertEqual(h.requests[-1]['machine_codes'], ['710', '719'])
        self.assertEqual(h.requests[-1]['item_actions'], {'710':OUTER, '719':OUTER})
        self.assertEqual(h.requests[-1]['shift'], 'شب')

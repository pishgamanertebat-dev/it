"""Exercise the staged real plugin and Hermes tools namespace, offline."""

from __future__ import annotations

import asyncio
import importlib
import importlib.util
import os
import subprocess
import sqlite3
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace


HERMES_CODE = Path(r"C:\Users\win-10\AppData\Local\hermes\hermes-agent")
STAGE_ROOT = Path(r"E:\KomatsoAI\runtime\work_order_bale_pilot")


async def verify(root: Path) -> None:
    os.environ["HERMES_HOME"] = str(root / "hermes_home")
    os.environ["BALE_ADMIN_IDS"] = "455740857"
    sys.path.insert(0, str(HERMES_CODE))
    # Use the installed host dependencies while retaining the approved project
    # interpreter; no package installation or environment-file edits are needed.
    sys.path.append(str(HERMES_CODE / "venv" / "Lib" / "site-packages"))
    import tools
    original_tools = tools
    assert Path(tools.__file__).resolve() == HERMES_CODE / "tools" / "__init__.py"

    plugin_path = Path(r"E:\KomatsoAI\runtime\staff_roster\plugin.py")
    spec = importlib.util.spec_from_file_location("_bale_pilot_plugin", plugin_path, submodule_search_locations=[str(plugin_path.parent)])
    plugin = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = plugin
    spec.loader.exec_module(plugin)
    plugin.DB_PATH = root / "registration.db"
    replies = []

    async def send(chat, text):
        replies.append((chat, text))
        return {"ok": True}

    documents = []

    async def send_document(**kwargs):
        documents.append((kwargs["chat_id"], kwargs["document"].read()))
        return SimpleNamespace(message_id="test-document")

    async def send_message(**kwargs):
        replies.append((kwargs['chat_id'], kwargs['text']))
        return SimpleNamespace(message_id='test-text')

    gateway = SimpleNamespace(adapters={"bale": SimpleNamespace(send=send, _bot=SimpleNamespace(send_document=send_document, send_message=send_message))})

    async def event(text, user="455740857", platform="bale", chat_type="dm"):
        source = SimpleNamespace(platform=platform, user_id=user, chat_id=user, chat_type=chat_type, user_name="TEST")
        result = plugin._handle_bale(SimpleNamespace(source=source, text=text), gateway)
        await asyncio.sleep(0)
        return result

    assert await event("ordinary chat") is None
    assert sys.modules["tools"] is original_tools
    handler_module = importlib.import_module("tools.fleet.work_orders.channels.bale.message_handler")
    handler_module._handler.db_path = root / "permissions.db"
    migration = importlib.import_module("tools.fleet.work_orders.migrations.002_create_work_order_permissions")
    con = sqlite3.connect(root / "permissions.db")
    try:
        migration.create_schema(con)
        con.executemany("INSERT INTO service_work_order_users (bale_id, role, active) VALUES (?, 'MAINTENANCE_MANAGER', 1)", [("455740857",), ("1006",), ("1008",)])
        con.commit()
    finally:
        con.close()
    core_schema = importlib.import_module("tools.fleet.work_orders.migrations.001_create_work_order_schema_v1")
    con = sqlite3.connect(root / "permissions.db")
    try:
        con.execute("CREATE TABLE machines (id INTEGER PRIMARY KEY, canonical_code TEXT)")
        con.executemany('INSERT INTO machines(canonical_code) VALUES (?)', [('HD714',),('HD465',),('EX231',)])
        core_schema.create_schema(con)
        importlib.import_module('tools.fleet.work_orders.migrations.003_staff_dispatch').migrate(con)
        importlib.import_module('tools.fleet.work_orders.migrations.004_staff_roster').migrate(con)
        con.commit()
    finally:
        con.close()
    importlib.import_module('tools.fleet.work_orders.core.db').DB_PATH = root / 'permissions.db'
    importlib.import_module('tools.fleet.work_orders.core.permissions').DB_PATH = root / 'permissions.db'

    async def isolated_worker(request):
        process = await asyncio.create_subprocess_exec(
            r"E:\KomatsoAI\.venv\Scripts\python.exe", "-E", "-s", "-B", "-X", "utf8", "-m",
            "tools.fleet.work_orders.channels.bale.create_worker",
            "--test-db", str(root / "permissions.db"), "--test-output", str(root / "orders"),
            cwd=r"E:\KomatsoAI", stdin=asyncio.subprocess.PIPE,
            stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
        import json
        output, errors = await process.communicate(json.dumps(request).encode("utf-8"))
        assert process.returncode == 0, errors.decode("utf-8", errors="replace")
        return json.loads(output.decode("utf-8"))

    handler_module._handler.worker = isolated_worker
    con = plugin._connect()
    try:
        con.executemany(
            "INSERT INTO channel_users (platform,user_id,chat_id,registration_status,first_seen_at,updated_at) VALUES ('bale',?,?,?,?,?)",
            [(user,user,status,"TEST","TEST") for user,status in (("1006","approved"),("1007","approved"),("1008","rejected"))],
        )
        con.commit()
    finally:
        con.close()

    assert (await event("حکم کار"))["reason"] == "work-order-menu"
    assert replies[-1][0] == "455740857" and "1) هواکش" in replies[-1][1]
    assert (await event("۲"))["reason"] == "work-order-selection-rejected"
    assert (await event("۱"))["reason"] == "work-order-type-selected"
    await asyncio.gather(*list(handler_module._handler.tasks))
    await asyncio.sleep(0)
    session = next(iter(handler_module._handler.pending.values()))
    assert session.proposal['plan_date'] == '1405/06/09'
    assert len(session.proposal['items']) == 6
    await event('حذف')
    await event('۱ ۳')
    assert len(session.proposal['items']) == 4
    await event('اضافه')
    await event('714')
    await asyncio.gather(*list(handler_module._handler.tasks))
    await event('۲')
    assert session.proposal['items'][-1]['action_code'] == 'AIR_FILTER_INNER_OUTER'
    await event('تایید')
    assert (await event("صبح"))["reason"] == "work-order-creating"
    await asyncio.gather(*list(handler_module._handler.tasks))
    await asyncio.sleep(0)
    assert "✅ حکم کار ساخته شد" in replies[-1][1]
    assert len(list((root / "orders").rglob("*.xlsx"))) == 1
    assert len(documents) == 1 and documents[0][0] == "455740857"
    assert documents[0][1].startswith(b"PK")
    handler_module._handler.pending.clear()
    assert (await event("ثبت تأیید AF-1405-06-09-001"))["reason"] == "work-order-review"
    await asyncio.gather(*list(handler_module._handler.tasks))
    await asyncio.sleep(0)
    assert "تایید بررسی فایل" in replies[-1][1]
    handler_module._handler.pending.clear()
    assert (await event("ویرایش AF-1405-06-09-001"))["reason"] == "work-order-review"
    await asyncio.gather(*list(handler_module._handler.tasks))
    await asyncio.sleep(0)
    assert "پیشنهاد حکم هواکش" in replies[-1][1]
    await event('حذف')
    await event('۱')
    await event('اضافه')
    await event('714')
    await asyncio.gather(*list(handler_module._handler.tasks))
    await asyncio.sleep(0)
    await event('۱')
    await event('تایید')
    await event("صبح ظهر")
    await asyncio.gather(*list(handler_module._handler.tasks))
    await asyncio.sleep(0)
    assert "AF-1405-06-09-002" in replies[-1][1]
    assert "ویرایش" in replies[-1][1]
    assert len(documents) == 2
    assert len(list((root / "orders").rglob("*.xlsx"))) == 2
    from openpyxl import load_workbook
    wb = load_workbook(next((root / "orders").rglob("*002.xlsx")), read_only=True)
    assert wb.active["A1"].value == "لیست هواکش شیفت صبح-ظهر"
    assert wb.active['F1'].value == '1405/06/09'
    assert any(row[3] == 'تعویض هواکش داخلی و بیرونی' for row in wb.active.iter_rows(min_row=3,values_only=True))
    wb.close()
    denied = await isolated_worker({"action": "edit", "bale_id": "1006", "work_order_no": "AF-1405-06-09-001"})
    assert not denied["ok"]
    await event('تایید')
    await asyncio.gather(*list(handler_module._handler.tasks))
    await asyncio.sleep(0)
    assert '1) محسن غضنفری' in replies[-1][1] and '2) حسین محمودی' in replies[-1][1] and '3) پوریا آسترکی' in replies[-1][1]
    roster_text = replies[-1][1].split('سرویسکار را انتخاب کنید:')[1]
    assert (await event('۴'))['reason'] == 'work-order-input-rejected'
    assert (await event('۲'))['reason'] == 'work-order-dispatch'
    await asyncio.gather(*list(handler_module._handler.tasks))
    await asyncio.sleep(0)
    assert len(documents) == 3 and documents[-1][0] == '85539397'
    assert documents[-1][1] == documents[-2][1]
    await event('۲')
    await asyncio.gather(*list(handler_module._handler.tasks))
    assert len(documents) == 3
    staff_flow = importlib.import_module('tools.fleet.work_orders.channels.bale.staff_flow')
    assert (await event('تایید', '85539397'))['reason'] == 'staff-receipt'
    await asyncio.gather(*list(staff_flow.tasks))
    assert any(c == '85539397' and 'توسط شما تایید شد' in t for c,t in replies)
    assert any(c == '455740857' and 'دریافت حکم' in t for c,t in replies)
    notices = len([t for c,t in replies if c == '455740857' and 'دریافت حکم' in t])
    await event('تایید', '85539397')
    await asyncio.gather(*list(staff_flow.tasks))
    assert len([t for c,t in replies if c == '455740857' and 'دریافت حکم' in t]) == notices
    con = sqlite3.connect(root / 'permissions.db')
    row = con.execute("SELECT status,assigned_staff_id,approved_at,sent_at,acknowledged_at FROM service_work_orders WHERE work_order_no='AF-1405-06-09-002'").fetchone()
    con.close()
    assert row[0] == 'SENT' and all(row[1:])
    await event('حکم کار')
    assert (await event('۳'))['reason'] == 'work-order-type-selected'
    await asyncio.gather(*list(handler_module._handler.tasks))
    session = next(iter(handler_module._handler.pending.values()))
    assert session.stage == 'PROPOSAL'
    assert session.proposal['cutoff'] == '1405/06/15 - روز'
    assert session.proposal['plan_date'] == '1405/06/16'
    assert 's1' in [i['machine_code'] for i in session.proposal['items']]
    assert 'S1' not in [i['machine_code'] for i in session.proposal['items']]
    await event('حذف')
    await event('۱ ۲')
    await event('اضافه')
    await event('714 231')
    await asyncio.gather(*list(handler_module._handler.tasks))
    assert session.stage == 'PROPOSAL'  # Greasing needs no inner/outer choice.
    assert all(i['action_code']=='GREASING_FULL' for i in session.proposal['items'])
    assert (await event('تایید'))['reason'] == 'work-order-creating'
    await asyncio.gather(*list(handler_module._handler.tasks))
    await asyncio.sleep(0)
    assert 'GR-1405-06-16-001' in replies[-1][1], replies[-1]
    con = sqlite3.connect(root / 'permissions.db')
    assert con.execute("SELECT shift FROM service_work_orders WHERE work_order_no='GR-1405-06-16-001'").fetchone()[0] == 'روزانه'
    con.close()
    grease_path = next((root / 'orders').rglob('GR-*.xlsx'))
    wb = load_workbook(grease_path, read_only=True)
    assert wb.active['A1'].value == 'لیست روزانه گریسکاری'
    assert wb.active['D3'].value == wb.active['D4'].value == 'گریسکاری کامل'
    assert wb.active['F1'].value == '1405/06/16'
    assert any(row[2]=='s1' for row in wb.active.iter_rows(min_row=3,values_only=True))
    wb.close()
    handler_module._handler.pending.clear()
    await event('ویرایش GR-1405-06-16-001')
    await asyncio.gather(*list(handler_module._handler.tasks))
    await asyncio.sleep(0)
    assert 'پیشنهاد حکم گریس‌کاری' in '\n'.join(t for _,t in replies[-4:])
    await event('اضافه')
    await event('465')
    await asyncio.gather(*list(handler_module._handler.tasks))
    assert (await event('تایید'))['reason'] == 'work-order-creating'
    await asyncio.gather(*list(handler_module._handler.tasks))
    await event('تایید')
    await asyncio.gather(*list(handler_module._handler.tasks))
    await asyncio.sleep(0)
    assert replies[-1][1].split('سرویسکار را انتخاب کنید:')[1] == roster_text
    manager_copy = documents[-1][1]
    await event('۳')
    await asyncio.gather(*list(handler_module._handler.tasks))
    assert documents[-1] == ('85539397', manager_copy)
    assert (await event('تایید', '85539397'))['reason'] == 'staff-receipt'
    await asyncio.gather(*list(staff_flow.tasks))
    assert any(c == '455740857' and 'گریس‌کاری' in t and 'دریافت حکم' in t for c,t in replies)
    assert (await event('تایید GR-1405-06-16-002', '85539397'))['reason'] == 'staff-receipt'
    await asyncio.gather(*list(staff_flow.tasks))
    assert (await event('تایید GR-1405-06-16-001', '85539397'))['reason'] == 'staff-receipt-denied'
    # Re-open two isolated receipts to exercise the real plugin's numeric bridge.
    con = sqlite3.connect(root / 'permissions.db')
    con.execute("UPDATE service_work_orders SET acknowledged_at=NULL WHERE status='SENT'")
    con.commit()
    con.close()
    assert (await event('تایید', '85539397'))['reason'] == 'staff-receipt-select'
    assert '1) حکم' in replies[-1][1] and '2) حکم' in replies[-1][1]
    selected_numbers = list(staff_flow.receipt_choices[('85539397','85539397')]['numbers'])
    assert (await event('۱', '85539397'))['reason'] == 'staff-receipt'
    await asyncio.gather(*list(staff_flow.tasks))
    assert (await event('۲', '85539397'))['reason'] == 'staff-receipt'
    await asyncio.gather(*list(staff_flow.tasks))
    con = sqlite3.connect(root / 'permissions.db')
    assert all(con.execute('SELECT acknowledged_at FROM service_work_orders WHERE work_order_no=?',(n,)).fetchone()[0] for n in selected_numbers)
    con.close()
    assert (await event("حکم کار", "1006"))["reason"] == "work-order-menu"
    assert (await event("انصراف", "1006"))["reason"] == "work-order-menu-cancelled"
    assert (await event("حکم کار", "1007"))["reason"] == "work-order-permission-denied"
    assert (await event("حکم کار", "1008"))["reason"] == "bale-registration-rejected"
    assert (await event("حکم کار", "1009"))["reason"] == "bale-registration-started"
    count = len(replies)
    assert await event("حکم کار", platform="telegram") is None
    assert await event("حکم کار", chat_type="group") is None
    assert await event("ordinary chat", "1006") is None
    assert len(replies) == count
    assert sys.modules["tools"] is original_tools
    print("PASS: real staged plugin imports with Hermes tools namespace preserved")
    print("PASS: pilot menu, disabled type, selection, approved manager, cancellation")
    print("PASS: full Bale form -> real project worker -> FILE_READY order and Excel")
    print("PASS: manager document upload, persisted confirmation, edit after session loss, combined shifts, owner check")
    print('PASS: staff selection, disabled placeholders, identical Excel dispatch, acknowledgement, manager notification, no repeated dispatch or notification')
    print('PASS: greasing real-source proposal, remove/add without action choice, case-sensitive s1, fixed Excel, edit after session loss, dispatch and receipt')
    print("PASS: permission denial, rejected registration, new-user onboarding")
    print("PASS: Telegram, group messages and ordinary chat are unchanged")
    print("All replies used the fake adapter; all database writes used temporary files.")


if __name__ == "__main__":
    STAGE_ROOT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=STAGE_ROOT) as directory:
        asyncio.run(verify(Path(directory)))

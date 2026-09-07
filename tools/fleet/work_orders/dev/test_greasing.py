from pathlib import Path
from openpyxl import load_workbook
from tools.fleet.work_orders.dev.test_bale_work_order_create import WorkOrderCreateTests
from tools.fleet.work_orders.dev.test_permissions import test_database
from tools.fleet.work_orders.core import service
from tools.fleet.work_orders.types.greasing import builder


class GreasingTests(WorkOrderCreateTests):
    def test_aliases_and_fixed_action_are_validated(self):
        self.assertEqual(builder.get_items(['714'])[0]['machine_code'], 'HD714')
        for codes in (['714','HD714'],['UNKNOWN'],[]):
            with self.assertRaises(ValueError):
                builder.get_items(codes)
        with self.assertRaises(ValueError):
            builder.get_items(['714'],actions={'HD714':'AIR_FILTER_OUTER'})
        with test_database(self.db_path) as con:
            con.execute("INSERT INTO machines(canonical_code) VALUES ('EX714')")
        with self.assertRaises(ValueError):
            builder.get_items(['714'])

    def test_greasing_reuses_layout_with_fixed_title_and_independent_numbering(self):
        air = self.create()
        with test_database(self.db_path) as con:
            con.execute("INSERT INTO machines(canonical_code) VALUES ('EX231')")
        order = service.create_work_order(work_order_type='GREASING',jalali_date='1405/06/15',shift='صبح-ظهر',machine_codes=['465','710','711','712','713','714','231'],created_by='bale:455740857')
        self.assertEqual(order['work_order_no'],'GR-1405-06-15-001')
        self.assertEqual(order['status'],'FILE_READY')
        self.assertEqual(order['shift'],'صبح-ظهر')
        self.assertTrue(all(i['action_code']=='GREASING_FULL' for i in order['items']))
        self.assertIn('greasing',Path(order['excel_path']).parts)
        wb=load_workbook(order['excel_path'])
        old=load_workbook(air['excel_path'])
        try:
            self.assertEqual(wb.active['A1'].value,'لیست روزانه گریسکاری')
            self.assertEqual(wb.active['F1'].value,'1405/06/15')
            self.assertEqual([wb.active.cell(r,4).value for r in range(3,10)],['گریسکاری کامل']*7)
            self.assertEqual(wb.active['D3']._style,old.active['D3']._style)
            self.assertEqual(wb.active.column_dimensions['D'].width,old.active.column_dimensions['D'].width)
            self.assertIn('هواکش',old.active['A1'].value)
        finally:
            wb.close()
            old.close()

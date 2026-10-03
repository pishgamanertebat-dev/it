"""Isolated creation and audit checks for explicit oil-service cycle overrides."""
import importlib
import json
from copy import deepcopy
from unittest.mock import patch

from openpyxl import load_workbook

from tools.fleet.oil_change.proposal import build_proposal, resolve_items, override_interval
from tools.fleet.oil_change.test_proposal import sample_source
from tools.fleet.work_orders.channels.bale import create_worker
from tools.fleet.work_orders.channels.bale.proposal_form import add_items, edit_oil_interval, creation_metadata
from tools.fleet.work_orders.core import service
from tools.fleet.work_orders.dev.test_permissions import PermissionDatabaseTestCase, test_database
from tools.fleet.work_orders.types.oil_change.builder import TEMPLATES


class OilServiceEditTests(PermissionDatabaseTestCase):
    def setUp(self):
        super().setUp()
        self.source = sample_source('PC1250-8', 'EX1252', last=1800, remaining=100)
        self.output_root = self.db_path.parent / 'orders'
        with test_database(self.db_path) as con:
            con.execute('CREATE TABLE machines (id INTEGER PRIMARY KEY, canonical_code TEXT)')
            con.execute("INSERT INTO machines (canonical_code) VALUES ('EX1252')")
            importlib.import_module('tools.fleet.work_orders.migrations.001_create_work_order_schema_v1').create_schema(con)
        for target, kwargs in (
            ('tools.fleet.work_orders.core.db.DB_PATH', {'new':self.db_path}),
            ('tools.fleet.work_orders.core.service.WORK_ORDER_OUTPUT_ROOT', {'new':self.output_root}),
            ('tools.fleet.oil_change.source.source_hash', {'return_value':'fixture'}),
            ('tools.fleet.oil_change.source.read_source', {'side_effect':lambda *a,**kw:self.source}),
            ('tools.fleet.oil_change.proposal.read_source', {'side_effect':lambda *a,**kw:self.source}),
        ):
            patcher = patch(target, **kwargs)
            patcher.start()
            self.addCleanup(patcher.stop)

    def request(self):
        proposal = build_proposal()
        self.assertNotIn('EX1252', [item['machine_code'] for item in proposal['items']])
        add_items(proposal, resolve_items(['1252'], self.source), '')
        edit_oil_interval(proposal, '۱۲۵۲', '۱۸۰۰')
        return dict(action='create', bale_id='455740857', work_order_type='OIL_CHANGE',
            machine_codes=['EX1252'], jalali_date=proposal['plan_date'], shift='روزانه',
            item_actions={item['machine_code']:item['action_code'] for item in proposal['items']},
            proposal=creation_metadata(proposal))

    def test_1252_1800_creates_correct_sheet_with_audit_and_unchanged_source(self):
        source_before = deepcopy(self.source)
        template = TEMPLATES['PC1250-8'][0]
        template_before = template.read_bytes()
        result = create_worker.execute_request(self.request(), db_path=self.db_path)
        self.assertTrue(result['ok'], result)
        order = service.get_work_order(result['orders'][0]['work_order_no'])
        self.assertEqual(order['items'][0]['action_code'], 'OIL_CHANGE_PC1250-8_1800')
        wb = load_workbook(order['excel_path'])
        try:
            self.assertEqual(wb.sheetnames, ['1800'])
            self.assertEqual(wb.active['F3'].value, 'EX1252')
        finally:
            wb.close()
        audit = json.loads(order['notes'].split(': ', 1)[1])
        self.assertEqual(audit['service_overrides'], {'EX1252':1800})
        component = audit['item']['components']['oil_change']
        self.assertEqual((component['planned_interval'], component['next_interval']), (2000,1800))
        self.assertEqual(component['last_interval'],1800)
        self.assertEqual(component['remaining'],100)
        self.assertEqual(order['created_by'], 'bale:455740857')
        self.assertIsNone(order['sent_at'])
        self.assertEqual(self.source, source_before)
        self.assertEqual(template.read_bytes(), template_before)
        edited = create_worker.execute_request(dict(action='edit', bale_id='455740857',
            work_order_no=order['work_order_no']), db_path=self.db_path)
        self.assertTrue(edited['ok'], edited)
        self.assertEqual(edited['proposal']['items'][0]['action_code'], 'OIL_CHANGE_PC1250-8_1800')
        self.assertEqual(edited['proposal']['service_overrides'], {'EX1252':1800})

    def test_override_requires_valid_cycle_selected_identity_and_matching_action(self):
        request = self.request()
        cases = [None, [], {'EX1252':True}, {'EX1252':1801}, {'EX1252':'1800.0'},
                 {'EX1252':2200}, {'UNKNOWN':1800}, {'EX1252':'=1800'}, {}]
        with patch.object(service, 'create_work_order') as create:
            for overrides in cases:
                with self.subTest(overrides=overrides):
                    request['proposal']['service_overrides'] = overrides
                    result = create_worker.execute_request(request, db_path=self.db_path)
                    self.assertEqual(result['error'], 'INVALID_INPUT', result)
            request['proposal']['service_overrides'] = {'EX1252':1800}
            request['item_actions']['EX1252'] = 'OIL_CHANGE_PC800-7_1800'
            self.assertEqual(create_worker.execute_request(request, db_path=self.db_path)['error'], 'INVALID_INPUT')
            create.assert_not_called()

    def test_changed_source_and_revoked_permission_block_even_explicit_override(self):
        request = self.request()
        with patch.object(service, 'create_work_order') as create:
            request['proposal']['source_sha256'] = 'old'
            self.assertEqual(create_worker.execute_request(request, db_path=self.db_path)['error'], 'INVALID_INPUT')
            request['proposal']['source_sha256'] = 'fixture'
            with test_database(self.db_path) as con:
                con.execute("UPDATE service_work_order_users SET active=0 WHERE bale_id='455740857'")
            self.assertEqual(create_worker.execute_request(request, db_path=self.db_path)['error'], 'DENIED')
            create.assert_not_called()

    def test_repeated_override_retains_original_evidence_for_all_cycles(self):
        item = resolve_items(['1252'], self.source)[0]
        original = deepcopy(item)
        for interval in range(200,2001,200):
            item = override_interval(item, interval)
            component = item['components']['oil_change']
            self.assertEqual(component['planned_interval'],2000)
            self.assertEqual(component['next_interval'],interval)
            self.assertTrue(item['action_code'].endswith('_'+str(interval)))
            for field in ('current_meter','target_meter','remaining','last_interval','last_service'):
                self.assertEqual(component[field], original['components']['oil_change'][field])

    def test_batch_overrides_apply_to_only_selected_machine(self):
        second = sample_source('HD785-7', 'HD708', last=200)
        self.source['machines'] += second['machines']
        self.source['plans'] += second['plans']
        request = self.request()
        request['machine_codes'].append('HD708')
        request['item_actions']['HD708'] = 'OIL_CHANGE_HD785-7_400'
        result = create_worker.execute_request(request, db_path=self.db_path)
        self.assertTrue(result['ok'],result)
        self.assertEqual(len(result['orders']),2)
        self.assertIn('1800 ساعتی',result['orders'][0]['item_summary'])
        self.assertIn('400 ساعتی',result['orders'][1]['item_summary'])

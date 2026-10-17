"""AF layout preserves all rows, metadata and independent completion columns."""
import tempfile
from pathlib import Path
import unittest
from openpyxl import load_workbook
from tools.fleet.work_orders.types.air_filter.builder import build_document


class AirFilterPrintTests(unittest.TestCase):
    def test_rows_metadata_wrapping_signature_and_unbounded_pagination(self):
        root=Path('E:/KomatsoAI/runtime/air-filter-fixes-20261017/fixtures');root.mkdir(parents=True,exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as tmp:
            for count in (1,2,5,6,7,20,40,100):
                with self.subTest(count=count):
                    items=[dict(machine_code=str(700+i),machine_name='دامپتراک کوماتسو با نام بلند آزمایشی',
                                action_text='تعویض هواکش داخلی و بیرونی') for i in range(count)]
                    path=Path(tmp)/f'AF-{count}.xlsx'
                    record=dict(work_order_no='AF-1405-07-25-002',jalali_date='1405/07/25',shift='صبح-ظهر',revision_of='AF-1405-07-25-001',created_at='2026-10-16 21:00:00')
                    build_document(output_path=path,jalali_date='1405/01/01',shift='شب',items=items,work_order=record)
                    wb=load_workbook(path);ws=wb.active
                    try:
                        self.assertEqual(ws['F1'].value,record['jalali_date'])
                        self.assertIn(record['shift'],ws['A1'].value)
                        self.assertIn(record['revision_of'],ws['A1'].value)
                        self.assertIn(record['work_order_no'],ws.oddHeader.center.text)
                        self.assertIn('تاریخ صدور: 1405/07/25',ws.oddHeader.center.text)
                        self.assertEqual(ws.page_setup.fitToHeight,0)
                        self.assertEqual(ws.print_title_rows,'$1:$2')
                        self.assertIn(f'E{count+3}:F{count+3}',{str(m) for m in ws.merged_cells.ranges})
                        for row,item in enumerate(items,3):
                            self.assertEqual(ws.cell(row,3).value,item['machine_code'])
                            self.assertEqual(ws.cell(row,4).value,item['action_text'])
                            self.assertTrue(ws.cell(row,2).alignment.wrap_text)
                            self.assertTrue(ws.cell(row,4).alignment.wrap_text)
                            self.assertGreaterEqual(ws.cell(row,2).font.sz,14)
                            for col in (5,6):
                                self.assertFalse(any(ws.cell(row,col).coordinate in m for m in ws.merged_cells.ranges))
                    finally:wb.close()

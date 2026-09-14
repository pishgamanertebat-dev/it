import asyncio
import hashlib
import importlib
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook
from tools.fleet.work_orders.types.oil_change import builder
from tools.fleet.work_orders.channels.bale import create_worker, staff_flow
from tools.fleet.work_orders.channels.bale.message_handler import WorkOrderMenuHandler
from tools.fleet.work_orders.core import service, staff_dispatch
from tools.fleet.work_orders.dev.test_bale_work_order_create import WorkOrderCreateTests
from tools.fleet.work_orders.dev.test_permissions import test_database


class OilChangeTests(WorkOrderCreateTests):
    def setUp(self):
        super().setUp()
        from tools.fleet.oil_change.test_proposal import sample_source
        self.source = sample_source()
        for target in ('tools.fleet.oil_change.source.read_source', 'tools.fleet.oil_change.proposal.read_source'):
            patcher = patch(target, side_effect=lambda *a,**kw:self.source)
            patcher.start();self.addCleanup(patcher.stop)
        patcher=patch('tools.fleet.oil_change.source.source_hash',return_value='fixture')
        patcher.start();self.addCleanup(patcher.stop)

    def oil_order(self, interval=200):
        return service.create_work_order(work_order_type='OIL_CHANGE', jalali_date='1405/06/18',
            shift='روزانه', machine_codes=['702'], created_by='bale:455740857',
            item_actions={'HD702': f'OIL_CHANGE_{interval}'})

    def test_multiple_selected_machines_get_independent_reviewable_orders(self):
        from tools.fleet.oil_change.test_proposal import sample_source
        self.source=sample_source('HD785-7','HD708',last=200)
        second=sample_source('R330-9','EX332',last=2000,remaining=24)
        self.source['machines'] += second['machines']
        self.source['plans'] += second['plans']
        async def scenario():
            documents, replies = [], []
            async def worker(request):
                return create_worker.execute_request(request,db_path=self.db_path)
            async def document(gateway,chat,order):
                documents.append(order)
            handler=WorkOrderMenuHandler(db_path=self.db_path,worker=worker,document_sender=document)
            async def enter(text):
                event=SimpleNamespace(text=text,source=SimpleNamespace(platform='bale',chat_type='dm',user_id='455740857',chat_id='455740857'))
                handler.handle(event,None,send=lambda g,c,t:replies.append(t))
                while handler.tasks: await asyncio.gather(*list(handler.tasks))
            await enter('حکم کار');await enter('۲')
            session=next(iter(handler.pending.values()))
            self.assertEqual([i['machine_code'] for i in session.proposal['items']],['HD708'])
            await enter('اضافه');await enter('EX332')
            self.assertEqual(session.stage,'PROPOSAL')
            await enter('حذف');await enter('1')
            self.assertEqual([i['machine_code'] for i in session.proposal['items']],['EX332'])
            await enter('اضافه');await enter('HD708')
            await enter('تایید')
            self.assertEqual(len(documents),2)
            self.assertNotEqual(documents[0]['work_order_no'],documents[1]['work_order_no'])
            for document in documents:
                order=service.get_work_order(document['work_order_no'])
                self.assertEqual(len(order['items']),1)
                self.assertEqual(order['status'],'FILE_READY')
                self.assertIsNone(order['sent_at'])
                self.assertTrue(any('ثبت تایید '+document['work_order_no'] in text for text in replies))
            self.assertIn('200 ساعتی',documents[0]['item_summary'])
            self.assertIn('400 ساعتی',documents[1]['item_summary'])
            handler.pending.clear()
            await enter('ارسال مجدد '+documents[0]['work_order_no'])
            self.assertEqual(documents[-1]['work_order_no'],documents[0]['work_order_no'])
        asyncio.run(scenario())

    def test_source_changes_and_wrong_cycle_block_oil_creation(self):
        request=dict(action='create',bale_id='455740857',work_order_type='OIL_CHANGE',
                     machine_codes=['HD708'],item_actions={'HD708':builder.action_for('HD785-7',400)},
                     jalali_date='1405/06/23',shift='روزانه',proposal={'source_sha256':'outdated'})
        with patch.object(service,'create_work_order') as create:
            self.assertEqual(create_worker.execute_request(request,db_path=self.db_path)['error'],'INVALID_INPUT')
            request['proposal']['source_sha256']='fixture'
            request['item_actions']['HD708']=builder.action_for('HD785-7',200)
            self.assertEqual(create_worker.execute_request(request,db_path=self.db_path)['error'],'INVALID_INPUT')
            create.assert_not_called()

    def test_partial_batch_failure_reports_created_orders_without_retry(self):
        from tools.fleet.oil_change.proposal import build_proposal
        from tools.fleet.oil_change.test_proposal import sample_source
        second=sample_source('HD785-7','HD709')
        self.source['machines']+=second['machines'];self.source['plans']+=second['plans']
        proposal=build_proposal()
        request=dict(action='create',bale_id='455740857',work_order_type='OIL_CHANGE',
                     machine_codes=['HD708','HD709'],item_actions={i['machine_code']:i['action_code'] for i in proposal['items']},
                     jalali_date=proposal['plan_date'],shift='روزانه',proposal=proposal)
        original=service.create_work_order
        def create(**kwargs):
            if kwargs['machine_codes']==['HD709']:raise ValueError('test failure')
            return original(**kwargs)
        with patch.object(service,'create_work_order',side_effect=create), self.assertLogs(level='ERROR'):
            result=create_worker.execute_request(request,db_path=self.db_path)
        self.assertTrue(result['ok'])
        self.assertEqual(len(result['orders']),1)
        self.assertIn('HD709',result['batch_error'])

    def test_all_ten_templates_keep_cells_styles_and_package_assets(self):
        before = hashlib.sha256(builder.SOURCE.read_bytes()).hexdigest()
        source = load_workbook(builder.SOURCE)
        try:
            for interval in builder.INTERVALS:
                with self.subTest(interval=interval):
                    order = self.oil_order(interval)
                    self.assertEqual(order['status'], 'FILE_READY')
                    self.assertTrue(order['work_order_no'].startswith('OC-'))
                    self.assertEqual(order['items'][0]['action_code'], f'OIL_CHANGE_{interval}')
                    wb = load_workbook(order['excel_path'])
                    try:
                        self.assertEqual(wb.sheetnames, [str(interval)])
                        expected, actual = source[str(interval)], wb.active
                        for row in expected:
                            for cell in row:
                                self.assertEqual(actual[cell.coordinate].value, 'HD702' if cell.coordinate == 'F3' else cell.value)
                                self.assertEqual(actual[cell.coordinate]._style, cell._style)
                        self.assertEqual(str(actual.merged_cells), str(expected.merged_cells))
                    finally:
                        wb.close()
                    with ZipFile(builder.SOURCE) as original, ZipFile(order['excel_path']) as output:
                        changed = [n for n in original.namelist() if original.read(n) != output.read(n)]
                        self.assertEqual(len(changed), 2)
                        self.assertIn('xl/workbook.xml', changed)
        finally:
            source.close()
        self.assertEqual(hashlib.sha256(builder.SOURCE.read_bytes()).hexdigest(), before)

    def test_invalid_input_cannot_create_orders(self):
        for value in ('0', '201', '4350', '2200', '200.0', '=200', True):
            with self.assertRaises(ValueError):
                builder.normalize_interval(value)
        self.assertEqual(builder.normalize_interval('١٢٠٠'), 1200)
        self.assertEqual(builder.normalize_code('۷۰۲'), 'HD702')
        for codes, actions in (([], {}), (['701', '702'], {}), (['=1+1'], {}), (['702'], None), (['702'], {'HD701':'OIL_CHANGE_200'})):
            with self.assertRaises(ValueError):
                builder.get_items(codes, actions)
        with test_database(self.db_path) as con:
            self.assertEqual(con.execute('SELECT COUNT(*) FROM service_work_orders').fetchone()[0], 0)

    def test_bale_automatic_review_dispatch_and_receipt_offline(self):
        for choice, code in [('1', '702'), ('2', '463'), ('3', '708'), ('4', '۸۰۱'), ('5', 'ex333'), ('6', '851')]:
            with self.subTest(model=choice):
                self.run_bale_model(choice, code)

    def run_bale_model(self, model_choice, machine_code):
        from tools.fleet.oil_change.test_proposal import sample_source
        from tools.fleet.work_orders.types.oil_change.form import MODELS
        self.source=sample_source(MODELS[model_choice],builder.normalize_code(machine_code,MODELS[model_choice]),last=1000)
        with test_database(self.db_path) as con:
            importlib.import_module('tools.fleet.work_orders.migrations.003_staff_dispatch').migrate(con)
        async def scenario():
            documents, messages, requests = [], [], []
            async def document(**kw):
                documents.append((str(kw['chat_id']), kw['document'].read()))
                return {'ok': True}
            async def message(**kw):
                messages.append((str(kw['chat_id']), kw['text']))
            gateway = SimpleNamespace(adapters={'bale': SimpleNamespace(_bot=SimpleNamespace(send_document=document, send_message=message))})
            async def worker(request):
                requests.append(request)
                return create_worker.execute_request(request, db_path=self.db_path)
            handler = WorkOrderMenuHandler(db_path=self.db_path, worker=worker)
            def event(text, actor='455740857'):
                return SimpleNamespace(text=text, source=SimpleNamespace(platform='bale', chat_type='dm', user_id=actor, chat_id=actor))
            def send(gateway, chat, reply):
                messages.append((str(chat), reply))
            async def enter(text):
                result = handler.handle(event(text), gateway, send=send)
                if handler.tasks:
                    await asyncio.gather(*list(handler.tasks))
                return result
            await enter('حکم کار')
            await enter('۲')
            self.assertEqual(requests[0]['action'], 'propose')
            self.assertEqual(next(iter(handler.pending.values())).stage,'PROPOSAL')
            self.assertEqual(len(documents),0)
            await enter('تایید')
            session = next(iter(handler.pending.values()))
            self.assertEqual(session.stage, 'REVIEW')
            self.assertEqual(len(documents), 1)
            self.assertEqual(documents[0][0], '455740857')
            number = session.order_no
            self.assertEqual(service.get_work_order(number)['status'], 'FILE_READY')
            from tools.fleet.work_orders.types.oil_change.form import MODELS
            saved = service.get_work_order(number)
            self.assertEqual(builder.parse_action(saved['items'][0]['action_code']), (MODELS[model_choice], 1200))
            self.assertEqual(saved['items'][0]['machine_code'], builder.normalize_code(machine_code, MODELS[model_choice]))
            await enter('تایید')
            self.assertEqual(session.stage, 'STAFF')
            self.assertEqual(len(documents), 1)
            await enter('1')
            self.assertEqual(service.get_work_order(number)['status'], 'SENT')
            self.assertEqual(len(documents), 2)
            actor = documents[1][0]
            result = staff_flow.handle_staff_receipt(event('تایید ' + number, actor), gateway, send=send)
            self.assertEqual(result['reason'], 'staff-receipt')
            await asyncio.gather(*list(staff_flow.tasks))
            self.assertIsNotNone(service.get_work_order(number)['acknowledged_at'])
            self.assertTrue(any(chat == '455740857' and 'دریافت حکم تعویض روغن' in text for chat, text in messages))
            # Commands survive expired in-memory forms and reconstruct oil editing.
            await enter('حکم کار')
            await enter('۲')
            await enter('تایید')
            editable = session = next(iter(handler.pending.values()))
            edit_number = editable.order_no
            handler.pending.clear()
            await enter('ویرایش ' + edit_number)
            self.assertEqual(next(iter(handler.pending.values())).stage, 'PROPOSAL')
            self.assertEqual(service.get_work_order(edit_number)['status'], 'FILE_READY')
        with patch('tools.fleet.work_orders.core.permissions.DB_PATH', self.db_path):
            asyncio.run(scenario())

    def test_loader_bale_review_pdf_dispatch_receipt(self):
        for choice, code in [('7', '۶۰۲'), ('8', 'w473')]:
            with self.subTest(model=choice):
                self.run_bale_model(choice, code)

    def test_final_five_models_create_review_files_from_bale(self):
        from tools.fleet.work_orders.types.oil_change.form import MODELS
        async def scenario():
            documents = []
            async def worker(request):
                return create_worker.execute_request(request, db_path=self.db_path)
            async def document_sender(gateway, chat, order):
                documents.append(order)
            handler = WorkOrderMenuHandler(db_path=self.db_path, worker=worker, document_sender=document_sender)
            async def enter(text):
                event = SimpleNamespace(text=text, source=SimpleNamespace(platform='bale', chat_type='dm', user_id='455740857', chat_id='455740857'))
                handler.handle(event, None, send=lambda *args: None)
                if handler.tasks:
                    await asyncio.gather(*list(handler.tasks))
            for choice, raw, canonical in [('۹','۱۵۳','D153'), ('۱۰','d154','D154'), ('۱۱','۳۲۲','EX322'), ('۱۲','ex522','EX522'), ('۱۳','602','EX602')]:
                from tools.fleet.oil_change.test_proposal import sample_source
                self.source=sample_source(MODELS[str(int(choice))],canonical,last=1800)
                await enter('حکم کار')
                await enter('۲')
                session = next(iter(handler.pending.values()))
                self.assertEqual(session.proposal['items'][0]['machine_code'],canonical)
                await enter('تایید')
                self.assertEqual(session.stage, 'REVIEW')
                order = service.get_work_order(session.order_no)
                self.assertEqual(order['status'], 'FILE_READY')
                self.assertIsNone(order['assigned_staff_id'])
                self.assertEqual(builder.parse_action(order['items'][0]['action_code']), (MODELS[str(int(choice))], 2000))
                self.assertEqual(order['items'][0]['machine_code'], canonical)
            self.assertEqual(len(documents), 5)
        asyncio.run(scenario())

    def test_new_model_templates_all_intervals(self):
        for model, code in [('HD465-7R', 'HD464'), ('HD785-7', 'HD709'), ('PC800-7', 'EX801'), ('R330-9', 'EX333'), ('PC850-8', 'EX851'), ('WA600-6', 'W602'), ('WA470-3', 'W473'),
                            ('D155A-2', 'D153'), ('D155A-6', 'D154'), ('R320-9', 'EX322'), ('R520-9', 'EX522'), ('PC600-8', 'EX602')]:
            source_path = builder.TEMPLATES[model][0]
            before = hashlib.sha256(source_path.read_bytes()).hexdigest()
            original = load_workbook(source_path)
            try:
                for interval in builder.INTERVALS:
                    with self.subTest(model=model, interval=interval):
                        order = service.create_work_order(work_order_type='OIL_CHANGE', jalali_date='1405/06/18',
                            shift='روزانه', machine_codes=[code], created_by='bale:455740857',
                            item_actions={code: builder.action_for(model, interval)})
                        wb = load_workbook(order['excel_path'])
                        try:
                            self.assertEqual(wb.sheetnames, [str(interval)])
                            expected, actual = original[str(interval)], wb.active
                            for row in expected:
                                for cell in row:
                                    expected_value = code if cell.coordinate == 'F3' else cell.value
                                    if model == 'PC600-8' and cell.coordinate == 'C2':
                                        expected_value = interval
                                    self.assertEqual(actual[cell.coordinate].value, expected_value)
                                    self.assertEqual(actual[cell.coordinate]._style, cell._style)
                            self.assertEqual(str(actual.merged_cells), str(expected.merged_cells))
                        finally:
                            wb.close()
                        with ZipFile(source_path) as source, ZipFile(order['excel_path']) as output:
                            changed = [n for n in source.namelist() if source.read(n) != output.read(n)]
                            self.assertEqual(len(changed), 2)
            finally:
                original.close()
            self.assertEqual(hashlib.sha256(source_path.read_bytes()).hexdigest(), before)
        for action in ['OIL_CHANGE_HD999_200', 'OIL_CHANGE_HD785-7_2200', 'OIL_CHANGE_HD465-7R_201']:
            with self.assertRaises(ValueError):
                builder.get_items(['463'], {'HD463': action})

    def test_prefix_follows_selected_model_and_rejects_mismatches(self):
        for model in ('D155A-2', 'D155A-6'):
            self.assertEqual(builder.normalize_code('۱۵۲', model), 'D152')
            self.assertEqual(builder.normalize_code('d152', model), 'D152')
            with self.assertRaises(ValueError):
                builder.normalize_code('EX152', model)
        for model in ('WA600-6', 'WA470-3'):
            for code in ('۶۰۱', 'w601', 'wa601'):
                self.assertEqual(builder.normalize_code(code, model), 'W601')
            for code in ('EX601', 'HD601'):
                with self.assertRaises(ValueError):
                    builder.normalize_code(code, model)
        for model in ('PC800-7', 'R330-9', 'PC850-8'):
            self.assertEqual(builder.normalize_code('۸۰۱', model), 'EX801')
            self.assertEqual(builder.normalize_code('ex801', model), 'EX801')
            with self.assertRaises(ValueError):
                builder.normalize_code('HD801', model)
            with self.assertRaises(ValueError):
                builder.get_items(['HD801'], {'HD801': builder.action_for(model, 200)})
        with self.assertRaises(ValueError):
            builder.normalize_code('EX701', 'HD785-5')


import sqlite3

import openpyxl



DB = r"E:\KomatsoAI\data\fleet\db\fleet_ops.db"



conn = sqlite3.connect(DB)

cur = conn.cursor()



cur.execute("""

SELECT id, raw_path

FROM ingest_sources

WHERE source_type = 'service_events' 

ORDER BY id DESC

LIMIT 1

""")



row = cur.fetchone()



if not row:

    raise SystemExit("No service_events source found")



source_id, raw_path = row



print("SERVICE FILE INSPECTION")

print("=" * 70)

print("Source ID:", source_id)

print("File:", raw_path)



wb = openpyxl.load_workbook(raw_path, data_only=True)



print()

print("SHEETS")

print("-" * 70)



for ws in wb.worksheets:

    print(f"{ws.title} | rows={ws.max_row} cols={ws.max_column}")



print()

print("SAMPLE FIRST SHEET")

print("-" * 70)



ws = wb.worksheets[0]



print("Sheet:", ws.title)



for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15), values_only=True):

    print(row)



conn.close()

"""Regression tests for year selection, import isolation and annual reports."""
import contextlib
import io
import json
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from normalize_service_year import normalize, select_year_columns
import service_history


class ServiceYearTests(unittest.TestCase):
    def setUp(self):
        root = Path(__file__).resolve().parents[2] / 'runtime' / 'service_year_fix'
        root.mkdir(parents=True, exist_ok=True)
        self.tmp = tempfile.TemporaryDirectory(dir=root)
        self.addCleanup(self.tmp.cleanup)
        self.db = str(Path(self.tmp.name) / 'fleet.db')
        self.xlsx = str(Path(self.tmp.name) / 'hours.xlsx')
        w = Workbook(); s = w.active; s.title = 'ساعت کاری'
        s.cell(4, 2, 'EX801'); s.cell(4, 5, 500)
        # Partial 1403, full-year boundary 1404, unmerged month, then 1405.
        for col, month, day, shift, value in [
            (6,'مرداد',1,'روز',4), (7,'اسفند',29,'روز',5),
            (8,'فروردین',1,'روز',6), (9,'تیر',31,'روز',7),
            (10,'مرداد',1,'روز',8), (11,None,2,'شب',9),
            (12,'اسفند',29,'روز',10), (13,'فروردین',1,'روز',11),
        ]:
            if month: s.cell(1,col,month)
            s.cell(2,col,day); s.cell(3,col,shift); s.cell(4,col,value)
        s.cell(3,15,'مانده به تعویض'); s.cell(4,14,450); s.cell(4,15,50)
        w.save(self.xlsx); w.close()
        with contextlib.closing(sqlite3.connect(self.db)) as c, c:
            c.executescript('''
                CREATE TABLE fleet_settings(setting_key TEXT, setting_value TEXT);
                INSERT INTO fleet_settings VALUES ('operational_jalali_year','1405');
                CREATE TABLE ingest_sources(id INTEGER, source_type TEXT, raw_path TEXT);
                CREATE TABLE machines(id INTEGER, canonical_code TEXT, machine_type_hint TEXT, model_key TEXT, identity_status TEXT);
                INSERT INTO machines VALUES(1,'EX801','excavator',NULL,'VERIFIED');
                CREATE TABLE service_shift_hours(id INTEGER PRIMARY KEY, ingest_source_id INTEGER, machine_id INTEGER, canonical_code TEXT, jalali_date TEXT, shift TEXT, work_hours REAL, note TEXT, raw_value TEXT, source_sheet TEXT, source_row INTEGER, source_col INTEGER, UNIQUE(ingest_source_id,source_sheet,source_row,source_col));
                CREATE TABLE service_pm_snapshots(id INTEGER PRIMARY KEY, ingest_source_id INTEGER, machine_id INTEGER, canonical_code TEXT, jalali_asof_date TEXT, next_service_hour REAL, current_meter_hour REAL, remaining_hours REAL, pm_issue_raw TEXT, oil_analysis_raw TEXT, source_sheet TEXT, source_row INTEGER, UNIQUE(ingest_source_id,source_sheet,source_row));
            ''')
            c.execute("INSERT INTO ingest_sources VALUES (1,'service_events',?)",(self.xlsx,))

    def run_import(self, year):
        with contextlib.redirect_stdout(io.StringIO()): normalize(year,self.db)

    def report(self, *args):
        out=io.StringIO()
        with patch.object(sys,'argv',['service_history','EX801','--db',self.db,'--json',*args]), contextlib.redirect_stdout(out):
            service_history.main()
        return json.loads(out.getvalue())

    def test_historical_lazy_import_and_annual_total(self):
        result=self.report('--year','۱۴۰۴','--limit','1')
        self.assertEqual(result['year_summary']['total_work_hours'],40)
        self.assertEqual(result['year_summary']['numeric_records'],5)
        self.assertEqual(len(result['recent_records']),1)
        self.assertIsNone(result['service_status'])
        self.assertEqual(result,self.report('--year','1404','--limit','1'))
        self.run_import(1404)
        self.assertEqual(self.report('--year','1404')['year_summary']['total_work_hours'],40)

    def test_default_and_partial_year_are_isolated(self):
        self.run_import(None)
        result=self.report()
        self.assertEqual(result['year_summary']['total_work_hours'],11)
        self.assertEqual(result['service_policy']['requested_year'],1405)
        self.assertEqual(result['service_status']['current_meter_hour'],450)
        self.assertEqual(self.report('--year','1403')['year_summary']['total_work_hours'],9)
        self.assertEqual(self.report()['year_summary'],result['year_summary'])

    def test_unsupported_year_never_wraps_or_claims_zero(self):
        for year in ['1402','1406']:
            result=self.report('--year',year)
            self.assertEqual(result['status'],'SOURCE_IMPORT_FAILED')
            self.assertIsNone(result['year_summary']['total_work_hours'])
        with contextlib.closing(sqlite3.connect(self.db)) as c, c:
            self.assertEqual(c.execute('SELECT count(*) FROM service_shift_hours').fetchone()[0],0)

    def test_undated_current_snapshot_not_used_for_history(self):
        self.run_import(None)
        with contextlib.closing(sqlite3.connect(self.db)) as c, c:
            c.execute('UPDATE service_pm_snapshots SET jalali_asof_date=NULL')
        self.assertIsNone(self.report('--year','1404')['service_status'])

    def test_missing_source_is_distinct_from_no_records(self):
        Path(self.xlsx).unlink()
        self.assertEqual(self.report('--year','1404')['status'],'SOURCE_IMPORT_FAILED')
        self.assertEqual(self.report()['status'],'NO_NORMALIZED_RECORDS')


if __name__ == '__main__':
    unittest.main()


from __future__ import annotations

import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


DB_PATH = Path(
    r"E:\KomatsoAI\data\fleet\db\fleet_ops.db"
)

PARSER_VERSION = "maintenance_v1"


def clean(value):
    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\u200c", " ")
        .replace("ي", "ی")
        .replace("ك", "ک")
        .split()
    )


def detect_columns(ws):
    max_scan = min(ws.max_row, 15)

    for row_number in range(1, max_scan + 1):

        headers = {}

        for cell in ws[row_number]:
            text = clean(cell.value)

            if text:
                headers[cell.column] = text

        if not headers:
            continue

        date_col = None
        code_col = None
        mechanic_col = None
        action_col = None
        parts_col = None

        for col, text in headers.items():

            if "تاریخ" in text:
                date_col = col

            if (
                "کد مکانیزم" in text
                or text == "حفارها"
            ):
                code_col = col

            if (
                "نام مکانیک" in text
                or "تعمیرکار" in text
            ):
                mechanic_col = col

            if "نوع خرابی" in text:
                action_col = col

            if "قطعات مصرفی" in text:
                parts_col = col

        # تاریخ و شرح تعمیر برای Parse کافی هستند.
        if date_col and action_col:
            return {
                "header_row": row_number,
                "date_col": date_col,
                "code_col": code_col,
                "mechanic_col": mechanic_col,
                "action_col": action_col,
                "parts_col": parts_col,
            }

    return None


def value_at(row, column_number):
    if not column_number:
        return ""

    index = column_number - 1

    if index >= len(row):
        return ""

    return clean(row[index])


def init_db(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS maintenance_events (
            id INTEGER PRIMARY KEY AUTOINCREMENT,

            raw_excel_row_id INTEGER NOT NULL UNIQUE,
            ingest_source_id INTEGER NOT NULL,

            parser_version TEXT NOT NULL,

            event_date_raw TEXT,

            machine_code_raw TEXT,
            mechanic_raw TEXT,

            action_raw TEXT,
            parts_raw TEXT,

            sheet_name TEXT NOT NULL,
            sheet_index INTEGER NOT NULL,
            excel_row INTEGER NOT NULL,

            normalized_at TEXT NOT NULL,

            FOREIGN KEY (raw_excel_row_id)
                REFERENCES raw_excel_rows(id),

            FOREIGN KEY (ingest_source_id)
                REFERENCES ingest_sources(id)
        )
        """
    )

    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS
        idx_maintenance_events_date
        ON maintenance_events(event_date_raw)
        """
    )

    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS
        idx_maintenance_events_code
        ON maintenance_events(machine_code_raw)
        """
    )

    conn.commit()


def main():
    if not DB_PATH.exists():
        raise SystemExit(
            f"ERROR: DB not found: {DB_PATH}"
        )

    conn = sqlite3.connect(DB_PATH)

    conn.execute(
        "PRAGMA foreign_keys = ON"
    )

    init_db(conn)

    source = conn.execute(
        """
        SELECT
            id,
            original_filename,
            raw_path

        FROM ingest_sources

        WHERE source_type = 'maintenance_events'

        ORDER BY id DESC

        LIMIT 1
        """
    ).fetchone()

    if not source:
        raise SystemExit(
            "ERROR: maintenance_events source not found"
        )

    source_id = source[0]
    filename = source[1]
    raw_path = Path(source[2])

    if not raw_path.exists():
        raise SystemExit(
            f"ERROR: raw file missing: {raw_path}"
        )

    # Raw row IDs for traceability
    raw_ids = {}

    rows = conn.execute(
        """
        SELECT
            id,
            sheet_index,
            excel_row

        FROM raw_excel_rows

        WHERE ingest_source_id = ?
        """,
        (source_id,),
    ).fetchall()

    for raw_id, sheet_index, excel_row in rows:
        raw_ids[
            (
                sheet_index,
                excel_row,
            )
        ] = raw_id

    wb = load_workbook(
        raw_path,
        read_only=True,
        data_only=False,
    )

    now = datetime.now(
        timezone.utc
    ).isoformat(timespec="seconds")

    parsed_sheets = 0
    unparsed_sheets = []

    candidate_rows = 0
    inserted = 0
    skipped = 0

    with conn:

        for sheet_index, sheet_name in enumerate(
            wb.sheetnames,
            start=1,
        ):
            ws = wb[sheet_name]

            columns = detect_columns(ws)

            if not columns:
                unparsed_sheets.append(
                    sheet_name
                )
                continue

            parsed_sheets += 1

            header_row = columns[
                "header_row"
            ]

            for excel_row, row in enumerate(
                ws.iter_rows(
                    values_only=True
                ),
                start=1,
            ):
                if excel_row <= header_row:
                    continue

                event_date = value_at(
                    row,
                    columns["date_col"],
                )

                machine_code = value_at(
                    row,
                    columns["code_col"],
                )

                mechanic = value_at(
                    row,
                    columns["mechanic_col"],
                )

                action = value_at(
                    row,
                    columns["action_col"],
                )

                parts = value_at(
                    row,
                    columns["parts_col"],
                )

                # Skip genuinely empty rows.
                if not any(
                    (
                        event_date,
                        machine_code,
                        mechanic,
                        action,
                        parts,
                    )
                ):
                    continue

                # Header-like/repeated junk rows
                if (
                    event_date == "تاریخ"
                    or action == "نوع خرابی"
                ):
                    continue

                raw_excel_row_id = raw_ids.get(
                    (
                        sheet_index,
                        excel_row,
                    )
                )

                # Only normalize rows that exist
                # in immutable raw layer.
                if raw_excel_row_id is None:
                    continue

                candidate_rows += 1

                cur = conn.execute(
                    """
                    INSERT OR IGNORE INTO maintenance_events (
                        raw_excel_row_id,
                        ingest_source_id,
                        parser_version,
                        event_date_raw,
                        machine_code_raw,
                        mechanic_raw,
                        action_raw,
                        parts_raw,
                        sheet_name,
                        sheet_index,
                        excel_row,
                        normalized_at
                    )
                    VALUES (
                        ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
                    )
                    """,
                    (
                        raw_excel_row_id,
                        source_id,
                        PARSER_VERSION,
                        event_date,
                        machine_code,
                        mechanic,
                        action,
                        parts,
                        sheet_name,
                        sheet_index,
                        excel_row,
                        now,
                    ),
                )

                if cur.rowcount == 1:
                    inserted += 1
                else:
                    skipped += 1

    wb.close()

    total = conn.execute(
        """
        SELECT COUNT(*)
        FROM maintenance_events
        WHERE ingest_source_id = ?
        """,
        (source_id,),
    ).fetchone()[0]

    print()
    print("MAINTENANCE NORMALIZATION COMPLETE")
    print(f"Source ID: {source_id}")
    print(f"File: {filename}")
    print(f"Sheets: {len(wb.sheetnames)}")
    print(f"Parsed sheets: {parsed_sheets}")
    print(
        f"Unparsed sheets: "
        f"{len(unparsed_sheets)}"
    )
    print(
        f"Candidate event rows: "
        f"{candidate_rows}"
    )
    print(f"Inserted: {inserted}")
    print(f"Skipped: {skipped}")
    print(
        f"Normalized total: {total}"
    )

    if unparsed_sheets:
        print()
        print("UNPARSED SHEETS")

        for name in unparsed_sheets:
            print(f"  {name}")

    print()
    print("SAMPLE: SHEET 714")

    rows = conn.execute(
        """
        SELECT
            event_date_raw,
            machine_code_raw,
            mechanic_raw,
            action_raw,
            parts_raw

        FROM maintenance_events

        WHERE
            ingest_source_id = ?
            AND sheet_name = '714'

        ORDER BY excel_row

        LIMIT 5
        """,
        (source_id,),
    ).fetchall()

    for row in rows:
        print(
            f"{row[0]} | "
            f"code={row[1]} | "
            f"mechanic={row[2]} | "
            f"action={row[3]} | "
            f"parts={row[4]}"
        )

    print()
    print("601 SHEET CHECK")

    rows = conn.execute(
        """
        SELECT
            sheet_name,
            event_date_raw,
            machine_code_raw,
            action_raw

        FROM maintenance_events

        WHERE
            ingest_source_id = ?
            AND sheet_name IN ('601', '601.')

        ORDER BY
            sheet_name,
            excel_row

        LIMIT 20
        """,
        (source_id,),
    ).fetchall()

    for row in rows:
        print(
            f"sheet={row[0]} | "
            f"date={row[1]} | "
            f"raw_code={row[2]} | "
            f"{row[3]}"
        )

    conn.close()


if __name__ == "__main__":
    main()
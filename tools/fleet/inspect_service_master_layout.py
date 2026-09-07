import sqlite3
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


DB = r"E:\KomatsoAI\data\fleet\db\fleet_ops.db"

MONTHS = (
    "فروردین",
    "اردیبهشت",
    "خرداد",
    "تیر",
    "مرداد",
    "شهریور",
    "مهر",
    "آبان",
    "آذر",
    "دی",
    "بهمن",
    "اسفند",
)


def clean(value):
    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\u200c", " ")
        .replace("ي", "ی")
        .replace("ك", "ک")
        .split()
    )


def build_header_map(ws, max_header_row=5):
    """
    Expands merged header cells so every column gets the
    visible header value belonging to it.
    """
    result = {}

    for r in range(1, max_header_row + 1):
        for c in range(1, ws.max_column + 1):
            value = ws.cell(r, c).value
            if value is not None:
                result[(r, c)] = clean(value)

    for merged in ws.merged_cells.ranges:
        if merged.min_row > max_header_row:
            continue

        top_value = clean(
            ws.cell(
                merged.min_row,
                merged.min_col
            ).value
        )

        if not top_value:
            continue

        for r in range(
            merged.min_row,
            min(
                merged.max_row,
                max_header_row
            ) + 1
        ):
            for c in range(
                merged.min_col,
                merged.max_col + 1
            ):
                result[(r, c)] = top_value

    return result


def find_machine_row(ws, code):
    wanted = code.upper()

    for r in range(4, ws.max_row + 1):
        value = clean(ws.cell(r, 2).value).upper()

        if value == wanted:
            return r

    return None


conn = sqlite3.connect(DB)

source = conn.execute("""
SELECT
    id,
    raw_path,
    source_modified_at
FROM ingest_sources
WHERE source_type='service_events'
ORDER BY id DESC
LIMIT 1
""").fetchone()

if not source:
    raise SystemExit(
        "ERROR: service_events source not found"
    )

source_id, raw_path, modified_at = source

print("SERVICE MASTER LAYOUT AUDIT")
print("=" * 90)
print(f"SOURCE ID: {source_id}")
print(f"SOURCE MODIFIED: {modified_at}")
print(f"FILE: {raw_path}")

wb = load_workbook(
    raw_path,
    data_only=True,
    read_only=False
)

targets = [
    "EX801",
    "HD701",
    "HD714",
]

for sheet_name in [
    "ساعت کاری (2)",
    "ساعت کاری",
]:

    if sheet_name not in wb.sheetnames:
        continue

    ws = wb[sheet_name]

    print()
    print("=" * 90)
    print(f"SHEET: {sheet_name}")
    print(
        f"SIZE: rows={ws.max_row}, "
        f"cols={ws.max_column}"
    )

    headers = build_header_map(
        ws,
        max_header_row=5
    )

    # --------------------------------------------------------
    # Find month/year-like markers in top rows
    # --------------------------------------------------------

    print()
    print("MONTH / YEAR HEADER MARKERS")
    print("-" * 90)

    markers = []

    for r in range(1, 6):
        for c in range(1, ws.max_column + 1):

            raw = ws.cell(r, c).value

            if raw is None:
                continue

            text = clean(raw)

            contains_month = any(
                m in text
                for m in MONTHS
            )

            contains_year = any(
                y in text
                for y in (
                    "1403",
                    "1404",
                    "1405",
                    "۱۴۰۳",
                    "۱۴۰۴",
                    "۱۴۰۵",
                )
            )

            if contains_month or contains_year:
                markers.append(
                    (
                        ws.cell(r, c).coordinate,
                        text
                    )
                )

    if markers:
        for coord, value in markers:
            print(
                f"{coord:<8} | {value}"
            )
    else:
        print(
            "No explicit month/year labels "
            "found in rows 1-5."
        )

    # --------------------------------------------------------
    # Show actual visible headers across top rows
    # Only non-empty text headers, avoiding thousands of blanks
    # --------------------------------------------------------

    print()
    print("TOP HEADER CELLS")
    print("-" * 90)

    shown = set()

    for r in range(1, 6):
        for c in range(1, ws.max_column + 1):

            value = clean(
                ws.cell(r, c).value
            )

            if not value:
                continue

            key = (r, c, value)

            if key in shown:
                continue

            shown.add(key)

            coord = ws.cell(
                r,
                c
            ).coordinate

            print(
                f"{coord:<8} | {value}"
            )

    # --------------------------------------------------------
    # Inspect selected machines
    # --------------------------------------------------------

    for machine_code in targets:

        row_no = find_machine_row(
            ws,
            machine_code
        )

        print()
        print("-" * 90)

        if row_no is None:
            print(
                f"{machine_code}: NOT FOUND"
            )
            continue

        machine_name = clean(
            ws.cell(row_no, 1).value
        )

        internal_code = clean(
            ws.cell(row_no, 3).value
        )

        print(
            f"MACHINE: {machine_code}"
        )

        print(
            f"ROW: {row_no}"
        )

        print(
            f"INTERNAL CODE: {internal_code}"
        )

        print(
            f"NAME: {machine_name}"
        )

        cells = []

        for c in range(
            4,
            ws.max_column + 1
        ):
            value = ws.cell(
                row_no,
                c
            ).value

            if value is None:
                continue

            if clean(value) == "":
                continue

            h1 = headers.get(
                (1, c),
                ""
            )
            h2 = headers.get(
                (2, c),
                ""
            )
            h3 = headers.get(
                (3, c),
                ""
            )
            h4 = headers.get(
                (4, c),
                ""
            )
            h5 = headers.get(
                (5, c),
                ""
            )

            cells.append(
                (
                    c,
                    value,
                    h1,
                    h2,
                    h3,
                    h4,
                    h5,
                )
            )

        print(
            f"NON-EMPTY DATA CELLS: "
            f"{len(cells)}"
        )

        print()
        print(
            "FIRST 20 NON-EMPTY CELLS"
        )

        for item in cells[:20]:

            c, value, h1, h2, h3, h4, h5 = item

            coord = (
                f"{get_column_letter(c)}"
                f"{row_no}"
            )

            print(
                f"{coord:<9} "
                f"value={value!r} | "
                f"H1={h1!r} | "
                f"H2={h2!r} | "
                f"H3={h3!r}"
            )

        print()
        print(
            "LAST 40 NON-EMPTY CELLS"
        )

        for item in cells[-40:]:

            c, value, h1, h2, h3, h4, h5 = item

            coord = (
                f"{get_column_letter(c)}"
                f"{row_no}"
            )

            print(
                f"{coord:<9} "
                f"value={value!r} | "
                f"H1={h1!r} | "
                f"H2={h2!r} | "
                f"H3={h3!r}"
            )

wb.close()
conn.close()

print()
print("=" * 90)
print("LAYOUT AUDIT COMPLETE")
import math
import sqlite3
from collections import Counter
from openpyxl import load_workbook


DB = r"E:\KomatsoAI\data\fleet\db\fleet_ops.db"

MONTH_NUM = {
    "فروردین": 1,
    "اردیبهشت": 2,
    "خرداد": 3,
    "تیر": 4,
    "مرداد": 5,
    "شهریور": 6,
    "مهر": 7,
    "آبان": 8,
    "آذر": 9,
    "دی": 10,
    "بهمن": 11,
    "اسفند": 12,
}

# Verified source-code correction.
# The Excel uses W601 while the canonical fleet identity is WA601.
SOURCE_CODE_OVERRIDES = {
    "W601": "WA601",
}


def clean(value):
    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\u200c", " ")
        .replace("ي", "ی")
        .replace("ك", "ک")
        .split()
    )


def is_number(value):
    if isinstance(value, bool):
        return False

    if isinstance(value, (int, float)):
        try:
            return math.isfinite(float(value))
        except Exception:
            return False

    return False


def header_map(ws, max_row=3):
    result = {}

    for r in range(1, max_row + 1):
        for c in range(1, ws.max_column + 1):
            value = clean(ws.cell(r, c).value)

            if value:
                result[(r, c)] = value

    for merged in ws.merged_cells.ranges:

        if merged.min_row > max_row:
            continue

        value = clean(
            ws.cell(
                merged.min_row,
                merged.min_col
            ).value
        )

        if not value:
            continue

        for r in range(
            merged.min_row,
            min(merged.max_row, max_row) + 1
        ):
            for c in range(
                merged.min_col,
                merged.max_col + 1
            ):
                result[(r, c)] = value

    return result


conn = sqlite3.connect(DB)
cur = conn.cursor()

source = cur.execute("""
SELECT id, raw_path
FROM ingest_sources
WHERE source_type='service_events'
ORDER BY id DESC
LIMIT 1
""").fetchone()

if not source:
    raise SystemExit(
        "ERROR: service_events source not found"
    )

source_id, raw_path = source


year_row = cur.execute("""
SELECT setting_value
FROM fleet_settings
WHERE setting_key='operational_jalali_year'
""").fetchone()

if not year_row:
    raise SystemExit(
        "ERROR: operational_jalali_year not found"
    )

operational_year = int(year_row[0])


machines = {
    str(code).upper(): {
        "id": machine_id,
        "code": code,
    }
    for machine_id, code
    in cur.execute("""
        SELECT
            id,
            canonical_code
        FROM machines
    """)
}


wb = load_workbook(
    raw_path,
    data_only=True,
    read_only=False
)

ws = wb["ساعت کاری"]

headers = header_map(ws)


# ============================================================
# Detect 1405 cycle
# ============================================================

farvardin_columns = []

for c in range(
    1,
    ws.max_column + 1
):
    if clean(
        ws.cell(1, c).value
    ) == "فروردین":
        farvardin_columns.append(c)


if not farvardin_columns:
    raise SystemExit(
        "ERROR: no Farvardin block found"
    )

# Latest Farvardin block corresponds to operational 1405
cycle_start = farvardin_columns[-1]


remaining_col = None

for c in range(
    1,
    ws.max_column + 1
):
    if clean(
        ws.cell(3, c).value
    ) == "مانده به تعویض":
        remaining_col = c


if remaining_col is None:
    raise SystemExit(
        "ERROR: remaining-service column not found"
    )


current_meter_col = remaining_col - 1
pm_col = remaining_col + 1
oil_analysis_col = remaining_col + 2
daily_end = remaining_col - 2


# ============================================================
# Counters
# ============================================================

shift_seen = 0
shift_inserted = 0
shift_existing = 0

snapshot_seen = 0
snapshot_inserted = 0
snapshot_existing = 0

numeric_hours = 0
text_notes = 0

suspicious_numeric = []

unresolved = Counter()

latest_date = None


# ============================================================
# Import
# ============================================================

for row_no in range(
    4,
    ws.max_row + 1
):

    source_code = clean(
        ws.cell(
            row_no,
            2
        ).value
    )

    if not source_code:
        continue

    source_code_upper = (
        source_code.upper()
    )

    canonical_lookup = (
        SOURCE_CODE_OVERRIDES.get(
            source_code_upper,
            source_code_upper
        )
    )

    machine = machines.get(
        canonical_lookup
    )

    if machine is None:
        unresolved[source_code] += 1
        continue


    row_latest_date = None


    # ========================================================
    # Shift / daily work-hour records
    # ========================================================

    for c in range(
        cycle_start,
        daily_end + 1
    ):

        month_name = headers.get(
            (1, c),
            ""
        )

        day_raw = headers.get(
            (2, c),
            ""
        )

        shift = headers.get(
            (3, c),
            ""
        )

        if month_name not in MONTH_NUM:
            continue

        if shift not in (
            "روز",
            "شب",
        ):
            continue

        try:
            day = int(
                float(day_raw)
            )
        except Exception:
            continue

        if not (
            1 <= day <= 31
        ):
            continue

        raw = ws.cell(
            row_no,
            c
        ).value

        if raw is None:
            continue

        raw_text = clean(raw)

        if not raw_text:
            continue


        month_num = MONTH_NUM[
            month_name
        ]

        jalali_date = (
            f"{operational_year}/"
            f"{month_num:02d}/"
            f"{day:02d}"
        )


        if (
            row_latest_date is None
            or jalali_date > row_latest_date
        ):
            row_latest_date = (
                jalali_date
            )


        if (
            latest_date is None
            or jalali_date > latest_date
        ):
            latest_date = jalali_date


        work_hours = None
        note = None


        if is_number(raw):

            value = float(raw)

            if 0 <= value <= 24:

                work_hours = value
                numeric_hours += 1

            else:

                # Do not treat an out-of-range value as work hours.
                # Preserve it as raw/note for review.
                note = raw_text

                suspicious_numeric.append(
                    (
                        machine["code"],
                        jalali_date,
                        shift,
                        raw_text,
                        row_no,
                        c,
                    )
                )

                text_notes += 1

        else:

            note = raw_text
            text_notes += 1


        shift_seen += 1

        cur.execute("""
        INSERT OR IGNORE INTO service_shift_hours (
            ingest_source_id,
            machine_id,
            canonical_code,
            jalali_date,
            shift,
            work_hours,
            note,
            raw_value,
            source_sheet,
            source_row,
            source_col
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            source_id,
            machine["id"],
            machine["code"],
            jalali_date,
            shift,
            work_hours,
            note,
            raw_text,
            ws.title,
            row_no,
            c,
        ))

        if cur.rowcount == 1:
            shift_inserted += 1
        else:
            shift_existing += 1


    # ========================================================
    # PM / service status snapshot
    # ========================================================

    target_change = ws.cell(
        row_no,
        5
    ).value

    current_meter = ws.cell(
        row_no,
        current_meter_col
    ).value

    remaining = ws.cell(
        row_no,
        remaining_col
    ).value

    pm_raw = clean(
        ws.cell(
            row_no,
            pm_col
        ).value
    )

    oil_raw = clean(
        ws.cell(
            row_no,
            oil_analysis_col
        ).value
    )


    has_snapshot = any(
        (
            target_change is not None,
            current_meter is not None,
            remaining is not None,
            bool(pm_raw),
            bool(oil_raw),
        )
    )


    if has_snapshot:

        snapshot_seen += 1

        target_value = (
            float(target_change)
            if is_number(target_change)
            else None
        )

        current_value = (
            float(current_meter)
            if is_number(current_meter)
            else None
        )

        remaining_value = (
            float(remaining)
            if is_number(remaining)
            else None
        )


        cur.execute("""
        INSERT OR IGNORE INTO service_pm_snapshots (
            ingest_source_id,
            machine_id,
            canonical_code,
            jalali_asof_date,
            next_service_hour,
            current_meter_hour,
            remaining_hours,
            pm_issue_raw,
            oil_analysis_raw,
            source_sheet,
            source_row
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            source_id,
            machine["id"],
            machine["code"],
            row_latest_date,
            target_value,
            current_value,
            remaining_value,
            pm_raw or None,
            oil_raw or None,
            ws.title,
            row_no,
        ))

        if cur.rowcount == 1:
            snapshot_inserted += 1
        else:
            snapshot_existing += 1


conn.commit()


# ============================================================
# Verification
# ============================================================

db_shift_count = cur.execute("""
SELECT COUNT(*)
FROM service_shift_hours
WHERE ingest_source_id=?
  AND jalali_date LIKE ?
""", (
    source_id,
    f"{operational_year}/%",
)).fetchone()[0]


db_snapshot_count = cur.execute("""
SELECT COUNT(*)
FROM service_pm_snapshots
WHERE ingest_source_id=?
""", (
    source_id,
)).fetchone()[0]


print("SERVICE 1405 NORMALIZATION")
print("=" * 78)

print(
    f"Operational year: "
    f"{operational_year}"
)

print(
    f"Source ID: "
    f"{source_id}"
)

print(
    f"Latest date: "
    f"{latest_date}"
)

print()

print("SHIFT HOURS")
print("-" * 78)

print(
    f"Seen:      {shift_seen}"
)

print(
    f"Inserted:  {shift_inserted}"
)

print(
    f"Existing:  {shift_existing}"
)

print(
    f"Numeric:   {numeric_hours}"
)

print(
    f"Notes:     {text_notes}"
)

print(
    f"DB total:  {db_shift_count}"
)


print()
print("PM SNAPSHOTS")
print("-" * 78)

print(
    f"Seen:      {snapshot_seen}"
)

print(
    f"Inserted:  {snapshot_inserted}"
)

print(
    f"Existing:  {snapshot_existing}"
)

print(
    f"DB total:  {db_snapshot_count}"
)


print()
print("SUSPICIOUS NUMERIC VALUES")
print("-" * 78)

if suspicious_numeric:

    for item in suspicious_numeric:

        (
            code,
            date,
            shift,
            raw,
            row_no,
            col_no,
        ) = item

        print(
            f"{code} | "
            f"{date} | "
            f"{shift} | "
            f"raw={raw} | "
            f"row={row_no} | "
            f"col={col_no}"
        )

else:
    print("NONE")


print()
print("UNRESOLVED / NOT IMPORTED")
print("-" * 78)

if unresolved:

    for code, count in sorted(
        unresolved.items()
    ):
        print(
            f"{code}: {count}"
        )

else:
    print("NONE")


print()
print("=" * 78)
print("NORMALIZATION COMPLETE")


wb.close()
conn.close()
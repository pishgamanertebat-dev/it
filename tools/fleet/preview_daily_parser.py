from __future__ import annotations

import re
import sqlite3
from pathlib import Path

from openpyxl import load_workbook


DB_PATH = Path(r"E:\KomatsoAI\data\fleet\db\fleet_ops.db")

DATE_RE = re.compile(
    r"(14\d{2})[./_-](\d{1,2})[./_-](\d{1,2})"
)


def clean(value) -> str:
    if value is None:
        return ""

    return (
        str(value)
        .replace("\u200c", " ")
        .strip()
    )


def normalize_header(value) -> str:
    text = clean(value)

    text = (
        text
        .replace("ي", "ی")
        .replace("ك", "ک")
    )

    return re.sub(r"\s+", " ", text)


def extract_sheet_date(title: str) -> str | None:
    match = DATE_RE.search(title)

    if not match:
        return None

    year, month, day = map(int, match.groups())

    return f"{year:04d}/{month:02d}/{day:02d}"


def latest_daily_snapshot() -> Path:
    if not DB_PATH.exists():
        raise SystemExit(
            f"ERROR: DB not found: {DB_PATH}"
        )

    with sqlite3.connect(DB_PATH) as conn:
        row = conn.execute(
            """
            SELECT raw_path
            FROM ingest_sources
            WHERE source_type = 'daily_fault_reports'
            ORDER BY id DESC
            LIMIT 1
            """
        ).fetchone()

    if not row:
        raise SystemExit(
            "ERROR: no daily_fault_reports snapshot registered"
        )

    path = Path(row[0])

    if not path.exists():
        raise SystemExit(
            f"ERROR: raw snapshot not found: {path}"
        )

    return path


def find_column(headers, predicate):
    for index, header in enumerate(headers, start=1):
        if predicate(header):
            return index

    return None


def parse_sheet(ws):
    header_row = None
    headers = []

    # Header ممکن است در ردیف اول یا دوم یا کمی پایین‌تر باشد.
    for row_number in range(
        1,
        min(ws.max_row, 12) + 1
    ):
        row_headers = [
            normalize_header(
                ws.cell(row_number, col).value
            )
            for col in range(1, ws.max_column + 1)
        ]

        has_machine_code = any(
            header == "کد"
            or "کد جدید" in header
            for header in row_headers
        )

        has_fault_column = any(
            "شرح معایب" in header
            for header in row_headers
        )

        if has_machine_code and has_fault_column:
            header_row = row_number
            headers = row_headers
            break

    if header_row is None:
        return None, []

    col_row_number = find_column(
        headers,
        lambda h: h == "ردیف",
    )

    col_machine_type = find_column(
        headers,
        lambda h: "نوع دستگاه" in h,
    )

    col_machine_code = find_column(
        headers,
        lambda h: h == "کد"
        or "کد جدید" in h,
    )

    col_mechanical = find_column(
        headers,
        lambda h: "شرح معایب مکانیکی" in h,
    )

    col_fabrication = find_column(
        headers,
        lambda h: "شرح معایب آهنگری" in h,
    )

    col_general = find_column(
        headers,
        lambda h:
            h.startswith("شرح معایب")
            and "مکانیکی" not in h
            and "آهنگری" not in h,
    )

    records = []

    for excel_row in range(
        header_row + 1,
        ws.max_row + 1
    ):
        machine_type = (
            clean(
                ws.cell(
                    excel_row,
                    col_machine_type
                ).value
            )
            if col_machine_type
            else ""
        )

        machine_code = (
            clean(
                ws.cell(
                    excel_row,
                    col_machine_code
                ).value
            )
            if col_machine_code
            else ""
        )

        # ردیف کاملاً خالی را نادیده بگیر.
        if not machine_type and not machine_code:
            continue

        record = {
            "excel_row": excel_row,

            "row_no": (
                clean(
                    ws.cell(
                        excel_row,
                        col_row_number
                    ).value
                )
                if col_row_number
                else ""
            ),

            "machine_type": machine_type,

            "machine_code": machine_code,

            "mechanical": (
                clean(
                    ws.cell(
                        excel_row,
                        col_mechanical
                    ).value
                )
                if col_mechanical
                else ""
            ),

            "fabrication": (
                clean(
                    ws.cell(
                        excel_row,
                        col_fabrication
                    ).value
                )
                if col_fabrication
                else ""
            ),

            "general": (
                clean(
                    ws.cell(
                        excel_row,
                        col_general
                    ).value
                )
                if col_general
                else ""
            ),
        }

        records.append(record)

    return header_row, records


def main():
    raw_path = latest_daily_snapshot()

    print(f"RAW: {raw_path}")
    print()

    wb = load_workbook(
        raw_path,
        read_only=True,
        data_only=True,
    )

    dated_sheets = []

    for ws in wb.worksheets:
        report_date = extract_sheet_date(ws.title)

        if report_date:
            dated_sheets.append(
                (report_date, ws.title)
            )

    dated_sheets.sort(reverse=True)

    print(
        f"DATED SHEETS FOUND: {len(dated_sheets)}"
    )

    print()

    # فعلاً فقط سه روز آخر را Preview می‌کنیم.
    for report_date, sheet_name in dated_sheets[:3]:
        ws = wb[sheet_name]

        header_row, records = parse_sheet(ws)

        print("=" * 70)
        print(f"DATE: {report_date}")
        print(f"SHEET: {sheet_name!r}")
        print(f"HEADER ROW: {header_row}")
        print(f"RECORDS: {len(records)}")
        print("=" * 70)

        for item in records:
            print(
                f"{item['machine_code']} | "
                f"{item['machine_type']}"
            )

            if item["mechanical"]:
                print(
                    f"  MECH: {item['mechanical']}"
                )

            if item["fabrication"]:
                print(
                    f"  FAB: {item['fabrication']}"
                )

            if item["general"]:
                print(
                    f"  GENERAL: {item['general']}"
                )

        print()


if __name__ == "__main__":
    main()
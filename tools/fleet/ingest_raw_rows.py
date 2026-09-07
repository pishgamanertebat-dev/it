from __future__ import annotations

import argparse
import hashlib
import json
import sqlite3
from datetime import date, datetime, time, timezone
from pathlib import Path

from openpyxl import load_workbook


DB_PATH = Path(r"E:\KomatsoAI\data\fleet\db\fleet_ops.db")


def json_safe(value):
    if value is None:
        return None

    if isinstance(value, (datetime, date, time)):
        return value.isoformat()

    if isinstance(value, (str, int, float, bool)):
        return value

    return str(value)


def row_to_json(values) -> str:
    safe_values = [json_safe(v) for v in values]

    return json.dumps(
        safe_values,
        ensure_ascii=False,
        separators=(",", ":"),
    )


def sha256_text(text: str) -> str:
    return hashlib.sha256(
        text.encode("utf-8")
    ).hexdigest()


def init_db(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS raw_excel_rows (
            id INTEGER PRIMARY KEY AUTOINCREMENT,

            ingest_source_id INTEGER NOT NULL,

            source_type TEXT NOT NULL,
            original_filename TEXT NOT NULL,
            source_imported_at TEXT NOT NULL,

            sheet_name TEXT NOT NULL,
            sheet_index INTEGER NOT NULL,
            excel_row INTEGER NOT NULL,

            row_sha256 TEXT NOT NULL,
            row_json TEXT NOT NULL,

            stored_at TEXT NOT NULL,

            UNIQUE (
                ingest_source_id,
                sheet_index,
                excel_row
            ),

            FOREIGN KEY (
                ingest_source_id
            )
            REFERENCES ingest_sources(id)
        )
        """
    )

    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS
        idx_raw_excel_rows_source
        ON raw_excel_rows (
            ingest_source_id
        )
        """
    )

    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS
        idx_raw_excel_rows_sheet
        ON raw_excel_rows (
            ingest_source_id,
            sheet_name
        )
        """
    )

    conn.commit()


def get_latest_source(
    conn: sqlite3.Connection,
    source_type: str,
):
    row = conn.execute(
        """
        SELECT
            id,
            source_type,
            original_filename,
            imported_at,
            raw_path
        FROM ingest_sources
        WHERE source_type = ?
        ORDER BY id DESC
        LIMIT 1
        """,
        (source_type,),
    ).fetchone()

    if not row:
        raise SystemExit(
            f"ERROR: no registered source found: {source_type}"
        )

    return row


def main() -> None:
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "source_type",
        help="Example: daily_fault_reports",
    )

    args = parser.parse_args()

    if not DB_PATH.exists():
        raise SystemExit(
            f"ERROR: DB not found: {DB_PATH}"
        )

    conn = sqlite3.connect(DB_PATH)

    conn.execute(
        "PRAGMA foreign_keys = ON"
    )

    init_db(conn)

    (
        ingest_source_id,
        source_type,
        original_filename,
        source_imported_at,
        raw_path_text,
    ) = get_latest_source(
        conn,
        args.source_type,
    )

    raw_path = Path(raw_path_text)

    if not raw_path.exists():
        raise SystemExit(
            f"ERROR: raw file not found: {raw_path}"
        )

    print(f"SOURCE ID: {ingest_source_id}")
    print(f"SOURCE: {source_type}")
    print(f"FILE: {original_filename}")
    print(f"RAW: {raw_path}")
    print()

    # data_only=False:
    # فرمول اصلی Excel را حفظ می‌کنیم، نه فقط نتیجه محاسبه‌شده.
    wb = load_workbook(
        raw_path,
        read_only=True,
        data_only=False,
    )

    stored_at = datetime.now(
        timezone.utc
    ).isoformat(timespec="seconds")

    total_nonempty = 0
    inserted = 0
    skipped = 0

    with conn:
        for sheet_index, ws in enumerate(
            wb.worksheets,
            start=1,
        ):
            sheet_count = 0

            for excel_row, row in enumerate(
                ws.iter_rows(values_only=True),
                start=1,
            ):
                # ردیف کاملاً خالی ارزش Raw ندارد؛
                # شماره ردیف Excel همچنان حفظ می‌شود.
                if all(
                    value is None or str(value).strip() == ""
                    for value in row
                ):
                    continue

                total_nonempty += 1
                sheet_count += 1

                row_json = row_to_json(row)
                row_sha256 = sha256_text(row_json)

                cur = conn.execute(
                    """
                    INSERT OR IGNORE INTO raw_excel_rows (
                        ingest_source_id,
                        source_type,
                        original_filename,
                        source_imported_at,
                        sheet_name,
                        sheet_index,
                        excel_row,
                        row_sha256,
                        row_json,
                        stored_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        ingest_source_id,
                        source_type,
                        original_filename,
                        source_imported_at,
                        ws.title,
                        sheet_index,
                        excel_row,
                        row_sha256,
                        row_json,
                        stored_at,
                    ),
                )

                if cur.rowcount == 1:
                    inserted += 1
                else:
                    skipped += 1

            print(
                f"{ws.title!r}: "
                f"{sheet_count} non-empty rows"
            )

    wb.close()
    conn.close()

    print()
    print("RAW ROW INGEST COMPLETE")
    print(f"Non-empty rows: {total_nonempty}")
    print(f"Inserted: {inserted}")
    print(f"Skipped: {skipped}")


if __name__ == "__main__":
    main()
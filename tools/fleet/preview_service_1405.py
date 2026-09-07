import math
import sqlite3
from collections import Counter
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


DB = r"E:\KomatsoAI\data\fleet\db\fleet_ops.db"

MONTH_NUM = {
    "فروردین": 1,
    "اردیبهشت": 2,
    "خرداد": 3,
    "تیر": 4,
    "مرداد": 5,
    "شهریور": 6,
    "مهر": 7,
    "آبان": 8,
    "آذر": 9,
    "دی": 10,
    "بهمن": 11,
    "اسفند": 12,
}


def clean(value):
    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\u200c", " ")
        .replace("ي", "ی")
        .replace("ك", "ک")
        .split()
    )


def numeric(value):
    if isinstance(value, bool):
        return False

    if isinstance(value, (int, float)):
        try:
            return math.isfinite(float(value))
        except Exception:
            return False

    return False


def header_map(ws, max_row=3):
    result = {}

    for r in range(1, max_row + 1):
        for c in range(1, ws.max_column + 1):
            value = clean(ws.cell(r, c).value)

            if value:
                result[(r, c)] = value

    for merged in ws.merged_cells.ranges:

        if merged.min_row > max_row:
            continue

        value = clean(
            ws.cell(
                merged.min_row,
                merged.min_col
            ).value
        )

        if not value:
            continue

        for r in range(
            merged.min_row,
            min(merged.max_row, max_row) + 1
        ):
            for c in range(
                merged.min_col,
                merged.max_col + 1
            ):
                result[(r, c)] = value

    return result


conn = sqlite3.connect(DB)
cur = conn.cursor()

source = cur.execute("""
SELECT id, raw_path
FROM ingest_sources
WHERE source_type='service_events'
ORDER BY id DESC
LIMIT 1
""").fetchone()

if not source:
    raise SystemExit(
        "ERROR: service_events source not found"
    )

source_id, raw_path = source

operational_year_row = cur.execute("""
SELECT setting_value
FROM fleet_settings
WHERE setting_key='operational_jalali_year'
""").fetchone()

operational_year = (
    int(operational_year_row[0])
    if operational_year_row
    else 1405
)

machines = {
    str(code).upper(): {
        "id": machine_id,
        "code": code,
        "status": identity_status,
    }
    for machine_id, code, identity_status
    in cur.execute("""
        SELECT
            id,
            canonical_code,
            identity_status
        FROM machines
    """)
}


wb = load_workbook(
    raw_path,
    data_only=True,
    read_only=False
)

ws = wb["ساعت کاری"]

headers = header_map(ws)


# ============================================================
# Locate operational 1405 cycle
# ============================================================

farvardin_starts = []

for c in range(1, ws.max_column + 1):

    value = clean(
        ws.cell(1, c).value
    )

    if value == "فروردین":
        farvardin_starts.append(c)


if not farvardin_starts:
    raise SystemExit(
        "ERROR: no Farvardin block found"
    )


# Latest Farvardin block = current operational cycle.
cycle_start = farvardin_starts[-1]


remaining_col = None

for c in range(1, ws.max_column + 1):

    value = clean(
        ws.cell(3, c).value
    )

    if value == "مانده به تعویض":
        remaining_col = c


if remaining_col is None:
    raise SystemExit(
        "ERROR: remaining-service column not found"
    )


current_meter_col = remaining_col - 1
pm_col = remaining_col + 1
oil_analysis_col = remaining_col + 2

daily_end = remaining_col - 2


print("SERVICE 1405 PREVIEW")
print("=" * 80)

print(
    f"Operational Jalali year: "
    f"{operational_year}"
)

print(
    f"Source ID: {source_id}"
)

print(
    f"Sheet: {ws.title}"
)

print(
    f"1405 cycle starts: "
    f"{get_column_letter(cycle_start)}1"
)

print(
    f"Daily data ends: "
    f"{get_column_letter(daily_end)}"
)

print(
    f"Current meter column: "
    f"{get_column_letter(current_meter_col)}"
)

print(
    f"Remaining column: "
    f"{get_column_letter(remaining_col)}"
)

print(
    f"PM column: "
    f"{get_column_letter(pm_col)}"
)

print(
    f"Oil-analysis flag column: "
    f"{get_column_letter(oil_analysis_col)}"
)


# ============================================================
# Show detected 1405 month blocks
# ============================================================

print()
print("DETECTED 1405 MONTH BLOCKS")
print("-" * 80)

month_starts = []

for c in range(
    cycle_start,
    daily_end + 1
):

    raw = clean(
        ws.cell(1, c).value
    )

    if raw in MONTH_NUM:
        month_starts.append(
            (
                c,
                raw,
                MONTH_NUM[raw]
            )
        )

for c, month_name, month_num in month_starts:

    print(
        f"{operational_year}/"
        f"{month_num:02d} "
        f"{month_name:<10} "
        f"starts at "
        f"{get_column_letter(c)}"
    )


# ============================================================
# Parse preview
# ============================================================

shift_rows = []
snapshots = []

unresolved = Counter()
text_values = Counter()

suspicious_hours = []

latest_global_date = None


for row_no in range(
    4,
    ws.max_row + 1
):

    source_code = clean(
        ws.cell(
            row_no,
            2
        ).value
    )

    if not source_code:
        continue

    source_code_upper = (
        source_code.upper()
    )

    machine = machines.get(
        source_code_upper
    )

    row_has_1405_data = False
    row_latest_date = None

    parsed_cells = []


    # --------------------------------------------------------
    # Daily / shift hours
    # --------------------------------------------------------

    for c in range(
        cycle_start,
        daily_end + 1
    ):

        month_name = headers.get(
            (1, c),
            ""
        )

        day_raw = headers.get(
            (2, c),
            ""
        )

        shift = headers.get(
            (3, c),
            ""
        )

        if month_name not in MONTH_NUM:
            continue

        if shift not in (
            "روز",
            "شب",
        ):
            continue

        try:
            day = int(
                float(day_raw)
            )
        except Exception:
            continue

        if not (
            1 <= day <= 31
        ):
            continue

        raw = ws.cell(
            row_no,
            c
        ).value

        if raw is None:
            continue

        raw_text = clean(raw)

        if not raw_text:
            continue

        row_has_1405_data = True

        month_num = MONTH_NUM[
            month_name
        ]

        jalali_date = (
            f"{operational_year}/"
            f"{month_num:02d}/"
            f"{day:02d}"
        )

        if (
            row_latest_date is None
            or jalali_date > row_latest_date
        ):
            row_latest_date = (
                jalali_date
            )

        if (
            latest_global_date is None
            or jalali_date > latest_global_date
        ):
            latest_global_date = (
                jalali_date
            )

        work_hours = None
        note = None

        if numeric(raw):

            work_hours = float(raw)

            if (
                work_hours < 0
                or work_hours > 24
            ):
                suspicious_hours.append(
                    (
                        source_code,
                        jalali_date,
                        shift,
                        raw
                    )
                )

        else:

            note = raw_text
            text_values[note] += 1

        parsed_cells.append({
            "date": jalali_date,
            "shift": shift,
            "hours": work_hours,
            "note": note,
            "raw": raw_text,
            "col": c,
        })


    # --------------------------------------------------------
    # PM snapshot
    # --------------------------------------------------------

    target_change = ws.cell(
        row_no,
        5
    ).value

    current_meter = ws.cell(
        row_no,
        current_meter_col
    ).value

    remaining = ws.cell(
        row_no,
        remaining_col
    ).value

    pm_raw = clean(
        ws.cell(
            row_no,
            pm_col
        ).value
    )

    oil_raw = clean(
        ws.cell(
            row_no,
            oil_analysis_col
        ).value
    )

    has_snapshot = any(
        (
            target_change is not None,
            current_meter is not None,
            remaining is not None,
            bool(pm_raw),
            bool(oil_raw),
        )
    )

    if (
        not row_has_1405_data
        and not has_snapshot
    ):
        continue


    if machine is None:

        unresolved[
            source_code
        ] += 1

        continue


    shift_rows.extend(
        {
            "machine_id": machine["id"],
            "canonical_code": machine["code"],
            **item,
        }
        for item in parsed_cells
    )


    if has_snapshot:

        check_diff = None

        if (
            numeric(target_change)
            and numeric(current_meter)
            and numeric(remaining)
        ):
            calculated = (
                float(target_change)
                - float(current_meter)
            )

            check_diff = (
                calculated
                - float(remaining)
            )

        snapshots.append({
            "machine_id":
                machine["id"],

            "canonical_code":
                machine["code"],

            "source_row":
                row_no,

            "asof":
                row_latest_date,

            "target_change":
                target_change,

            "current_meter":
                current_meter,

            "remaining":
                remaining,

            "pm_raw":
                pm_raw or None,

            "oil_raw":
                oil_raw or None,

            "check_diff":
                check_diff,
        })


# ============================================================
# Results
# ============================================================

print()
print("=" * 80)
print("PREVIEW COUNTS")

print(
    f"Shift/service cells: "
    f"{len(shift_rows)}"
)

numeric_count = sum(
    1
    for x in shift_rows
    if x["hours"] is not None
)

note_count = sum(
    1
    for x in shift_rows
    if x["note"] is not None
)

print(
    f"Numeric work-hour cells: "
    f"{numeric_count}"
)

print(
    f"Text/note cells: "
    f"{note_count}"
)

print(
    f"PM snapshots: "
    f"{len(snapshots)}"
)

print(
    f"Latest detected 1405 date: "
    f"{latest_global_date}"
)

print(
    f"Unresolved source codes: "
    f"{len(unresolved)}"
)

print(
    f"Suspicious numeric cells "
    f"(hours <0 or >24): "
    f"{len(suspicious_hours)}"
)


# ============================================================
# PM arithmetic checks
# ============================================================

valid_checks = []
bad_checks = []

for s in snapshots:

    diff = s["check_diff"]

    if diff is None:
        continue

    if abs(diff) <= 0.01:
        valid_checks.append(s)
    else:
        bad_checks.append(s)


print()
print("PM ARITHMETIC CHECK")
print("-" * 80)

print(
    f"Target - Current = Remaining PASS: "
    f"{len(valid_checks)}"
)

print(
    f"Mismatch: "
    f"{len(bad_checks)}"
)

for s in bad_checks[:10]:

    print(
        f"MISMATCH | "
        f"{s['canonical_code']} | "
        f"target={s['target_change']} | "
        f"current={s['current_meter']} | "
        f"remaining={s['remaining']} | "
        f"diff={s['check_diff']}"
    )


# ============================================================
# Specific machine examples
# ============================================================

for wanted in (
    "HD714",
    "HD701",
    "EX801",
):

    print()
    print("=" * 80)
    print(
        f"MACHINE PREVIEW: {wanted}"
    )

    rows = [
        x
        for x in shift_rows
        if x[
            "canonical_code"
        ].upper() == wanted
    ]

    snap = next(
        (
            x
            for x in snapshots
            if x[
                "canonical_code"
            ].upper() == wanted
        ),
        None
    )

    if not rows and not snap:
        print("NO 1405 DATA")
        continue

    if rows:

        print(
            f"Shift cells: "
            f"{len(rows)}"
        )

        dates = [
            x["date"]
            for x in rows
        ]

        print(
            f"Date range: "
            f"{min(dates)} -> "
            f"{max(dates)}"
        )

        numeric_hours = sum(
            x["hours"]
            for x in rows
            if x["hours"] is not None
        )

        print(
            f"Total numeric hours "
            f"in detected 1405 block: "
            f"{numeric_hours:g}"
        )

        print(
            "Last 12 records:"
        )

        for x in sorted(
            rows,
            key=lambda y: (
                y["date"],
                0
                if y["shift"] == "روز"
                else 1
            )
        )[-12:]:

            print(
                f"  {x['date']} | "
                f"{x['shift']} | "
                f"hours={x['hours']} | "
                f"note={x['note']}"
            )

    if snap:

        print()
        print("SERVICE STATUS")

        print(
            f"  as-of: "
            f"{snap['asof']}"
        )

        print(
            f"  target/change hour: "
            f"{snap['target_change']}"
        )

        print(
            f"  current meter: "
            f"{snap['current_meter']}"
        )

        print(
            f"  remaining: "
            f"{snap['remaining']}"
        )

        print(
            f"  PM raw: "
            f"{snap['pm_raw']}"
        )

        print(
            f"  oil-analysis raw: "
            f"{snap['oil_raw']}"
        )


# ============================================================
# Notes
# ============================================================

print()
print("=" * 80)
print("TEXT VALUES / NOTES")

for text, count in (
    text_values.most_common(20)
):
    print(
        f"{count:4} | {text}"
    )


print()
print("=" * 80)
print("UNRESOLVED CODES")

if unresolved:

    for code, count in sorted(
        unresolved.items()
    ):
        print(
            f"{code}: {count}"
        )

else:
    print("NONE")


print()
print("=" * 80)

print(
    "READ-ONLY PREVIEW COMPLETE - "
    "DATABASE NOT MODIFIED"
)

wb.close()
conn.close()
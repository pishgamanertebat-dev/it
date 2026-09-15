import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from .source import FormulaReader, is_yellow, read_source, source_hash
from .proposal import MODEL_MAP, build_proposal, evaluate, next_interval, resolve_items


def sample_source(model='HD785-7', code='HD708', last=200, remaining=-2):
    kind, label = next(k for k,v in MODEL_MAP.items() if v==model)
    return dict(machines=[dict(row=4,code=code,legacy_code='',name=kind,remaining=remaining,
                              current_meter=1002,target_meter=1000,last_service=(6,1,0),last_work=(6,2,0),errors=[])],
                plans=[dict(row=2,kind=kind,model=label,code=code,last_interval=last)],
                cutoff=(6,2,0),as_of=(6,22,1),warnings=['اطلاعات قدیمی'],sha256='fixture')


class ProposalTests(unittest.TestCase):
    def test_cycle_advances_once_including_wrap(self):
        for last in range(200,2001,200):
            self.assertEqual(next_interval(last),200 if last==2000 else last+200)
        self.assertEqual(next_interval('۲۰۰'),400)
        for bad in (None,True,0,201,2200,'=200'):
            with self.assertRaises(ValueError): next_interval(bad)

    def test_warning_boundaries_and_source_plan_date(self):
        for remaining, state, selected, warned in [(-1,'DUE',1,0),(0,'DUE',1,0),(.5,'NEAR_DUE',0,1),(24,'NEAR_DUE',0,1),(24.5,'NEAR_DUE',0,1),(30,'NEAR_DUE',0,1),(30.5,'OK',0,0)]:
            with patch('tools.fleet.oil_change.proposal.read_source',return_value=sample_source(remaining=remaining)):
                proposal=build_proposal()
            self.assertEqual(proposal['plan_date'],'1405/06/23')
            self.assertEqual(proposal['evaluations'][0]['components']['oil_change']['state'],state)
            self.assertEqual(len(proposal['items']),selected)
            self.assertEqual(len(proposal['warnings']),warned)
            if warned:
                self.assertEqual(proposal['warnings'], [f'HD708 : تا موعد بعدی {remaining:g} ساعت، برای سرویس 400 ساعتی'])

    def test_confirmed_bulldozer_alias_and_legacy_identity(self):
        for raw,model,canonical in [('W151','D155A-6','D151'),('W152','D155A-2','D152')]:
            source=sample_source(model,raw)
            source['machines'][0].update(code='D155',legacy_code=raw[1:])
            item=resolve_items([raw],source)[0]
            self.assertEqual(item['machine_code'],canonical)
            self.assertIn(model,item['action_code'])

    def test_pc1250_is_supported_by_automatic_planning(self):
        source = sample_source('PC1250-8', 'EX1252', last=1800, remaining=30)
        items, review = evaluate(source)
        self.assertEqual(review, [])
        self.assertEqual(items[0]['machine_code'], 'EX1252')
        self.assertEqual(items[0]['components']['oil_change']['state'], 'NEAR_DUE')
        self.assertTrue(items[0]['action_code'].endswith('_2000'))

    def test_missing_duplicate_unsupported_and_invalid_rows_require_review(self):
        for mutate in (lambda s:s['plans'][0].update(model='9999-9'),
                       lambda s:s['plans'].append(s['plans'][0].copy()),
                       lambda s:s['machines'].clear(),
                       lambda s:s['machines'][0].update(errors=['bad formula'])):
            source=sample_source(); mutate(source)
            items,review=evaluate(source)
            self.assertEqual(items,[])
            self.assertTrue(review)
            with self.assertRaises(ValueError):resolve_items(['HD708'],source)

    def test_manual_add_keeps_automatic_action_and_rejects_duplicates(self):
        source=sample_source(last=2000,remaining=100)
        self.assertTrue(resolve_items(['708'],source)[0]['action_code'].endswith('_200'))
        with self.assertRaises(ValueError):resolve_items(['708','HD708'],source)

    def test_excluded_units_never_enter_orders_warnings_or_review(self):
        for with_plans in (False, True):
            source = sample_source()
            for code, kind, model in [('W471','لودر','470-6'), ('EX231','بیل مکانیکی','230')]:
                source['machines'].append({**source['machines'][0], 'code':code, 'remaining':-100, 'errors':['bad formula']})
                if with_plans:
                    source['plans'].append(dict(row=9,code=code,kind=kind,model=model,last_interval=None))
            items, review = evaluate(source)
            self.assertEqual([i['machine_code'] for i in items], ['HD708'])
            self.assertEqual(review, [])
            for code in ('W471', 'EX231', '471', '231'):
                with self.assertRaises(ValueError):
                    resolve_items([code], source)


class SourceTests(unittest.TestCase):
    def test_only_exact_yellow_proves_service_even_on_empty_cell(self):
        cell=Workbook().active['A1']
        for color,expected in [('FFFF00',True),('FFFFFF00',True),('FFC000',False),('FFCE3C',False),('FFFF01',False)]:
            cell.fill=PatternFill('solid',fgColor=color)
            self.assertEqual(is_yellow(cell),expected)
        cell.fill=PatternFill('solid',fgColor='FFFF00');cell.fill.fgColor.tint=.1
        self.assertFalse(is_yellow(cell))

    def test_safe_formula_calculation_no_cached_values_required(self):
        s=Workbook().active
        for address,value in {'E4':200,'F4':'=100+20','G4':10,'H4':'خراب','I4':5,'J4':'=SUM(F4:I4)','K4':'=E4-J4'}.items():
            s[address]=value
        reader=FormulaReader(s)
        self.assertEqual(reader.value('K4'),65)
        self.assertEqual(FormulaReader(s,[9]).value('K4'),70)
        for formula in ('=__import__("os")','=SUM(F5:I5)','=K4','=1/0','=E5+1','=NOW()'):
            s['K4']=formula
            with self.assertRaises(ValueError):FormulaReader(s).value('K4')

    def test_workbooks_yellow_orange_future_and_recalculation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);hours=root/'hours.xlsx';plans=root/'plans.xlsx'
            w=Workbook();s=w.active;s.title='ساعت کاری'
            s['G1']='فروردین'
            for col,day in [(7,1),(9,2),(11,3)]:
                s.cell(2,col,day);s.cell(3,col,'روز');s.cell(3,col+1,'شب')
            s['M3']='ساعت کار';s['N3']='مانده به تعویض'
            for addr,val in {'A4':'دامپتراک','B4':'HD708','E4':200,'F4':170,'G4':10,'I4':5,'K4':24,'M4':'=SUM(F4:L4)','N4':'=E4-M4'}.items():s[addr]=val
            s['H4'].fill=PatternFill('solid',fgColor='FFFF00')
            s['I4'].fill=PatternFill('solid',fgColor='FFC000')
            s['K4'].fill=PatternFill('solid',fgColor='FFFF00')
            for row, code, name in [(5,'W471','لودر'), (6,'EX231','بیل مکانیکی')]:
                s.cell(row,1,name);s.cell(row,2,code)
                s.cell(row,14,'=1/0')
                s.cell(row,10,12).fill=PatternFill('solid',fgColor='FFFF00')
            w.save(hours);w.close()
            p=Workbook();p.active.append(['ردیف','نوع دستگاه','مدل دستگاه','کد دستگاه','نوع سرویس']);p.active.append([1,'دامپتراک','785-7','HD708',200])
            p.active.append([2,'لودر','470-6','W471',None]);p.active.append([3,'بیل مکانیکی',230,'EX231',None]);p.save(plans);p.close()
            before=source_hash(hours,plans)
            source=read_source(hours,plans,as_of=(1,2,1))
            machine=source['machines'][0]
            self.assertEqual([m['code'] for m in source['machines']], ['HD708'])
            self.assertEqual([p['code'] for p in source['plans']], ['HD708'])
            self.assertEqual(machine['remaining'],15)
            self.assertEqual(machine['current_meter'],185)
            self.assertEqual(machine['last_service'],(1,1,1))
            self.assertEqual(source['cutoff'],(1,2,0))
            self.assertEqual(before,source_hash(hours,plans))
            p=Workbook();p.save(plans);p.close()
            self.assertNotEqual(before,source_hash(hours,plans))

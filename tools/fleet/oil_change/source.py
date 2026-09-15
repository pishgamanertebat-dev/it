"""Read source workbooks without changing them or trusting stale formula caches."""
import ast
import hashlib
from io import BytesIO
import math
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple
from tools.fleet.greasing.source import (
    SOURCE, SHEET, clean, code_text, date_columns, completed_cutoff, format_shift,
)
from .scope import excluded_codes, included_plans

PLANNING_SOURCE = Path('E:/Function/برنامه ریزی سرویس.xlsx')


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def source_hash(path=SOURCE, planning_path=PLANNING_SOURCE):
    return hashlib.sha256(Path(path).read_bytes() + b'\0' + Path(planning_path).read_bytes()).hexdigest()


class FormulaReader:
    """Only numeric literals, cell references, SUM and arithmetic; never execute Excel text."""
    def __init__(self, sheet, excluded_columns=()):
        self.sheet = sheet
        self.excluded_columns = set(excluded_columns)
        self.cache = {}
        self.visiting = set()

    def value(self, address):
        address = address.replace('$', '').upper()
        row, col = coordinate_to_tuple(address)
        if col in self.excluded_columns:
            return 0
        if address in self.cache:
            return self.cache[address]
        if address in self.visiting:
            raise ValueError('فرمول دوری در سلول ' + address)
        cell = self.sheet[address]
        value = cell.value
        if cell.data_type == 'e':
            raise ValueError('خطای اکسل در سلول ' + address)
        if not isinstance(value, str) or not value.startswith('='):
            return value
        self.visiting.add(address)
        try:
            expression = value[1:].upper()
            def sum_range(match):
                left, top, right, bottom = range_boundaries(match[1].replace('$', ''))
                if bottom != row or top != row:
                    raise ValueError('فرمول جمع باید به همان ردیف دستگاه اشاره کند: ' + address)
                values = [self.value(self.sheet.cell(row, c).coordinate) for c in range(left, right + 1)]
                return repr(sum(v for v in values if number(v)))
            expression = re.sub(r'SUM\(\s*([A-Z$]+\d+:[A-Z$]+\d+)\s*\)', sum_range, expression)
            def reference(match):
                target = match[0]
                if coordinate_to_tuple(target.replace('$', ''))[0] != row:
                    raise ValueError('ارجاع فرمول به دستگاه دیگر: ' + address)
                result = self.value(target)
                if result is None:
                    return '0'
                if not number(result):
                    raise ValueError('مقدار غیرعددی در فرمول ' + address)
                return repr(result)
            expression = re.sub(r'\$?[A-Z]{1,3}\$?\d+', reference, expression)
            def calculate(node):
                if isinstance(node, ast.Constant) and number(node.value):
                    return node.value
                if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
                    return calculate(node.operand) * (-1 if isinstance(node.op, ast.USub) else 1)
                if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub, ast.Mult, ast.Div)):
                    a, b = calculate(node.left), calculate(node.right)
                    if isinstance(node.op, ast.Add): return a + b
                    if isinstance(node.op, ast.Sub): return a - b
                    if isinstance(node.op, ast.Mult): return a * b
                    return a / b
                raise ValueError('فرمول پشتیبانی‌نشده در سلول ' + address)
            result = calculate(ast.parse(expression, mode='eval').body)
            if not number(result):
                raise ValueError('نتیجهٔ نامعتبر فرمول در سلول ' + address)
            self.cache[address] = result
            return result
        except (SyntaxError, ZeroDivisionError, OverflowError) as exc:
            raise ValueError('فرمول نامعتبر در سلول ' + address) from exc
        finally:
            self.visiting.remove(address)


def is_yellow(cell):
    fill = cell.fill
    return (fill.patternType == 'solid' and fill.fgColor.type == 'rgb'
            and fill.fgColor.tint == 0 and str(fill.fgColor.rgb)[-6:].upper() == 'FFFF00')


def read_source(path=SOURCE, planning_path=PLANNING_SOURCE, as_of=None):
    hours_data, planning_data = Path(path).read_bytes(), Path(planning_path).read_bytes()
    hours = load_workbook(BytesIO(hours_data), data_only=False)
    planning = load_workbook(BytesIO(planning_data), data_only=False)
    limit = as_of or completed_cutoff()
    try:
        sheet = planning.active
        expected = ('نوع دستگاه','مدل دستگاه','کد دستگاه','نوع سرویس')
        if tuple(clean(sheet.cell(1,c).value) for c in range(2,6)) != expected:
            raise ValueError('سرستون‌های فایل برنامه‌ریزی سرویس تغییر کرده‌اند.')
        plans = [dict(row=r, kind=clean(sheet.cell(r,2).value), model=clean(sheet.cell(r,3).value),
                      code=code_text(sheet.cell(r,4).value), last_interval=sheet.cell(r,5).value)
                 for r in range(2,sheet.max_row+1) if any(sheet.cell(r,c).value is not None for c in range(2,6))]
        excluded = excluded_codes(plans)
        plans = included_plans(plans)
        ws = hours[SHEET]
        columns = date_columns(ws)
        remaining_columns = [c for c in range(1, ws.max_column + 1) if clean(ws.cell(3,c).value) == 'مانده به تعویض']
        if len(remaining_columns) != 1:
            raise ValueError('ستون یکتای مانده به تعویض پیدا نشد.')
        remaining_col = remaining_columns[0]
        reader = FormulaReader(ws, [c for stamp,c in columns if stamp > limit])
        machines, populated = [], []
        for row in range(4, ws.max_row + 1):
            name, code = clean(ws.cell(row,1).value), code_text(ws.cell(row,2).value)
            if not name or code.upper() in excluded:
                continue
            events, work, errors = [], [], []
            for stamp, col in columns:
                if stamp > limit:
                    continue
                cell = ws.cell(row,col)
                if is_yellow(cell):
                    events.append(stamp)
                    populated.append(stamp)
                try:
                    value = reader.value(cell.coordinate)
                except ValueError:
                    value = cell.value
                if value is not None and clean(value) not in ('', '-'):
                    populated.append(stamp)
                if number(value) and value > 0:
                    work.append(stamp)
            values = {}
            for key, col in [('remaining',remaining_col), ('current_meter',remaining_col-1), ('target_meter',5)]:
                try:
                    value = reader.value(ws.cell(row,col).coordinate)
                    if not number(value):
                        raise ValueError('مقدار عددی معتبر در ' + ws.cell(row,col).coordinate + ' وجود ندارد.')
                    values[key] = value
                except ValueError as exc:
                    values[key] = None
                    errors.append(str(exc))
            machines.append(dict(row=row, code=code, legacy_code=code_text(ws.cell(row,3).value), name=name,
                                 last_service=max(events) if events else None,
                                 last_work=max(work) if work else None, errors=errors, **values))
        cutoff = max(populated) if populated else None
        warnings = []
        if cutoff is None:
            warnings.append('هیچ ثبت کارکرد یا تعویض روغن در سال عملیاتی تا تاریخ محاسبه پیدا نشد.')
        elif cutoff < limit:
            warnings.append('اطلاعات کارکرد قدیمی است؛ آخرین ثبت: ' + format_shift(cutoff) + '؛ مانده فقط از داده‌های ثبت‌شده محاسبه شده است.')
        return dict(machines=machines, plans=plans, cutoff=cutoff, as_of=limit, warnings=warnings,
                    sha256=hashlib.sha256(hours_data+b'\0'+planning_data).hexdigest())
    finally:
        hours.close()
        planning.close()

from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import load_workbook


DB_PATH = Path(
    r"E:\KomatsoAI\data\fleet\db\fleet_ops.db"
)


HEADER_WORDS = (
    "تاریخ",
    "شرح",
    "عیب",
    "خرابی",
    "تعمیر",
    "مکانیک",
    "کد",
    "قطعه",
    "اقدام",
    "ساعت",
    "وضعیت",
    "توضیح",
)


REPRESENTATIVE_SHEETS = (
    "متفرقه",
    "801",
    "601",
    "601.",
    "1254",
    "152",
    "702",
    "714",
    "تعمیرگاه",
)


def clean(value):
    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\u200c", " ")
        .replace("ي", "ی")
        .replace("ك", "ک")
        .split()
    )


def row_values(ws, row_number):
    values = []

    for cell in ws[row_number]:
        value = clean(cell.value)

        if value:
            values.append(
                f"{cell.column_letter}={value}"
            )

    return values


def detect_header(ws):
    best_row = None
    best_score = 0
    best_values = []

    max_scan = min(
        ws.max_row,
        15,
    )

    for row_number in range(
        1,
        max_scan + 1,
    ):
        values = row_values(
            ws,
            row_number,
        )

        joined = " | ".join(values)

        score = sum(
            1
            for word in HEADER_WORDS
            if word in joined
        )

        if score > best_score:
            best_score = score
            best_row = row_number
            best_values = values

    return (
        best_row,
        best_score,
        best_values,
    )


def next_nonempty_rows(
    ws,
    start_row,
    count=3,
):
    result = []

    if start_row is None:
        start_row = 0

    for row_number in range(
        start_row + 1,
        ws.max_row + 1,
    ):
        values = row_values(
            ws,
            row_number,
        )

        if not values:
            continue

        result.append(
            (
                row_number,
                values,
            )
        )

        if len(result) >= count:
            break

    return result


def main():
    if not DB_PATH.exists():
        raise SystemExit(
            f"ERROR: DB not found: {DB_PATH}"
        )

    conn = sqlite3.connect(DB_PATH)

    source = conn.execute(
        """
        SELECT
            id,
            original_filename,
            raw_path,
            imported_at

        FROM ingest_sources

        WHERE source_type = 'maintenance_events'

        ORDER BY id DESC

        LIMIT 1
        """
    ).fetchone()

    conn.close()

    if not source:
        raise SystemExit(
            "ERROR: maintenance_events source not found"
        )

    source_id = source[0]
    filename = source[1]
    raw_path = Path(source[2])
    imported_at = source[3]

    if not raw_path.exists():
        raise SystemExit(
            f"ERROR: raw file not found: {raw_path}"
        )

    print("MAINTENANCE PARSER PREVIEW")
    print("=" * 70)

    print(f"Source ID: {source_id}")
    print(f"File: {filename}")
    print(f"Imported: {imported_at}")
    print(f"Raw: {raw_path}")

    wb = load_workbook(
        raw_path,
        read_only=True,
        data_only=False,
    )

    print()
    print(
        f"SHEETS: {len(wb.sheetnames)}"
    )

    print()
    print("HEADER DETECTION")
    print("=" * 70)

    for sheet_name in wb.sheetnames:

        ws = wb[sheet_name]

        (
            header_row,
            score,
            header_values,
        ) = detect_header(ws)

        print()
        print(
            f"[{sheet_name}] "
            f"rows={ws.max_row} "
            f"cols={ws.max_column} "
            f"header_row={header_row} "
            f"score={score}"
        )

        if header_values:
            print(
                "  HEADER: "
                + " | ".join(
                    header_values
                )
            )
        else:
            print(
                "  HEADER: NOT DETECTED"
            )

    print()
    print()
    print("REPRESENTATIVE SHEETS")
    print("=" * 70)

    for sheet_name in REPRESENTATIVE_SHEETS:

        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]

        (
            header_row,
            score,
            header_values,
        ) = detect_header(ws)

        print()
        print("#" * 70)
        print(
            f"SHEET: {sheet_name}"
        )
        print(
            f"SIZE: rows={ws.max_row}, "
            f"cols={ws.max_column}"
        )
        print(
            f"HEADER ROW: {header_row}"
        )
        print(
            f"HEADER SCORE: {score}"
        )

        if header_values:
            print(
                "HEADER:"
            )

            for item in header_values:
                print(
                    f"  {item}"
                )

        print()
        print("FIRST DATA ROWS:")

        samples = next_nonempty_rows(
            ws,
            header_row,
            count=3,
        )

        if not samples:
            print("  None")

        for row_number, values in samples:

            print(
                f"  ROW {row_number}:"
            )

            for item in values:
                print(
                    f"    {item}"
                )

    wb.close()


if __name__ == "__main__":
    main()
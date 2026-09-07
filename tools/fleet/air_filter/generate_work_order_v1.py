import argparse
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


TEMPLATE_SHEET = "Sheet1 (486)"

DATA_START_ROW = 3
TEMPLATE_DATA_ROWS = 6

DEFAULT_ACTION = "تعویض هواکش بیرونی"


def normalize_code(value):
    return str(value).strip().upper()


def machine_name_from_code(code):
    c = normalize_code(code)

    if c.isdigit():
        n = int(c)

        if 461 <= n <= 469:
            return "دامپتراک"

        if 701 <= n <= 716:
            return "دامپتراک"

        if n == 231:
            return "بیل"

    if c.startswith("MZ"):
        return "مزدا"

    if c in {"S1", "S3"}:
        return "کامیون سهند زرد"

    if c == "TA1":
        return "کامیون آب پاش"

    if c == "TR1":
        return "خاور"

    if c == "DG1":
        return "ژنراتور"

    if c in {"PR2", "PR3"}:
        return "پیکاپ ریچ"

    raise ValueError(
        f"کد دستگاه ناشناخته است: {code}"
    )


def copy_row_style(ws, source_row, target_row):
    for col in range(1, 7):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)

        if src.has_style:
            dst._style = copy(src._style)

        if src.number_format:
            dst.number_format = src.number_format

        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)

    ws.row_dimensions[target_row].height = (
        ws.row_dimensions[source_row].height
    )


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--template",
        required=True,
        help="Original air-filter workbook"
    )

    parser.add_argument(
        "--date",
        required=True,
        help="Example: 1405/06/09"
    )

    parser.add_argument(
        "--output",
        required=True
    )

    parser.add_argument(
        "--codes",
        nargs="+",
        required=True
    )

    args = parser.parse_args()

    template_path = Path(args.template)
    output_path = Path(args.output)

    if not template_path.exists():
        raise SystemExit(
            f"TEMPLATE NOT FOUND: {template_path}"
        )

    if not args.codes:
        raise SystemExit(
            "NO MACHINE CODES"
        )

    wb = load_workbook(template_path)

    if TEMPLATE_SHEET not in wb.sheetnames:
        raise SystemExit(
            f"SHEET NOT FOUND: {TEMPLATE_SHEET}"
        )

    # فقط همان شیت روز 9 باقی بماند.
    ws = wb[TEMPLATE_SHEET]

    for other_ws in list(wb.worksheets):
        if other_ws.title != TEMPLATE_SHEET:
            wb.remove(other_ws)

    machine_count = len(args.codes)

    # --------------------------------------------------------
    # تعداد ردیف‌های فرم را با تعداد دستگاه‌ها هماهنگ کن
    # --------------------------------------------------------

    if machine_count > TEMPLATE_DATA_ROWS:
        extra = machine_count - TEMPLATE_DATA_ROWS

        # ردیف امضا در تمپلیت اصلی ردیف 9 است.
        ws.insert_rows(
            DATA_START_ROW + TEMPLATE_DATA_ROWS,
            amount=extra
        )

        source_style_row = (
            DATA_START_ROW
            + TEMPLATE_DATA_ROWS
            - 1
        )

        for row in range(
            DATA_START_ROW + TEMPLATE_DATA_ROWS,
            DATA_START_ROW + machine_count
        ):
            copy_row_style(
                ws,
                source_style_row,
                row
            )

    elif machine_count < TEMPLATE_DATA_ROWS:
        remove_count = (
            TEMPLATE_DATA_ROWS
            - machine_count
        )

        ws.delete_rows(
            DATA_START_ROW + machine_count,
            amount=remove_count
        )

    # --------------------------------------------------------
    # تاریخ
    # --------------------------------------------------------

    ws["F1"] = args.date

    # --------------------------------------------------------
    # اطلاعات دستگاه‌ها
    # --------------------------------------------------------

    for index, raw_code in enumerate(
        args.codes,
        start=1
    ):
        row = DATA_START_ROW + index - 1

        machine_name = machine_name_from_code(
            raw_code
        )

        # ردیف
        ws.cell(row, 1).value = index

        # نام دستگاه
        ws.cell(row, 2).value = machine_name

        # کد دستگاه
        #
        # همان چیزی که مسئول نت وارد کرده نمایش داده می‌شود.
        ws.cell(row, 3).value = str(raw_code).strip()

        # شرح اقدام
        ws.cell(row, 4).value = DEFAULT_ACTION

        # انجام شد / انجام نشد باید برای فرم چاپی خالی بماند.
        ws.cell(row, 5).value = None
        ws.cell(row, 6).value = None

    # --------------------------------------------------------
    # محدوده چاپ
    # --------------------------------------------------------

    signature_row = (
        DATA_START_ROW
        + machine_count
    )

    ws.print_area = (
        f"A1:F{signature_row}"
    )

    # --------------------------------------------------------
    # ذخیره
    # --------------------------------------------------------

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    wb.save(output_path)
    wb.close()

    print("=" * 70)
    print("AIR FILTER WORK ORDER V1 CREATED")
    print("=" * 70)

    print(
        f"Date: {args.date}"
    )

    print(
        f"Machines: {machine_count}"
    )

    for index, raw_code in enumerate(
        args.codes,
        start=1
    ):
        print(
            f"{index}. "
            f"{raw_code} | "
            f"{machine_name_from_code(raw_code)} | "
            f"{DEFAULT_ACTION}"
        )

    print()
    print(
        f"OUTPUT: {output_path}"
    )

    print("=" * 70)


if __name__ == "__main__":
    main()
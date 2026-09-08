"""Deterministic proposals from observed service events; never writes source data."""
from collections import Counter
from io import BytesIO
from pathlib import Path
import hashlib
import math

from openpyxl import load_workbook
from tools.fleet.preview_air_filter_due import (
    clean, clean_code, parse_header_date, jalali_ordinal, next_jalali_day, fmt_date,
)

SOURCE = Path('E:/Function/هواکش.xlsx')
MARKS = {'صبح', 'ظهر', 'عصر', 'شب', '*'}
from tools.fleet.air_filter.rules import OUTER, BOTH, ACTIONS, rule_for


def _sheet_name(workbook):
    preferred = 'Sheet1 (2)'
    if preferred in workbook.sheetnames:
        return preferred
    candidates = []
    for sheet in workbook:
        first = tuple(clean(sheet.cell(1, col).value) for col in range(1, 4))
        fields = tuple(clean(sheet.cell(2, col).value) for col in range(4, 7))
        if first == ('ردیف', 'نام دستگاه', 'کد دستگاه') and fields == ('کارکرد', 'درونی', 'بیرونی'):
            candidates.append(sheet.title)
    if len(candidates) != 1:
        raise ValueError('شیت عملیاتی هواکش به‌صورت یکتا قابل تشخیص نیست.')
    return candidates[0]


def _source_code(name, value):
    code = clean_code(value)
    if code.isdecimal() and name == 'بیل مکانیکی':
        return 'EX' + code
    if code.isdecimal() and name == 'لودر':
        return ('WA' if code == '601' else 'W') + code
    if code.isdecimal() and name == 'بلدوزر':
        return 'D' + code
    return code


def has_month_activity(daily, cutoff):
    for entry in daily:
        if entry['date'][0] != cutoff[0] or entry['date'] > cutoff:
            continue
        value = entry['hours']
        if clean(value) in {'', '-'}:
            continue
        if not isinstance(value, bool) and isinstance(value, (int, float)) and value == 0:
            continue
        # Invalid recorded hours still require review; they are not inactivity.
        return True
    return False


def read_source(path=SOURCE, target=None):
    payload = Path(path).read_bytes()
    cached = load_workbook(BytesIO(payload), data_only=True, read_only=True)
    formulas = load_workbook(BytesIO(payload), data_only=False, read_only=True)
    try:
        sheet_name = _sheet_name(cached)
        rows = list(cached[sheet_name].iter_rows(values_only=True))
        raw = list(formulas[sheet_name].iter_rows(values_only=True))
    finally:
        cached.close()
        formulas.close()
    machines = [(i, clean(r[1]), _source_code(clean(r[1]), r[2])) for i,r in enumerate(rows[2:],2) if r[1] or r[2]]
    blocks = []
    issues = []
    col = 3
    while col < len(rows[0]):
        md = parse_header_date(rows[0][col])
        if md is None and clean(rows[1][col]) == 'کارکرد':
            raise ValueError(f'تاریخ ستون {col+1} قابل تشخیص نیست؛ ساختار اکسل بررسی شود.')
        if md and (target is None or md < target):
            fields = tuple(clean(v) for v in rows[1][col:col+5])
            if fields[:3] != ('کارکرد','درونی','بیرونی'):
                issues.append({'code':'FILE','reason':f'ساختار تاریخ {fmt_date(md)} نامعتبر است.'})
                col += 1
                continue
            width = 5 if fields == ('کارکرد','درونی','بیرونی','درونی','بیرونی') and all(
                parse_header_date(rows[0][col+offset]) == md for offset in range(5)
            ) else 3
            blocks.append((md,col,width))
            col += width
        else:
            col += 1
    counts = Counter(md for md,c,width in blocks)
    if any(n > 1 for n in counts.values()):
        raise ValueError('تاریخ تکراری در اکسل؛ پیش از محاسبه ساختار فایل بررسی شود.')
    if issues:
        raise ValueError(' / '.join(i['reason'] for i in issues))
    populated = [(md,c) for md,c,width in blocks if any(
        any(clean(v) for v in rows[i][c:c+width]) for i,n,k in machines
    )]
    if not populated:
        raise ValueError('هیچ دادهٔ عملیاتی پیش از روز هدف وجود ندارد.')
    cutoff = max(md for md,c in populated)
    plan = target or next_jalali_day(*cutoff)
    if plan[0] > 12:
        raise ValueError('روز هدف خارج از سال عملیاتی ۱۴۰۵ است.')
    blocks = sorted((md,c,width) for md,c,width in blocks if md <= cutoff)
    if jalali_ordinal(*plan) - jalali_ordinal(*cutoff) > 1:
        issues.append({'code':'FILE','reason':f'داده تا {fmt_date(cutoff)} است و با روز هدف فاصله دارد.'})
    result = []
    code_counts = Counter(k for i,n,k in machines)
    for i,name,code in machines:
        daily = []
        for md,c,width in blocks:
            values = list(rows[i][c:c+width])
            for offset in range(width):
                if values[offset] is None and isinstance(raw[i][c+offset],str) and raw[i][c+offset].startswith('='):
                    values[offset] = 'فرمول بدون مقدار محاسبه‌شده'
            inner_values = [values[1]] + ([values[3]] if width == 5 else [])
            outer_values = [values[2]] + ([values[4]] if width == 5 else [])
            daily.append({'date':md,'hours':values[0],
                          'inner':next((v for v in inner_values if clean(v)),None),
                          'outer':next((v for v in outer_values if clean(v)),None),
                          'inner_values':inner_values,'outer_values':outer_values})
        if has_month_activity(daily, cutoff):
            result.append({'code':code,'name':name,'daily':daily,'duplicate':code_counts[code] > 1})
    return result, cutoff, plan, issues, hashlib.sha256(payload).hexdigest()


def evaluate_machine(machine, cutoff, plan):
    code,name = machine['code'],machine['name']
    result = {'code':code,'name':name,'components':{},'issues':[], 'action_code':None}
    if not code or machine.get('duplicate'):
        result['issues'].append('کد خالی یا تکراری است؛ نیازمند بررسی')
        return result
    try:
        rule = rule_for(code,name)
    except ValueError as exc:
        result['issues'].append(str(exc))
        return result
    if rule is None:
        result['excluded'] = True
        return result
    events = {'inner':None,'outer':None}
    unknown = {'inner':[], 'outer':[]}
    hours = []
    for d in sorted(machine['daily'],key=lambda d:d['date']):
        md = d['date']
        if md > cutoff or md >= plan:
            continue
        mark_values = {k:d.get(k + '_values',[d[k]]) for k in events}
        marks = {k:[clean(value) for value in values] for k,values in mark_values.items()}
        inner,outer = any(text in MARKS for text in marks['inner']),any(text in MARKS for text in marks['outer'])
        if rule.get('together') and (inner or outer):
            inner = outer = True
        if inner:
            events['inner'] = events['outer'] = md
        elif outer:
            events['outer'] = md
        for k,texts in marks.items():
            for text in texts:
                if text and text not in MARKS:
                    unknown[k].append(md)
                    result['issues'].append(f"{fmt_date(md)} {'داخلی' if k == 'inner' else 'بیرونی'}: {text!r} علامت سرویس معتبر نیست")
        wh = d['hours']
        numeric = not isinstance(wh,bool) and isinstance(wh,(int,float)) and math.isfinite(wh) and wh >= 0
        if numeric:
            hours.append((md,float(wh)))
        elif clean(wh) not in {'','-'}:
            hours.append((md,None))
            result['issues'].append(f'{fmt_date(md)} کارکرد نامشخص: {wh!r}')
    fields = ['outer'] if 'calendar' in rule else ['inner','outer']
    for k in fields:
        base = events[k]
        label = 'داخلی' if k == 'inner' else 'بیرونی'
        component = {'last_service':fmt_date(base), 'state':'NEEDS_REVIEW','value':None}
        result['components'][k] = component
        if base is None:
            result['issues'].append(f'{label}: مبنای آخرین سرویس مشخص نیست')
            continue
        relevant_unknown = unknown[k] + (unknown['inner'] if k == 'outer' else [])
        if any(md >= base for md in relevant_unknown):
            continue
        if 'calendar' in rule:
            value = jalali_ordinal(*plan) - jalali_ordinal(*base)
            threshold = rule['calendar']
            alert = threshold
            unit = 'روز'
        else:
            if any(wh is None for md,wh in hours if md > base):
                continue
            value = sum(wh for md,wh in hours if md > base and wh is not None and wh > 0)
            threshold,alert = rule[k]
            unit = 'ساعت'
        component.update(value=value,threshold=threshold,unit=unit,state='DUE' if value >= threshold else 'NEAR_DUE' if value >= alert else 'OK')
    if result['components'].get('inner',{}).get('state') == 'DUE':
        result['action_code'] = BOTH
    elif result['components'].get('outer',{}).get('state') == 'DUE':
        result['action_code'] = BOTH if rule.get('together') else OUTER
    return result


def build_proposal(path=SOURCE, target=None):
    machines,cutoff,plan,issues,digest = read_source(path,target)
    evaluated = [evaluate_machine(m,cutoff,plan) for m in machines]
    selected = []
    warnings = []
    for m in evaluated:
        if m['action_code']:
            evidence = '; '.join(f"{'داخلی' if k == 'inner' else 'بیرونی'}: {c['value']:g}/{c['threshold']} {c['unit']}؛ آخرین انجام {c['last_service']}" for k,c in m['components'].items() if c['value'] is not None)
            selected.append({'machine_code':m['code'],'machine_name':m['name'],'action_code':m['action_code'],'action_text':ACTIONS[m['action_code']], 'evidence':evidence, 'components':m['components']})
        for field,comp in m['components'].items():
            if comp['state'] == 'NEAR_DUE':
                warnings.append(f"{m['code']} {'داخلی' if field == 'inner' else 'بیرونی'}: {comp['value']:g} از {comp['threshold']} {comp['unit']}")
        issues.extend({'code':m['code'],'reason':reason} for reason in m['issues'])
    return {'cutoff':fmt_date(cutoff),'plan_date':fmt_date(plan),'source_sha256':digest,
            'items':selected,'warnings':warnings,'review':issues,'machines':evaluated}

import argparse
from pathlib import Path
from openpyxl import load_workbook


OPERATIONAL_YEAR = 1405
SOURCE_SHEET = "Sheet1 (2)"

# دامپتراک
DUMP_OUTER_INTERVAL = 20
DUMP_OUTER_ALERT = 17

DUMP_INNER_INTERVAL = 100
DUMP_INNER_ALERT = 90

# بیل 231
EXCAVATOR_CODE = "231"
EXCAVATOR_INTERVAL = 10
EXCAVATOR_ALERT = 7

# ژنراتور
GENERATOR_CODE = "DG1"
GENERATOR_INTERVAL = 20
GENERATOR_ALERT = 17

# سایر خودروها
CALENDAR_INTERVAL_DAYS = 2

VALID_MARKS = {
    "صبح",
    "ظهر",
    "عصر",
    "*",
}


def clean(v):
    if v is None:
        return ""

    return " ".join(
        str(v)
        .replace("\u200c", " ")
        .replace("ي", "ی")
        .replace("ك", "ک")
        .split()
    )


def clean_code(v):
    if v is None:
        return ""

    if isinstance(v, float) and v.is_integer():
        v = int(v)

    return clean(v).upper()


def parse_sheet_date(v):
    if v is None or isinstance(v, bool):
        return None

    if isinstance(v, (int, float)):
        text = f"{float(v):.2f}"
    else:
        text = clean(v).replace("/", ".")

    parts = text.split(".")

    if len(parts) != 2:
        return None

    try:
        month = int(parts[0])
        day = int(parts[1])
    except ValueError:
        return None

    if not (1 <= month <= 12):
        return None

    max_day = 31 if month <= 6 else 30

    if month == 12:
        max_day = 29

    if not (1 <= day <= max_day):
        return None

    return month, day


def parse_target(text):
    text = (
        clean(text)
        .replace("-", "/")
        .replace(".", "/")
    )

    parts = text.split("/")

    if len(parts) == 3:
        year, month, day = map(int, parts)

    elif len(parts) == 2:
        year = OPERATIONAL_YEAR
        month, day = map(int, parts)

    else:
        raise ValueError(
            "date must be 1405/06/03 or 06/03"
        )

    if year != OPERATIONAL_YEAR:
        raise ValueError(
            f"Only year {OPERATIONAL_YEAR} supported"
        )

    return month, day


def ordinal(md):
    month, day = md

    month_days = [
        31,31,31,31,31,31,
        30,30,30,30,30,29
    ]

    return sum(
        month_days[:month - 1]
    ) + day


def fmt(md):
    if md is None:
        return "-"

    return (
        f"{OPERATIONAL_YEAR}/"
        f"{md[0]:02d}/"
        f"{md[1]:02d}"
    )


def work_hours(v):
    if isinstance(v, bool):
        return None

    if isinstance(v, (int, float)):
        v = float(v)

        if v >= 0:
            return v

    return None


def service_mark(v):
    return clean(v) in VALID_MARKS


def detect_blocks(ws):

    blocks = []

    col = 4

    while col <= ws.max_column - 2:

        md = parse_sheet_date(
            ws.cell(1, col).value
        )

        if md is None:
            col += 1
            continue

        if (
            clean(ws.cell(2, col).value)
            == "کارکرد"

            and clean(
                ws.cell(2, col + 1).value
            )
            == "درونی"

            and clean(
                ws.cell(2, col + 2).value
            )
            == "بیرونی"
        ):

            blocks.append({
                "date": md,
                "work": col,
                "inner": col + 1,
                "outer": col + 2,
            })

            col += 3

        else:
            col += 1

    unique = {}

    for b in blocks:
        unique.setdefault(
            b["date"],
            b
        )

    return sorted(
        unique.values(),
        key=lambda x: ordinal(
            x["date"]
        )
    )


def main():
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parents[3]))
    from tools.fleet.air_filter.proposal_cli import run
    run(backtest=True)


if __name__ == '__main__':
    main()

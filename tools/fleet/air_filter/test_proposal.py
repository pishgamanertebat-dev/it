import unittest
import tempfile
from pathlib import Path
from openpyxl import Workbook
from tools.fleet.air_filter.proposal import evaluate_machine, build_proposal
from tools.fleet.air_filter.rules import BOTH, OUTER


def day(md,hours=0,inner=None,outer=None):
    return dict(date=md,hours=hours,inner=inner,outer=outer)


class ProposalTests(unittest.TestCase):
    def evaluate(self, name, daily, code='710', cutoff=(6,8),plan=(6,9)):
        return evaluate_machine(dict(code=code,name=name,daily=daily),cutoff,plan)

    def test_all_hour_thresholds_and_alert_boundaries(self):
        for name,code,outer,oa,inner,ia in [('دامپتراک','710',20,17,100,90),('کامیون سهند زرد','S1',10,7,50,45),('کامیون آب پاش','TA1',10,7,50,45),('خاور','TR1',10,7,10,7)]:
            for field,due,alert in [('outer',outer,oa),('inner',inner,ia)]:
                for hours,state in [(alert-1,'OK'),(alert,'NEAR_DUE'),(due-1,'NEAR_DUE'),(due,'DUE')]:
                    with self.subTest(name=name,field=field,hours=hours):
                        r=self.evaluate(name,[day((6,1),24,inner='صبح'),day((6,8),hours)],code)
                        self.assertEqual(r['components'][field]['state'],state)

    def test_cross_month_excludes_service_day_and_future(self):
        r=self.evaluate('دامپتراک',[day((5,31),24,inner='صبح'),day((6,1),11),day((6,8),9),day((6,9),100,inner='ظهر')])
        self.assertEqual(r['components']['outer']['value'],20)
        self.assertEqual(r['action_code'],OUTER)

    def test_inner_service_resets_outer_and_due_includes_both(self):
        r=self.evaluate('دامپتراک',[day((5,1),inner='صبح'),day((5,2),100),day((6,8),outer='صبح')])
        self.assertEqual(r['action_code'],BOTH)
        r=self.evaluate('دامپتراک',[day((5,1),outer='صبح'),day((6,7),inner='شب'),day((6,8),3)])
        self.assertEqual(r['components']['outer']['value'],3)

    def test_calendar_overdue_repeats_and_no_baseline_is_visible(self):
        for name,code in [('مزدا','MZ1'),('پیکاپ ریچ','PR2')]:
            for elapsed in [2,3,5]:
                r=self.evaluate(name,[day((6,9-elapsed),outer='ظهر')],code)
                self.assertEqual(r['action_code'],OUTER)
            r=self.evaluate(name,[],code)
            self.assertTrue(r['issues'])
            self.assertIsNone(r['action_code'])

    def test_unknown_service_and_invalid_hours_are_not_events_or_zero(self):
        r=self.evaluate('دامپتراک',[day((5,8),inner='صبح'),day((5,9),4,outer='نامشخص'),day((6,8),20)])
        self.assertEqual(r['components']['outer']['last_service'],'1405/05/08')
        self.assertEqual(r['components']['outer']['state'],'NEEDS_REVIEW')
        self.assertTrue(r['issues'])
        for bad in [-1,True,float('nan'),'نامشخص']:
            r=self.evaluate('دامپتراک',[day((6,1),inner='صبح'),day((6,8),bad)])
            self.assertEqual(r['components']['outer']['state'],'NEEDS_REVIEW')

    def test_excavator_loader_bulldozer_rules(self):
        excavator=self.evaluate('بیل مکانیکی',[day((6,1),inner='صبح'),day((6,8),4)],'EX231')
        self.assertEqual(excavator['components']['outer']['state'],'NEAR_DUE')
        excavator=self.evaluate('بیل مکانیکی',[day((6,1),inner='صبح'),day((6,8),5)],'EX231')
        self.assertEqual(excavator['action_code'],OUTER)
        excavator=self.evaluate('بیل مکانیکی',[day((6,1),inner='صبح'),day((6,8),10)],'EX231')
        self.assertEqual(excavator['action_code'],BOTH)
        for name,code in [('لودر','W471'),('بلدوزر','D151')]:
            near=self.evaluate(name,[day((6,1),outer='شب'),day((6,8),7)],code)
            self.assertEqual({c['state'] for c in near['components'].values()},{'NEAR_DUE'})
            due=self.evaluate(name,[day((6,1),outer='شب'),day((6,8),10)],code)
            self.assertEqual(due['action_code'],BOTH)

    def test_new_codes_are_valid_work_order_items(self):
        from tools.fleet.work_orders.types.air_filter.builder import get_items
        actions={'EX601':OUTER,'WA601':BOTH,'W471':BOTH,'D151':BOTH,'D152':BOTH}
        items=get_items(list(actions),actions=actions)
        self.assertEqual([i['machine_name'] for i in items],
                         ['بیل مکانیکی','لودر','لودر','بلدوزر','بلدوزر'])
        bad=dict(actions); bad['W471']=OUTER
        with self.assertRaises(ValueError):
            get_items(list(bad),actions=bad)

    def test_five_column_day_night_layout_is_combined(self):
        from tools.fleet.air_filter.proposal import read_source
        root=Path('runtime/air_filter_proposal'); root.mkdir(parents=True,exist_ok=True)
        with tempfile.TemporaryDirectory(dir=root) as directory:
            path=Path(directory)/'source.xlsx'
            workbook=Workbook(); sheet=workbook.active; sheet.title='1'
            for col,value in enumerate(('ردیف','نام دستگاه','کد دستگاه'),1):
                sheet.cell(1,col,value)
            for col,field in enumerate(('کارکرد','درونی','بیرونی','درونی','بیرونی'),4):
                sheet.cell(1,col,6.01); sheet.cell(2,col,field)
            sheet.cell(3,2,'بیل مکانیکی'); sheet.cell(3,3,801)
            sheet.cell(3,4,4); sheet.cell(3,8,'شب')
            workbook.save(path); workbook.close()
            machines,cutoff,plan,issues,_=read_source(path,target=(6,2))
        self.assertEqual((cutoff,plan,issues),((6,1),(6,2),[]))
        self.assertEqual(machines[0]['code'],'EX801')
        self.assertEqual(machines[0]['daily'][0]['outer_values'],[None,'شب'])
        result=evaluate_machine(machines[0],cutoff,plan)
        self.assertEqual(result['components']['outer']['last_service'],'1405/06/01')

    def test_inner_warning_does_not_replace_outer_due(self):
        r=self.evaluate('دامپتراک',[day((5,1),inner='صبح'),day((5,2),70),day((6,1),outer='صبح'),day((6,8),20)])
        self.assertEqual(r['components']['inner']['state'],'NEAR_DUE')
        self.assertEqual(r['action_code'],OUTER)

    def test_real_source_target_and_quality(self):
        p=build_proposal()
        self.assertEqual(p['cutoff'],'1405/06/15')
        self.assertEqual(p['plan_date'],'1405/06/16')
        self.assertFalse({'S3','DG1','W471','462'} & {r['code'] for r in p['review']})
        self.assertIn('EX231',{i['machine_code'] for i in p['items']})
        codes={r['code'] for r in p['machines']}
        self.assertTrue({'EX601','WA601','D151','D152'}.issubset(codes))
        self.assertEqual(len(codes),len(p['machines']))

    def test_month_activity_ignores_old_work_and_future_work(self):
        from tools.fleet.air_filter.proposal import has_month_activity
        daily=[day((5,31),100,inner='صبح'),day((6,1),None),day((6,2),0),day((6,3),'-'),day((6,9),10)]
        self.assertFalse(has_month_activity(daily,(6,8)))
        self.assertTrue(has_month_activity(daily+[day((6,4),1)],(6,8)))
        self.assertTrue(has_month_activity(daily+[day((6,4),'نامشخص')],(6,8)))

    def test_generator_is_excluded_even_with_work(self):
        from tools.fleet.work_orders.types.air_filter.builder import get_items
        result=self.evaluate('ژنراتور',[day((6,1),inner='صبح'),day((6,8),100)],'DG1')
        self.assertTrue(result['excluded'])
        self.assertEqual(result['issues'],[])
        with self.assertRaises(ValueError):
            get_items(['DG1'])

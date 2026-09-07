import sqlite3
import openpyxl

DB = r"E:\KomatsoAI\data\fleet\db\fleet_ops.db"

conn = sqlite3.connect(DB)
cur = conn.cursor()

cur.execute("""
SELECT raw_path
FROM ingest_sources
WHERE source_type = 'service_events'
ORDER BY id DESC
LIMIT 1
""")

row = cur.fetchone()

if not row:
    raise SystemExit("No service_events found")

raw_path = row[0]

print("SERVICE HEADER INSPECTION")
print("=" * 70)
print("FILE:", raw_path)

wb = openpyxl.load_workbook(raw_path, data_only=True)


for sheet in ["بیل 101", "HD714"]:

    if sheet not in wb.sheetnames:
        print("Missing sheet:", sheet)
        continue

    ws = wb[sheet]

    print()
    print("=" * 70)
    print("SHEET:", sheet)
    print("=" * 70)

    for r in range(1, 4):

        values = []

        for c in range(1, ws.max_column + 1):

            value = ws.cell(r, c).value

            if value is not None:
                values.append(
                    f"{c-1}: {value}"
                )

        print()
        print("HEADER ROW", r)
        print("-" * 70)

        for item in values:
            print(item)


conn.close()
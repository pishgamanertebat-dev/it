"""One-file mine import: replay confirmed fault edits through the existing writer.

The edits journal is the only authority for Bale changes. Import events live in
that same journal, allowing hash-based recovery after an uncertain replacement.
No upstream write, second operational workbook, or general sync engine exists.
"""
from contextlib import closing
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import uuid

import openpyxl
from .sync_cells import local_changes, extend_changes, apply_changes, verify_changes, refresh_confirmed_changes

from .entry_service import (
    CONFIG, RUNTIME, EntryError, SECTIONS, writer_lock, source_reader, open_audit,
    recover_prepared, publish_candidate, today_sheet, layout, find_row, text_value,
    set_description, clone_day, key,
)

KIND = 'mine_fault_merge'


class SyncConflict(EntryError):
    pass


def digest(data):
    return hashlib.sha256(data).hexdigest()


def sync_schema(db):
    db.execute('''CREATE TABLE IF NOT EXISTS fault_sync_sources (
        source TEXT PRIMARY KEY, start_rowid INTEGER NOT NULL,
        initial_hash TEXT NOT NULL, upstream_hash TEXT NOT NULL, enabled_at TEXT NOT NULL)''')
    db.execute('''CREATE TABLE IF NOT EXISTS fault_sync_failures (
        id INTEGER PRIMARY KEY, source TEXT, upstream_hash TEXT,
        created TEXT, error TEXT)''')
    db.commit()


def initialize(upstream, *, config=CONFIG, runtime=RUNTIME):
    """Enable at a verified equal snapshot; historical overwritten edits are excluded."""
    settings = json.loads(Path(config).read_text(encoding='utf8'))
    source = Path(settings['source']).resolve(strict=True)
    incoming_hash = digest(Path(upstream).read_bytes())
    with writer_lock(Path(runtime)), closing(open_audit(runtime)) as db:
        sync_schema(db)
        with source_reader(source) as stream:
            current_data = stream.read()
            current_hash = digest(current_data)
            recover_prepared(db, source, current_hash)
            row = db.execute('SELECT * FROM fault_sync_sources WHERE source=?', (str(source),)).fetchone()
            if row:
                return dict(status='already_enabled', enabled_at=row['enabled_at'])
            if current_hash != incoming_hash:
                raise SyncConflict('FAULT_SYNC_INITIAL_MISMATCH: local and mine must match before activation')
            baseline = Path(runtime) / 'backups' / (current_hash + '.xlsx')
            baseline.parent.mkdir(exist_ok=True)
            if not baseline.exists():
                baseline.write_bytes(current_data)
            cutoff = db.execute('SELECT COALESCE(MAX(rowid),0) FROM edits').fetchone()[0]
            timestamp = datetime.now(timezone.utc).isoformat()
            with db:
                db.execute('INSERT INTO fault_sync_sources VALUES (?,?,?,?,?)',
                           (str(source), cutoff, current_hash, incoming_hash, timestamp))
            return dict(status='enabled', start_rowid=cutoff, enabled_at=timestamp, hash=current_hash)


def journal_state(db, source, state, runtime=RUNTIME):
    overlays = {}
    local_overlays = []
    latest_hash, upstream_hash = state['initial_hash'], state['upstream_hash']
    rows = db.execute("SELECT rowid,* FROM edits WHERE source=? AND rowid>? AND status='succeeded' ORDER BY rowid",
                      (str(source), state['start_rowid'])).fetchall()
    for row in rows:
        request = json.loads(row['request'])
        if request.get('_kind') == KIND:
            result = json.loads(row['result'])
            overlays = {(v['date'], key(v['code']), v['section']): v for v in result['overlays']}
            upstream_hash = result['upstream_hash']
            local_overlays = result.get('local_overlays', [])
        else:
            if row['before_hash'] != latest_hash:
                baseline = Path(runtime) / 'backups' / (latest_hash + '.xlsx')
                original = Path(runtime) / 'backups' / (row['before_hash'] + '.xlsx')
                if (not baseline.is_file() or digest(baseline.read_bytes()) != latest_hash
                        or not original.is_file() or digest(original.read_bytes()) != row['before_hash']):
                    raise SyncConflict('FAULT_SYNC_BASELINE_MISSING: pre-Bale local changes need review')
                local_overlays = extend_changes(local_overlays, local_changes(baseline.read_bytes(), original.read_bytes()))
            section = request.get('section')
            if section not in SECTIONS or not isinstance(request.get('description'), str):
                raise SyncConflict('FAULT_SYNC_INVALID_JOURNAL: unrecognized confirmed edit')
            field = (request['date'], key(request['code']), section)
            previous = overlays.get(field)
            if previous and request['expected'] != previous['description']:
                raise SyncConflict('FAULT_SYNC_JOURNAL_CHAIN_CONFLICT: ' + repr(field))
            overlays[field] = dict(date=request['date'], code=request['code'], name=request['name'],
                section=section, base=previous['base'] if previous else request['expected'],
                description=request['description'], operation=row['operation'], actor=row['actor'])
            if local_overlays:
                snapshot = Path(runtime) / 'backups' / (row['after_hash'] + '.xlsx')
                if not snapshot.is_file() or digest(snapshot.read_bytes()) != row['after_hash']:
                    raise SyncConflict('FAULT_SYNC_BASELINE_MISSING: confirmed edit snapshot unavailable')
                local_overlays = refresh_confirmed_changes(local_overlays, snapshot.read_bytes())
        latest_hash = row['after_hash']
    return overlays, latest_hash, upstream_hash, local_overlays


def combine(incoming, overlays, template, local_overlays=()):
    """Only journal-addressed cells change; new rows/days use the existing template."""
    book = openpyxl.load_workbook(io.BytesIO(incoming))
    try:
        remaining, applied = apply_changes(book, local_overlays)
        for field, edit in sorted(overlays.items()):
            day, code, section = field
            sheet = today_sheet(book, day)
            cols = layout(sheet) if sheet is not None else None
            row = find_row(sheet, code, cols) if sheet is not None else None
            mine = text_value(sheet.cell(row, cols[section]).value) if row else ''
            if row and (sheet.title, sheet.cell(row, cols[section]).coordinate) in applied and mine != edit['description']:
                raise SyncConflict(f'FAULT_SYNC_BALE_LOCAL_CONFLICT: date={day} code={code} section={section}; local file preserved')
            if mine not in (edit['base'], edit['description']):
                raise SyncConflict(f'FAULT_SYNC_CELL_CONFLICT: date={day} code={code} section={section}; local file preserved')
            # A matching manual entry acknowledges the local value as the new baseline.
            edit['base'] = mine
            if row:
                cell = sheet.cell(row, cols[section])
                cell.value = edit['description'] or None
                if edit['description']:
                    cell.data_type = 's'
            elif edit['description']:
                if sheet is None:
                    clone_day(book, day, template, restoring_confirmed=True)
                set_description(book, edit, section, edit['description'], day, '', template)
        output = io.BytesIO()
        book.save(output)
        candidate = output.getvalue()
    finally:
        book.close()
    # Validate the serialized workbook and every addressed field before publication.
    verified = openpyxl.load_workbook(io.BytesIO(candidate), read_only=not any(v.get('kind') == 'row_height' for v in local_overlays))
    try:
        verify_changes(verified, applied)
        for (day, code, section), edit in overlays.items():
            sheet = today_sheet(verified, day)
            cols = layout(sheet) if sheet is not None else None
            row = find_row(sheet, code, cols) if sheet is not None else None
            field_value = text_value(sheet.cell(row, cols[section]).value) if row else ''
            if field_value != edit['description']:
                raise SyncConflict('FAULT_SYNC_VALIDATION_FAILED: ' + repr((day, code, section)))
    finally:
        verified.close()
    return candidate, remaining


def merge_snapshot(staged, destination, *, config=CONFIG, runtime=RUNTIME, reconciliation=None):
    settings = json.loads(Path(config).read_text(encoding='utf8'))
    source = Path(settings['source']).resolve(strict=True)
    if Path(destination).resolve(strict=True) != source:
        raise SyncConflict('FAULT_SYNC_DESTINATION_DENIED')
    incoming = Path(staged).read_bytes()
    incoming_hash = digest(incoming)
    runtime = Path(runtime)
    with writer_lock(runtime), closing(open_audit(runtime)) as db:
        sync_schema(db)
        try:
            with source_reader(source) as stream:
                original = stream.read()
                current_hash = digest(original)
                recover_prepared(db, source, current_hash)
                state = db.execute('SELECT * FROM fault_sync_sources WHERE source=?', (str(source),)).fetchone()
                if state is None:
                    raise SyncConflict('FAULT_SYNC_NOT_INITIALIZED: publication denied')
                overlays, latest_hash, upstream_hash, local_overlays = journal_state(db, source, state, runtime)
                reconciled = False
                if current_hash != latest_hash:
                    if reconciliation is not None:
                        # One-time, hash-pinned operator recovery of the legacy
                        # deployment, which did not retain published baselines.
                        if any(reconciliation.get(k) != v for k, v in dict(local_hash=current_hash,
                                upstream_hash=incoming_hash, journal_hash=latest_hash).items()):
                            raise SyncConflict('FAULT_SYNC_RECONCILIATION_CHANGED: review snapshots again')
                        changes = local_changes(incoming, original, incoming_superset=True)
                        if changes != reconciliation.get('changes'):
                            raise SyncConflict('FAULT_SYNC_RECONCILIATION_CHANGED: cell review differs')
                    else:
                        baseline = runtime / 'backups' / (latest_hash + '.xlsx')
                        if not baseline.is_file() or digest(baseline.read_bytes()) != latest_hash:
                            raise SyncConflict('FAULT_SYNC_BASELINE_MISSING: reviewed reconciliation required; local file preserved')
                        changes = local_changes(baseline.read_bytes(), original)
                    local_overlays = extend_changes(local_overlays, changes)
                    reconciled = True
                if incoming_hash == upstream_hash and not reconciled:
                    return dict(status='unchanged', upstream_hash=incoming_hash, protected_cells=len(overlays),
                                protected_local_cells=len(local_overlays))
                candidate, local_overlays = combine(incoming, overlays, settings['template'], local_overlays)
                result = dict(status='merged', upstream_hash=incoming_hash, overlays=list(overlays.values()),
                              local_overlays=local_overlays, protected_local_cells=len(local_overlays),
                              reconciled_local_change=reconciled, protected_cells=len(overlays))
                request = dict(_kind=KIND, upstream_hash=incoming_hash, operation=uuid.uuid4().hex)
                if reconciliation is not None:
                    request['reviewed_reconciliation'] = reconciliation
                return publish_candidate(source, original, candidate, db, 'mine-file-sync', request,
                                         result, runtime=runtime)
        except Exception as exc:
            with db:
                db.execute('INSERT INTO fault_sync_failures(source,upstream_hash,created,error) VALUES (?,?,?,?)',
                           (str(source), incoming_hash, datetime.now(timezone.utc).isoformat(), str(exc)))
            raise


def main():
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--staged', required=True, type=Path)
    parser.add_argument('--destination', type=Path)
    parser.add_argument('--initialize', action='store_true')
    args = parser.parse_args()
    try:
        if args.initialize:
            result = initialize(args.staged)
        else:
            if args.destination is None:
                parser.error('--destination is required for publication')
            result = merge_snapshot(args.staged, args.destination)
        print(json.dumps({k: v for k, v in result.items() if k not in {'overlays', 'local_overlays'}}, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps(dict(status='error', error=str(exc)), ensure_ascii=False))
        raise SystemExit(1)


if __name__ == '__main__':
    main()

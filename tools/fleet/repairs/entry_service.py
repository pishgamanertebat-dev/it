"""Confirmed daily-report edits, with source locking, backups and a durable journal."""
from __future__ import annotations

from contextlib import contextmanager
from copy import copy, deepcopy
from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import re
import sqlite3
import tempfile
import time
import uuid

import openpyxl

from tools.fleet.overflow.report import normalize, validate_date
from tools.fleet.report_caption import jalali_today
from .report import ROOT, SECTIONS

CONFIG = ROOT / 'settings/repairs_entry.json'
RUNTIME = ROOT / 'runtime/repairs-entry'
FLEET_DB = ROOT / 'data/fleet/db/fleet_ops.db'


class EntryError(ValueError):
    pass


def permission(actor, config=CONFIG):
    settings = json.loads(Path(config).read_text(encoding='utf-8'))
    if not actor or str(actor) not in settings['allowed_users']:
        raise PermissionError('اجازهٔ ثبت شرح خرابی را ندارید.')
    if not settings.get('template'):
        raise EntryError('قالب خالی گزارش روزانه تنظیم نشده است.')
    return settings


def key(value):
    return normalize(value).replace('آ', 'ا').replace(' ', '').upper()


def layout(sheet):
    labels = {'number': 'ردیف', 'name': 'نوع دستگاه', 'code': 'کد جدید', **SECTIONS}
    cells = list(sheet[2])
    result = {}
    for field, label in labels.items():
        matches = [c.column for c in cells if key(c.value) == key(label)]
        if len(matches) != 1:
            raise EntryError('ستون‌های گزارش شناخته نشدند؛ فایل باید بررسی شود.')
        result[field] = matches[0]
    return result


def dated_sheets(book):
    result = []
    for sheet in book.worksheets:
        found = set()
        for cell in sheet[1]:
            for raw in re.findall(r'(?<!\d)1[34]\d{2}[/.-]\d{1,2}[/.-]\d{1,2}(?!\d)', normalize(cell.value)):
                found.add(validate_date(raw))
        if len(found) == 1:
            result.append((found.pop(), sheet))
        elif found:
            raise EntryError('تاریخ سربرگ یکی از شیت‌ها مبهم است.')
    if not result:
        raise EntryError('شیت روزانهٔ معتبری در فایل پیدا نشد.')
    return sorted(result, key=lambda item: item[0], reverse=True)


def today_sheet(book, day):
    matches = [s for d, s in dated_sheets(book) if d == day]
    if len(matches) > 1:
        raise EntryError('چند شیت برای امروز وجود دارد؛ ابتدا فایل باید اصلاح شود.')
    return matches[0] if matches else None


def machine(book, entered, fleet_db=FLEET_DB):
    entered = key(entered)
    if not entered or len(entered) > 30:
        raise EntryError('کد دستگاه معتبر نیست.')
    known = {}
    for _, sheet in dated_sheets(book):
        try:
            cols = layout(sheet)
        except EntryError:
            continue
        for row in sheet.iter_rows(min_row=3, values_only=True):
            code = row[cols['code'] - 1]
            name = row[cols['name'] - 1]
            if code is not None and name and key(code) not in {'-', 'جمع'}:
                known.setdefault(key(code), {'code': str(code).strip(), 'name': str(name).strip()})
    if entered in known:
        return known[entered]
    candidates = {code: value for code, value in known.items()
                  if entered.isdecimal() and re.sub(r'^[A-Z]+', '', code) == entered}
    if Path(fleet_db).exists():
        conn = sqlite3.connect(Path(fleet_db).resolve().as_uri() + '?mode=ro', uri=True)
        try:
            rows = conn.execute('''SELECT DISTINCT m.canonical_code, m.machine_type_hint, m.model_key
                FROM machines m LEFT JOIN machine_aliases a ON a.machine_id=m.id AND a.verified=1
                WHERE UPPER(m.canonical_code)=? OR UPPER(a.alias_code)=?''', (entered, entered)).fetchall()
            for code, name, model in rows:
                if name or model:
                    candidates.setdefault(key(code), known.get(key(code), {'code': code, 'name': name or model}))
        finally:
            conn.close()
    if len(candidates) == 1:
        return next(iter(candidates.values()))
    if candidates:
        raise EntryError('کد کوتاه مبهم است؛ کد کامل دستگاه را وارد کنید: ' + '، '.join(sorted(candidates)))
    raise EntryError('کد دستگاه در فهرست موجود نیست؛ کد کامل و دقیق دستگاه را وارد کنید.')


def find_row(sheet, code, cols):
    rows = [r for r in range(3, sheet.max_row + 1) if key(sheet.cell(r, cols['code']).value) == key(code)]
    if len(rows) > 1:
        raise EntryError('این دستگاه در شیت امروز تکراری است؛ ابتدا فایل باید بررسی شود.')
    return rows[0] if rows else None


def text_value(value):
    return '' if value is None else str(value)


def preview(actor, code, section, *, config=CONFIG, day=None, fleet_db=FLEET_DB):
    settings = permission(actor, config)
    if section not in SECTIONS:
        raise EntryError('بخش انتخاب‌شده معتبر نیست.')
    day = day or jalali_today()
    book = openpyxl.load_workbook(settings['source'])
    try:
        selected = machine(book, code, fleet_db)
        sheet = today_sheet(book, day)
        current = ''
        if sheet:
            cols = layout(sheet)
            row = find_row(sheet, selected['code'], cols)
            if row:
                current = text_value(sheet.cell(row, cols[section]).value)
        return {**selected, 'date': day, 'section': section, 'expected': current}
    finally:
        book.close()


@contextmanager
def source_reader(path):
    """Deny other writers while allowing our atomic replacement on Windows."""
    if os.name == 'nt':
        import ctypes
        import msvcrt
        from ctypes import wintypes
        kernel = ctypes.WinDLL('kernel32', use_last_error=True)
        kernel.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                      wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        kernel.CreateFileW.restype = wintypes.HANDLE
        handle = kernel.CreateFileW(str(path), 0x80000000, 1 | 4, None, 3, 0, None)
        if handle == wintypes.HANDLE(-1).value:
            raise EntryError('فایل اکسل باز یا در حال ویرایش است؛ آن را ببندید و دوباره تأیید کنید.')
        fd = msvcrt.open_osfhandle(handle, os.O_RDONLY | os.O_BINARY)
        with os.fdopen(fd, 'rb') as stream:
            yield stream
    else:
        with Path(path).open('rb') as stream:
            yield stream


@contextmanager
def writer_lock(runtime):
    runtime.mkdir(parents=True, exist_ok=True)
    with (runtime / 'writer.lock').open('a+b') as stream:
        stream.seek(0)
        if not stream.read(1):
            stream.write(b'0')
            stream.flush()
        if os.name == 'nt':
            import msvcrt
            deadline = time.monotonic() + 30
            while True:
                stream.seek(0)
                try:
                    msvcrt.locking(stream.fileno(), msvcrt.LK_NBLCK, 1)
                    break
                except OSError:
                    if time.monotonic() >= deadline:
                        raise EntryError('ثبت دیگری در حال انجام است؛ کمی بعد دوباره تأیید کنید.')
                    time.sleep(.1)
            try:
                yield
            finally:
                stream.seek(0)
                msvcrt.locking(stream.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl
            fcntl.flock(stream, fcntl.LOCK_EX)
            try:
                yield
            finally:
                fcntl.flock(stream, fcntl.LOCK_UN)


def replace_source(temporary, source):
    if os.name == 'nt':
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL('kernel32', use_last_error=True)
        kernel.ReplaceFileW.argtypes = [wintypes.LPCWSTR, wintypes.LPCWSTR, wintypes.LPCWSTR,
                                        wintypes.DWORD, wintypes.LPVOID, wintypes.LPVOID]
        kernel.ReplaceFileW.restype = wintypes.BOOL
        if not kernel.ReplaceFileW(str(source), str(temporary), None, 0, None, None):
            raise ctypes.WinError(ctypes.get_last_error())
    else:
        os.replace(temporary, source)


def copy_style(source, target):
    for attribute in ('font', 'fill', 'border', 'alignment', 'protection'):
        setattr(target, attribute, copy(getattr(source, attribute)))
    target.number_format = source.number_format


def wrapped_lines(text):
    text = text_value(text)
    if not text.strip():
        return 0
    return sum(max(1, (len(line) + 69) // 70) for line in text.splitlines())


def fitted_height(texts, default=20):
    lines = max((wrapped_lines(text) for text in texts), default=0)
    return default if not lines else min(409, max(default, lines * 18 + 8))


def apply_template_row_style(sheet, row, template):
    """Copy cell and row style from the blank form row, never from a previous data row."""
    for col in range(1, max(sheet.max_column, template.max_column) + 1):
        copy_style(template.cell(3, col), sheet.cell(row, col))
    source, target = template.row_dimensions[3], sheet.row_dimensions[row]
    target.height, target.hidden, target.outlineLevel = source.height, source.hidden, source.outlineLevel
    copy_style(source, target)


def machine_rows(sheet, cols):
    return [r for r in range(3, sheet.max_row + 1) if sheet.cell(r, cols['code']).value is not None]


def row_texts(sheet, row, cols):
    return [text_value(sheet.cell(row, cols[name]).value) for name in SECTIONS]


def copy_template(template, destination, *, day=None):
    """Copy the fixed blank form across workbooks without reusing style IDs."""
    for row in template.iter_rows():
        for source in row:
            target = destination.cell(source.row, source.column)
            target.value = source.value.replace('{date}', day) if day and isinstance(source.value, str) else source.value
            copy_style(source, target)
    for row, dimension in template.row_dimensions.items():
        target = destination.row_dimensions[row]
        target.height, target.hidden, target.outlineLevel = dimension.height, dimension.hidden, dimension.outlineLevel
        copy_style(dimension, target)
    for col, dimension in template.column_dimensions.items():
        target = destination.column_dimensions[col]
        target.width, target.hidden, target.outlineLevel = dimension.width, dimension.hidden, dimension.outlineLevel
        target.min, target.max = dimension.min, dimension.max
        copy_style(dimension, target)
    for merged in template.merged_cells.ranges:
        destination.merge_cells(str(merged))
    for attr in ('sheet_format', 'sheet_properties', 'page_margins', 'page_setup', 'print_options', 'HeaderFooter', 'views'):
        setattr(destination, attr, deepcopy(getattr(template, attr)))
    destination.print_area = str(template.print_area).split('!', 1)[-1] if template.print_area else 'A1:E3'
    destination.print_title_rows = template.print_title_rows
    destination.print_title_cols = template.print_title_cols


def create_blank_template(source, destination):
    """One-time setup: retain headers and one formatted EMPTY entry row."""
    destination = Path(destination)
    if destination.exists():
        raise FileExistsError('Template already exists')
    source_book = openpyxl.load_workbook(source)
    output = openpyxl.Workbook()
    try:
        latest = dated_sheets(source_book)[0][1]
        layout(latest)
        # Isolate only the form before copying to the dedicated workbook.
        if latest.max_row > 3:
            latest.delete_rows(4, latest.max_row - 3)
        for cell in latest[3]:
            cell.value = None
        for cell in latest[1]:
            if cell.value is not None:
                cell.value = re.sub(r'1[34]\d{2}[/.-]\d{1,2}[/.-]\d{1,2}', '{date}', normalize(cell.value))
        latest.print_area = f'A1:{openpyxl.utils.get_column_letter(latest.max_column)}3'
        copy_template(latest, output.active)
        output.active.title = 'Daily report template'
        destination.parent.mkdir(parents=True, exist_ok=True)
        output.save(destination)
    finally:
        source_book.close()
        output.close()


def clone_day(book, day, template_path):
    if dated_sheets(book)[0][0] > day:
        raise EntryError('فایل شامل تاریخی بعد از امروز است؛ ابتدا تاریخ‌ها بررسی شوند.')
    title = 'گزارش روزانه ' + day.replace('/', '.')
    if title in book.sheetnames:
        raise EntryError('نام شیت امروز موجود است ولی تاریخ سربرگ آن تطابق ندارد.')
    template_book = openpyxl.load_workbook(template_path)
    try:
        template = template_book.active
        layout(template)
        if any(cell.value is not None for row in template.iter_rows(min_row=3) for cell in row):
            raise EntryError('قالب گزارش باید خالی از دستگاه و شرح خرابی باشد.')
        sheet = book.create_sheet(title, 0)
        copy_template(template, sheet, day=day)
    finally:
        template_book.close()
    return sheet


def set_description(book, selected, section, description, day, expected, template_path):
    result = {'date': day, 'sheet': '', 'code': selected['code'], 'name': selected['name'], 'section': section}
    clearing = not description
    sheet = today_sheet(book, day)
    if sheet is None and clearing:
        return result
    sheet = sheet or clone_day(book, day, template_path)
    result['sheet'] = sheet.title
    cols = layout(sheet)
    row = find_row(sheet, selected['code'], cols)
    old = text_value(sheet.cell(row, cols[section]).value) if row else ''
    if old != expected:
        raise EntryError('شرح این دستگاه پس از نمایش شما تغییر کرده است؛ «شرح خرابی» را دوباره باز کنید و متن جدید را بررسی کنید.')
    if row is None and clearing:
        return result
    template_book = openpyxl.load_workbook(template_path)
    try:
        template = template_book.active
        default_height = template.row_dimensions[3].height or 20
        if row is None:
            occupied = machine_rows(sheet, cols)
            row = max(occupied) + 1 if occupied else 3
            apply_template_row_style(sheet, row, template)
            sheet.cell(row, cols['number'], row - 2)
            sheet.cell(row, cols['code'], selected['code']).data_type = 's'
            sheet.cell(row, cols['name'], selected['name']).data_type = 's'
            sheet.print_area = f'A1:{openpyxl.utils.get_column_letter(max(cols.values()))}{row}'
        cell = sheet.cell(row, cols[section])
        cell.value = description or None
        if description:
            cell.data_type = 's'  # Descriptions beginning with '=' remain literal text.
        remaining = row_texts(sheet, row, cols)
        if not any(text.strip() for text in remaining):
            sheet.delete_rows(row)
            occupied = machine_rows(sheet, cols)
            last = occupied[-1] if occupied else 3
            if occupied:
                for index, current in enumerate(occupied, 1):
                    sheet.cell(current, cols['number'], index)
                    sheet.row_dimensions[current].height = fitted_height(row_texts(sheet, current, cols), default_height)
            else:
                apply_template_row_style(sheet, 3, template)
            for extra in [idx for idx in list(sheet.row_dimensions) if idx > last]:
                del sheet.row_dimensions[extra]
            sheet.print_area = f'A1:{openpyxl.utils.get_column_letter(max(cols.values()))}{last}'
        else:
            cell.alignment = copy(cell.alignment)
            cell.alignment = openpyxl.styles.Alignment(horizontal=cell.alignment.horizontal or 'right',
                vertical='center', wrap_text=True, readingOrder=2)
            sheet.row_dimensions[row].height = fitted_height(remaining, default_height)
    finally:
        template_book.close()
    return result


def commit(actor, request, *, config=CONFIG, runtime=RUNTIME, day=None, fleet_db=FLEET_DB):
    settings = permission(actor, config)
    operation = str(request.get('operation', ''))
    if not re.fullmatch(r'[a-f0-9]{32}', operation):
        raise EntryError('شناسهٔ تأیید معتبر نیست.')
    section, description = request.get('section'), request.get('description')
    if section not in SECTIONS or not isinstance(description, str) or len(description) > 1800:
        raise EntryError('شرح باید حداکثر ۱۸۰۰ نویسه باشد.')
    def mutate(book):
        selected = machine(book, request['code'], fleet_db)
        return set_description(book, selected, section, description, request['date'],
                               request['expected'], settings['template'])
    return commit_workbook(actor, request, settings, mutate, runtime=runtime, day=day)


def commit_workbook(actor, request, settings, mutate, *, runtime, day=None, command='شرح خرابی'):
    """Journal and atomically replace an existing workbook for a confirmed operation."""
    operation = str(request.get('operation', ''))
    if not re.fullmatch(r'[a-f0-9]{32}', operation):
        raise EntryError('شناسهٔ تأیید معتبر نیست.')
    source = Path(settings['source']).resolve(strict=True)
    runtime = Path(runtime)
    fingerprint = hashlib.sha256(json.dumps([str(actor), str(source), request], sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    with writer_lock(runtime):
        db = sqlite3.connect(runtime / 'audit.sqlite3')
        db.row_factory = sqlite3.Row
        try:
            db.execute('''CREATE TABLE IF NOT EXISTS edits (operation TEXT PRIMARY KEY, fingerprint TEXT,
                actor TEXT, source TEXT, request TEXT, before_hash TEXT, after_hash TEXT, backup TEXT,
                status TEXT, result TEXT, created TEXT)''')
            with source_reader(source) as locked:
                original = locked.read()
                digest = hashlib.sha256(original).hexdigest()
                for pending in db.execute("SELECT * FROM edits WHERE status='prepared' AND source=?", (str(source),)).fetchall():
                    if digest == pending['after_hash']:
                        status = 'succeeded'
                    elif digest == pending['before_hash']:
                        status = 'not_applied'
                    else:
                        raise EntryError('نتیجهٔ ثبت قبلی نیاز به بررسی دارد؛ برای جلوگیری از بازنویسی، ثبت متوقف شد.')
                    db.execute('UPDATE edits SET status=? WHERE operation=?', (status, pending['operation']))
                db.commit()
                previous = db.execute('SELECT * FROM edits WHERE operation=?', (operation,)).fetchone()
                if previous:
                    if previous['fingerprint'] != fingerprint:
                        raise EntryError('اطلاعات تأیید تغییر کرده است؛ فرم را دوباره باز کنید.')
                    if previous['status'] == 'succeeded':
                        return json.loads(previous['result'])
                if request.get('date') != (day or jalali_today()):
                    raise EntryError(f'روز عوض شده است؛ «{command}» را دوباره برای تاریخ امروز باز کنید.')
                book = openpyxl.load_workbook(io.BytesIO(original))
                try:
                    result = mutate(book)
                    output = io.BytesIO()
                    book.save(output)
                    candidate = output.getvalue()
                finally:
                    book.close()
                backup = runtime / 'backups' / (digest + '.xlsx')
                backup.parent.mkdir(exist_ok=True)
                if not backup.exists():
                    backup.write_bytes(original)
                after = hashlib.sha256(candidate).hexdigest()
                with db:
                    db.execute('INSERT OR REPLACE INTO edits VALUES (?,?,?,?,?,?,?,?,?,?,?)',
                        (operation, fingerprint, str(actor), str(source), json.dumps(request, ensure_ascii=False),
                         digest, after, str(backup), 'prepared', json.dumps(result, ensure_ascii=False),
                         datetime.now(timezone.utc).isoformat()))
                temporary = source.with_name('.repairs-' + uuid.uuid4().hex + '.xlsx')
                try:
                    with temporary.open('xb') as stream:
                        stream.write(candidate)
                        stream.flush()
                        os.fsync(stream.fileno())
                    if hashlib.sha256(source.read_bytes()).hexdigest() != digest:
                        raise EntryError('فایل هم‌زمان تغییر کرد؛ فرم را دوباره باز کنید.')
                    replace_source(temporary, source)
                except OSError as exc:
                    raise EntryError('ذخیرهٔ اکسل انجام نشد؛ فایل را ببندید و دوباره تأیید کنید.') from exc
                finally:
                    temporary.unlink(missing_ok=True)
                with db:
                    db.execute("UPDATE edits SET status='succeeded' WHERE operation=?", (operation,))
                return result
        finally:
            db.close()


def main():
    import sys
    request = json.load(sys.stdin)
    try:
        if request['action'] == 'preview':
            result = preview(request['actor'], request['code'], request['section'])
        elif request['action'] == 'commit':
            result = commit(request['actor'], request['request'])
        else:
            raise EntryError('درخواست معتبر نیست.')
        response = {'ok': True, 'result': result}
    except (EntryError, PermissionError) as exc:
        response = {'ok': False, 'message': str(exc)}
    except Exception:
        import traceback
        traceback.print_exc(file=sys.stderr)
        response = {'ok': False, 'message': 'دسترسی یا ذخیرهٔ فایل ناموفق بود؛ دوباره تلاش کنید. اطلاعاتی تأیید نشده است.'}
    print(json.dumps(response, ensure_ascii=False))


if __name__ == '__main__':
    main()

"""Exercise every live workbook sheet on a copy; never write the source."""
import hashlib
import json
from pathlib import Path
import shutil
import uuid

import openpyxl

from .maintenance_service import CONFIG, ROOT, machines, layout, append_record, preview, commit
from tools.fleet.report_caption import jalali_today


def snapshot(book):
    return {s.title: {c.coordinate: (c.value, c.data_type, str(c._style))
                     for row in s for c in row if c.value is not None} for s in book}


def main():
    source = Path(json.loads(CONFIG.read_text(encoding='utf-8'))['source'])
    original = source.read_bytes()
    directory = ROOT/'artifacts/maintenance-entry-preview'/uuid.uuid4().hex
    directory.mkdir(parents=True)
    target = directory/source.name
    shutil.copy2(source, target)
    book = openpyxl.load_workbook(target)
    before = snapshot(book)
    titles = book.sheetnames[:]
    results = []
    for selected in machines(book):
        request = {**selected, 'date': jalali_today(), 'mechanic': 'تعمیرکار آزمایشی',
                   'description': 'ثبت آزمایشی روی کپی\nخط دوم شرح کار', 'parts': 'قطعه آزمایشی'}
        result = append_record(book, request)
        _, columns = layout(book[selected['sheet']])
        results.append((result, columns))
    book.save(target)
    book.close()
    book = openpyxl.load_workbook(target)
    after = snapshot(book)
    assert titles == book.sheetnames
    for title, cells in before.items():
        for coordinate, value in cells.items():
            assert after[title][coordinate] == value, (title, coordinate)
    for result, columns in results:
        sheet = book[result['sheet']]
        assert sheet.cell(result['row'], columns['description']).value == 'ثبت آزمایشی روی کپی\nخط دوم شرح کار'
        assert sheet.row_dimensions[result['row']].height > 24
    book.close()
    config = directory/'config.json'
    config.write_text(json.dumps({'source': str(target), 'allowed_users': ['455740857', '654806764']}))
    request = {**preview('654806764', 'HD710', config=config), 'mechanic': 'تعمیرکار آزمایشی',
               'description': 'آزمون ذخیرهٔ اتمیک', 'parts': 'مصرف نشده', 'operation': uuid.uuid4().hex}
    saved = commit('654806764', request, config=config, runtime=directory/'journal')
    content = target.read_bytes()
    assert commit('654806764', request, config=config, runtime=directory/'journal') == saved
    assert target.read_bytes() == content
    assert source.read_bytes() == original, 'Live workbook changed during rehearsal'
    report = {'sheets_checked': len(results), 'source_unchanged': True,
              'source_sha256': hashlib.sha256(original).hexdigest(), 'output': str(target),
              'historical_values_and_styles_preserved': True, 'duplicate_confirmation_safe': True}
    (directory/'result.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False))


if __name__ == '__main__':
    main()

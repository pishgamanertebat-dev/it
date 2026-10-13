"""Conservative cell-value reconciliation for the single driver workbook.

Sheet/row structure and non-cell assets are never guessed. Blank-string rewrites
are treated as Excel serialization, while formulas and nonblank values are exact.
"""
from copy import copy
from datetime import date, datetime, time
import io
import zipfile

import openpyxl

from .entry_service import EntryError, layout, key


def value(cell):
    raw = cell.value
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return dict(value=None, type='n')
    if isinstance(raw, (datetime, date, time)):
        return dict(value=raw.isoformat(), type=cell.data_type, kind=type(raw).__name__)
    if not isinstance(raw, (str, int, float, bool)):
        raise EntryError('FAULT_SYNC_UNSUPPORTED_CELL_VALUE: ' + cell.coordinate)
    return dict(value=raw, type=cell.data_type)


def assign(cell, payload):
    raw = payload['value']
    if payload.get('kind'):
        raw = {'datetime': datetime, 'date': date, 'time': time}[payload['kind']].fromisoformat(raw)
    cell.value = raw
    cell.data_type = payload['type']


def styles(sheet):
    result = {}
    for cell in sheet._cells.values():
        if cell.style_id not in result:
            result[cell.style_id] = tuple(copy(getattr(cell, prop)) for prop in
                ('font', 'fill', 'border', 'alignment', 'protection', 'number_format'))
    return result


def assets(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()
                if not (name.startswith('docProps/') or
                    (name.startswith('xl/worksheets/') and '/_rels/' not in name) or name in
                    {'xl/workbook.xml', 'xl/styles.xml', 'xl/sharedStrings.xml', 'xl/calcChain.xml'})}


def dimensions(sheet, prop):
    if prop == 'row_dimensions':
        return {k: dict(v) for k, v in sheet.row_dimensions.items()}
    # Excel may split one identical H:XFD column group into H:H and I:XFD.
    # Compare effective contiguous ranges, not this serialization difference.
    ranges = []
    for dim in sorted(sheet.column_dimensions.values(), key=lambda d: d.min or openpyxl.utils.column_index_from_string(d.index)):
        attributes = dict(dim)
        start = int(attributes.pop('min', None) or openpyxl.utils.column_index_from_string(dim.index))
        end = int(attributes.pop('max', None) or start)
        if ranges and ranges[-1][1] + 1 == start and ranges[-1][2] == attributes:
            ranges[-1] = (ranges[-1][0], end, attributes)
        else:
            ranges.append((start, end, attributes))
    return ranges


def property_value(sheet, prop):
    if prop == 'page_setup':
        attributes = dict(sheet.page_setup)
        for dpi in ('verticalDpi', 'horizontalDpi'):
            if attributes.get(dpi) == '0':
                attributes.pop(dpi)
        return attributes
    return str(getattr(sheet, prop))


def local_changes(baseline, current, *, incoming_superset=False):
    # An appended payload is not an ordinary Excel save and must not be ignored.
    end = current.rfind(b'PK\x05\x06')
    if end < 0 or end + 22 + int.from_bytes(current[end+20:end+22], 'little') != len(current):
        raise EntryError('FAULT_SYNC_UNJOURNALED_LOCAL_CHANGE: unsupported trailing data')
    books = []
    try:
        books = [openpyxl.load_workbook(io.BytesIO(data)) for data in (baseline, current)]
        before, after = books
        if (not incoming_superset and before.sheetnames != after.sheetnames) or any(
                name not in before.sheetnames for name in after.sheetnames):
            raise EntryError('FAULT_SYNC_LOCAL_STRUCTURE_CONFLICT: sheet addition/removal/reorder')
        if not incoming_superset and assets(baseline) != assets(current):
            raise EntryError('FAULT_SYNC_LOCAL_STRUCTURE_CONFLICT: workbook assets changed')
        if not incoming_superset and before.defined_names != after.defined_names:
            raise EntryError('FAULT_SYNC_LOCAL_STRUCTURE_CONFLICT: defined names changed')
        changes = []
        for sheet in after:
            old = before[sheet.title]
            try:
                code_col = layout(old)['code']
            except EntryError:
                code_col = None
            for prop in ('merged_cells', 'print_area', 'print_title_rows', 'print_title_cols',
                         'freeze_panes', 'sheet_format', 'sheet_properties', 'sheet_state',
                         'page_setup', 'page_margins', 'print_options', 'protection'):
                if property_value(old, prop) != property_value(sheet, prop):
                    raise EntryError(f'FAULT_SYNC_LOCAL_STRUCTURE_CONFLICT: sheet={sheet.title} property={prop}')
            if dimensions(old, 'column_dimensions') != dimensions(sheet, 'column_dimensions'):
                raise EntryError(f'FAULT_SYNC_LOCAL_STRUCTURE_CONFLICT: sheet={sheet.title} property=column_dimensions')
            for row in set(old.row_dimensions) | set(sheet.row_dimensions):
                a, b = old.row_dimensions.get(row), sheet.row_dimensions.get(row)
                attributes = []
                for dim in (a, b):
                    attrs = dict(dim) if dim else {}
                    attrs.pop('ht', None)
                    attrs.pop('customHeight', None)
                    attributes.append(attrs)
                if attributes[0] != attributes[1]:
                    raise EntryError(f'FAULT_SYNC_LOCAL_STRUCTURE_CONFLICT: sheet={sheet.title} row={row}')
                if (a.height if a else None) != (b.height if b else None):
                    edit = dict(sheet=sheet.title, cell=f'@row{row}', kind='row_height', row=row,
                                base=a.height if a else None, local=b.height if b else None)
                    if code_col and row >= 3:
                        edit['anchor'] = dict(column=code_col, code=key(old.cell(row, code_col).value))
                    changes.append(edit)
            old_styles, new_styles = styles(old), styles(sheet)
            for coordinate in set(old._cells) | set(sheet._cells):
                a, b = old.cell(*coordinate), sheet.cell(*coordinate)
                for cell, cache in ((a, old_styles), (b, new_styles)):
                    if cell.style_id not in cache:
                        cache[cell.style_id] = tuple(copy(getattr(cell, prop)) for prop in
                            ('font', 'fill', 'border', 'alignment', 'protection', 'number_format'))
                if (old_styles.get(a.style_id) != new_styles.get(b.style_id)
                        or a.comment != b.comment or a.hyperlink != b.hyperlink):
                    raise EntryError(f'FAULT_SYNC_LOCAL_STRUCTURE_CONFLICT: sheet={sheet.title} cell={b.coordinate} style/comment/link')
                if value(a) == value(b):
                    continue
                edit = dict(sheet=sheet.title, cell=b.coordinate, base=value(a), local=value(b))
                if code_col and b.row >= 3:
                    edit['anchor'] = dict(column=code_col, code=key(old.cell(b.row, code_col).value))
                changes.append(edit)
        return sorted(changes, key=lambda v: (v['sheet'], v['cell']))
    finally:
        for book in books:
            book.close()


def extend_changes(previous, changes):
    merged = {(v['sheet'], v['cell']): dict(v) for v in previous}
    for edit in changes:
        address = edit['sheet'], edit['cell']
        if address in merged:
            if edit['base'] != merged[address]['local']:
                raise EntryError('FAULT_SYNC_LOCAL_JOURNAL_CONFLICT: ' + repr(address))
            merged[address]['local'] = edit['local']
        else:
            merged[address] = edit
    return list(merged.values())


def apply_changes(book, changes):
    pending, applied = [], {}
    for edit in changes:
        name, address = edit['sheet'], edit['cell']
        if name not in book.sheetnames:
            raise EntryError(f'FAULT_SYNC_LOCAL_CELL_CONFLICT: missing sheet={name} cell={address}')
        if edit.get('kind') == 'row_height':
            anchor = edit.get('anchor')
            if anchor and key(book[name].cell(edit['row'], anchor['column']).value) != anchor['code']:
                raise EntryError(f'FAULT_SYNC_LOCAL_CELL_CONFLICT: row identity changed sheet={name} row={edit["row"]}')
            row = book[name].row_dimensions.get(edit['row'])
            mine = row.height if row else None
            if mine not in (edit['base'], edit['local']):
                raise EntryError(f'FAULT_SYNC_LOCAL_CELL_CONFLICT: sheet={name} row={edit["row"]} height')
            if mine != edit['local']:
                book[name].row_dimensions[edit['row']].height = edit['local']
                pending.append(edit)
            applied[(name, address)] = edit['local']
            continue
        cell = book[name][address]
        anchor = edit.get('anchor')
        if anchor and key(book[name].cell(cell.row, anchor['column']).value) != anchor['code']:
            raise EntryError(f'FAULT_SYNC_LOCAL_CELL_CONFLICT: row identity changed sheet={name} cell={address}')
        mine = value(cell)
        if mine not in (edit['base'], edit['local']):
            raise EntryError(f'FAULT_SYNC_LOCAL_CELL_CONFLICT: sheet={name} cell={address} '
                             f'base={edit["base"]!r} local={edit["local"]!r} mine={mine!r}; local file preserved')
        if mine != edit['local']:
            assign(cell, edit['local'])
            pending.append(edit)
        applied[(name, address)] = edit['local']
    return pending, applied


def verify_changes(book, applied):
    for (name, address), expected in applied.items():
        if address.startswith('@row'):
            row = book[name].row_dimensions.get(int(address[4:]))
            actual = row.height if row else None
        else:
            actual = value(book[name][address])
        if actual != expected:
            raise EntryError(f'FAULT_SYNC_LOCAL_VALIDATION_FAILED: sheet={name} cell={address}')


def refresh_confirmed_changes(previous, snapshot):
    """A later confirmed Bale write may explicitly update a protected value."""
    book = openpyxl.load_workbook(io.BytesIO(snapshot))
    try:
        refreshed = []
        for entry in previous:
            edit = dict(entry)
            if edit['sheet'] not in book.sheetnames:
                raise EntryError('FAULT_SYNC_LOCAL_JOURNAL_CONFLICT: confirmed edit removed protected sheet')
            sheet = book[edit['sheet']]
            if edit.get('kind') == 'row_height':
                dim = sheet.row_dimensions.get(edit['row'])
                edit['local'] = dim.height if dim else None
            else:
                edit['local'] = value(sheet[edit['cell']])
            refreshed.append(edit)
        return refreshed
    finally:
        book.close()

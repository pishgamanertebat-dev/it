from datetime import date
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

from tools.fleet.report_caption import report_caption, jalali_today
from tools.scheduler.tasks import repairs, BaleSender, validate_repairs
from .report import ROOT, latest_sheet, section_columns, build_report


class RepairsTests(unittest.TestCase):
    def setUp(self):
        runtime = ROOT / 'runtime/repairs-tests'
        runtime.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=runtime)
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / 'source.xlsx'

    def workbook(self, items):
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, header in items:
            sheet = book.create_sheet(name)
            sheet.append(['گزارش روزانه', header])
        book.save(self.source)
        book.close()

    def test_latest_by_header_not_position_or_tab_name(self):
        self.workbook([('1405.06.17', 'تاریخ: ۱۴۰۵/۰۶/۱۷'),
                       ('1405.06.30', 'تاریخ: 1405/06/15'), ('old', '1404/12/29')])
        self.assertEqual(latest_sheet(self.source), dict(date='1405/06/17', sheet='1405.06.17'))

    def test_reordered_sheets_and_new_year(self):
        self.workbook([('old', '1404/12/29'), ('new', '1405/01/01'), ('notes', 'توضیحات')])
        self.assertEqual(latest_sheet(self.source)['sheet'], 'new')

    def test_ambiguous_and_missing_dates_fail(self):
        for items in [[('a', '1405/06/17'), ('b', '1405/06/17')],
                      [('a', '1405/06/17 و 1405/06/18')], [('a', 'بدون تاریخ')]]:
            self.workbook(items)
            with self.assertRaises(ValueError):
                latest_sheet(self.source)

    def test_jalali_calendar_year_boundaries_and_today(self):
        for gregorian, jalali in [(date(2026, 9, 16), '1405/06/25'),
                                  (date(2025, 3, 20), '1403/12/30'),
                                  (date(2025, 3, 21), '1404/01/01'),
                                  (date(2026, 3, 21), '1405/01/01')]:
            self.assertEqual(jalali_today(gregorian), jalali)

    def test_current_overflow_has_only_one_line(self):
        self.assertEqual(report_caption('سرریز روزانه', '1405/06/25', today=date(2026, 9, 16)),
                         'سرریز روزانه 1405/06/25')

    def test_stale_and_future_dates_warn(self):
        for value in ['1405/06/17', '1405/06/26']:
            caption = report_caption('سرریز روزانه', value, today=date(2026, 9, 16))
            self.assertEqual(len(caption.splitlines()), 2)
            self.assertIn('⚠️', caption)
            self.assertIn('1405/06/25', caption)
            self.assertNotIn('بنویسید', caption)
            self.assertNotIn('تعداد ردیف', caption)

    def test_repairs_caption_states_latest_day(self):
        caption = report_caption('گزارش روزانه تعمیرات', '1405/06/17', today=date(2026, 9, 16), latest=True)
        self.assertIn('آخرین روز ثبت‌شده: 1405/06/17', caption)
        self.assertIn('⚠️', caption)

    def test_scheduled_repairs_sends_pdf_with_caption(self):
        def build(directory, source, *, section):
            self.assertEqual(section, 'mechanical')
            pdf = Path(directory) / 'report.pdf'
            pdf.write_bytes(b'%PDF-test')
            return dict(pdf=str(pdf), caption='latest-day-warning')
        with patch('tools.fleet.repairs.report.build_report', side_effect=build), patch('tools.scheduler.tasks.BaleSender') as sender:
            repairs('654806764', {})
            args = sender.return_value.document.call_args.args
            self.assertEqual(args[0], '654806764')
            self.assertEqual(args[2], 'latest-day-warning')
            self.assertEqual(Path(args[1]).suffix, '.pdf')
            sender.return_value.close.assert_called_once()
            self.assertFalse(Path(args[1]).exists())

    def test_section_columns_follow_headers_when_reordered(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(['گزارش', '1405/06/17'])
        sheet.append(['شرح معایب اهنگری', 'کد جدید', 'شرح معایب مکانیکی', 'نوع دستگاه', 'ردیف'])
        book.save(self.source)
        book.close()
        self.assertEqual(section_columns(self.source, sheet.title, 'mechanical'), [5, 4, 2, 3])
        self.assertEqual(section_columns(self.source, sheet.title, 'metalwork'), [5, 4, 2, 1])

    def test_missing_section_never_falls_back_to_full_sheet(self):
        self.workbook([('report', '1405/06/17')])
        with self.assertRaises(ValueError):
            section_columns(self.source, 'report', 'metalwork')
        for value in ['all', '', None]:
            with self.assertRaises(ValueError):
                validate_repairs({'section': value})

    def test_build_selects_columns_and_preserves_source_bytes(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(['گزارش', '1405/06/17'])
        sheet.append(['ردیف', 'نوع دستگاه', 'کد جدید', 'شرح معایب مکانیکی', 'شرح معایب آهنگری'])
        sheet.append([1, 'لودر', 'WA601', 'mechanical-only', 'metalwork-only'])
        book.save(self.source)
        book.close()
        original = self.source.read_bytes()
        for section, columns in [('mechanical', [1, 2, 3, 4]), ('metalwork', [1, 2, 3, 5])]:
            with patch('tools.fleet.repairs.report.export_pdf') as exporter:
                result = build_report(Path(self.temp.name) / section, self.source, section=section)
            self.assertEqual(exporter.call_args.args[-1], columns)
            self.assertIn(section, Path(result['pdf']).name)
            self.assertEqual(result['section'], section)
            self.assertEqual(self.source.read_bytes(), original)

    def test_metalwork_schedule_passes_section_and_recipient(self):
        with patch('tools.fleet.repairs.report.build_report', return_value={'pdf': 'test.pdf', 'caption': 'metalwork'}) as build, patch('tools.scheduler.tasks.BaleSender') as sender:
            repairs('455740857', {'section': 'metalwork'})
            self.assertEqual(build.call_args.kwargs, {'section': 'metalwork'})
            sender.return_value.document.assert_called_once_with('455740857', 'test.pdf', 'metalwork')

    def test_document_transport_multipart_and_redaction(self):
        pdf = Path(self.temp.name) / 'report.pdf'
        pdf.write_bytes(b'%PDF-test')
        with patch.dict('os.environ', {'BALE_BOT_TOKEN': 'private-token'}):
            sender = BaleSender()
        self.addCleanup(sender.close)
        with patch.object(sender.session, 'post') as post:
            post.return_value.status_code = 200
            post.return_value.json.return_value = {'ok': True}
            sender.document('654806764', pdf, 'caption')
            self.assertEqual(post.call_args.kwargs['files']['document'][2], 'application/pdf')
            self.assertTrue(post.call_args.args[0].endswith('/sendDocument'))
            post.side_effect = RuntimeError('private-token')
            with self.assertRaises(RuntimeError) as caught:
                sender.document('654806764', pdf, 'caption')
            self.assertNotIn('private-token', str(caught.exception))


if __name__ == '__main__':
    unittest.main()

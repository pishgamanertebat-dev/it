"""Real-layout rehearsal on a copy, with no production workbook writes or sends."""
from contextlib import closing
import hashlib
import json
from pathlib import Path
import shutil
import sqlite3
import uuid

import openpyxl

from .entry_service import CONFIG, ROOT, preview, commit, machine, clone_day, layout, find_row, today_sheet
from .sync_service import initialize, merge_snapshot, SyncConflict
from tools.fleet.report_caption import jalali_today


def signature(sheet):
    digest = hashlib.sha256()
    for position, cell in sorted(sheet._cells.items()):
        if cell.value is not None or cell.has_style:
            digest.update(repr((position, cell.value, cell.data_type, tuple(cell._style) if cell.has_style else None)).encode())
    digest.update(repr(([(r, d.height, d.hidden) for r, d in sheet.row_dimensions.items()],
        [(c, d.width, d.hidden) for c, d in sheet.column_dimensions.items()],
        sorted(map(str, sheet.merged_cells.ranges)), str(sheet.print_area),
        str(sheet.page_setup), str(sheet.page_margins), sheet.print_title_rows, sheet.print_title_cols)).encode())
    return digest.hexdigest()


def main():
    out = ROOT/'runtime/department-fault-entry-20261010/real-copy'
    out.mkdir(parents=True, exist_ok=True)
    settings = json.loads(CONFIG.read_text(encoding='utf8'))
    live = Path(settings['source']); original_hash = hashlib.sha256(live.read_bytes()).hexdigest()
    source = out/'local.xlsx'; mine = out/'mine.xlsx'
    shutil.copy2(live, source); shutil.copy2(live, mine)
    settings['source'] = str(source)
    cfg = out/'config.json'; cfg.write_text(json.dumps(settings, ensure_ascii=False), encoding='utf8')
    kw = dict(config=cfg, runtime=out/'journal')
    initialize(mine, **kw)
    before = openpyxl.load_workbook(mine)
    try:
        history = {sheet.title: signature(sheet) for sheet in before}
    finally:
        before.close()
    day = jalali_today()
    for actor, section, description in [('654806764', 'mechanical', 'mechanical rehearsal'),
            ('1732374823', 'mechanical', 'mechanical deputy edit'),
            ('387679249', 'metalwork', 'metalwork rehearsal')]:
        selection = preview(actor, '469', section, config=cfg)
        request = dict(selection, description=description, operation=uuid.uuid4().hex)
        commit(actor, request, **kw)
        unchanged = source.read_bytes()
        commit(actor, request, **kw)
        assert unchanged == source.read_bytes()
    incoming = openpyxl.load_workbook(mine)
    try:
        selected = machine(incoming, '469')
        sheet = clone_day(incoming, day, settings['template'])
        sheet.append([1, selected['name'], selected['code'], None, None, 'mine driver information'])
        # Row three already exists as a formatted empty template row.
        sheet['A3'], sheet['B3'], sheet['C3'] = 1, selected['name'], selected['code']
        sheet['D3'], sheet['E3'], sheet['F3'] = None, None, 'mine driver information'
        for cell in sheet[4]:
            cell.value = None
        sheet['F2'] = 'Other mine information'
        incoming.save(mine)
    finally:
        incoming.close()
    first = merge_snapshot(mine, source, **kw)
    assert first['protected_cells'] == 2
    repeat_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    assert merge_snapshot(mine, source, **kw)['status'] == 'unchanged'
    assert repeat_hash == hashlib.sha256(source.read_bytes()).hexdigest()
    incoming = openpyxl.load_workbook(mine)
    try:
        today_sheet(incoming, day)['F3'] = 'updated mine driver information'
        incoming.save(mine)
    finally:
        incoming.close()
    merge_snapshot(mine, source, **kw)
    merged = openpyxl.load_workbook(source)
    try:
        for title, expected in history.items():
            assert signature(merged[title]) == expected, title
        sheet = today_sheet(merged, day); cols = layout(sheet)
        row = find_row(sheet, 'HD469', cols)
        assert sheet.cell(row, cols['mechanical']).value == 'mechanical deputy edit'
        assert sheet.cell(row, cols['metalwork']).value == 'metalwork rehearsal'
        assert sheet['F3'].value == 'updated mine driver information'
        assert len(merged.sheetnames) == len(history)+1
    finally:
        merged.close()
    incoming = openpyxl.load_workbook(mine)
    try:
        today_sheet(incoming, day)['D3'] = 'different mine edit'
        incoming.save(mine)
    finally:
        incoming.close()
    safe = source.read_bytes()
    try:
        merge_snapshot(mine, source, **kw)
    except SyncConflict:
        pass
    else:
        raise AssertionError('Expected same-cell conflict')
    assert safe == source.read_bytes()
    assert hashlib.sha256(live.read_bytes()).hexdigest() == original_hash
    with closing(sqlite3.connect(out/'journal/audit.sqlite3')) as con:
        assert con.execute('PRAGMA integrity_check').fetchone() == ('ok',)
    result = dict(success=True, historical_sheets_preserved=len(history),
        live_source_unchanged=True, alias_469='HD469', both_sections_preserved=True,
        mine_information_preserved=True, repeat_no_write=True, conflict_preserved_local=True,
        production_test_entries=0, original_sha256=original_hash)
    (out/'evidence.json').write_text(json.dumps(result, indent=2), encoding='utf8')
    print(json.dumps(result))


if __name__ == '__main__':
    main()

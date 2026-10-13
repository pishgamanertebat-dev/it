"""Controlled one-workbook sync; all writes target disposable copies."""
from contextlib import closing
from concurrent.futures import ThreadPoolExecutor
import importlib.util
import json
import os
from pathlib import Path
import shutil
import sqlite3
import unittest
from unittest.mock import patch

import openpyxl

from .test_department_entry import DepartmentWriterTests
from .sync_service import initialize, merge_snapshot, SyncConflict
from .entry_service import ROOT, EntryError, replace_source


class FaultSyncTests(unittest.TestCase):
    book = DepartmentWriterTests.book
    request = DepartmentWriterTests.request
    save = DepartmentWriterTests.save

    def setUp(self):
        DepartmentWriterTests.setUp(self)
        self.mine = self.root/'mine.xlsx'
        shutil.copy2(self.source, self.mine)
        self.sync_kw = dict(config=self.config, runtime=self.root/'state')
        initialize(self.mine, **self.sync_kw)

    def incoming_day(self, *, mechanical=None, metalwork=None, other='driver information'):
        book = openpyxl.load_workbook(self.mine)
        try:
            sheet = book.copy_worksheet(book.worksheets[0])
            sheet.title = 'mine-day-26'
            sheet['C1'] = '1405/06/26'
            sheet['D3'], sheet['E3'], sheet['H3'] = mechanical, metalwork, other
            sheet['H2'] = 'Driver data'
            sheet['H3'].number_format = '@'
            book.save(self.mine)
        finally:
            book.close()

    def modify_mine(self, coordinate, value, sheet='mine-day-26'):
        book = openpyxl.load_workbook(self.mine)
        try:
            book[sheet][coordinate] = value
            book.save(self.mine)
        finally:
            book.close()

    def merge(self):
        return merge_snapshot(self.mine, self.source, **self.sync_kw)

    def values(self):
        book = openpyxl.load_workbook(self.source)
        try:
            sheet = book['mine-day-26']
            return sheet['D3'].value, sheet['E3'].value, sheet['H3'].value
        finally:
            book.close()

    def test_mine_new_day_both_sections_other_data_and_repeat(self):
        self.save('654806764', self.request('654806764', 'mechanical', 'engine'))
        self.save('387679249', self.request('387679249', 'metalwork', 'welding'))
        self.incoming_day()
        result = self.merge()
        self.assertEqual(result['protected_cells'], 2)
        self.assertEqual(self.values(), ('engine', 'welding', 'driver information'))
        before = self.source.read_bytes()
        self.assertEqual(self.merge()['status'], 'unchanged')
        self.assertEqual(self.source.read_bytes(), before)
        self.modify_mine('H3', 'new driver information')
        self.merge()
        self.assertEqual(self.values(), ('engine', 'welding', 'new driver information'))
        with closing(sqlite3.connect(self.root/'state/audit.sqlite3')) as con:
            rows = con.execute("SELECT status,backup FROM edits WHERE actor='mine-file-sync'").fetchall()
        self.assertEqual(len(rows), 2)
        for status, backup in rows:
            self.assertEqual(status, 'succeeded')
            self.assertTrue(Path(backup).is_file())

    def test_same_cell_conflict_preserves_current_then_recovery(self):
        self.save('654806764', self.request('654806764', 'mechanical', 'engine'))
        self.incoming_day(mechanical='manual mine edit')
        before = self.source.read_bytes()
        with self.assertRaisesRegex(SyncConflict, 'FAULT_SYNC_CELL_CONFLICT'):
            self.merge()
        self.assertEqual(self.source.read_bytes(), before)
        with closing(sqlite3.connect(self.root/'state/audit.sqlite3')) as con:
            self.assertEqual(con.execute('SELECT COUNT(*) FROM fault_sync_failures').fetchone()[0], 1)
        self.modify_mine('D3', None)
        self.merge()
        self.assertEqual(self.values()[0], 'engine')

    def test_matching_manual_edit_then_new_bale_edit_uses_updated_baseline(self):
        self.save('654806764', self.request('654806764', 'mechanical', 'engine'))
        self.incoming_day(mechanical='engine')
        self.merge()
        self.save('1732374823', self.request('1732374823', 'mechanical', 'engine v2'))
        self.modify_mine('H3', 'mine update')
        self.merge()
        self.assertEqual(self.values(), ('engine v2', None, 'mine update'))

    def test_missing_machine_preserves_mine_footer_and_print_columns(self):
        self.save('654806764', self.request('654806764', 'mechanical', 'engine'))
        self.incoming_day()
        book = openpyxl.load_workbook(self.mine)
        try:
            sheet = book['mine-day-26']
            sheet['C3'] = 'HD710'  # Bale device is missing from this incoming day.
            sheet['A4'], sheet['H4'] = 'mine footer', 'other mine information'
            sheet.merge_cells('H4:I4')
            sheet.print_area = 'A1:I4'
            book.save(self.mine)
        finally:
            book.close()
        self.merge()
        book = openpyxl.load_workbook(self.source)
        try:
            sheet = book['mine-day-26']
            self.assertEqual(sheet['A4'].value, 'mine footer')
            self.assertEqual(sheet['H4'].value, 'other mine information')
            self.assertIn('H4:I4', list(map(str, sheet.merged_cells.ranges)))
            self.assertEqual(sheet['C5'].value, 'EX332')
            self.assertEqual(sheet['D5'].value, 'engine')
            self.assertIn('I$5', sheet.print_area)
        finally:
            book.close()

    def test_uncertain_publication_recovers_without_duplicate(self):
        self.save('387679249', self.request('387679249', 'metalwork', 'welding'))
        self.incoming_day()
        def replaced_then_lost(*args):
            replace_source(*args)
            raise OSError('simulated lost result after atomic replacement')
        with patch('tools.fleet.repairs.entry_service.replace_source', side_effect=replaced_then_lost):
            with self.assertRaises(ValueError):
                self.merge()
        before = self.source.read_bytes()
        self.assertEqual(self.merge()['status'], 'unchanged')
        self.assertEqual(self.source.read_bytes(), before)
        with closing(sqlite3.connect(self.root/'state/audit.sqlite3')) as con:
            self.assertEqual(con.execute("SELECT COUNT(*),MIN(status) FROM edits WHERE actor='mine-file-sync'").fetchone(), (1, 'succeeded'))

    def test_unjournaled_local_change_is_not_discarded(self):
        self.incoming_day()
        with self.source.open('ab') as stream:
            stream.write(b'external local change')
        before = self.source.read_bytes()
        with self.assertRaisesRegex(EntryError, 'UNJOURNALED_LOCAL_CHANGE'):
            self.merge()
        self.assertEqual(self.source.read_bytes(), before)

    def test_open_excel_handle_preserves_file_and_retry_succeeds(self):
        self.incoming_day()
        before = self.source.read_bytes()
        with self.source.open('r+b'):
            with self.assertRaises(EntryError):
                self.merge()
        self.assertEqual(self.source.read_bytes(), before)
        self.assertEqual(self.merge()['status'], 'merged')

    def edit_local(self, cell, text, sheet=None):
        book = openpyxl.load_workbook(self.source)
        try:
            (book[sheet] if sheet else book.worksheets[0])[cell] = text
            book.save(self.source)
        finally:
            book.close()

    def test_manual_header_fix_and_new_mine_sheet_are_both_preserved(self):
        self.edit_local('C1', 'corrected header: 1405/06/17')
        self.incoming_day()
        result = self.merge()
        self.assertTrue(result['reconciled_local_change'])
        self.assertEqual(result['protected_local_cells'], 1)
        book = openpyxl.load_workbook(self.source)
        try:
            self.assertEqual(book.worksheets[0]['C1'].value, 'corrected header: 1405/06/17')
            self.assertIn('mine-day-26', book.sheetnames)
        finally:
            book.close()
        before = self.source.read_bytes()
        self.assertEqual(self.merge()['status'], 'unchanged')
        self.assertEqual(before, self.source.read_bytes())
        self.modify_mine('H3', 'new upstream change')
        self.merge()
        self.assertEqual(self.values()[2], 'new upstream change')

    def test_true_manual_cell_conflict_preserves_local_and_reports_values(self):
        self.edit_local('D3', 'manual local repair')
        self.incoming_day()
        self.modify_mine('D3', 'different mine repair', sheet='گزارش روزانه 1405.06.17')
        before = self.source.read_bytes()
        with self.assertRaisesRegex(EntryError, 'LOCAL_CELL_CONFLICT.*manual local repair.*different mine repair'):
            self.merge()
        self.assertEqual(before, self.source.read_bytes())

    def test_baseline_after_bale_edit_preserves_both_manual_and_bale(self):
        self.incoming_day()
        self.merge()
        self.save('654806764', self.request('654806764', 'mechanical', 'bale fault'))
        self.edit_local('C1', 'corrected date 1405/06/17')
        self.merge()
        self.assertEqual(self.values()[0], 'bale fault')

    def test_manual_edit_cannot_erase_confirmed_bale_value(self):
        self.incoming_day()
        self.merge()
        self.save('654806764', self.request('654806764', 'mechanical', 'bale fault'))
        self.edit_local('D3', None, sheet='mine-day-26')
        before = self.source.read_bytes()
        with self.assertRaisesRegex(SyncConflict, 'BALE_LOCAL_CONFLICT'):
            self.merge()
        self.assertEqual(before, self.source.read_bytes())

    def test_unrecognized_local_structure_preserved(self):
        book = openpyxl.load_workbook(self.source)
        book.worksheets[0].title = 'renamed manually'
        book.save(self.source)
        book.close()
        self.incoming_day()
        before = self.source.read_bytes()
        with self.assertRaisesRegex(EntryError, 'LOCAL_STRUCTURE_CONFLICT'):
            self.merge()
        self.assertEqual(before, self.source.read_bytes())

    def test_manual_correction_acknowledged_by_share_and_later_update_allowed(self):
        self.edit_local('D3', 'correction')
        self.incoming_day()
        self.merge()
        self.modify_mine('D3', 'correction', sheet='گزارش روزانه 1405.06.17')
        self.assertEqual(self.merge()['protected_local_cells'], 0)
        self.modify_mine('D3', 'later shared repair', sheet='گزارش روزانه 1405.06.17')
        self.merge()
        book = openpyxl.load_workbook(self.source)
        self.assertEqual(book.worksheets[0]['D3'].value, 'later shared repair')
        book.close()

    def test_manual_change_followed_by_bale_commit_is_preserved(self):
        self.incoming_day()
        self.merge()
        self.edit_local('C1', 'manually corrected 1405/06/17')
        self.save('654806764', self.request('654806764', 'mechanical', 'bale after manual'))
        self.modify_mine('H3', 'new upstream driver data')
        self.merge()
        self.assertEqual(self.values(), ('bale after manual', None, 'new upstream driver data'))
        book = openpyxl.load_workbook(self.source)
        self.assertEqual(book.worksheets[0]['C1'].value, 'manually corrected 1405/06/17')
        book.close()

    def test_confirmed_bale_edit_updates_a_manual_value_without_false_conflict(self):
        self.incoming_day()
        self.merge()
        self.edit_local('D3', 'manual fault', sheet='mine-day-26')
        self.merge()
        self.save('654806764', self.request('654806764', 'mechanical', 'confirmed next fault'))
        self.modify_mine('H3', 'driver update')
        self.merge()
        self.assertEqual(self.values(), ('confirmed next fault', None, 'driver update'))

    def test_local_row_height_and_equivalent_column_groups_are_preserved(self):
        book = openpyxl.load_workbook(self.source)
        sheet = book.worksheets[0]
        sheet.row_dimensions[3].height = 63
        sheet.column_dimensions['H'].min = 8
        sheet.column_dimensions['H'].max = 9
        sheet.column_dimensions['H'].width = 12
        book.save(self.source)
        book.save(self.mine)
        book.close()
        # Establish an equal baseline with a grouped column range.
        import hashlib
        with closing(sqlite3.connect(self.root/'state/audit.sqlite3')) as db:
            data = self.source.read_bytes()
            sha = hashlib.sha256(data).hexdigest()
            db.execute('UPDATE fault_sync_sources SET initial_hash=?, upstream_hash=?', (sha, sha))
            db.commit()
            (self.root/'state/backups'/ (sha+'.xlsx')).write_bytes(data)
        book = openpyxl.load_workbook(self.source)
        sheet = book.worksheets[0]
        sheet.row_dimensions[3].height = 67.5
        sheet.column_dimensions['H'].max = 8
        sheet.column_dimensions['I'].min = sheet.column_dimensions['I'].max = 9
        sheet.column_dimensions['I'].width = 12
        book.save(self.source)
        book.close()
        self.incoming_day()
        self.merge()
        book = openpyxl.load_workbook(self.source)
        self.assertEqual(book.worksheets[0].row_dimensions[3].height, 67.5)
        book.close()

    def test_changed_row_identity_is_reported_instead_of_applying_to_another_unit(self):
        self.edit_local('D3', 'manual fault')
        self.incoming_day()
        self.modify_mine('C3', 'HD710', sheet='گزارش روزانه 1405.06.17')
        before = self.source.read_bytes()
        with self.assertRaisesRegex(EntryError, 'row identity changed'):
            self.merge()
        self.assertEqual(before, self.source.read_bytes())

    def test_parallel_writers_and_sync_use_the_same_lock(self):
        mechanical = self.request('654806764', 'mechanical', 'engine')
        metalwork = self.request('387679249', 'metalwork', 'welding')
        with ThreadPoolExecutor(max_workers=2) as pool:
            first = pool.submit(self.save, '654806764', mechanical)
            second = pool.submit(self.save, '387679249', metalwork)
            first.result(); second.result()
        self.incoming_day()
        edit = self.request('1732374823', 'mechanical', 'new engine')
        with ThreadPoolExecutor(max_workers=2) as pool:
            imported = pool.submit(self.merge)
            edited = pool.submit(self.save, '1732374823', edit)
            imported.result(); edited.result()
        self.assertEqual(self.values(), ('new engine', 'welding', 'driver information'))
        self.modify_mine('H3', 'next mine update')
        self.merge()
        self.assertEqual(self.values(), ('new engine', 'welding', 'next mine update'))

    def test_actual_sync_hook_other_files_disconnect_reconnect(self):
        script = Path(os.environ.get('KOMATSO_SYNC_SCRIPT',
            'C:/Users/win-10/AppData/Local/hermes/scripts/mine_file_sync.py'))
        spec = importlib.util.spec_from_file_location('fault_sync_fixture', script)
        module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
        upstream = self.root/'share'; upstream.mkdir()
        destination = self.root/'local'; destination.mkdir()
        # Keep the fixture writer source in its existing path, using a flat mine source.
        fault_incoming = upstream/self.source.name
        shutil.copy2(self.mine, fault_incoming)
        other = upstream/'other.pdf'; other.write_bytes(b'other mine document')
        self.save('654806764', self.request('654806764', 'mechanical', 'engine'))
        self.incoming_day()
        shutil.copy2(self.mine, fault_incoming)
        # Sync destination is the fixture root, not E:\Function.
        def publish(staged, target, **kwargs):
            merge_snapshot(staged, target, **self.sync_kw)
        pulse = lambda **kwargs: None
        with patch.object(module, 'FAULT_REPORT', self.source), patch.object(module, 'publish_fault_report', side_effect=publish):
            result = module.empty_result()
            module.sync_tree(upstream, self.root, result, pulse)
            self.assertEqual(result['status'], 'success')
            self.assertEqual((self.root/'other.pdf').read_bytes(), other.read_bytes())
            self.assertEqual(self.values()[0], 'engine')
            before = self.source.read_bytes()
            offline = self.root/'share-offline'; upstream.rename(offline)
            with self.assertRaises(OSError):
                module.sync_tree(upstream, self.root, module.empty_result(), pulse)
            self.assertEqual(self.source.read_bytes(), before)
            offline.rename(upstream)
            self.modify_mine('H3', 'after reconnect')
            shutil.copy2(self.mine, fault_incoming)
            other.write_bytes(b'new other mine document')
            result = module.empty_result()
            module.sync_tree(upstream, self.root, result, pulse)
            self.assertEqual(result['status'], 'success')
            self.assertEqual(self.values(), ('engine', None, 'after reconnect'))
            self.assertEqual((self.root/'other.pdf').read_bytes(), other.read_bytes())


if __name__ == '__main__':
    unittest.main()

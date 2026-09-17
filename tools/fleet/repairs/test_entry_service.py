from contextlib import closing
from copy import copy
import hashlib
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
import uuid

import openpyxl
from openpyxl.styles import Font

from .entry_service import ROOT, EntryError, commit, preview, writer_lock, create_blank_template


class EntryServiceTests(unittest.TestCase):
    def setUp(self):
        parent = ROOT / 'runtime/repairs-entry-tests'
        parent.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=parent)
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'drivers.xlsx'
        self.template = self.root / 'template.xlsx'
        self.config = self.root / 'config.json'
        self.config.write_text(json.dumps({'allowed_users': ['641220453', '455740857'],
            'source': str(self.source), 'template': str(self.template)}), encoding='utf-8')
        self.kw = dict(config=self.config, fleet_db=self.root/'missing.db', day='1405/06/26')
        self.book()
        create_blank_template(self.source, self.template)

    def book(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'گزارش روزانه 1405.06.17'
        sheet.append([None, 'گزارش روزانه', 'تاریخ: 1405/06/17'])
        sheet.append(['ردیف','نوع دستگاه','کد جدید','شرح معایب مکانیکی','شرح معایب آهنگری'])
        sheet.append([1,'بیل مکانیکی','EX332','mechanical-old','metalwork-old'])
        sheet.append([2,'دامپتراک','HD710','old-710','old-metal-710'])
        sheet['D3'].font = Font(name='Arial', size=14, bold=True)
        sheet.column_dimensions['D'].width = 70
        sheet.row_dimensions[3].height = 40
        sheet.print_area = 'A1:E4'
        sheet.page_setup.orientation = 'landscape'
        archive = book.copy_worksheet(sheet)
        archive.title = 'archive'
        archive['C1'] = 'تاریخ: 1405/06/16'
        archive['C4'] = 'HD999'
        book.save(self.source)
        book.close()

    def request(self, code='EX332', section='mechanical', text='new fault'):
        selection = preview('641220453', code, section, **self.kw)
        return dict(selection, description=text, operation=uuid.uuid4().hex)

    def save(self, request, actor='641220453'):
        return commit(actor, request, runtime=self.root/'state', **self.kw)

    def test_new_day_and_preserve_history_and_other_section(self):
        original = self.source.read_bytes()
        request = self.request()
        result = self.save(request)
        book = openpyxl.load_workbook(self.source)
        try:
            today, old = book.worksheets[:2]
            self.assertEqual(today['C1'].value, 'تاریخ: 1405/06/26')
            self.assertEqual(today['D3'].value, 'new fault')
            self.assertIsNone(today['E3'].value)
            self.assertIsNone(today['D4'].value)
            self.assertEqual(old['D3'].value, 'mechanical-old')
            self.assertEqual(old['E3'].value, 'metalwork-old')
            self.assertEqual(copy(today['D3'].font), copy(old['D3'].font))
            self.assertEqual(today.column_dimensions['D'].width, 70)
            self.assertEqual(today.page_setup.orientation, 'landscape')
            self.assertIn('A$1:$E$3', str(today.print_area))
            self.assertEqual(result['date'], '1405/06/26')
        finally:
            book.close()
        backup = self.root/'state/backups'/ (hashlib.sha256(original).hexdigest()+'.xlsx')
        self.assertEqual(backup.read_bytes(), original)

    def test_two_sections_same_day_and_edit_and_clear(self):
        self.save(self.request())
        self.save(self.request(section='metalwork', text='welding'), actor='455740857')
        req = self.request(text='replacement')
        self.assertEqual(req['expected'], 'new fault')
        self.save(req)
        self.save(self.request(section='metalwork', text=''))
        book = openpyxl.load_workbook(self.source)
        try:
            self.assertEqual(len(book.worksheets), 3)
            self.assertEqual(book.worksheets[0]['D3'].value, 'replacement')
            self.assertIsNone(book.worksheets[0]['E3'].value)
        finally:
            book.close()

    def test_clear_only_description_removes_machine_row(self):
        self.save(self.request(text='temporary'))
        self.save(self.request(text=''))
        book = openpyxl.load_workbook(self.source)
        try:
            today = book.worksheets[0]
            self.assertIsNone(today['A3'].value)
            self.assertIsNone(today['C3'].value)
            self.assertIsNone(today['D3'].value)
            self.assertEqual(today.max_row, 3)
        finally:
            book.close()

    def test_clear_without_existing_row_does_not_create_sheet_or_row(self):
        self.save(self.request('HD999', text=''))
        book = openpyxl.load_workbook(self.source)
        try:
            self.assertEqual(book.worksheets[0].title, 'گزارش روزانه 1405.06.17')
            self.assertEqual(len(book.worksheets), 2)
            self.assertEqual(book.worksheets[0]['C3'].value, 'EX332')
        finally:
            book.close()
        self.save(self.request())
        self.save(self.request('HD999', text=''))
        book = openpyxl.load_workbook(self.source)
        try:
            today = book.worksheets[0]
            self.assertEqual(today['C3'].value, 'EX332')
            self.assertIsNone(today['C4'].value)
        finally:
            book.close()

    def test_clear_middle_row_renumbers_remaining(self):
        self.save(self.request('EX332'))
        self.save(self.request('HD710'))
        self.save(self.request('HD999'))
        self.save(self.request('HD710', text=''))
        book = openpyxl.load_workbook(self.source)
        try:
            today = book.worksheets[0]
            self.assertEqual(today['A3'].value, 1)
            self.assertEqual(today['C3'].value, 'EX332')
            self.assertEqual(today['A4'].value, 2)
            self.assertEqual(today['C4'].value, 'HD999')
            self.assertIsNone(today['C5'].value)
            self.assertIn('E$4', str(today.print_area))
        finally:
            book.close()

    def test_row_style_and_height_come_from_template_and_current_text(self):
        long_text = 'x' * 350
        expected_tall = min(409, max(40, 5 * 18 + 8))
        self.save(self.request(text=long_text))
        book = openpyxl.load_workbook(self.source)
        try:
            self.assertEqual(book.worksheets[0].row_dimensions[3].height, expected_tall)
        finally:
            book.close()
        self.save(self.request('HD999', text='short'))
        book = openpyxl.load_workbook(self.source)
        try:
            today = book.worksheets[0]
            self.assertEqual(today.row_dimensions[3].height, expected_tall)
            self.assertEqual(today.row_dimensions[4].height, 40)
            self.assertEqual(copy(today['D4'].font), copy(today['D3'].font))
        finally:
            book.close()
        self.save(self.request(section='metalwork', text=long_text))
        self.save(self.request(text='short now'))
        book = openpyxl.load_workbook(self.source)
        try:
            today = book.worksheets[0]
            self.assertEqual(today.row_dimensions[3].height, expected_tall)
            self.assertEqual(today['D3'].value, 'short now')
        finally:
            book.close()
        self.save(self.request(section='metalwork', text='short metal'))
        book = openpyxl.load_workbook(self.source)
        try:
            self.assertEqual(book.worksheets[0].row_dimensions[3].height, 40)
        finally:
            book.close()

    def test_append_known_machine_and_print_area(self):
        self.save(self.request('EX332'))
        self.save(self.request('HD999'))
        book = openpyxl.load_workbook(self.source)
        try:
            self.assertEqual(book.worksheets[0]['C4'].value, 'HD999')
            self.assertEqual(book.worksheets[0]['D4'].value, 'new fault')
            self.assertIn('E$4', str(book.worksheets[0].print_area))
        finally:
            book.close()

    def test_duplicate_confirmation_and_mutated_operation(self):
        req = self.request()
        self.save(req)
        saved = self.source.read_bytes()
        self.save(req)
        self.assertEqual(self.source.read_bytes(), saved)
        with self.assertRaises(EntryError):
            self.save(dict(req, description='tampered'))

    def test_concurrent_same_cell_is_rejected_other_column_is_kept(self):
        first = self.request(text='first')
        second = self.request(text='second')
        metal = self.request(section='metalwork', text='metal')
        self.save(first)
        with self.assertRaisesRegex(EntryError, 'تغییر کرده'):
            self.save(second, actor='455740857')
        self.save(metal, actor='455740857')

    def test_denied_unknown_machine_midnight(self):
        original = self.source.read_bytes()
        with self.assertRaises(PermissionError):
            preview('654806764', 'EX332', 'mechanical', **self.kw)
        with self.assertRaises(EntryError):
            self.request('UNKNOWN')
        req = self.request()
        with self.assertRaises(EntryError):
            self.save(dict(req, date='1405/06/25'))
        with self.assertRaises(PermissionError):
            self.save(req, actor='654806764')
        self.assertEqual(self.source.read_bytes(), original)

    def test_formula_like_description_is_literal(self):
        self.save(self.request(text='=HYPERLINK("http://example.invalid", "text")'))
        book = openpyxl.load_workbook(self.source)
        try:
            self.assertEqual(book.worksheets[0]['D3'].data_type, 's')
        finally:
            book.close()

    def test_failed_replace_then_retry_and_prepared_recovery(self):
        request = self.request()
        original = self.source.read_bytes()
        with patch('tools.fleet.repairs.entry_service.replace_source', side_effect=PermissionError('locked')):
            with self.assertRaises(EntryError):
                self.save(request)
        self.assertEqual(self.source.read_bytes(), original)
        self.save(request)
        with closing(sqlite3.connect(self.root/'state/audit.sqlite3')) as conn:
            conn.execute("UPDATE edits SET status='prepared'")
            conn.commit()
        saved = self.source.read_bytes()
        self.save(request)
        self.assertEqual(self.source.read_bytes(), saved)

    def test_revocation_checked_again_at_commit(self):
        request = self.request()
        config = json.loads(self.config.read_text())
        config['allowed_users'] = []
        self.config.write_text(json.dumps(config))
        with self.assertRaises(PermissionError):
            self.save(request)

    def test_new_day_has_only_devices_entered_that_day(self):
        request = self.request(text='updated-mechanical')
        self.assertEqual(request['expected'], '')
        self.save(request)
        self.kw['day'] = '1405/06/27'
        self.save(self.request('HD710', section='metalwork', text='new metal'))
        book = openpyxl.load_workbook(self.source)
        try:
            today = book.worksheets[0]
            self.assertEqual(today.max_row, 3)
            self.assertEqual(today['C3'].value, 'HD710')
            self.assertIsNone(today['D3'].value)
            self.assertEqual(today['E3'].value, 'new metal')
            self.assertEqual(book.worksheets[1]['D3'].value, 'updated-mechanical')
        finally:
            book.close()


if __name__ == '__main__':
    unittest.main()

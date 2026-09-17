"""Exercise the real workbook layout on an isolated copy; never send messages."""
from copy import copy
import hashlib
import json
from pathlib import Path
import shutil
import uuid

import openpyxl

from .entry_service import ROOT, CONFIG, preview, commit
from .report import DEFAULT_SOURCE, latest_sheet, section_columns


def signature(sheet):
    return {
        'cells': [(c.coordinate, c.value, str(copy(c.font)), str(copy(c.fill)), str(copy(c.border)),
                   str(copy(c.alignment)), c.number_format) for row in sheet for c in row],
        'rows': [(r, d.height, d.hidden) for r, d in sheet.row_dimensions.items()],
        'cols': [(c, d.width, d.hidden) for c, d in sheet.column_dimensions.items()],
        'merges': list(map(str, sheet.merged_cells.ranges)), 'area': str(sheet.print_area),
        'setup': str(sheet.page_setup),
    }


def main():
    directory = ROOT / 'artifacts/repairs-entry-preview'
    directory.mkdir(parents=True, exist_ok=True)
    before_hash = hashlib.sha256(DEFAULT_SOURCE.read_bytes()).hexdigest()
    original = openpyxl.load_workbook(DEFAULT_SOURCE)
    try:
        before = {s.title: signature(s) for s in original.worksheets}
    finally:
        original.close()
    target = directory / ('report-' + uuid.uuid4().hex + '.xlsx')
    shutil.copyfile(DEFAULT_SOURCE, target)
    settings = json.loads(CONFIG.read_text(encoding='utf-8'))
    settings['source'] = str(target)
    config = directory / (target.stem + '.json')
    config.write_text(json.dumps(settings, ensure_ascii=False), encoding='utf-8')
    options = {'config': config}
    journal = ROOT / 'runtime/repairs-entry-rehearsal' / target.stem
    results = []
    for actor, section, description in [
        ('641220453', 'mechanical', 'آزمایش ثبت مکانیکی روی نسخهٔ آزمایشی'),
        ('455740857', 'metalwork', 'آزمایش ثبت آهنگری روی نسخهٔ آزمایشی'),
        ('641220453', 'mechanical', 'آزمایش ویرایش مکانیکی روی نسخهٔ آزمایشی'),
    ]:
        selected = preview(actor, 'HD710', section, **options)
        request = {**selected, 'description': description, 'operation': uuid.uuid4().hex}
        result = commit(actor, request, runtime=journal, **options)
        assert result == commit(actor, request, runtime=journal, **options)
        results.append(result)
    updated = openpyxl.load_workbook(target)
    try:
        for title, expected in before.items():
            assert signature(updated[title]) == expected, title
        day_sheet = updated[results[-1]['sheet']]
        assert day_sheet.max_row == 3
        assert day_sheet['C3'].value == 'HD710'
        assert 'ویرایش' in day_sheet['D3'].value
        assert 'آهنگری' in day_sheet['E3'].value
    finally:
        updated.close()
    assert hashlib.sha256(DEFAULT_SOURCE.read_bytes()).hexdigest() == before_hash
    report = latest_sheet(target)
    assert report['date'] == results[-1]['date']
    assert section_columns(target, report['sheet'], 'mechanical') == [1, 2, 3, 4]
    assert section_columns(target, report['sheet'], 'metalwork') == [1, 2, 3, 5]
    evidence = {'workbook': str(target), 'original_sha256': before_hash,
                'historical_sheets_preserved': len(before), 'date': report['date'], 'messages_sent': 0}
    (directory/'evidence.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(evidence, ensure_ascii=False))


if __name__ == '__main__':
    main()

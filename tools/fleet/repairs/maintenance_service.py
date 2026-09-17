"""Append repair events to existing machine worksheets, preserving their history."""
from copy import copy
import json
import math
from pathlib import Path
import re

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from .entry_service import EntryError, ROOT, key, copy_style, commit_workbook
from tools.fleet.report_caption import jalali_today

CONFIG = ROOT / 'settings/maintenance_entry.json'
RUNTIME = ROOT / 'runtime/maintenance-entry'
FIELDS = {'date': ('تاریخ',), 'mechanic': ('نام مکانیک', 'تعمیرکار'),
          'code': ('کد مکانیزم', 'حفارها'), 'description': ('نوع خرابی',), 'parts': ('قطعات مصرفی',)}


def permission(actor, config=CONFIG):
    settings = json.loads(Path(config).read_text(encoding='utf-8'))
    if not actor or str(actor) not in settings['allowed_users']:
        raise PermissionError('اجازهٔ ثبت تعمیرات را ندارید.')
    return settings


def layout(sheet):
    for row in sheet.iter_rows(max_row=min(10, sheet.max_row)):
        columns = {}
        for field, labels in FIELDS.items():
            matches = [c.column for c in row if key(c.value) in {key(v) for v in labels}]
            if len(matches) == 1:
                columns[field] = matches[0]
        if len(columns) == len(FIELDS):
            return row[0].row, columns
    raise EntryError('ستون‌های شیت دستگاه شناخته نشدند؛ فایل باید بررسی شود.')


def machines(book):
    result = []
    for sheet in book:
        header, _ = layout(sheet)
        name = next((str(c.value).strip() for row in sheet.iter_rows(max_row=header-1)
                     for c in row if c.value is not None), sheet.title)
        short = key(sheet.title).rstrip('.')
        prefix = next((prefix for label, prefix in [('بیل', 'EX'), ('لودر', 'WA'),
                      ('بلدوزر', 'D'), ('تراک', 'HD'), ('دامپ', 'HD')] if label in name), '')
        canonical = prefix + short if short.isdecimal() else short
        result.append({'sheet': sheet.title, 'code': short, 'canonical': canonical, 'name': name})
    return result


def machine(book, entered):
    entered = key(entered)
    if not entered or len(entered) > 60:
        raise EntryError('کد دستگاه معتبر نیست.')
    choices = [m for m in machines(book) if entered in
               {key(m['name']), m['canonical'], m['code'], key(m['sheet'])}]
    if len(choices) == 1:
        return choices[0]
    if choices:
        raise EntryError('کد مشترک است؛ کد کامل دستگاه را وارد کنید: ' +
                         '، '.join(f"{m['canonical']} ({m['name']})" for m in choices))
    raise EntryError('شیتی برای این دستگاه وجود ندارد؛ کد دقیق دستگاه را وارد کنید.')


def preview(actor, code, *, config=CONFIG, day=None):
    settings = permission(actor, config)
    book = openpyxl.load_workbook(settings['source'])
    try:
        return {**machine(book, code), 'date': day or jalali_today()}
    finally:
        book.close()


def validate_text(value, limit=1800):
    if not isinstance(value, str) or not value.strip() or len(value) > limit or ILLEGAL_CHARACTERS_RE.search(value):
        raise EntryError(f'متن را در یک پیام، بین ۱ تا {limit} نویسه و بدون نویسهٔ کنترلی وارد کنید.')
    return value.strip()


def append_record(book, request):
    selected = machine(book, request['canonical'])
    if selected['sheet'] != request['sheet'] or selected['code'] != request['code']:
        raise EntryError('شیت دستگاه تغییر کرده است؛ «تعمیرات» را دوباره باز کنید.')
    sheet = book[selected['sheet']]
    header, columns = layout(sheet)
    occupied = [r for r in range(header + 1, sheet.max_row + 1)
                if any(c.value is not None for c in sheet[r])]
    row = max(occupied, default=header) + 1
    if any(rng.min_row <= row <= rng.max_row for rng in sheet.merged_cells.ranges):
        raise EntryError('ردیف بعدی سلول ادغام‌شده دارد؛ قالب شیت باید بررسی شود.')
    template = max(occupied, default=header + 1)
    height = 24
    for field, col in columns.items():
        cell = sheet.cell(row, col)
        copy_style(sheet.cell(template, col), cell)
        value = request[field]
        cell.value = value
        cell.data_type = 's'
        cell.alignment = copy(cell.alignment)
        cell.alignment = openpyxl.styles.Alignment(horizontal=cell.alignment.horizontal or 'right',
                            vertical='center', wrap_text=True, readingOrder=2)
        # Excel has no stored automatic row-height calculation. Estimate using
        # this column's width and font size, including explicit newlines.
        width = sheet.column_dimensions[cell.column_letter].width or 13
        size = cell.font.sz or 11
        capacity = max(1, int((width - 2) * 11 / size))
        lines = sum(max(1, math.ceil(len(line) / capacity)) for line in str(value).split('\n'))
        height = max(height, lines * size * 1.5 + 8)
    if height > 409:
        raise EntryError('متن در یک ردیف اکسل جا نمی‌شود؛ آن را کوتاه‌تر یا در چند ثبت جدا وارد کنید.')
    sheet.row_dimensions[row].height = height
    sheet.row_dimensions[row].hidden = False
    if sheet.print_area:
        # Keep existing print ranges and extend only their bottom edge.
        sheet.print_area = re.sub(r'(\$[A-Z]+\$)(\d+)(?=,|$)',
                                 lambda m: m[1] + str(max(int(m[2]), row)), str(sheet.print_area))
    return {**selected, 'date': request['date'], 'row': row}


def commit(actor, request, *, config=CONFIG, runtime=RUNTIME, day=None):
    settings = permission(actor, config)
    for field in ('mechanic', 'description', 'parts'):
        validate_text(request.get(field), 200 if field == 'mechanic' else 1800)
    return commit_workbook(actor, request, settings, lambda book: append_record(book, request),
                           runtime=runtime, day=day, command='تعمیرات')


def main():
    import sys
    payload = json.load(sys.stdin)
    try:
        if payload['action'] == 'preview':
            result = preview(payload['actor'], payload['code'])
        elif payload['action'] == 'commit':
            result = commit(payload['actor'], payload['request'])
        else:
            raise EntryError('درخواست معتبر نیست.')
        response = {'ok': True, 'result': result}
    except (EntryError, PermissionError) as exc:
        response = {'ok': False, 'message': str(exc)}
    except Exception:
        import traceback
        traceback.print_exc(file=sys.stderr)
        response = {'ok': False, 'message': 'خواندن یا ذخیرهٔ اکسل ناموفق بود؛ فایل را ببندید و دوباره تلاش کنید.'}
    print(json.dumps(response, ensure_ascii=False))


if __name__ == '__main__':
    main()

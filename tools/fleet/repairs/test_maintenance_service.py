from contextlib import closing
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
import uuid

import openpyxl

from .maintenance_service import ROOT, EntryError, preview, commit, layout


class MaintenanceServiceTests(unittest.TestCase):
    def setUp(self):
        parent = ROOT / 'runtime/maintenance-tests'
        parent.mkdir(parents=True, exist_ok=True)
        temp = tempfile.TemporaryDirectory(dir=parent)
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.source = self.root / 'repairs.xlsx'
        self.config = self.root / 'config.json'
        self.config.write_text(json.dumps({'source': str(self.source),
                                          'allowed_users': ['455740857', '654806764']}))
        book = openpyxl.Workbook()
        book.remove(book.active)
        for title, name, offset, reverse in [('710', 'دامپ 710', 0, False),
                                             ('601', 'بیل 601', 0, True),
                                             ('601.', 'لودر 601', 0, False),
                                             ('851', 'بیل 851', 1, True)]:
            sheet = book.create_sheet(title)
            sheet.cell(1, 1+offset, name)
            headers = ['تاریخ', 'نام مکانیک', 'کد مکانیزم', 'نوع خرابی', 'قطعات مصرفی']
            if reverse:
                headers[1:3] = reversed(headers[1:3])
            for col, value in enumerate(headers, 1+offset):
                sheet.cell(2, col, value)
            sheet.cell(3, 1+offset, '1405/01/01')
            sheet.cell(3, 4+offset, 'existing')
            sheet.cell(40, 1+offset).number_format = '@'
            for col in range(1, 7):
                sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 40
            sheet.print_area = f'A1:{openpyxl.utils.get_column_letter(5+offset)}3'
        book.save(self.source)
        book.close()
        self.kw = {'config': self.config, 'day': '1405/06/26'}

    def request(self, code='710', **changes):
        return {**preview('455740857', code, **self.kw), 'mechanic': 'تعمیرکار',
                'description': 'خط اول\nخط دوم\nخط سوم', 'parts': 'مصرف نشده',
                'operation': uuid.uuid4().hex, **changes}

    def save(self, request, actor='455740857'):
        return commit(actor, request, runtime=self.root/'journal', **self.kw)

    def test_append_both_users_column_orders_multiline_and_history(self):
        for code, actor in [('710', '455740857'), ('EX601', '654806764'), ('851', '455740857')]:
            request = self.request(code)
            result = self.save(request, actor)
            self.assertEqual(result['row'], 4)
            book = openpyxl.load_workbook(self.source)
            try:
                sheet = book[result['sheet']]
                _, cols = layout(sheet)
                self.assertEqual(sheet.cell(3, cols['description']).value, 'existing')
                for field in cols:
                    self.assertEqual(sheet.cell(4, cols[field]).value, request[field])
                self.assertTrue(sheet.cell(4, cols['description']).alignment.wrap_text)
                self.assertGreater(sheet.row_dimensions[4].height, 40)
                self.assertIn('$4', str(sheet.print_area))
            finally:
                book.close()
        self.assertEqual(len(list((self.root/'journal/backups').glob('*.xlsx'))), 3)

    def test_ambiguous_unknown_and_disallowed_leave_source_untouched(self):
        before = self.source.read_bytes()
        for code in ['601', 'HD999', 'WA710']:
            with self.assertRaises(EntryError):
                self.request(code)
        self.assertEqual(self.request('WA601')['sheet'], '601.')
        self.assertEqual(self.request('۷۱۰')['sheet'], '710')
        with self.assertRaises(PermissionError):
            self.save(self.request(), '641220453')
        with self.assertRaises(PermissionError):
            preview('641220453', '710', **self.kw)
        self.assertEqual(before, self.source.read_bytes())

    def test_duplicate_and_distinct_same_day_requests(self):
        first, second = self.request(), self.request()
        self.save(first)
        before = self.source.read_bytes()
        self.save(first)
        self.assertEqual(before, self.source.read_bytes())
        with self.assertRaises(EntryError):
            self.save(dict(first, parts='changed'))
        self.assertEqual(self.save(second, '654806764')['row'], 5)

    def test_midnight_invalid_text_and_changed_sheet(self):
        before = self.source.read_bytes()
        for changes in ({'date': '1405/06/25'}, {'mechanic': ''}, {'description': '\x00'},
                        {'sheet': '601.'}, {'parts': 'x'*1801}, {'description': '\n'.join(['line']*50)}):
            with self.assertRaises(EntryError):
                self.save(self.request(**changes))
        self.assertEqual(before, self.source.read_bytes())

    def test_formula_text_and_empty_device(self):
        book = openpyxl.load_workbook(self.source)
        for cell in book['710'][3]:
            cell.value = None
        book.save(self.source)
        book.close()
        request = self.request(description='=1+2', mechanic='=name', parts='=part')
        self.assertEqual(self.save(request)['row'], 3)
        book = openpyxl.load_workbook(self.source)
        try:
            for address in ('B3', 'D3', 'E3'):
                self.assertEqual(book['710'][address].data_type, 's')
        finally:
            book.close()

    def test_failed_replace_retry_and_crash_recovery(self):
        request = self.request()
        original = self.source.read_bytes()
        with patch('tools.fleet.repairs.entry_service.replace_source', side_effect=PermissionError('locked')):
            with self.assertRaises(EntryError):
                self.save(request)
        self.assertEqual(original, self.source.read_bytes())
        self.save(request)
        with closing(sqlite3.connect(self.root/'journal/audit.sqlite3')) as db:
            db.execute("UPDATE edits SET status='prepared'")
            db.commit()
        saved = self.source.read_bytes()
        self.save(request)
        self.assertEqual(saved, self.source.read_bytes())


if __name__ == '__main__':
    unittest.main()

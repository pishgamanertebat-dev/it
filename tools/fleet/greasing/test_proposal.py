import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from .source import read_source, date_columns, color_kind, SOURCE, completed_cutoff
from .proposal import evaluate, build_proposal, resolve_items, rule_for, catalog


def entry(day, hours, color='NONE', month=6, shift=0):
    return {'stamp':(month,day,shift),'hours':hours,'color':color,'cell':f'test-{month}-{day}-{shift}'}


def machine(code='HD715', entries=None):
    return {'code':code,'name':'test','entries':entries or [],'row':4,'legacy_code':'','duplicate':False}


class RuleTests(unittest.TestCase):
    def test_hours_after_service_include_night_and_cross_month(self):
        m=machine(entries=[entry(31,11,'ORANGE',month=5),entry(31,3,month=5,shift=1),entry(1,11),entry(2,11),entry(3,11),entry(4,13)])
        c,reason=evaluate(m,(6,4,1),(6,5))
        self.assertIsNone(reason)
        self.assertEqual((c['value'],c['remaining'],c['state']),(49,11,'OK'))

    def test_hour_thresholds(self):
        for code,threshold,warning in [('HD715',60,54),('EX231',10,7),('W472',10,7),('D155',10,7)]:
            for hours,state in [(warning-1,'OK'),(warning,'NEAR_DUE'),(threshold,'DUE')]:
                m=machine(code,[entry(1,None,'ORANGE')]+[entry(2+i,min(10,hours-i*10)) for i in range((hours+9)//10)])
                c,reason=evaluate(m,(6,15,1),(6,16))
                self.assertIsNone(reason)
                self.assertEqual(c['state'],state)

    def test_orange_blank_and_dash_are_real_resets_yellow_is_not(self):
        for value in (None,'-',0):
            m=machine(entries=[entry(1,10,'ORANGE'),entry(2,11),entry(3,value,'ORANGE'),entry(4,9,'OTHER')])
            c,_=evaluate(m,(6,4,1),(6,5))
            self.assertEqual(c['value'],9)
        c,reason=evaluate(machine('EX231',[entry(1,10,'OTHER')]),(6,2,1),(6,3))
        self.assertIsNone(c)
        self.assertIn('مشخص نیست',reason)

    def test_calendar_cycles_follow_latest_real_service(self):
        for days,state in [(5,'OK'),(6,'NEAR_DUE'),(7,'NEAR_DUE'),(8,'DUE'),(9,'DUE')]:
            c,_=evaluate(machine('s1',[entry(1,24,'ORANGE')]),(6,days,1),(6,1+days))
            self.assertEqual(c['state'],state)
        c,_=evaluate(machine('s1',[entry(1,24,'ORANGE'),entry(5,None,'ORANGE')]),(6,8,1),(6,9))
        self.assertEqual(c['value'],4)

    def test_future_event_invalid_hours_and_suspicious_color(self):
        base=[entry(1,0,'ORANGE'),entry(2,11)]
        c,_=evaluate(machine(entries=base+[entry(4,0,'ORANGE')]),(6,3,1),(6,4))
        self.assertEqual(c['value'],11)
        for bad in (entry(3,-2),entry(3,'خراب'),entry(3,4,'SUSPECT'),entry(3,25)):
            c,reason=evaluate(machine(entries=base+[bad]),(6,3,1),(6,4))
            self.assertIsNone(c)
            self.assertTrue(reason)

    def test_exact_orange_only(self):
        ws=Workbook().active
        for rgb,expected in [('FFFFC000','ORANGE'),('00FFC000','ORANGE'),('FFFFFF00','OTHER'),('FFFFCE3C','SUSPECT')]:
            ws['A1'].fill=PatternFill('solid',fgColor=rgb)
            self.assertEqual(color_kind(ws['A1']),expected)

    def test_case_sensitive_identity_and_duplicate_model_codes(self):
        lower=machine('s1'); upper=machine('S1'); upper['row']=5
        source={'machines':[lower,upper]}
        self.assertEqual(resolve_items(['s1'],source)[0]['machine_code'],'s1')
        self.assertEqual(resolve_items(['S1'],source)[0]['machine_code'],'S1')
        source['machines']=[lower]
        with self.assertRaises(ValueError): resolve_items(['S1'],source)
        for row,legacy in [(6,'151'),(7,'152')]:
            source['machines'].append({**machine('D155'),'row':row,'legacy_code':legacy,'duplicate':True})
        self.assertEqual([i['machine_code'] for i in resolve_items(['151','152'],source)],['151','152'])
        with self.assertRaises(ValueError): resolve_items(['D155'],source)

    def test_generator_and_ex1251_are_fully_outside_greasing(self):
        excluded=[]
        for row,code,name in ((4,'EX1251','بیل مکانیکی آوردن موتور'),(5,'','ژنراتور')):
            excluded.append({'row':row,'code':code,'legacy_code':'','name':name,'duplicate':False,
                             'entries':[entry(1,None,'ORANGE'),entry(2,20,'SUSPECT')]})
        source={'machines':excluded,'cutoff':(6,2,1),'as_of':(6,2,1),'warnings':[],
                'sha256':'test','excluded_count':0}
        self.assertEqual(catalog(source),[])
        with self.assertRaises(ValueError):
            resolve_items(['EX1251'],source)


class ReaderTests(unittest.TestCase):
    def workbook(self):
        w=Workbook(); s=w.active; s.title='ساعت کاری'; col=7
        for month in range(1,7):
            for day in (range(1,32) if month < 6 else [1,2,3]):
                if day == 1 and month < 6:
                    s.cell(1,col,['فروردین','اردیبهشت','خرداد','تیر','مرداد'][month-1])
                s.cell(2,col,day); s.cell(3,col,'روز'); s.cell(3,col+1,'شب'); col+=2
        for row,code in [(4,'S1'),(5,'s1'),(6,'S2')]:
            s.cell(row,1,'کامیون'); s.cell(row,2,code)
        return w

    def test_rollover_inactive_rows_and_future_cells(self):
        w=self.workbook(); s=w.active; columns=dict(date_columns(s))
        s.cell(4,columns[(5,31,0)],10).fill=PatternFill('solid',fgColor='FFFFC000')
        s.cell(5,columns[(5,31,0)]).fill=PatternFill('solid',fgColor='FFFFC000')
        s.cell(5,columns[(6,1,0)],4)
        s.cell(5,columns[(6,3,0)],24).fill=PatternFill('solid',fgColor='FFFFC000')
        with tempfile.TemporaryDirectory(dir='runtime/greasing_proposal') as temp:
            p=Path(temp)/'source.xlsx'; w.save(p); w.close()
            source=read_source(p,as_of=(6,2,1))
        self.assertEqual(source['cutoff'],(6,1,0))
        self.assertEqual([m['code'] for m in source['machines']],['s1'])
        self.assertTrue(source['warnings'])
        c,_=evaluate(source['machines'][0],source['cutoff'],(6,2))
        self.assertEqual(c['value'],2)

    def test_real_source_snapshot(self):
        p=build_proposal(as_of=(6,16,1))
        self.assertEqual(p['cutoff'],'1405/06/15 - روز')
        self.assertEqual(p['plan_date'],'1405/06/16')
        codes={i['machine_code'] for i in p['items']}
        self.assertTrue({'s1','TR1','TA1'}.issubset(codes))
        self.assertFalse({'S1','S2','HD715','HD468'} & codes)
        truck=next(i for i in p['evaluations'] if i['machine_code']=='HD715')['components']['greasing']
        self.assertEqual((truck['value'],truck['remaining']),(49,11))
        self.assertTrue(any(i.startswith('EX801: 9 از 10 ساعت') for i in p['warnings']))
        self.assertFalse(any(i['code']=='EX1251' or 'ژنراتور' in i['code'] for i in p['review']))
        self.assertEqual({i['code'] for i in p['review'] if 'رنگ مشکوک' in i['reason']},{'HD712','HD714'})
        self.assertEqual(next(i for i in p['evaluations'] if i['machine_code']=='EX231')['components']['greasing']['last_service'],'1405/05/18 - روز')


if __name__=='__main__': unittest.main()

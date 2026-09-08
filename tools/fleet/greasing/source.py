"""Read only greasing evidence from the operational workbook; preserve source identities."""
from collections import Counter
from datetime import datetime, date, timedelta, timezone
from io import BytesIO
from pathlib import Path
import hashlib
import math
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

SOURCE = Path('E:/Function/ساعت و مصرف روغن مکانیزم ها .xlsx')
SHEET = 'ساعت کاری'
MONTHS = ('فروردین','اردیبهشت','خرداد','تیر','مرداد','شهریور','مهر','آبان','آذر','دی','بهمن','اسفند')
LENGTHS = (31,31,31,31,31,31,30,30,30,30,30,29)


def clean(value):
    return '' if value is None else ' '.join(str(value).replace('\u200c',' ').translate(str.maketrans('كي۰۱۲۳۴۵۶۷۸۹','کی0123456789')).split())


def code_text(value):
    if isinstance(value,float) and value.is_integer():
        value=int(value)
    return clean(value)


def ordinal(md):
    return sum(LENGTHS[:md[0]-1])+md[1]


def from_ordinal(day):
    if not 1 <= day <= 365:
        raise ValueError('تاریخ خارج از سال عملیاتی ۱۴۰۵ است.')
    for month,length in enumerate(LENGTHS,1):
        if day <= length:
            return month,day
        day -= length


def format_date(md):
    return f'1405/{md[0]:02d}/{md[1]:02d}'


def format_shift(stamp):
    return format_date(stamp[:2]) + (' - روز' if stamp[2] == 0 else ' - شب')


def completed_cutoff(now=None):
    now=now or datetime.now(timezone(timedelta(hours=3,minutes=30)))
    # Operational year 1405 begins on Gregorian 2026-03-21.
    day=(now.date()-date(2026,3,21)).days+1
    # Morning proposals use completed calendar days, not today's partial shifts.
    return (*from_ordinal(day-1),1)


def numeric(value):
    return isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(value) and 0 <= value <= 24


def color_kind(cell, neutral_themes=(0,1)):
    fill=cell.fill
    if fill.patternType is None:
        return 'NONE'
    color=fill.fgColor
    if fill.patternType == 'solid' and color.type == 'rgb' and color.tint == 0:
        rgb=str(color.rgb)[-6:].upper()
        if rgb in {'FFC000', 'FFCE3C'}:
            return 'ORANGE'
        if rgb in {'FFFF00','FFFFFF','000000','92D050','0099FF','00FF99','00B0F0'}:
            return 'OTHER'
    # Workbook's neutral row shading, verified separately from event colors.
    if fill.patternType == 'solid' and color.type == 'theme' and color.theme in neutral_themes:
        return 'OTHER'
    return 'SUSPECT'


def date_columns(ws):
    anchors=[c for c in range(7,ws.max_column+1) if clean(ws.cell(1,c).value)=='فروردین']
    if not anchors:
        raise ValueError('شروع سال عملیاتی در سرستون‌ها پیدا نشد.')
    start=anchors[-1]
    month=1
    previous=None
    columns=[]
    seen=set()
    for col in range(start,ws.max_column+1):
        if clean(ws.cell(3,col).value) != 'روز':
            continue
        if clean(ws.cell(3,col+1).value) != 'شب':
            raise ValueError(f'ستون شب پس از ستون {col} معتبر نیست.')
        raw=ws.cell(2,col).value
        if isinstance(raw,bool) or not clean(raw).isdigit():
            raise ValueError(f'روز ستون {col} قابل تشخیص نیست.')
        day=int(clean(raw))
        label=clean(ws.cell(1,col).value)
        if previous is not None and day < previous:
            if day != 1 or previous != LENGTHS[month-1]:
                raise ValueError(f'توالی روزها پیش از ستون {col} نیازمند بررسی است.')
            month += 1
        if not 1 <= month <= 12 or not 1 <= day <= LENGTHS[month-1]:
            raise ValueError('تاریخ غیرمعتبر در تاریخچهٔ گریس‌کاری')
        if label in MONTHS and MONTHS.index(label)+1 != month:
            raise ValueError(f'عنوان ماه و توالی روزها در ستون {col} هم‌خوان نیستند.')
        if (month,day) in seen:
            raise ValueError('تاریخ تکراری در تاریخچهٔ گریس‌کاری')
        seen.add((month,day))
        for shift in (0,1):
            columns.append(((month,day,shift),col+shift))
        previous=day
    return columns


def read_source(path=SOURCE, as_of=None):
    payload=Path(path).read_bytes()
    workbook=load_workbook(BytesIO(payload),data_only=True)
    rawbook=load_workbook(BytesIO(payload),data_only=False,read_only=True)
    limit=as_of or completed_cutoff()
    try:
        ws=workbook[SHEET]
        neutral_themes=set()
        if workbook.loaded_theme:
            theme=ET.fromstring(workbook.loaded_theme)
            scheme=theme.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme')
            # Excel theme-index ordering is light1, dark1, light2, dark2.
            colors={e.tag.rsplit('}',1)[-1]: list(e)[0].attrib for e in scheme}
            for index,key in enumerate(('lt1','dk1','lt2','dk2')):
                color=colors[key].get('lastClr',colors[key].get('val'))
                if color in {'FFFFFF','000000','E7E6E6','44546A'}:
                    neutral_themes.add(index)
        columns=date_columns(ws)
        allowed=[(stamp,c) for stamp,c in columns if stamp <= limit]
        if not allowed:
            raise ValueError('پیش از تاریخ محاسبه، داده‌ای وجود ندارد.')
        rows=[r for r in range(4,ws.max_row+1) if clean(ws.cell(r,1).value)]
        lastrow=max(rows)
        start=columns[0][1]
        end=allowed[-1][1]
        rawrows=list(rawbook[SHEET].iter_rows(min_row=4,max_row=lastrow,min_col=start,max_col=end,values_only=True))
        machines=[]
        populated=[]
        for r in rows:
            entries=[]
            for stamp,c in allowed:
                cell=ws.cell(r,c)
                value=cell.value
                if value is None and str(rawrows[r-4][c-start]).startswith('='):
                    value='فرمول بدون مقدار محاسبه‌شده'
                kind=color_kind(cell, neutral_themes)
                entries.append({'stamp':stamp,'hours':value,'color':kind,'cell':cell.coordinate})
                if clean(value) not in {'','-'} or kind == 'ORANGE':
                    populated.append(stamp)
            machines.append({'row':r,'code':code_text(ws.cell(r,2).value),'legacy_code':code_text(ws.cell(r,3).value),
                             'name':clean(ws.cell(r,1).value),'entries':entries})
        if not populated:
            raise ValueError('هیچ کارکرد ثبت‌شده‌ای پیش از تاریخ محاسبه پیدا نشد.')
        cutoff=max(populated)
        period=cutoff[0]
        active=[m for m in machines if any(e['stamp'][0]==period and e['stamp']<=cutoff and numeric(e['hours']) and e['hours']>0 for e in m['entries'])]
        counts=Counter(m['code'] for m in machines if m['code'])
        for m in active:
            m['duplicate']=bool(m['code'] and counts[m['code']]>1)
            m['entries']=[e for e in m['entries'] if e['stamp']<=cutoff]
        warnings=[]
        if cutoff[2] == 0:
            warnings.append('آخرین ثبت مربوط به شیفت روز است؛ دادهٔ شیفت شب پس از آن ثبت نشده است.')
        if cutoff < allowed[-1][0]:
            warnings.append(f'ستون‌های بعد از {format_shift(cutoff)} تا {format_shift(allowed[-1][0])} دادهٔ ثبت‌شده ندارند.')
        if ordinal(limit[:2])-ordinal(cutoff[:2]) >= 1:
            warnings.append('آخرین داده قدیمی‌تر از آخرین شیفت قابل محاسبه است؛ جمع‌بندی کارکرد بررسی شود.')
        return {'machines':active,'cutoff':cutoff,'as_of':limit,'warnings':warnings,
                'sha256':hashlib.sha256(payload).hexdigest(),'excluded_count':len(machines)-len(active)}
    finally:
        workbook.close()
        rawbook.close()

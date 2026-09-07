from pathlib import Path
import sqlite3
from openpyxl import load_workbook


DB_PATH = Path(
    r"E:\KomatsoAI\data\fleet\db\fleet_ops.db"
)


def clean(value):
    if value is None:
        return ""

    return " ".join(
        str(value)
        .replace("\u200c", " ")
        .replace("ي", "ی")
        .replace("ك", "ک")
        .split()
    )


conn = sqlite3.connect(DB_PATH)

source = conn.execute(
    """
    SELECT id, raw_path
    FROM ingest_sources
    WHERE source_type = 'service_events'
    ORDER BY id DESC
    LIMIT 1
    """
).fetchone()

if not source:
    raise SystemExit(
        "ERROR: service_events source not found"
    )

source_id, raw_path = source

wb = load_workbook(
    raw_path,
    read_only=True,
    data_only=True,
)

print("SERVICE MASTER MAPPING")
print("=" * 70)
print(f"Source ID: {source_id}")
print()

for sheet_name in [
    "ساعت کاری (2)",
    "ساعت کاری",
]:
    if sheet_name not in wb.sheetnames:
        continue

    ws = wb[sheet_name]

    print("=" * 70)
    print(f"SHEET: {sheet_name}")
    print()

    # From inspection:
    # A = machine description
    # B = new/canonical code
    # C = old/internal code

    count = 0

    for excel_row, row in enumerate(
        ws.iter_rows(
            min_row=4,
            values_only=True,
        ),
        start=4,
    ):
        machine_name = clean(
            row[0] if len(row) > 0 else None
        )

        canonical_code = clean(
            row[1] if len(row) > 1 else None
        )

        internal_code = clean(
            row[2] if len(row) > 2 else None
        )

        if not any(
            (
                machine_name,
                canonical_code,
                internal_code,
            )
        ):
            continue

        # Ignore rows that are clearly not equipment rows.
        if not internal_code:
            continue

        print(
            f"row={excel_row} | "
            f"internal={internal_code} | "
            f"canonical={canonical_code or '(empty)'} | "
            f"name={machine_name}"
        )

        count += 1

    print()
    print(f"Rows found: {count}")
    print()

print("=" * 70)
print("CHECK AGAINST MACHINE REGISTRY")
print()

rows = conn.execute(
    """
    SELECT
        canonical_code,
        machine_type_hint,
        identity_status
    FROM machines
    ORDER BY canonical_code
    """
).fetchall()

registered = {
    str(row[0]).upper(): row
    for row in rows
}

# Check mappings from primary master sheet.
if "ساعت کاری (2)" in wb.sheetnames:

    ws = wb["ساعت کاری (2)"]

    for excel_row, row in enumerate(
        ws.iter_rows(
            min_row=4,
            values_only=True,
        ),
        start=4,
    ):
        machine_name = clean(
            row[0] if len(row) > 0 else None
        )

        canonical_code = clean(
            row[1] if len(row) > 1 else None
        )

        internal_code = clean(
            row[2] if len(row) > 2 else None
        )

        if not internal_code:
            continue

        if not canonical_code:
            print(
                f"NO NEW CODE | "
                f"internal={internal_code} | "
                f"name={machine_name}"
            )
            continue

        match = registered.get(
            canonical_code.upper()
        )

        if match:
            print(
                f"OK | "
                f"{internal_code} -> "
                f"{canonical_code} | "
                f"{match[2]}"
            )
        else:
            print(
                f"NOT REGISTERED | "
                f"{internal_code} -> "
                f"{canonical_code} | "
                f"name={machine_name}"
            )

wb.close()
conn.close()
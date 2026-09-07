import sqlite3
import openpyxl

DB = r"E:\KomatsoAI\data\fleet\db\fleet_ops.db"

conn = sqlite3.connect(DB)
cur = conn.cursor()

cur.execute("""
SELECT id, raw_path
FROM ingest_sources
WHERE original_filename LIKE '%ساعت و مصرف روغن%'
ORDER BY id DESC
LIMIT 1
""")

row = cur.fetchone()

if not row:
    raise SystemExit("No service file found")

source_id, raw_path = row

print("SERVICE SHEET INSPECTION")
print("=" * 70)
print("Source ID:", source_id)
print("File:", raw_path)

wb = openpyxl.load_workbook(raw_path, data_only=True)

ws = wb["بیل 101"]

print()
print("SHEET:", ws.title)
print("ROWS:", ws.max_row)
print("COLS:", ws.max_column)

print()
print("FIRST 25 ROWS")
print("-" * 70)

for r in ws.iter_rows(min_row=1, max_row=25, values_only=True):
    print(r)

conn.close()
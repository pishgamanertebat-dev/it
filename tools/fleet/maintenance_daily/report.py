"""Exact-date extraction over an inspected, explicit workbook schema, and vector A4 PDF."""
from __future__ import annotations
from pathlib import Path
from contextlib import contextmanager
import hashlib,html,io,json,os,re,unicodedata
import openpyxl
from integrations.hermes.function_domain.scoped import ScopedReader,MAX_PARSE_BYTES
from integrations.hermes.function_domain.extract import zip_safe
from tools.fleet.overflow.report import validate_date
SOURCE_NAME='تعمیرات 1405.xlsx'
LAYOUT=Path(__file__).with_name('layout.json')
DIGITS=str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩','01234567890123456789')
COMPANY='شرکت صنعتی و معدنی پیشگامان ارتباط هشت بهشت'
TITLE='گزارش روزانه تعمیرات ماشین‌آلات'
REPORT_COLUMNS=['تاریخ','کد مکانیزم','نام مکانیک / تعمیرکار','نوع خرابی','قطعات مصرفی']

def output_columns(columns, rule):
    """One human-facing order from explicit inspected source coordinates.

    Runtime never classifies cell values or guesses which cell is a person/code.
    Exact header aliases are available for trusted fixture layouts without a map.
    """
    mapped=rule.get('report_columns')
    if mapped is None:
        aliases=[('تاریخ',),('کد مکانیزم',),('نام مکانیک','تعمیرکار'),('نوع خرابی',),('قطعات مصرفی',)]
        mapped=[]
        for labels in aliases:
            found=[c for c,label in columns if label in labels]
            if len(found)!=1:raise ValueError('Missing or ambiguous report field')
            mapped.append(found[0])
    if (not isinstance(mapped,list) or len(mapped)!=5 or any(type(c)!=int for c in mapped)
            or len(set(mapped))!=5 or set(mapped)!={c for c,_ in columns}
            or mapped[0]!=next(c for c,label in columns if label=='تاریخ')):
        raise ValueError('Invalid explicit report column mapping')
    return mapped

def display(v):
    return '' if v is None else str(v).strip()

def exact_date(v):
    if not isinstance(v,str):raise ValueError('Workbook date must be an explicit Jalali string')
    v=v.strip().translate(DIGITS)
    if not re.fullmatch(r'1[34]\d{2}/\d{1,2}/\d{1,2}',v):raise ValueError('Unsupported or ambiguous workbook date')
    return validate_date(v)

@contextmanager
def source_snapshot(reader=None):
    reader=reader or ScopedReader(allowed_files=[SOURCE_NAME])
    # Reuse pinned parent and READ-only Windows handle; active writers cause failure.
    with reader.open(SOURCE_NAME) as (_,f):
        info=os.fstat(f.fileno())
        if info.st_size>MAX_PARSE_BYTES:raise ValueError('Workbook too large')
        data=f.read(MAX_PARSE_BYTES+1)
        if len(data)>MAX_PARSE_BYTES:raise ValueError('Workbook too large')
        digest=hashlib.sha256(data).hexdigest()
        try:yield data,digest
        finally:
            f.seek(0)
            after=hashlib.sha256(f.read(MAX_PARSE_BYTES+1)).hexdigest()
            if after!=digest:raise RuntimeError('Source workbook changed during reporting')

def scan(data, target, *, layout=None):
    """Scan every sheet and validate even non-matching rows; never carry a blank date."""
    target=exact_date(target)
    rules=(layout or json.loads(LAYOUT.read_text(encoding='utf-8')))['sheets']
    zip_safe(data)
    book=openpyxl.load_workbook(io.BytesIO(data),data_only=False,keep_links=False)
    devices=[];dates=set();excluded=[];undated=0
    try:
        if not 1<=len(book.sheetnames)<=256:raise ValueError('Sheet count exceeds limit')
        if set(book.sheetnames)!=set(rules):raise ValueError('Workbook sheets differ from inspected layout; inspect before updating manifest')
        for s in book:
            rule=rules[s.title]
            if s.max_row>10000 or s.max_column>32:raise ValueError('Worksheet dimensions exceed reporting bounds')
            if rule.get('exclude'):
                excluded.append(s.title);continue
            h=rule['header_row'];columns=rule['columns'];indices=[c for c,_ in columns]
            expected={c:label for c,label in columns}
            actual={c.column:c.value for c in s[h] if c.value is not None}
            if actual!=expected:raise ValueError('Header differs from inspected schema: '+s.title)
            titles=[[c.coordinate,c.value] for r in range(1,h) for c in s[r] if c.value is not None]
            if titles!=rule['title_cells']:raise ValueError('Device header changed: '+s.title)
            date_cols=[c for c,label in columns if label=='تاریخ']
            if len(date_cols)!=1:raise ValueError('One exact date column is required')
            dc=date_cols[0];merged_dates={}
            report_indices=output_columns(columns,rule)
            for region in s.merged_cells.ranges:
                if region.max_row<=h:continue
                if not (rule.get('date_merged_blocks') and region.min_col==region.max_col==dc and region.min_row>h and region.max_row-region.min_row<1000):
                    raise ValueError('Uninspected data-cell merge: '+s.title)
                day=exact_date(s.cell(region.min_row,dc).value)
                for r in range(region.min_row,region.max_row+1):merged_dates[r]=day
            matched=[]
            for r in range(h+1,s.max_row+1):
                cells=list(s[r]);values={c.column:c.value for c in cells if c.value is not None}
                if not values:continue
                if values==expected:continue  # Exact repeated header only.
                if set(values)-set(indices):raise ValueError('Uninspected nonempty helper column: '+s.title)
                if any(c.data_type in {'f','e'} for c in cells if c.value is not None):raise ValueError('Formula/error requires separate inspection')
                raw=s.cell(r,dc).value
                if raw is None and r not in merged_dates:
                    placeholders=rule.get('undated_placeholders',[])
                    if any(values=={p['column']:p['value']} for p in placeholders):undated+=1;continue
                    raise ValueError('Undated activity cannot be associated safely: '+s.title)
                day=merged_dates[r] if r in merged_dates else exact_date(raw)
                # A date-only template row is not a recorded activity.
                if not any(display(values.get(c)) for c in indices if c!=dc):continue
                dates.add(day)
                if day==target:
                    vals=[day if c==dc else display(values.get(c)) for c in report_indices]
                    if any(len(v)>20000 for v in vals):raise ValueError('Cell text exceeds rendering bounds')
                    matched.append({'source_row':r,'values':vals})
            if matched:
                devices.append({'sheet':s.title,'device':rule['device'],'identity_note':rule.get('identity_note',''),
                    'columns':list(REPORT_COLUMNS),'rows':matched})
        return {'date':target,'devices':devices,'devices_matched':len(devices),'rows_matched':sum(len(d['rows']) for d in devices),
            'sheets_scanned':len(book.sheetnames),'excluded_sheets':excluded,'undated_placeholders_skipped':undated,
            'source_sha256':hashlib.sha256(data).hexdigest(),'available_dates':sorted(dates)}
    finally:book.close()

FONT_DIR=Path('C:/Windows/Fonts')
CSS='''@font-face{font-family:Maintenance;src:url(tahoma.ttf)} @font-face{font-family:Maintenance;src:url(tahomabd.ttf);font-weight:bold} body{font-family:Maintenance;font-size:10px;color:#173448;margin:0} h1{font-size:18px;margin:4px 0} h2{font-size:13px;margin:4px 0} p{margin:3px 0} table{border-collapse:collapse;width:100%;table-layout:fixed} td,th{border:0.6px solid #c3ced5;padding:6px 4px;vertical-align:top;text-align:right} th{background:#173448;color:white;font-size:10px}'''
WIDTHS=[17,16,14,36,17]
def render(report, output, *, test_sample=False):
    import fitz
    if not report['devices'] or report['rows_matched']<1:raise ValueError('Empty PDF prohibited')
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True)
    temporary=output.with_suffix('.partial.pdf')
    if not all((FONT_DIR/n).is_file() for n in ["tahoma.ttf","tahomabd.ttf"]):raise ValueError("Required Persian fonts unavailable")
    archive=fitz.Archive(str(FONT_DIR))
    doc=fitz.open();page=None;y=0;bottom=790
    def insert(p,body,top,height,left=30,right=565):
        spare,scale=p.insert_htmlbox(fitz.Rect(left,top,right,top+height),body,css=CSS,archive=archive,scale_low=1)
        if spare<0 or scale!=1:raise ValueError('PDF text cannot fit without truncation/scaling')
        return height-spare
    def height(body):
        with fitz.open() as probe:
            p=probe.new_page(width=595.276,height=841.89)
            return insert(p,body,0,680)
    def cell_body(value,header=False):
        return '<p dir="auto" style="margin:0;text-align:right;'+('color:white;font-weight:bold' if header else '')+'">'+(html.escape(value).replace(chr(10),'<br>') or '—')+'</p>'
    height_cache={}
    def row_height(values,header=False):
        cache_key=(tuple(values),header)
        if cache_key in height_cache:return height_cache[cache_key]
        heights=[]
        for value,width in zip(values,WIDTHS):
            with fitz.open() as probe:
                p=probe.new_page(width=595.276,height=841.89)
                heights.append(insert(p,cell_body(value,header),0,680,left=0,right=535*width/100-8))
        height_cache[cache_key]=max(heights)+12
        return height_cache[cache_key]
    def draw_row(values,top,rh,header=False,shaded=False):
        x=565
        for value,width in zip(values,WIDTHS):
            left=x-535*width/100
            fill=(.09,.20,.28) if header else ((.94,.96,.98) if shaded else (1,1,1))
            page.draw_rect(fitz.Rect(left,top,x,top+rh),color=(.76,.81,.84),fill=fill,width=.5)
            insert(page,cell_body(value,header),top+5,rh-10,left=left+4,right=x-4)
            x=left
        return rh
    def new_page():
        nonlocal page,y
        page=doc.new_page(width=595.276,height=841.89)
        header=f'<p dir="rtl">{COMPANY}</p><h1 dir="rtl">{TITLE}</h1><p dir="rtl">تاریخ: {report["date"]}</p><p dir="rtl">تعداد دستگاه‌های دارای فعالیت تعمیراتی: {report["devices_matched"]}</p>'
        if test_sample:header+='<p dir="rtl" style="color:#a33">نمونه آزمایشی قالب گزارش</p><p dir="rtl">تاریخ داده: '+report['date']+'</p>'
        y=30+insert(page,header,30,150)+12
    try:
        new_page()
        for device in report['devices']:
            title='<h2 dir="rtl">دستگاه: '+html.escape(device['device'])+'</h2><p dir="rtl">Sheet: '+html.escape(device['sheet'])+'</p>'
            if device['identity_note']:title+='<p dir="rtl">'+html.escape(device['identity_note'])+'</p>'
            th=height(title);hh=row_height(device['columns'],True)
            first=row_height(device['rows'][0]['values'])
            if y+th+hh+first>bottom:new_page()
            y+=insert(page,title,y,th+0.2);y+=draw_row(device['columns'],y,hh,header=True)
            for i,row in enumerate(device['rows']):
                rh=row_height(row['values'])
                if y+rh>bottom:
                    new_page();y+=insert(page,title,y,th+0.2);y+=draw_row(device['columns'],y,hh,header=True)
                if y+rh>bottom:raise ValueError('Single record exceeds A4 page; no data truncated')
                y+=draw_row(row['values'],y,rh,shaded=i%2==0)
            y+=12
        for i,p in enumerate(doc):
            insert(p,f'<p dir="rtl">صفحه {i+1} از {len(doc)} | {report["date"]}</p>',810,20)
        doc.set_metadata({'title':TITLE+' '+report['date'],'subject':'test sample' if test_sample else 'maintenance daily','author':COMPANY})
        doc.subset_fonts();canonical_unicode_maps(doc);doc.save(temporary,garbage=4,deflate=True,no_new_id=True);doc.close()
        validation=validate_pdf(temporary,report,test_sample=test_sample)
        os.replace(temporary,output);return validation
    finally:
        if not doc.is_closed:doc.close()
        temporary.unlink(missing_ok=True)

def canonical_unicode_maps(doc):
    """Normalize Arabic presentation glyphs in ToUnicode, leaving drawing glyphs intact.

    Tahoma exposes contextual forms and common Persian/Arabic glyph aliases.
    Expand the generated mappings deterministically to base Persian Unicode so
    normal Persian queries work in PDF readers, including lam-alef ligatures.
    """
    seen=set()
    def normalized(hexvalue):
        value=bytes.fromhex(hexvalue).decode('utf-16-be')
        value=unicodedata.normalize('NFKC',value).replace('ي','ی').replace('ك','ک').replace('\u00ad','-')
        return value.encode('utf-16-be').hex().upper()
    for page in doc:
        for font in page.get_fonts():
            kind,ref=doc.xref_get_key(font[0],'ToUnicode')
            if kind!='xref':raise ValueError('Searchable font mapping missing')
            xref=int(ref.split()[0])
            if xref in seen:continue
            seen.add(xref);cmap=doc.xref_stream(xref).decode('ascii')
            def expand(match):
                typ=match.group(1);body=match.group(2);pairs=[]
                for line in body.strip().splitlines():
                    tokens=re.findall(r'<([0-9A-Fa-f]+)>',line)
                    if typ=='bfchar' and len(tokens)==2:pairs.append((tokens[0],normalized(tokens[1])))
                    elif typ=='bfrange' and len(tokens)==3 and '[' not in line:
                        first,last,base=[int(t,16) for t in tokens]
                        if last-first>10000:raise ValueError('Font mapping range exceeds bound')
                        for i in range(first,last+1):pairs.append((format(i,'04X'),normalized(format(base+i-first,'04X'))))
                    else:raise ValueError('Unsupported font Unicode mapping')
                return '\n'.join(str(len(pairs[i:i+100]))+' beginbfchar\n'+'\n'.join('<'+a+'> <'+b+'>' for a,b in pairs[i:i+100])+'\nendbfchar' for i in range(0,len(pairs),100))
            cmap=re.sub(r'\d+ begin(bfchar|bfrange)\s*([\s\S]*?)end(?:bfchar|bfrange)',expand,cmap)
            doc.update_stream(xref,cmap.encode('ascii'))

def text_key(v):
    v=unicodedata.normalize('NFKC',v).replace('ي','ی').replace('ك','ک').replace('\u00ad','-')
    return re.sub(r'[\s()\u200c\u200e\u200f]+','',v)

def validate_pdf(path,report,*,test_sample=False):
    import fitz
    path=Path(path)
    if not path.is_file() or not 0<path.stat().st_size<45*1024*1024:raise ValueError('Unreadable or oversized PDF')
    with fitz.open(path) as pdf:
        if pdf.is_repaired or pdf.needs_pass or not pdf.page_count:raise ValueError('Malformed PDF')
        texts=[]
        for p in pdf:
            if abs(p.rect.width-595.276)>1 or abs(p.rect.height-841.89)>1 or p.get_images():raise ValueError('A4 vector PDF required')
            t=p.get_text()
            if not t.strip() or report['date'] not in t:raise ValueError('Searchable date missing')
            # Force renderer to validate page resources as well as searchable content.
            p.get_pixmap(matrix=fitz.Matrix(.5,.5));texts.append(t)
        combined=text_key('\n'.join(texts))
        for d in report['devices']:
            if text_key(d['device']) not in combined or text_key('Sheet: '+d['sheet']) not in combined:raise ValueError('Device section missing')
            for row in d['rows']:
                for value in row['values']:
                    if value and text_key(value) not in combined:raise ValueError('Source text missing from searchable PDF')
        if text_key(TITLE) not in combined:raise ValueError('Persian text is not searchable')
        if test_sample and text_key('نمونه آزمایشی قالب گزارش') not in combined:raise ValueError('Test label missing')
        return {'pages':pdf.page_count,'bytes':path.stat().st_size,'searchable':True,'vector':True,'devices':report['devices_matched'],'rows':report['rows_matched'],'date':report['date'],'pdf_sha256':hashlib.sha256(path.read_bytes()).hexdigest()}

def build(date, output_dir, reader=None):
    date=exact_date(date)
    with source_snapshot(reader) as (data,digest):
        report=scan(data,date)
        if not report['devices']:return {'status':'skipped','reason':'no_target_date_records','report':report}
        path=Path(output_dir)/('maintenance-daily-'+date.replace('/','-')+'.pdf')
        validation=render(report,path)
        return {'status':'ready','report':report,'pdf':str(path),'validation':validation}

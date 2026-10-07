"""Maintenance extraction, offline vector PDF, receive-only auth and scheduler tests.

Synthetic workbooks/identities and fake transports; no real messages.
"""
from pathlib import Path
from contextlib import closing
from datetime import datetime,timezone
import ast,copy,hashlib,io,json,sqlite3,unittest
from unittest.mock import Mock,patch
from zoneinfo import ZoneInfo
import openpyxl,yaml,fitz
from tools.admin.test_admin1 import AdminFixture
from tools.authorization.maintenance import ROLE,CAPABILITY,RESOURCE,migrate
from tools.fleet.maintenance_daily import report,delivery
from tools.scheduler.tasks import maintenance_daily_report,overflow_report_date
from tools.scheduler.runner import tick,load_config

DATE='1405/07/13'
HEAD=['تاریخ','کد مکانیزم','نام مکانیک','نوع خرابی','قطعات مصرفی']

def workbook(spec):
    w=openpyxl.Workbook();w.remove(w.active);rules={}
    for name,rows in spec:
        s=w.create_sheet(name);s.append(['دستگاه '+name]);s.merge_cells('A1:E1');s.append(HEAD)
        for row in rows:s.append(row)
        rules[name]={'header_row':2,'columns':[[i+1,h] for i,h in enumerate(HEAD)],'title_cells':[['A1','دستگاه '+name]],'device':'دستگاه '+name,'exclude':False,'date_merged_blocks':False}
    return w,{'sheets':rules}

def package(w):
    buf=io.BytesIO();w.save(buf);return buf.getvalue()

class ExtractionTests(unittest.TestCase):
    def extract(self,spec,date=DATE):
        w,layout=workbook(spec);return report.scan(package(w),date,layout=layout)
    def test_target_single_row_and_nonmatching_sheet(self):
        r=self.extract([('801',[[DATE,801,'رحمان','تعویض فیلتر','فیلتر']]),('802',[['1405/07/12',802,'شاهین','OLD',None]])])
        self.assertEqual([d['sheet'] for d in r['devices']],['801']);self.assertEqual(r['rows_matched'],1);self.assertEqual(r['sheets_scanned'],2)
    def test_multiple_rows_devices_exact_duplicates_and_no_leakage(self):
        row=[DATE,801,'مکانیک','NEW','قطعه'];r=self.extract([('801',[row,row,['1405/07/12',801,'مکانیک','OLD',None],row]),('802',[[DATE,802,'مکانیک','SECOND',None]])])
        self.assertEqual(r['devices_matched'],2);self.assertEqual(r['rows_matched'],4)
        self.assertEqual(len(r['devices'][0]['rows']),3);self.assertNotIn('OLD',json.dumps(r))
    def test_exact_date_no_latest_fallback_zero_result(self):
        r=self.extract([('801',[['1405/07/12',801,None,'OLDER',None]])]);self.assertEqual(r['devices'],[]);self.assertEqual(r['date'],DATE)
    def test_persian_digits_are_exact_not_guess(self):
        r=self.extract([('801',[['۱۴۰۵/۰۷/۱۳',801,None,'عملیات',None]])]);self.assertEqual(r['rows_matched'],1)
        for invalid in ['yesterday','1405/07','date: 1405/07/13','1405/13/13',123]:
            with self.subTest(invalid=invalid),self.assertRaises(ValueError):report.exact_date(invalid)
    def test_unmerged_blank_activity_never_inherits_date(self):
        w,l=workbook([('801',[[DATE,801,None,'FIRST',None],[None,801,None,'ORPHAN',None]])])
        with self.assertRaises(ValueError):report.scan(package(w),DATE,layout=l)
    def test_proven_merged_date_block_bounded_only_with_explicit_rule(self):
        w,l=workbook([('801',[[DATE,801,None,'FIRST',None],[None,801,None,'SECOND',None],['1405/07/12',801,None,'OLD',None]])]);w['801'].merge_cells('A3:A4')
        with self.assertRaises(ValueError):report.scan(package(w),DATE,layout=l)
        l['sheets']['801']['date_merged_blocks']=True;r=report.scan(package(w),DATE,layout=l)
        self.assertEqual(r['rows_matched'],2);self.assertEqual([row['values'][0] for row in r['devices'][0]['rows']],[DATE,DATE]);self.assertNotIn('OLD',json.dumps(r))
    def test_hidden_device_rows_are_included_explicit_helper_sheet_excluded(self):
        w,l=workbook([('801',[[DATE,801,None,'HIDDEN',None]]),('helper',[[DATE,0,None,'HELPER',None]])]);w['801'].row_dimensions[3].hidden=True;w['801'].sheet_state='hidden';l['sheets']['helper']['exclude']=True
        r=report.scan(package(w),DATE,layout=l);self.assertEqual(r['devices_matched'],1);self.assertEqual(r['excluded_sheets'],['helper'])
    def test_exact_repeated_headers_skip_without_record_creation(self):
        r=self.extract([('801',[[DATE,801,None,'FIRST',None],HEAD,[DATE,801,None,'SECOND',None]])]);self.assertEqual(r['rows_matched'],2)
    def test_unexpected_formula_header_identity_extra_column_unknown_sheet_fail(self):
        for mutation in ['formula','header','title','column','sheet']:
            w,l=workbook([('801',[[DATE,801,None,'FIRST',None]])]);s=w['801']
            if mutation=='formula':s['E3']='=1+1'
            if mutation=='header':s['C2']='غير معروف'
            if mutation=='title':s['A1']='OTHER'
            if mutation=='column':s['F3']='SECRET'
            if mutation=='sheet':w.create_sheet('unexpected')
            with self.subTest(mutation=mutation),self.assertRaises(ValueError):report.scan(package(w),DATE,layout=l)
    def test_only_explicit_placeholder_is_skipped(self):
        w,l=workbook([('801',[[None,None,None,'.',None],[DATE,801,None,'FIRST',None]])]);l['sheets']['801']['undated_placeholders']=[{'column':4,'value':'.'}]
        r=report.scan(package(w),DATE,layout=l);self.assertEqual(r['rows_matched'],1);self.assertEqual(r['undated_placeholders_skipped'],1)
    def test_real_manifest_identity_and_exclusions(self):
        l=json.loads(report.LAYOUT.read_text(encoding='utf-8'))['sheets'];self.assertEqual(len(l),56)
        self.assertEqual([n for n,r in l.items() if r['exclude']],['متفرقه','تعمیرگاه']);self.assertEqual(l['705']['device'],'705');self.assertEqual(l['S1']['device'],'S1')

    def test_fixed_report_order_for_different_source_column_orders(self):
        w,l=workbook([('801',[[DATE,801,'رحمان','FIRST','فیلتر']]),('1252',[[DATE,'شایسته',1252,'SECOND',None]])])
        source_headers=['تاریخ','تعمیرکار','کد مکانیزم','نوع خرابی','قطعات مصرفی']
        for i,label in enumerate(source_headers,1):w['1252'].cell(2,i,label)
        l['sheets']['1252']['columns']=[[i+1,label] for i,label in enumerate(source_headers)]
        r=report.scan(package(w),DATE,layout=l)
        self.assertEqual([d['columns'] for d in r['devices']],[report.REPORT_COLUMNS,report.REPORT_COLUMNS])
        self.assertEqual(r['devices'][0]['rows'][0]['values'],[DATE,'801','رحمان','FIRST','فیلتر'])
        self.assertEqual(r['devices'][1]['rows'][0]['values'],[DATE,'1252','شایسته','SECOND',''])

    def test_explicit_mappings_correct_reversed_headers_without_guessing_values(self):
        w,l=workbook([('802',[[DATE,'شاهین',802,'REPAIR_802','PART_802']]),('705',[[DATE,705,'رحمان-جباری','REPAIR_705','PART_705']])])
        reversed_headers=['تاریخ','تعمیرکار','کد مکانیزم','نوع خرابی','قطعات مصرفی']
        for i,label in enumerate(reversed_headers,1):w['705'].cell(2,i,label)
        l['sheets']['705']['columns']=[[i+1,label] for i,label in enumerate(reversed_headers)]
        l['sheets']['802']['report_columns']=[1,3,2,4,5]
        l['sheets']['705']['report_columns']=[1,2,3,4,5]
        r=report.scan(package(w),DATE,layout=l)
        self.assertEqual(r['devices'][0]['rows'][0]['values'],[DATE,'802','شاهین','REPAIR_802','PART_802'])
        self.assertEqual(r['devices'][1]['rows'][0]['values'],[DATE,'705','رحمان-جباری','REPAIR_705','PART_705'])
        # An alphanumeric code and a missing name remain source values; no invention.
        w['802'].cell(3,3,'MZ10');w['802'].cell(3,2).value=None
        r=report.scan(package(w),DATE,layout=l)
        self.assertEqual(r['devices'][0]['rows'][0]['values'][1:3],['MZ10',''])

    def test_invalid_explicit_map_does_not_duplicate_drop_or_move_date(self):
        for mapping in [[1,2,2,4,5],[1,2,3,4,6],[2,1,3,4,5],[1,2,3,4],[True,2,3,4,5]]:
            w,l=workbook([('801',[[DATE,801,'رحمان','REPAIR','PART']])]);l['sheets']['801']['report_columns']=mapping
            with self.subTest(mapping=mapping),self.assertRaises(ValueError):report.scan(package(w),DATE,layout=l)

class MaintenanceFixture(AdminFixture):
    def setUp(self):
        super().setUp();migrate(self.auth,self.directory/'backups',assignments=['101','202'])
        self.r={'date':DATE,'devices':[{'device':'دستگاه 801','sheet':'801','identity_note':'','columns':HEAD,'rows':[{'source_row':3,'values':[DATE,'801','رحمان','تعویض فیلتر روغن','فیلتر'] }]}], 'devices_matched':1,'rows_matched':1,'sheets_scanned':2,'source_sha256':'fixture'}
        self.outputs=[];self.sender=Mock()
    def fake_build(self,date,directory):
        self.assertEqual(date,DATE);p=Path(directory)/'report.pdf';p.write_bytes(b'%PDF fixture');self.outputs.append(p)
        return {'status':'ready','report':self.r,'pdf':str(p)}
    def run_delivery(self,builder=None):
        with patch.object(delivery,'build',side_effect=builder or self.fake_build) as b,patch.object(delivery,'validate_pdf',return_value={'pages':1,'bytes':12}),patch.object(delivery,'verify_source_digest'),patch('tools.scheduler.tasks.BaleSender',return_value=self.sender):
            r=maintenance_daily_report(None,{'date':DATE},authorization=self.auth)
            return r,b

class AuthorizationTests(MaintenanceFixture,unittest.TestCase):
    def test_receive_only_role_no_profile_scope_or_extra_privilege(self):
        self.assertEqual(self.auth.capabilities_for_role(ROLE),(CAPABILITY,));self.assertEqual(self.sql('SELECT resource FROM auth_capabilities WHERE capability=?',(CAPABILITY,)),[(RESOURCE,)])
        self.assertEqual(self.sql('SELECT * FROM auth_role_profiles WHERE role=?',(ROLE,)),[])
        for u in ['101','202']:
            self.assertTrue(self.auth.has_capability(u,CAPABILITY));self.assertIsNone(self.auth.resolve_profile(u,u));self.assertFalse(self.auth.function_scope(u))
        self.assertFalse(self.auth.has_capability('303',CAPABILITY))
    def test_idempotent_migration_and_existing_profiles_capabilities_unchanged(self):
        self.assign();before=self.auth.resolve_profile('101','101');scope=self.auth.function_scope('101');events=self.sql('SELECT * FROM auth_events')
        migrate(self.auth,self.directory/'backups',assignments=['101','202','101']);self.assertEqual(self.sql('SELECT * FROM auth_events'),events);self.assertEqual(self.auth.resolve_profile('101','101'),before);self.assertEqual(self.auth.function_scope('101'),scope)
    def test_conflicting_profile_extra_grant_or_resource_rolls_back(self):
        for stmt in ["INSERT INTO auth_role_profiles VALUES ('maintenance_daily_report_recipient','admin',100,1)","INSERT INTO auth_role_capabilities VALUES ('maintenance_daily_report_recipient','function.read_all')","UPDATE auth_capabilities SET resource='E:\\Function' WHERE capability='reports.maintenance.daily_receive'"]:
            with self.subTest(stmt=stmt):
                self.sql(stmt)
                with self.assertRaises(ValueError):migrate(self.auth,self.directory/'backups',assignments=['303'])
                self.sql("DELETE FROM auth_role_profiles WHERE role=?",(ROLE,));self.sql('DELETE FROM auth_role_capabilities WHERE role=? AND capability!=?',(ROLE,CAPABILITY));self.sql('UPDATE auth_capabilities SET resource=? WHERE capability=?',(RESOURCE,CAPABILITY))
    def test_pending_identity_refused_transaction_rolled_back(self):
        before=self.sql('SELECT * FROM auth_user_roles')
        with self.assertRaises(ValueError):migrate(self.auth,self.directory/'backups',assignments=['303'])
        self.assertEqual(before,self.sql('SELECT * FROM auth_user_roles'))
    def test_canonical_resolver_dedupe_private_approval_and_revocation(self):
        self.assertEqual(self.auth.resolve_active_recipients(CAPABILITY).recipients,('101','202'))
        self.sql("UPDATE channel_users SET chat_id='999' WHERE user_id='202'");self.assertEqual(self.auth.resolve_active_recipients(CAPABILITY).recipients,('101',))
        self.sql("UPDATE channel_users SET registration_status='revoked' WHERE user_id='101'");self.assertEqual(self.auth.resolve_active_recipients(CAPABILITY).recipients,())

class DeliveryTests(MaintenanceFixture,unittest.TestCase):
    def test_multi_recipient_one_build_and_cleanup(self):
        r,b=self.run_delivery();self.assertEqual(r['sent_count'],2);self.assertEqual(b.call_count,1);self.assertEqual([c.args[0] for c in self.sender.document.call_args_list],['101','202']);self.assertTrue(all(not p.exists() for p in self.outputs))
    def test_zero_recipients_no_generation_or_send(self):
        self.sql('UPDATE auth_user_roles SET active=0');r,b=self.run_delivery();self.assertEqual(r['status'],'skipped');b.assert_not_called();self.sender.document.assert_not_called()
    def test_revoke_before_generation(self):
        real=self.auth.resolve_active_recipients;count=0
        def resolve(cap):
            nonlocal count
            count+=1
            if count==2:self.sql('UPDATE auth_user_roles SET active=0')
            return real(cap)
        with patch.object(self.auth,'resolve_active_recipients',side_effect=resolve):r,b=self.run_delivery()
        self.assertEqual(r['reason'],'recipient_revoked_before_generation');b.assert_not_called()
    def test_revoke_during_generation_and_after_sender_init(self):
        def builder(date,directory):
            r=self.fake_build(date,directory);self.sql("UPDATE auth_user_roles SET active=0 WHERE user_id='202'");return r
        r,b=self.run_delivery(builder);self.assertEqual(r['sent_count'],1);self.assertEqual(r['revoked_count'],1)
        self.sender.reset_mock()
        def factory():self.sql('UPDATE auth_user_roles SET active=0');return self.sender
        with patch.object(delivery,'build',side_effect=self.fake_build),patch.object(delivery,'validate_pdf',return_value={'pages':1,'bytes':12}),patch.object(delivery,'verify_source_digest'),patch('tools.scheduler.tasks.BaleSender',side_effect=factory):
            r=maintenance_daily_report(None,{'date':DATE},authorization=self.auth)
        self.assertEqual(r['sent_count'],0);self.sender.document.assert_not_called()
    def test_failure_does_not_mark_all_successful_or_retry(self):
        self.sender.document.side_effect=[RuntimeError('uncertain'),None];r,b=self.run_delivery();self.assertEqual(r['status'],'failed');self.assertEqual(r['sent_count'],1);self.assertEqual(r['failed_count'],1);self.assertEqual(self.sender.document.call_count,2);self.assertTrue(all(not p.exists() for p in self.outputs))
    def test_zero_records_no_empty_pdf_send(self):
        r,b=self.run_delivery(lambda date,directory:{'status':'skipped','reason':'no_target_date_records','report':dict(self.r,devices=[],devices_matched=0,rows_matched=0)})
        self.assertEqual(r['status'],'waiting_for_data');self.assertEqual(r['reason'],'no_target_date_records');self.sender.message.assert_called();self.sender.document.assert_not_called()
    def test_invalid_pdf_date_source_or_generation_failure_never_send_cleanup(self):
        for failure in ['date','pdf','source','generation']:
            with self.subTest(failure=failure):
                self.sender.reset_mock()
                def builder(date,directory):
                    result=self.fake_build(date,directory)
                    if failure=='date':result['report']=dict(self.r,date='1405/07/12')
                    if failure=='generation':raise ValueError('failed')
                    return result
                with patch.object(delivery,'build',side_effect=builder),patch.object(delivery,'validate_pdf',side_effect=ValueError('invalid') if failure=='pdf' else None,return_value={'pages':1,'bytes':12}),patch.object(delivery,'verify_source_digest',side_effect=ValueError('changed') if failure=='source' else None),patch('tools.scheduler.tasks.BaleSender',return_value=self.sender):
                    with self.assertRaises(ValueError):maintenance_daily_report(None,{'date':DATE},authorization=self.auth)
                self.sender.document.assert_not_called();self.assertTrue(all(not p.exists() for p in self.outputs))
    def test_exact_required_date_occurrence_no_fallback_or_source_override(self):
        for params in [{},{'date':DATE,'source':'other'},{'date':DATE,'latest':True},{'date':DATE,'occurrence':'2026-10-06T09:00:00'},{'date':DATE,'occurrence':'2026-10-05T09:00:00+03:30'}]:
            with self.subTest(params=params),self.assertRaises(ValueError):maintenance_daily_report(None,params,authorization=self.auth)

class SchedulerTests(MaintenanceFixture,unittest.TestCase):
    def config(self,job=None):
        job=job or {'id':'maintenance_daily_report','enabled':True,'task':'maintenance_daily_report','recipient_capability':CAPABILITY,'timezone':'Asia/Tehran','trigger':{'type':'cron','hour':9,'minute':0},'params':{}}
        p=self.directory/'schedule.yaml';p.write_text(yaml.safe_dump({'version':1,'timezone':'Asia/Tehran','misfire_grace_seconds':900,'schedules':[job]}),encoding='utf-8');return p,job
    def test_daily_previous_tehran_occurrence_and_dispatch_dedupe(self):
        p,job=self.config();handler=Mock(return_value={'status':'succeeded'});registry={'maintenance_daily_report':(delivery.validate_params,handler)};state=self.directory/'runs.db';now=datetime(2026,10,6,9,5,tzinfo=ZoneInfo('Asia/Tehran'))
        tick(p,state,now=now,registry=registry,authorization=self.auth);tick(p,state,now=now,registry=registry,authorization=self.auth)
        self.assertEqual(handler.call_count,1);self.assertEqual(handler.call_args.args[1],{'date':DATE,'occurrence':'2026-10-06T09:00:00+03:30'})
        _,_,jobs=load_config(p,registry);due=jobs[0]['trigger'].get_next_fire_time(None,now);self.assertEqual(due.hour,9);self.assertEqual(due.day,7)
        self.assertEqual(overflow_report_date(datetime(2026,10,5,20,29,tzinfo=timezone.utc)),'1405/07/12')
        self.assertEqual(overflow_report_date(datetime(2026,10,5,20,30,tzinfo=timezone.utc)),DATE)
    def test_duplicate_wrong_time_profile_fixed_recipient_or_params_refused(self):
        p,base=self.config()
        cases=[dict(base,trigger={'type':'cron','hour':10,'minute':0}),dict(base,recipient='101'),dict(base,recipient_role='business_admin'),dict(base,params={'date':DATE}),dict(base,timezone='UTC'),dict(base,id='other')]
        for job in cases:
            with self.subTest(job=job):
                p,_=self.config(job)
                with self.assertRaises(ValueError):load_config(p)
        doc=yaml.safe_load(p.read_text());doc['schedules']=[base,dict(base,id='other')];p.write_text(yaml.safe_dump(doc),encoding='utf-8')
        with self.assertRaises(ValueError):load_config(p)
    def test_failed_dispatch_persisted_and_not_replayed(self):
        p,_=self.config();handler=Mock(return_value={'status':'failed','reason':'delivery_failed'});registry={'maintenance_daily_report':(delivery.validate_params,handler)};state=self.directory/'runs.db';now=datetime(2026,10,6,9,5,tzinfo=ZoneInfo('Asia/Tehran'))
        self.assertEqual(tick(p,state,now=now,registry=registry,authorization=self.auth),1);self.assertEqual(tick(p,state,now=now,registry=registry,authorization=self.auth),0);self.assertEqual(handler.call_count,1)
        with closing(sqlite3.connect(state)) as c:self.assertEqual(c.execute('SELECT status,error FROM runs').fetchall(),[('failed','delivery_failed')])

class PDFTests(AdminFixture,unittest.TestCase):
    def test_offline_combined_a4_searchable_vector_multiple_pages_source_unchanged(self):
        w,l=workbook([('801',[[DATE,801,'رحمان',('هیدرولیک گلدسته فیلتر قالب ماشین‌آلات '*5)+str(i),'قطعه'] for i in range(13)]),('802',[[DATE,802,'شاهین','تعویض قطعه',None]])]);data=package(w);h=hashlib.sha256(data).hexdigest();r=report.scan(data,DATE,layout=l);p=self.directory/'combined.pdf'
        with patch('requests.sessions.Session.request',side_effect=AssertionError('Network forbidden')):v=report.render(r,p,test_sample=True)
        self.assertGreater(v['pages'],1);self.assertEqual(v['devices'],2);self.assertEqual(v['rows'],14);self.assertTrue(v['searchable']);self.assertEqual(hashlib.sha256(data).hexdigest(),h)
        with fitz.open(p) as d:
            self.assertTrue(any(page.search_for('هیدرولیک') for page in d));self.assertTrue(all(page.get_text() for page in d));self.assertTrue(all(not page.get_images() for page in d))
        with self.assertRaises(ValueError):report.validate_pdf(p,dict(r,date='1405/07/12'))
    def test_empty_pdf_prohibited_and_safe_source_snapshot_hash(self):
        with self.assertRaises(ValueError):report.render({'devices':[],'rows_matched':0},self.directory/'empty.pdf')
        p=self.data/report.SOURCE_NAME;p.write_bytes(b'fixture')
        with report.source_snapshot(self.reader) as (data,h):self.assertEqual(data,b'fixture');self.assertEqual(h,hashlib.sha256(b'fixture').hexdigest())
        self.assertEqual(p.read_bytes(),b'fixture')
        if __import__('os').name=='nt':
            with report.source_snapshot(self.reader):
                with self.assertRaises(OSError):p.write_bytes(b'writer')
    def test_scheduled_code_no_model_web_delegation_or_subprocess(self):
        for module in [report,delivery]:
            text=Path(module.__file__).read_text(encoding='utf-8-sig');tree=ast.parse(text)
            names=[n.func.id for n in ast.walk(tree) if isinstance(n,ast.Call) and isinstance(n.func,ast.Name)]
            self.assertFalse(set(names)&{'delegate_task','web_search','web_extract','execute_code','run_agent'})
            imports=[n.module or '' for n in ast.walk(tree) if isinstance(n,ast.ImportFrom)]
            imports += [a.name for n in ast.walk(tree) if isinstance(n,ast.Import) for a in n.names]
            self.assertFalse(any(x.split('.')[0] in {'openai','anthropic','litellm','subprocess','model_tools','agent'} for x in imports))
if __name__=='__main__':unittest.main()

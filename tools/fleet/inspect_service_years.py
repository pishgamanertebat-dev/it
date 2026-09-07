import re
import sqlite3
from collections import Counter
from openpyxl import load_workbook


DB = r"E:\KomatsoAI\data\fleet\db\fleet_ops.db"

DATE_RE = re.compile(
    r"(?<!\d)"
    r"(13\d{2}|14\d{2})"
    r"[/-]"
    r"(\d{1,2})"
    r"[/-]"
    r"(\d{1,2})"
    r"(?!\d)"
)


def extract_date(value):
    if value is None:
        return None

    text = str(value).strip()
    m = DATE_RE.search(text)

    if not m:
        return None

    year = int(m.group(1))
    month = int(m.group(2))
    day = int(m.group(3))

    if not (1 <= month <= 12):
        return None

    if not (1 <= day <= 31):
        return None

    return f"{year:04d}/{month:02d}/{day:02d}"


conn = sqlite3.connect(DB)

row = conn.execute("""
SELECT id, raw_path
FROM ingest_sources
WHERE source_type='service_events'
ORDER BY id DESC
LIMIT 1
""").fetchone()

if not row:
    raise SystemExit("ERROR: service_events source not found")

source_id, raw_path = row

wb = load_workbook(
    raw_path,
    data_only=True,
    read_only=True
)

print("SERVICE YEAR AUDIT")
print("=" * 78)
print(f"SOURCE ID: {source_id}")
print(f"FILE: {raw_path}")
print()


# ------------------------------------------------------------
# Detailed machine sheets
# ------------------------------------------------------------

master_sheets = {
    "ساعت کاری",
    "ساعت کاری (2)",
    "Sheet1",
}

overall_years = Counter()

print("PER-MACHINE SHEET DATE COVERAGE")
print("=" * 78)

for ws in wb.worksheets:

    if ws.title in master_sheets:
        continue

    dates = []

    # Per-machine service sheets use column B as date.
    for row_no, values in enumerate(
        ws.iter_rows(
            min_row=4,
            values_only=True
        ),
        start=4
    ):
        if len(values) < 2:
            continue

        d = extract_date(values[1])

        if d:
            dates.append(d)

    years = Counter(
        d[:4]
        for d in dates
    )

    overall_years.update(years)

    if dates:
        print(
            f"{ws.title:<22} | "
            f"rows={len(dates):3d} | "
            f"years={dict(sorted(years.items()))} | "
            f"min={min(dates)} | "
            f"max={max(dates)}"
        )
    else:
        print(
            f"{ws.title:<22} | "
            f"NO DATED SERVICE ROWS"
        )


print()
print("OVERALL DETAILED SERVICE YEARS")
print("=" * 78)

for year, count in sorted(overall_years.items()):
    print(f"{year}: {count} dated rows")


# ------------------------------------------------------------
# Master sheets:
# find any visible 1403 / 1404 / 1405 markers
# ------------------------------------------------------------

for sheet_name in [
    "ساعت کاری (2)",
    "ساعت کاری",
]:
    if sheet_name not in wb.sheetnames:
        continue

    ws = wb[sheet_name]

    print()
    print("=" * 78)
    print(f"MASTER SHEET: {sheet_name}")
    print(
        f"SIZE: rows={ws.max_row}, "
        f"cols={ws.max_column}"
    )

    year_hits = {
        "1403": [],
        "1404": [],
        "1405": [],
    }

    full_dates = []

    for row in ws.iter_rows():

        for cell in row:

            value = cell.value

            if value is None:
                continue

            text = str(value)

            for year in year_hits:
                if year in text:
                    year_hits[year].append(
                        (cell.coordinate, text)
                    )

            d = extract_date(value)

            if d:
                full_dates.append(
                    (cell.coordinate, d, text)
                )

    for year in [
        "1403",
        "1404",
        "1405",
    ]:
        hits = year_hits[year]

        print()
        print(
            f"{year} MARKERS: {len(hits)}"
        )

        for coordinate, value in hits[:20]:
            print(
                f"  {coordinate}: {value}"
            )

        if len(hits) > 20:
            print(
                f"  ... +{len(hits)-20} more"
            )

    print()
    print(
        f"FULL JALALI DATE CELLS: "
        f"{len(full_dates)}"
    )

    if full_dates:
        normalized_dates = [
            item[1]
            for item in full_dates
        ]

        print(
            f"MIN DATE: "
            f"{min(normalized_dates)}"
        )

        print(
            f"MAX DATE: "
            f"{max(normalized_dates)}"
        )

        print("LATEST DATE CELLS:")

        for coordinate, d, raw in sorted(
            full_dates,
            key=lambda x: x[1],
            reverse=True
        )[:20]:
            print(
                f"  {coordinate}: "
                f"{d} | raw={raw}"
            )


wb.close()
conn.close()

print()
print("=" * 78)
print("AUDIT COMPLETE")